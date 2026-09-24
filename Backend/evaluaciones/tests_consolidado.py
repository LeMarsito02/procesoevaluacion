"""Pruebas del informe consolidado (motor/consolidado) con datos inventados."""
from __future__ import annotations

import io

import openpyxl
from django.test import SimpleTestCase

from motor.consolidado import generar_informe
from motor.esquemas.proceso import ResultadoRequisito
from motor.tecnica.informe import PENDIENTE, ResultadoInforme

LOTES = [(0, "LOTE 1")]
PROPONENTES = [("P-01", "CONSORCIO UNO"), ("P-02", "CONSORCIO DOS"), ("P-03", "CONSORCIO TRES")]
# Jurídica: requisitos 1 y 2; técnica: experiencia 101 y factores; financiera: 201 y 221.
REQUISITOS = {
    "juridica": {"generales": [1, 2]},
    # Los factores de puntaje (121 a 130) no habilitan: solo suman puntos.
    "tecnica": {"generales": [], "lote_0": [101]},
    "financiera": {"generales": [201], "lote_0": [221]},
}
PUNTOS = {121: 5, 122: 5, 123: 20, 124: 20, 125: 1, 126: 0.25, 127: 0.25}


def _r(hoja: str, numero: int, cumple: bool | None, decidido: bool = False, detalle: dict | None = None) -> ResultadoInforme:
    return ResultadoInforme(
        ResultadoRequisito(hoja=hoja, numero_orden=1, nombre_proponente=hoja, requisito=numero,
                           cumple=bool(cumple), detalle=detalle),
        decidido,
    )


def _area_completa(hoja: str, numeros: list[int], cumple: bool = True) -> dict:
    return {(hoja, n): _r(hoja, n, cumple, decidido=not cumple) for n in numeros}


class ConsolidadoTests(SimpleTestCase):
    def _informe(self, resultados):
        contenido = generar_informe("PRUEBA-1", "Objeto de prueba", LOTES, PROPONENTES, resultados, REQUISITOS, borrador=True)
        libro = openpyxl.load_workbook(io.BytesIO(contenido))
        return libro["Consolidado"], libro["Pendientes"]

    def test_puntaje_y_orden_de_elegibilidad(self):
        resultados = {"juridica": {}, "tecnica": {}, "financiera": {}}
        for hoja, _ in PROPONENTES:
            resultados["juridica"].update(_area_completa(hoja, [1, 2]))
            resultados["financiera"].update(_area_completa(hoja, [201, 221]))
            resultados["tecnica"].update(_area_completa(hoja, [101]))
            for n, puntos in PUNTOS.items():
                # P-02 no acredita discapacidad ni mujeres: menos puntaje.
                otorga = not (hoja == "P-02" and n in (125, 126))
                resultados["tecnica"][(hoja, n)] = _r(hoja, n, otorga, decidido=not otorga, detalle={"puntaje_maximo": puntos})
            resultados["tecnica"][(hoja, 130)] = _r(hoja, 130, True)  # sin obras inconclusas
        # P-03 no cumple la financiera: no entra en el orden.
        resultados["financiera"][("P-03", 221)] = _r("P-03", 221, False, decidido=True)
        consolidado, _ = self._informe(resultados)
        filas = {f[0].value: [c.value for c in f] for f in consolidado.iter_rows(min_row=6, max_row=8)}
        self.assertEqual(filas["P-01"][2:], ["CUMPLE", "CUMPLE", "CUMPLE", "CUMPLE", 51.5, 1])
        self.assertEqual(filas["P-02"][2:], ["CUMPLE", "CUMPLE", "CUMPLE", "CUMPLE", 50.25, 2])
        self.assertEqual(filas["P-03"][2:6], ["CUMPLE", "CUMPLE", "NO CUMPLE", "NO CUMPLE"])

    def test_lo_pendiente_no_recibe_orden_y_se_lista(self):
        resultados = {"juridica": {}, "tecnica": {}, "financiera": {}}
        for hoja, _ in PROPONENTES:
            resultados["juridica"].update(_area_completa(hoja, [1, 2]))
            resultados["financiera"].update(_area_completa(hoja, [201, 221]))
            resultados["tecnica"].update(_area_completa(hoja, [101]))
            for n, puntos in PUNTOS.items():
                resultados["tecnica"][(hoja, n)] = _r(hoja, n, True, detalle={"puntaje_maximo": puntos})
            resultados["tecnica"][(hoja, 130)] = _r(hoja, 130, True)
        # A P-02 le falta que una persona revise su experiencia.
        resultados["tecnica"][("P-02", 101)] = _r("P-02", 101, False, decidido=False)
        consolidado, pendientes = self._informe(resultados)
        filas = {f[0].value: [c.value for c in f] for f in consolidado.iter_rows(min_row=6, max_row=8)}
        self.assertEqual(filas["P-02"][3], PENDIENTE)
        self.assertEqual(filas["P-02"][5], PENDIENTE)
        self.assertEqual(filas["P-02"][7], PENDIENTE)
        self.assertEqual(filas["P-01"][7], "1 (empate entre 2)")
        self.assertEqual(filas["P-03"][7], "1 (empate entre 2)")
        self.assertIn("P-02", [f[0].value for f in pendientes.iter_rows(min_row=2)])

    def test_los_empatados_comparten_puesto_y_el_informe_lo_dice(self):
        """En ICCU-LP-027 más de veinte proponentes quedaron con 51,50 puntos:
        el empate es la norma. Inventar un orden entre ellos sería afirmar algo
        que el programa no sabe; el desempate lo resuelve la entidad con los
        criterios del pliego."""
        resultados = {"juridica": {}, "tecnica": {}, "financiera": {}}
        for hoja, _ in PROPONENTES:
            resultados["juridica"].update(_area_completa(hoja, [1, 2]))
            resultados["financiera"].update(_area_completa(hoja, [201, 221]))
            resultados["tecnica"].update(_area_completa(hoja, [101]))
            for n, puntos in PUNTOS.items():
                resultados["tecnica"][(hoja, n)] = _r(hoja, n, True, detalle={"puntaje_maximo": puntos})
            resultados["tecnica"][(hoja, 130)] = _r(hoja, 130, True)
        consolidado, _ = self._informe(resultados)
        filas = {f[0].value: [c.value for c in f] for f in consolidado.iter_rows(min_row=6, max_row=8)}
        ordenes = {filas[h][7] for h, _ in PROPONENTES}
        self.assertEqual(ordenes, {"1 (empate entre 3)"})


class ConsolidadoNoAplicaTests(SimpleTestCase):
    def test_el_patrimonio_que_no_aplica_no_vuelve_na_el_lote(self):
        resultados = {"juridica": {}, "tecnica": {}, "financiera": {}}
        hoja = "P-01"
        resultados["juridica"] = _area_completa(hoja, [1, 2])
        resultados["tecnica"] = {(hoja, 101): _r(hoja, 101, True)}
        for n, puntos in PUNTOS.items():
            resultados["tecnica"][(hoja, n)] = _r(hoja, n, True, detalle={"puntaje_maximo": puntos})
        resultados["tecnica"][(hoja, 130)] = _r(hoja, 130, True)
        resultados["financiera"] = {
            (hoja, 201): _r(hoja, 201, True),
            (hoja, 204): _r(hoja, 204, True, detalle={"financiera": {"no_aplica": True}}),
            (hoja, 221): _r(hoja, 221, False, decidido=True),
        }
        requisitos = {**REQUISITOS, "financiera": {"generales": [201, 204], "lote_0": [221]}}
        contenido = generar_informe("PRUEBA-1", "Objeto", LOTES, [(hoja, "UNO")], resultados, requisitos, borrador=True)
        libro = openpyxl.load_workbook(io.BytesIO(contenido))["Consolidado"]
        fila = [c.value for c in libro[6]]
        self.assertEqual(fila[4], "NO CUMPLE")  # financiera
        self.assertEqual(fila[5], "NO CUMPLE")  # habilitado
