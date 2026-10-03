"""Acta de revisión cruzada de las evaluaciones de un proceso.

Antes de adjudicar, los comités jurídico, técnico y financiero revisan juntos
las tres evaluaciones y dejan constancia firmada (plan de mejoramiento del
ICCU, auditoría a la etapa precontractual, hallazgo 3: dos adjudicaciones
revocadas por un cálculo con proponentes que no cumplían y por adjudicar a
quien no tenía el mayor puntaje).

El acta junta en un documento controlado (ISO 9001, 7.5):
- el estado de cada evaluación, su dependencia y su comité;
- el resultado consolidado por lote, con las mismas cifras del informe
  consolidado (habilitación en las tres áreas, puntaje y orden de
  elegibilidad entre los habilitados);
- los puntos de control que la revisión debe confirmar, con lo que el sistema
  ya puede verificar marcado como tal;
- espacio para observaciones y las firmas de todos los integrantes.

La genera el jefe de una de las dependencias del proceso o el administrador.
"""
from __future__ import annotations

import hashlib
import io
from datetime import datetime

from django.conf import settings
from django.core.files.base import ContentFile
from django.db import IntegrityError, transaction
from django.utils import timezone

from cuentas.models import Usuario
from evaluaciones.models import ActaRevisionCruzada, EstadoEvaluacion, Proceso, siguiente_consecutivo

CODIGO_FORMATO = "MEV-FOR-RCR-01"
VERSION_FORMATO = "1"
ORDEN_AREAS = ("juridica", "tecnica", "financiera")


class ErrorActa(ValueError):
    pass


def puede_generar(usuario: Usuario, proceso: Proceso) -> bool:
    """Quien gestiona alguna de las evaluaciones del proceso: el jefe de su
    dependencia (o del área) o el administrador."""
    from evaluaciones.permisos import puede_gestionar

    return any(puede_gestionar(usuario, e) for e in proceso.evaluaciones.select_related("dependencia"))


def _resumen(proceso: Proceso) -> dict:
    from evaluaciones.estructura import comite_activo
    from evaluaciones.servicios import avances

    evaluaciones = {e.tipo: e for e in proceso.evaluaciones.select_related("dependencia", "aprobada_por")}
    avance = avances([e.id for e in evaluaciones.values()])
    areas = []
    for tipo in ORDEN_AREAS:
        e = evaluaciones.get(tipo)
        if e is None:
            continue
        a = avance[e.id]
        areas.append({
            "tipo": tipo,
            "nombre": e.get_tipo_display(),
            "dependencia": e.dependencia.nombre if e.dependencia_id else "—",
            "estado": e.get_estado_display(),
            "aprobada": e.estado == EstadoEvaluacion.APROBADA,
            "aprobada_por": e.aprobada_por.nombre_completo if e.aprobada_por_id else None,
            "aprobada_en": e.aprobada_en.isoformat() if e.aprobada_en else None,
            "evaluados": a.evaluados,
            "proponentes": a.proponentes,
            "pendientes": a.pendientes,
            # Sin comité registrado (procesos anteriores), firma el responsable.
            "comite": [{"nombre": u.nombre_completo, "email": u.email}
                       for u in (comite_activo(e) or ([e.responsable] if e.responsable_id else []))],
        })
    return {"areas": areas}


def generar(proceso: Proceso, actor: Usuario) -> ActaRevisionCruzada:
    if not proceso.evaluaciones.exists():
        raise ErrorActa("El proceso no tiene evaluaciones.")
    resumen = _resumen(proceso)
    for _ in range(3):  # dos actas a la vez pueden pedir el mismo consecutivo
        try:
            with transaction.atomic():
                consecutivo = siguiente_consecutivo(proceso.entidad_id, "RCR")
                version = proceso.actas_revision_cruzada.count() + 1
                contenido = documento(proceso, actor, consecutivo, version, resumen)
                acta = ActaRevisionCruzada(
                    entidad_id=proceso.entidad_id, proceso=proceso, consecutivo=consecutivo, version=version,
                    resumen=resumen, generada_por=actor, sha256=hashlib.sha256(contenido).hexdigest(),
                )
                acta.archivo.save(f"{consecutivo}.docx", ContentFile(contenido), save=False)
                acta.save()
                return acta
        except IntegrityError:
            continue
    raise ErrorActa("No se pudo numerar el acta; inténtelo de nuevo.")


def documento(proceso: Proceso, actor: Usuario, consecutivo: str, version: int, resumen: dict) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    from evaluaciones.reporte import _parrafo, _tabla, _titulo, fecha_larga
    from evaluaciones.servicios import datos_consolidado
    from motor.consolidado import tabla_por_lote

    ahora = timezone.now()
    areas = resumen["areas"]
    todas_aprobadas = all(a["aprobada"] for a in areas)
    pendientes = sum(a["pendientes"] for a in areas)
    doc = Document()
    seccion = doc.sections[0]
    # Horizontal: la tabla consolidada tiene ocho columnas.
    seccion.orientation = WD_ORIENT.LANDSCAPE
    seccion.page_width, seccion.page_height = seccion.page_height, seccion.page_width
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)

    tabla = seccion.header.add_table(rows=2, cols=4, width=seccion.page_width - seccion.left_margin - seccion.right_margin)
    tabla.style = "Table Grid"
    for i, (etiqueta, valor) in enumerate([
        ("Entidad", proceso.entidad.nombre), ("Código del formato", CODIGO_FORMATO),
        ("Documento", "Acta de revisión cruzada de las evaluaciones"), ("Versión del formato", VERSION_FORMATO),
    ]):
        for j, texto in enumerate((etiqueta, valor)):
            celda = tabla.cell(i // 2, (i % 2) * 2 + j)
            celda.text = texto
            celda.paragraphs[0].runs[0].font.size = Pt(8)
            celda.paragraphs[0].runs[0].bold = j == 0

    _titulo(doc, "Acta de revisión cruzada de las evaluaciones")
    _tabla(doc, ["Dato", "Valor"], [
        ["Consecutivo", consecutivo],
        ["Versión del acta", f"{version}" + (" (reemplaza el acta anterior del proceso)" if version > 1 else "")],
        ["Fecha", fecha_larga(ahora)],
        ["Proceso", proceso.codigo],
        ["Objeto", proceso.objeto or (proceso.documento_base or {}).get("objeto_general", "")],
        ["Convoca y genera", f"{actor.nombre_completo} ({actor.get_rol_display()})"],
    ])
    _parrafo(doc, "")
    _parrafo(
        doc,
        "Los comités evaluadores jurídico, técnico y financiero se reúnen para revisar de forma cruzada e "
        "interdisciplinaria las evaluaciones del proceso antes de la adjudicación: verifican que el resultado "
        "consolidado, las fórmulas de ponderación y el orden de elegibilidad correspondan a lo evaluado por cada "
        "comité y a lo que establece el pliego de condiciones.",
    )
    if not todas_aprobadas or pendientes:
        _parrafo(
            doc,
            "Atención: al generar esta acta hay evaluaciones sin aprobar o requisitos pendientes de revisión (ver "
            "numeral 1). El resultado consolidado es preliminar y no debe usarse para proyectar la adjudicación.",
            negrita=True,
        )

    _titulo(doc, "1. Estado de las evaluaciones", 2)
    _tabla(doc, ["Área", "Dependencia", "Comité", "Estado", "Evaluados", "Pendientes", "Aprobó"], anchos=[2.2, 4.0, 5.4, 2.6, 2.1, 2.5, 4.0], filas=[
        [
            a["nombre"], a["dependencia"], ", ".join(m["nombre"] for m in a["comite"]) or "Sin designar", a["estado"],
            f"{a['evaluados']}/{a['proponentes']}", str(a["pendientes"]),
            f"{a['aprobada_por']} ({fecha_larga(datetime.fromisoformat(a['aprobada_en']))})" if a["aprobada_por"] else "—",
        ]
        for a in areas
    ])

    _titulo(doc, "2. Resultado consolidado por lote", 2)
    _parrafo(
        doc,
        "Mismas cifras del informe consolidado de MiEvaluador: habilitación en cada área con las decisiones de los "
        "comités, puntaje preliminar y orden de elegibilidad calculado solo entre los proponentes habilitados en las "
        "tres áreas. PENDIENTE indica algo que ni el sistema confirmó ni una persona decidió.",
        tam=9,
    )
    datos = datos_consolidado(proceso)
    for nombre_lote, filas in tabla_por_lote(datos["lotes"], datos["proponentes"], datos["resultados"],
                                             datos["requisitos"], datos["puntajes_adoptados"]):
        _parrafo(doc, nombre_lote.upper(), negrita=True, espacio=2)
        _tabla(doc, ["Prop.", "Proponente", "Jurídica", "Técnica", "Financiera", "Habilitado", "Puntaje", "Orden"],
               anchos=[1.6, 6.6, 2.4, 2.4, 2.4, 2.4, 2.0, 2.9],
               filas=[[str(c) if c != "" else "—" for c in f] for f in filas])
        _parrafo(doc, "")

    _titulo(doc, "3. Puntos de control de la revisión cruzada", 2)
    sistema_orden = "Verificado por MiEvaluador: el orden se calcula solo entre los habilitados en las tres áreas."
    _tabla(doc, ["Punto de control", "Cumple (Sí/No)", "Observación"], anchos=[11.0, 3.0, 8.7], filas=[
        ["Todas las evaluaciones están aprobadas y sin requisitos pendientes.",
         "Sí" if todas_aprobadas and not pendientes else "No", "Según el numeral 1, al generar el acta."],
        ["El orden de elegibilidad incluye únicamente a proponentes habilitados en las tres áreas.", "", sistema_orden],
        ["Las fórmulas de ponderación (media geométrica u otro método del pliego) se aplicaron solo con propuestas "
         "habilitadas y, en experiencia, con contratos válidos no subsanados; los datos de entrada coinciden.", "", ""],
        ["El orden de elegibilidad del informe definitivo publicado coincide con el proyecto de resolución de "
         "adjudicación, lote por lote.", "", ""],
        ["Se verificaron, a la fecha de esta acta, inhabilidades e incompatibilidades (incluidas las sobrevinientes, "
         "como adjudicaciones recientes en otros procesos) de quien ocupa el primer orden de elegibilidad en cada lote.", "", ""],
    ])
    _parrafo(doc, "")
    _titulo(doc, "4. Observaciones y conclusiones", 2)
    for _ in range(4):
        _parrafo(doc, "_" * 110, tam=9)

    _titulo(doc, "5. Firmas de los integrantes de los comités", 2)
    firmantes = [[m["nombre"], a["nombre"], "", ""] for a in areas for m in a["comite"]]
    firmantes.append([actor.nombre_completo, "Convoca", "", ""])
    _tabla(doc, ["Nombre", "Comité", "Firma", "Fecha"], firmantes, anchos=[7.0, 3.5, 8.2, 4.0])
    _parrafo(doc, "")
    _titulo(doc, "Control del documento", 2)
    _parrafo(
        doc,
        f"Generado por MiEvaluador {settings.MIEVALUADOR_VERSION} el {fecha_larga(ahora)}. La huella SHA-256 de este "
        "archivo queda registrada en la plataforma y en la auditoría: cualquier modificación posterior la cambia. Una "
        "nueva revisión genera una nueva versión con nuevo consecutivo; las anteriores se conservan. Conservación según "
        "las tablas de retención documental de la entidad.",
        tam=9,
    )
    pie = seccion.footer.paragraphs[0]
    pie.text = f"{CODIGO_FORMATO} · v{VERSION_FORMATO} · {consecutivo} · {proceso.codigo} · MiEvaluador by LeMarTek"
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()
