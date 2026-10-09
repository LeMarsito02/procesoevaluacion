"""Contratación de prestación de servicios (OPS) en la plataforma.

El trabajador de la fila analiza los documentos (`analizar`) y guarda lo que
leyó. Lo que decide la persona —qué periodos cuentan, cuáles son relacionados,
correcciones al perfil, documentos que revisó— se guarda aparte y se aplica
cada vez que se consulta (`idoneidad`): un análisis nuevo nunca lo pisa.
"""
from __future__ import annotations

import io
import logging
import re
import time
from datetime import date, timedelta

from django.conf import settings
from django.core.files.base import ContentFile
from django.db import transaction
from django.db.models import Q
from django.utils import timezone

from evaluaciones.models import ContratacionOps, DocumentoOps, EstadoOps, OrigenDocumentoOps, TablaHonorariosOps
from motor.ops import documentos as motor_documentos
from motor.ops.expediente import leer_expediente
from motor.ops.hoja_de_vida import EN_SIGEP, ILEGIBLE, NO_ESTA, SIN_HOJA, Declarada, ExperienciaDeclarada, en_el_sigep, misma_entidad
from motor.ops.identidad import identificar
from motor.ops.experiencia import Periodo
from motor.ops.honorarios import ESPECIALIZACION, MAESTRIA, SIN_POSGRADO, FranjaProfesional, FranjaReconocimiento, TablaHonorarios
from motor.ops.idoneidad import Idoneidad, Perfil, evaluar_idoneidad
from motor.ops.tiempo import anios_meses_dias, texto_duracion
from motor.ops.titulos import Titulo
from motor.procesamiento.zip_utils import extraer_pdfs

log = logging.getLogger("mievaluador.ops")
# Un análisis sin terminar en este tiempo se da por caído y vuelve a la fila.
ANALISIS_CAIDO = timedelta(minutes=20)
MAX_DOCUMENTOS = 120
POSGRADOS = {SIN_POSGRADO: "Sin posgrado", ESPECIALIZACION: "Especialización", MAESTRIA: "Maestría o más"}
NIVELES = {"profesional": "Profesional", "tecnologo": "Tecnólogo", "tecnico": "Técnico", "bachiller": "Bachiller", **POSGRADOS}



class ErrorOps(ValueError):
    """Algo que la persona puede corregir (se le muestra tal cual)."""


# --- Documentos --------------------------------------------------------------
def guardar_archivos(contratacion: ContratacionOps, origen: str, archivos: list[tuple[str, bytes]], usuario) -> int:
    """Guarda los PDF subidos; de un .zip se guardan los PDF que trae."""
    pdfs: list[tuple[str, bytes]] = []
    for nombre, contenido in archivos:
        if not contenido:
            continue
        if nombre.lower().endswith(".zip"):
            pdfs.extend((ruta.rsplit("/", 1)[-1], datos) for ruta, datos in extraer_pdfs(contenido).items())
        elif nombre.lower().endswith(".pdf") or contenido[:5] == b"%PDF-":
            pdfs.append((nombre, contenido))
        else:
            raise ErrorOps(f"«{nombre}» no es un PDF ni un .zip.")
    if contratacion.documentos.count() + len(pdfs) > MAX_DOCUMENTOS:
        raise ErrorOps(f"Son demasiados documentos (máximo {MAX_DOCUMENTOS}).")
    for nombre, contenido in pdfs:
        doc = DocumentoOps(
            entidad_id=contratacion.entidad_id, contratacion=contratacion, origen=origen,
            nombre_original=nombre[:255], tamano=len(contenido), subido_por=usuario,
        )
        doc.archivo.save("documento.pdf", ContentFile(contenido), save=True)
    return len(pdfs)


def _leer(doc: DocumentoOps) -> bytes:
    with doc.archivo.open("rb") as f:
        return f.read()


def _nombres_unicos(docs: list[DocumentoOps]) -> dict[str, DocumentoOps]:
    """El motor reconoce cada archivo por su nombre: dos con el mismo nombre se distinguen."""
    por_nombre: dict[str, DocumentoOps] = {}
    for doc in docs:
        nombre, n = doc.nombre_original, 2
        while nombre in por_nombre:
            nombre, n = f"{doc.nombre_original} ({n})", n + 1
        por_nombre[nombre] = doc
    return por_nombre


# --- Análisis (trabajador) ---------------------------------------------------
def analizar(contratacion: ContratacionOps) -> dict:
    """Lee los documentos y devuelve lo que se guarda en `resultado`."""
    docs = list(contratacion.documentos.all())
    de_la_entidad = _nombres_unicos([d for d in docs if d.origen == OrigenDocumentoOps.ENTIDAD])
    del_contratista = _nombres_unicos([d for d in docs if d.origen == OrigenDocumentoOps.CONTRATISTA])
    pdfs = {n: _leer(d) for n, d in del_contratista.items()}

    # El nombre y la cédula salen de los propios antecedentes: si la persona
    # no los escribió se toman de ahí, y si los escribió distinto se avisa.
    nombre_leido, cedula_leida = identificar(pdfs)
    avisos: list[str] = []
    if not contratacion.contratista_nombre or not contratacion.contratista_cedula:
        if not (nombre_leido and cedula_leida):
            raise ErrorOps(
                "No se pudo leer el nombre y la cédula del contratista en sus certificados de antecedentes: escríbalos en los datos de la contratación."
            )
        contratacion.contratista_nombre = contratacion.contratista_nombre or nombre_leido.title()
        contratacion.contratista_cedula = contratacion.contratista_cedula or cedula_leida
        contratacion.save(update_fields=["contratista_nombre", "contratista_cedula"])
    elif cedula_leida and cedula_leida != contratacion.contratista_cedula:
        avisos.append(
            f"Los certificados de antecedentes son de la cédula {cedula_leida} ({nombre_leido or 'sin nombre'}), no de la {contratacion.contratista_cedula}: "
            "revise los datos del contratista."
        )

    exp = leer_expediente(
        {n: _leer(d) for n, d in de_la_entidad.items()}, pdfs,
        contratacion.contratista_nombre, contratacion.contratista_cedula, contratacion.fecha_referencia, contratacion.exige_libreta,
    )

    def doc_id(nombre: str | None, de: dict[str, DocumentoOps] = del_contratista) -> str | None:
        return str(de[nombre].id) if nombre in de else None

    estudio, perfil, m = exp.estudio, exp.estudio.perfil if exp.estudio else None, exp.matricula
    return {
        "estudio": None if estudio is None else {
            "documento_id": doc_id(exp.estudio_archivo, de_la_entidad), "objeto": estudio.objeto, "plazo_meses": estudio.plazo_meses,
            "valor": estudio.valor,
        },
        "cdp": None if exp.cdp is None else {"documento_id": doc_id(exp.cdp_archivo, de_la_entidad), "valor": exp.cdp},
        "perfil": None if perfil is None else {
            "anios_minimos": perfil.anios_minimos, "anios_maximos": perfil.anios_maximos, "posgrado": perfil.posgrado,
            "honorarios_mensuales": perfil.honorarios_mensuales, "obligaciones": perfil.obligaciones,
            "especifica_minima": perfil.especifica_minima, "especifica_maxima": perfil.especifica_maxima, "descripcion": perfil.descripcion,
        },
        "sigep": None if exp.sigep is None else {
            "bloques": exp.sigep.bloques,
            "experiencias": [
                {"entidad": d.entidad, "inicio": d.inicio.isoformat(), "fin": d.fin.isoformat() if d.fin else None, "cargo": d.cargo}
                for d in exp.sigep.experiencias
            ],
        },
        "matricula": None if m is None else {
            "profesion": m.profesion, "numero": m.numero, "fecha": m.fecha.isoformat() if m.fecha else None, "sin_sanciones": m.sin_sanciones,
        },
        "titulos": [
            {"nivel": t.nivel, "nombre": t.nombre, "fecha": t.fecha.isoformat() if t.fecha else None,
             "documento_id": doc_id(t.archivo), "archivo": t.archivo, "pagina": t.pagina, "declarado": t.declarado}
            for t in exp.titulos
        ],
        "periodos": [
            {"id": pid, "inicio": p.inicio.isoformat(), "fin": p.fin.isoformat(), "entidad": p.entidad, "referencia": p.referencia,
             "suspensiones": [[a.isoformat(), b.isoformat()] for a, b in p.suspensiones], "abierto": p.abierto, "nota": p.nota,
             "documento_id": doc_id(p.archivo), "archivo": p.archivo, "pagina": p.pagina, "texto": p.texto}
            for pid, p in zip(_ids_de_periodos(exp.periodos), exp.periodos)
        ],
        "documentos": [
            {"clave": r.clave, "nombre": r.nombre, "estado": r.estado, "motivo": r.motivo, "documento_id": doc_id(r.archivo),
             "archivo": r.archivo, "fecha": r.fecha.isoformat() if r.fecha else None,
             "adicionales": [{"documento_id": doc_id(a), "archivo": a} for a in r.adicionales]}
            for r in exp.documentos
        ],
        "sin_reconocer": [{"documento_id": doc_id(n), "archivo": n} for n in exp.sin_reconocer],
        "avisos": [*avisos, *exp.avisos],
    }


def _ids_de_periodos(periodos: list[Periodo]) -> list[str]:
    """El identificador de un periodo son sus fechas: lo que la persona decidió
    sobre él lo sigue aunque un análisis nuevo lea más periodos o en otro
    orden, y deja de aplicar si sus fechas cambian. Nunca cae en otro periodo."""
    vistos: dict[str, int] = {}
    ids = []
    for p in periodos:
        base = f"p{p.inicio:%Y%m%d}-{p.fin:%Y%m%d}"
        vistos[base] = vistos.get(base, 0) + 1
        ids.append(base if vistos[base] == 1 else f"{base}-{vistos[base]}")
    return ids


def reclamar() -> ContratacionOps | None:
    limite = timezone.now() - ANALISIS_CAIDO
    with transaction.atomic():
        c = (
            ContratacionOps.objects.select_for_update(skip_locked=True)
            .filter(Q(estado=EstadoOps.PENDIENTE) | Q(estado=EstadoOps.ANALIZANDO, analisis_iniciado__lt=limite))
            .order_by("creada_en")
            .first()
        )
        if c is None:
            return None
        c.estado, c.error, c.analisis_iniciado = EstadoOps.ANALIZANDO, "", timezone.now()
        c.save(update_fields=["estado", "error", "analisis_iniciado"])
        return c


def atender_pendientes() -> int:
    """Analiza las contrataciones pendientes (lo llama el trabajador de la fila)."""
    n = 0
    while (c := reclamar()) is not None:
        inicio = time.monotonic()
        try:
            resultado = analizar(c)
            ContratacionOps.objects.filter(pk=c.pk).update(
                estado=EstadoOps.LISTA, resultado=resultado, analizada_en=timezone.now(),
                segundos=round(time.monotonic() - inicio, 1), version_sistema=settings.MIEVALUADOR_VERSION,
            )
            log.info("OPS %s analizada en %.0fs", c.pk, time.monotonic() - inicio)
        except ErrorOps as exc:
            ContratacionOps.objects.filter(pk=c.pk).update(estado=EstadoOps.ERROR, error=str(exc)[:2000])
        except Exception as exc:  # noqa: BLE001
            log.exception("Falló el análisis de la OPS %s", c.pk)
            ContratacionOps.objects.filter(pk=c.pk).update(estado=EstadoOps.ERROR, error=f"No se pudo analizar: {exc}"[:2000])
        n += 1
    return n


def pedir_analisis(contratacion: ContratacionOps) -> None:
    contratacion.estado, contratacion.error = EstadoOps.PENDIENTE, ""
    contratacion.save(update_fields=["estado", "error", "actualizada_en"])


# --- Tabla de honorarios -----------------------------------------------------
def tabla_de(entidad_id, vigencia: int) -> TablaHonorarios | None:
    fila = TablaHonorariosOps.objects.filter(entidad_id=entidad_id, vigencia=vigencia).first()
    if fila is None or not fila.profesional:
        return None
    return TablaHonorarios(
        vigencia=fila.vigencia, norma=fila.norma,
        profesional=[FranjaProfesional(f["desde"], f.get("hasta"), f["sin_especializacion"], f["con_especializacion"], f["con_maestria"]) for f in fila.profesional],
        reconocimiento=[FranjaReconocimiento(f["desde"], f.get("hasta"), f["valor"]) for f in fila.reconocimiento],
    )


def validar_tabla(profesional: list[dict], reconocimiento: list[dict]) -> None:
    def franjas(filas: list[dict], campos: tuple[str, ...], nombre: str) -> None:
        anterior = None
        for f in filas:
            desde, hasta = f.get("desde"), f.get("hasta")
            if not isinstance(desde, (int, float)) or desde < 0 or (hasta is not None and (not isinstance(hasta, (int, float)) or hasta <= desde)):
                raise ErrorOps(f"En {nombre}, cada franja necesita «desde» y un «hasta» mayor (o vacío en la última).")
            if anterior is not None and desde < anterior:
                raise ErrorOps(f"En {nombre}, las franjas deben ir de menor a mayor y sin cruzarse.")
            if hasta is None and f is not filas[-1]:
                raise ErrorOps(f"En {nombre}, solo la última franja puede quedar sin «hasta».")
            anterior = hasta
            if any(not isinstance(f.get(c), int) or f[c] < 0 for c in campos):
                raise ErrorOps(f"En {nombre}, los valores deben ser pesos enteros.")

    if not profesional:
        raise ErrorOps("La tabla necesita al menos una franja de experiencia profesional.")
    franjas(profesional, ("sin_especializacion", "con_especializacion", "con_maestria"), "experiencia profesional")
    franjas(reconocimiento, ("valor",), "reconocimiento por experiencia específica")


# --- Idoneidad con las decisiones de la persona ------------------------------
def perfil_vigente(contratacion: ContratacionOps) -> Perfil | None:
    """El perfil leído del estudio previo con las correcciones de la persona."""
    leido = (contratacion.resultado or {}).get("perfil") or {}
    corregido = {**leido, **(contratacion.decisiones.get("perfil") or {})}
    if corregido.get("anios_minimos") is None:
        return None
    return Perfil(
        anios_minimos=float(corregido["anios_minimos"]),
        anios_maximos=None if corregido.get("anios_maximos") is None else float(corregido["anios_maximos"]),
        posgrado=corregido.get("posgrado") or SIN_POSGRADO,
        honorarios_mensuales=int(corregido.get("honorarios_mensuales") or 0),
        obligaciones=list(leido.get("obligaciones") or []),
        especifica_minima=float(corregido.get("especifica_minima") or 0),
        especifica_maxima=None if corregido.get("especifica_maxima") is None else float(corregido["especifica_maxima"]),
        descripcion=leido.get("descripcion") or "",
    )


def _sigep(contratacion: ContratacionOps) -> ExperienciaDeclarada | None:
    guardado = (contratacion.resultado or {}).get("sigep")
    if guardado is None:
        return None
    return ExperienciaDeclarada(
        experiencias=[
            Declarada(d["entidad"], date.fromisoformat(d["inicio"]), date.fromisoformat(d["fin"]) if d["fin"] else None, d.get("cargo", ""))
            for d in guardado["experiencias"]
        ],
        bloques=guardado["bloques"],
    )


DUDAS_SIGEP = {
    SIN_HOJA: "No está la hoja de vida del SIGEP entre los documentos: confirme que este periodo está relacionado en ella.",
    ILEGIBLE: "La hoja de vida del SIGEP no se dejó leer completa: confirme que este periodo está relacionado en ella.",
}
OTRA_ENTIDAD = "En el SIGEP, en esas fechas figura otra entidad ({entidad}): confirme que es la misma experiencia."
FUERA_DEL_SIGEP = "No está relacionado en la hoja de vida del SIGEP: no es experiencia válida."


def _todos_los_periodos(contratacion: ContratacionOps) -> list[dict]:
    """Los periodos leídos y los que agregó la persona, con sus correcciones aplicadas."""
    decisiones = contratacion.decisiones.get("periodos") or {}
    declarada = _sigep(contratacion)
    periodos = []
    for p in [*((contratacion.resultado or {}).get("periodos") or []), *(contratacion.decisiones.get("agregados") or [])]:
        d = decisiones.get(p["id"]) or {}
        corregido = any(c in d for c in ("inicio", "fin", "entidad", "referencia"))
        periodos.append({
            "suspensiones": [], "abierto": False, "nota": "", "documento_id": None, "archivo": "", "pagina": None, "texto": "",
            **p, **{c: d[c] for c in ("inicio", "fin", "entidad", "referencia") if c in d},
            "decision": d, "manual": p["id"].startswith("m"), "corregido": corregido,
            # Si la persona corrigió las fechas o confirmó la lectura, lo dudoso ya no aplica.
            **({"abierto": False, "nota": ""} if "inicio" in d or "fin" in d or d.get("confirmado") else {}),
            "confirmado": bool(d.get("confirmado")),
            "duda": "" if ("inicio" in d or "fin" in d or d.get("confirmado")) else (p.get("nota") or ("Sin fecha de terminación" if p.get("abierto") else "")),
        })
        # La experiencia que cuenta debe estar relacionada en la hoja de vida del
        # SIGEP. Se cruza con las fechas ya corregidas; lo que la persona vio allí
        # con sus ojos lo marca ella.
        actual = periodos[-1]
        estado, cubren = en_el_sigep(date.fromisoformat(actual["inicio"]), date.fromisoformat(actual["fin"]), declarada, contratacion.fecha_referencia)
        actual["sigep"] = EN_SIGEP if d.get("en_sigep") else estado
        actual["sigep_confirmado"] = bool(d.get("en_sigep"))
        if not d.get("en_sigep"):
            if estado in DUDAS_SIGEP:
                actual["duda"] = (actual["duda"] + " " if actual["duda"] else "") + DUDAS_SIGEP[estado]
            elif estado == EN_SIGEP and misma_entidad(actual["entidad"], cubren) is False:
                actual["duda"] = (actual["duda"] + " " if actual["duda"] else "") + OTRA_ENTIDAD.format(entidad=cubren[0].entidad)
        if estado == EN_SIGEP and not actual["entidad"] and len(cubren) == 1:
            actual["entidad"] = cubren[0].entidad
    return periodos


def _periodos(contratacion: ContratacionOps) -> tuple[list[Periodo], dict[int, dict], list[dict]]:
    """Periodos que cuentan (con lo que la persona corrigió), sus datos y los que retiró."""
    cuentan: list[Periodo] = []
    datos: dict[int, dict] = {}
    retirados: list[dict] = []
    for p in _todos_los_periodos(contratacion):
        d = p["decision"]
        if d.get("incluir") is False:
            retirados.append({**p, "estado": "retirado", "motivo": "Retirado por decisión de la persona."})
            continue
        if p["sigep"] == NO_ESTA:
            retirados.append({**p, "estado": "sin_sigep", "motivo": FUERA_DEL_SIGEP})
            continue
        relacionada_decidida = "relacionada" in d or (p["manual"] and "relacionada" in p)
        periodo = Periodo(
            date.fromisoformat(p["inicio"]), date.fromisoformat(p["fin"]), entidad=p["entidad"], referencia=p["referencia"],
            suspensiones=[(date.fromisoformat(a), date.fromisoformat(b)) for a, b in p["suspensiones"]],
            abierto=p["abierto"], nota=p["nota"], archivo=p["archivo"], pagina=p["pagina"], fijo=d.get("incluir") is True,
            texto="" if relacionada_decidida else p.get("texto", ""),
            relacionada=bool(d["relacionada"] if "relacionada" in d else p.get("relacionada")),
        )
        cuentan.append(periodo)
        datos[id(periodo)] = p
    return cuentan, datos, retirados


def idoneidad(contratacion: ContratacionOps) -> tuple[Idoneidad, dict[int, dict], list[dict]]:
    resultado = contratacion.resultado or {}
    perfil = perfil_vigente(contratacion) or Perfil(0, None)
    periodos, datos, retirados = _periodos(contratacion)
    titulos = [
        Titulo(t["nivel"], t["nombre"], date.fromisoformat(t["fecha"]) if t["fecha"] else None, t["archivo"], t["pagina"], t.get("declarado", False))
        for t in resultado.get("titulos") or []
    ]
    if grado := contratacion.decisiones.get("grado"):
        titulos = [t for t in titulos if t.nivel != "profesional"] + [Titulo("profesional", "", date.fromisoformat(grado))]
    if posgrado := contratacion.decisiones.get("posgrado"):
        titulos = [t for t in titulos if t.nivel not in (ESPECIALIZACION, MAESTRIA)]
        if posgrado != SIN_POSGRADO:
            titulos.append(Titulo(posgrado, "", None))
    tabla = tabla_de(contratacion.entidad_id, contratacion.fecha_referencia.year)
    r = evaluar_idoneidad(periodos, titulos, perfil, tabla)
    if perfil_vigente(contratacion) is None:
        r.revisiones.insert(0, "Falta el perfil: registre la experiencia que exige el estudio previo.")
    if tabla is None:
        r.revisiones.append(f"La entidad no ha registrado su tabla de honorarios de {contratacion.fecha_referencia.year}: no se verificó el tope.")
    return r, datos, retirados


def detalle(contratacion: ContratacionOps) -> dict:
    """Todo lo que muestra la pantalla de una contratación."""
    resultado = contratacion.resultado or {}
    datos: dict = {
        "id": str(contratacion.id), "referencia": contratacion.referencia, "estado": contratacion.estado, "error": contratacion.error,
        "contratista_nombre": contratacion.contratista_nombre, "contratista_cedula": contratacion.contratista_cedula,
        "exige_libreta": contratacion.exige_libreta, "fecha_referencia": contratacion.fecha_referencia.isoformat(),
        "creada_en": contratacion.creada_en.isoformat(), "creada_por": contratacion.creada_por.nombre_completo,
        "analizada_en": contratacion.analizada_en.isoformat() if contratacion.analizada_en else None, "segundos": contratacion.segundos,
        "confirmada_en": contratacion.confirmada_en.isoformat() if contratacion.confirmada_en else None,
        "confirmada_por": contratacion.confirmada_por.nombre_completo if contratacion.confirmada_por_id else None,
        "documentos_eliminados_en": contratacion.documentos_eliminados_en.isoformat() if contratacion.documentos_eliminados_en else None,
        "archivos": [
            {"id": str(d.id), "origen": d.origen, "nombre": d.nombre_original, "tamano": d.tamano} for d in contratacion.documentos.all()
        ],
    }
    if contratacion.resultado is None:
        return datos
    r, de, retirados = idoneidad(contratacion)
    perfil = perfil_vigente(contratacion)
    tabla = tabla_de(contratacion.entidad_id, contratacion.fecha_referencia.year)
    sobran = {id(t) for t in r.retirar}

    def fila(p: dict, **extra) -> dict:
        return {"id": p["id"], "entidad": p["entidad"], "referencia": p["referencia"], "inicio": p["inicio"], "fin": p["fin"],
                "abierto": p["abierto"], "nota": p["nota"], "documento_id": p["documento_id"], "pagina": p["pagina"],
                "manual": p["manual"], "corregido": p["corregido"], "confirmado": p["confirmado"], "por_confirmar": bool(p["duda"]),
                "duda": p["duda"], "sigep": p["sigep"], "sigep_confirmado": p["sigep_confirmado"], **extra}

    experiencia = []
    for t in r.lineal.tramos:
        experiencia.append(fila(
            de[id(t.periodo)], cuenta_desde=t.inicio.isoformat(), dias=t.dias, duracion=texto_duracion(t.dias), recortado=t.recortado,
            relacionada=t.periodo.relacionada, obligaciones_iguales=t.periodo.obligaciones_iguales,
            estado="sobra" if id(t) in sobran else "cuenta", fijo=t.periodo.fijo,
            motivo="Sobra para la franja del perfil: se propone retirarlo." if id(t) in sobran else "",
        ))
    for d in r.lineal.descartados:
        experiencia.append(fila(de[id(d.periodo)], dias=0, duracion="", estado="descartado", motivo=d.motivo))
    for p in retirados:
        experiencia.append(fila(p, dias=0, duracion="", estado=p["estado"], motivo=p["motivo"]))
    experiencia.sort(key=lambda f: (f["inicio"], f["fin"]))

    decisiones_doc = contratacion.decisiones.get("documentos") or {}
    documentos = []
    for d in resultado.get("documentos") or []:
        decision = decisiones_doc.get(d["clave"])
        documentos.append({"adicionales": [], **d, "decision": decision, "estado_final": decision["estado"] if decision else d["estado"]})
    pendientes = sum(1 for d in documentos if d["estado_final"] in (motor_documentos.FALTA, motor_documentos.REVISION, motor_documentos.NO_CUMPLE))

    datos.update({
        "estudio": resultado.get("estudio"), "cdp": resultado.get("cdp"),
        "perfil": None if perfil is None else {
            "anios_minimos": perfil.anios_minimos, "anios_maximos": perfil.anios_maximos, "posgrado": perfil.posgrado,
            "honorarios_mensuales": perfil.honorarios_mensuales, "obligaciones": perfil.obligaciones,
            "especifica_minima": perfil.especifica_minima, "especifica_maxima": perfil.especifica_maxima, "descripcion": perfil.descripcion,
            "corregido": bool(contratacion.decisiones.get("perfil")),
        },
        "titulos": formacion_acreditada(resultado.get("titulos") or [], resultado.get("matricula")), "matricula": resultado.get("matricula"),
        "grado": r.grado.isoformat() if r.grado else None, "grado_corregido": bool(contratacion.decisiones.get("grado")),
        "posgrado": r.posgrado, "posgrado_corregido": bool(contratacion.decisiones.get("posgrado")),
        "experiencia": experiencia,
        "total_dias": r.dias, "total": texto_duracion(r.dias), "total_leido": texto_duracion(r.lineal.dias),
        "relacionada_dias": r.dias_relacionada, "relacionada": texto_duracion(r.dias_relacionada),
        "franja": r.franja.nombre if r.franja else None, "tope": r.tope, "tabla_honorarios": (tabla.norma or None) if tabla else None,
        "documentos": documentos, "documentos_pendientes": pendientes,
        "sin_reconocer": resultado.get("sin_reconocer") or [],
        "sigep": None if resultado.get("sigep") is None else {
            "empleos": resultado["sigep"]["bloques"], "leidos": len(resultado["sigep"]["experiencias"]),
            "confiable": bool(resultado["sigep"]["bloques"]) and resultado["sigep"]["bloques"] == len(resultado["sigep"]["experiencias"]),
        },
        "revisiones": list(dict.fromkeys(r.revisiones)),
    })
    vistos = set(contratacion.decisiones.get("avisos_vistos") or [])
    datos["avisos"] = [{"texto": a, "visto": a in vistos} for a in dict.fromkeys(resultado.get("avisos") or [])]
    datos["por_confirmar"] = _por_confirmar(contratacion, datos, tabla is not None)
    datos["cumple"] = conclusion(datos)
    return datos


def _por_confirmar(contratacion: ContratacionOps, datos: dict, hay_tabla: bool) -> list[str]:
    """Lo que impide dar una conclusión: mientras quede algo aquí, el sistema
    no dice ni «cumple» ni «no cumple». Un aprobado no puede apoyarse en algo
    que no se leyó completo o que nadie confirmó."""
    resultado, decisiones = contratacion.resultado or {}, contratacion.decisiones
    faltan: list[str] = []
    if datos.get("perfil") is None:
        faltan.append("Registrar el perfil que exige el estudio previo.")
    if datos["documentos_pendientes"]:
        faltan.append(f"Decidir {datos['documentos_pendientes']} documentos que faltan o están por revisar.")
    dudosos = sum(1 for f in datos["experiencia"] if f["por_confirmar"] and f["estado"] in ("cuenta", "sobra"))
    if dudosos:
        faltan.append(f"Confirmar o corregir {dudosos} periodos de experiencia cuya lectura quedó en duda.")
    titulos = resultado.get("titulos") or []
    if not decisiones.get("grado"):
        leido = any(t["nivel"] == "profesional" and t["fecha"] and not t.get("declarado") for t in titulos)
        if not leido:
            faltan.append("Confirmar la fecha de grado: no se leyó del diploma.")
    if not decisiones.get("posgrado") and datos["posgrado"] != SIN_POSGRADO:
        # El posgrado sube el tope de honorarios: debe constar en un diploma leído.
        leido = any(t["nivel"] == datos["posgrado"] and not t.get("declarado") and (t["fecha"] or t["nombre"]) for t in titulos)
        if not leido:
            faltan.append("Confirmar el posgrado: no se leyó del diploma.")
    if not hay_tabla:
        faltan.append(f"Registrar la tabla de honorarios de {contratacion.fecha_referencia.year}: sin ella no se verifica el tope.")
    sin_ver = sum(1 for a in datos["avisos"] if not a["visto"])
    if sin_ver:
        faltan.append(f"Revisar {sin_ver} avisos de la lectura de los documentos.")
    return faltan


def formacion_acreditada(titulos: list[dict], matricula: dict | None) -> list[dict]:
    """Un renglón por nivel: el diploma y su acta dicen lo mismo, y a veces
    uno trae el nombre y el otro la fecha. El de bachiller no va."""
    profesion = (matricula or {}).get("profesion") or ""
    por_nivel: dict[str, dict] = {}
    for t in titulos:
        if t["nivel"] == "bachiller":
            continue
        actual = por_nivel.get(t["nivel"])
        if actual is None:
            por_nivel[t["nivel"]] = dict(t)
            continue
        if not actual["nombre"] and t["nombre"]:
            actual["nombre"] = t["nombre"]
        if not actual["fecha"] and t["fecha"]:
            actual.update(fecha=t["fecha"], documento_id=t.get("documento_id"), pagina=t.get("pagina"), declarado=t.get("declarado", False))
    for t in por_nivel.values():
        if t["nivel"] == "profesional" and not t["nombre"]:
            t["nombre"] = profesion
    if "profesional" not in por_nivel and profesion:
        por_nivel = {"profesional": {"nivel": "profesional", "nombre": profesion, "fecha": None, "documento_id": None, "pagina": None,
                                     "declarado": False, "por_matricula": True}, **por_nivel}
    return list(por_nivel.values())


def conclusion(datos: dict) -> bool | None:
    """True/False cuando ya no queda nada por confirmar; None mientras tanto."""
    perfil = datos.get("perfil")
    if perfil is None or datos["por_confirmar"]:
        return None
    if datos["total_dias"] < perfil["anios_minimos"] * 360:
        return False
    if datos["relacionada_dias"] < perfil["especifica_minima"] * 360:
        return False
    if datos["tope"] is not None and perfil["honorarios_mensuales"] > datos["tope"]:
        return False
    orden = [SIN_POSGRADO, ESPECIALIZACION, MAESTRIA]
    return orden.index(datos["posgrado"]) >= orden.index(perfil["posgrado"])


def _fecha_iso(valor, nombre: str) -> str:
    try:
        return date.fromisoformat(valor).isoformat()
    except (TypeError, ValueError) as exc:
        raise ErrorOps(f"La fecha de {nombre} no es válida.") from exc


def aplicar_decisiones(contratacion: ContratacionOps, cambios: dict) -> None:
    """Mezcla las decisiones nuevas con las guardadas, validando lo que llega."""
    resultado = contratacion.resultado or {}
    decisiones = {**contratacion.decisiones}
    agregados = [dict(a) for a in decisiones.get("agregados") or []]
    if "agregar" in cambios:
        nuevo = cambios["agregar"] or {}
        inicio, fin = _fecha_iso(nuevo.get("inicio"), "inicio"), _fecha_iso(nuevo.get("fin"), "terminación")
        if fin < inicio:
            raise ErrorOps("La fecha de terminación es anterior a la de inicio.")
        if not str(nuevo.get("entidad") or "").strip():
            raise ErrorOps("Escriba la entidad o la empresa del periodo.")
        if len(agregados) >= 60:
            raise ErrorOps("Son demasiados periodos agregados a mano.")
        numero = max((int(a["id"][1:]) for a in agregados), default=0) + 1
        agregados.append({
            "id": f"m{numero}", "inicio": inicio, "fin": fin, "entidad": str(nuevo["entidad"]).strip()[:200],
            "referencia": str(nuevo.get("referencia") or "").strip()[:200], "relacionada": bool(nuevo.get("relacionada")),
        })
        decisiones["agregados"] = agregados
    if "quitar" in cambios:
        if not any(a["id"] == cambios["quitar"] for a in agregados):
            raise ErrorOps("Solo se quitan los periodos agregados a mano; los leídos se retiran.")
        decisiones["agregados"] = [a for a in agregados if a["id"] != cambios["quitar"]]
        decisiones["periodos"] = {k: v for k, v in (decisiones.get("periodos") or {}).items() if k != cambios["quitar"]}
    if "periodos" in cambios:
        validos = {p["id"]: p for p in [*(resultado.get("periodos") or []), *(decisiones.get("agregados") or [])]}
        guardadas = {**(decisiones.get("periodos") or {})}
        for pid, d in (cambios["periodos"] or {}).items():
            if pid not in validos:
                raise ErrorOps("Ese periodo no existe.")
            limpia = {}
            for campo in ("incluir", "relacionada"):
                if d.get(campo) is not None:
                    limpia[campo] = bool(d[campo])
            if d.get("confirmado"):
                limpia["confirmado"] = True
            if d.get("en_sigep"):
                limpia["en_sigep"] = True
            for campo in ("entidad", "referencia"):
                if isinstance(d.get(campo), str):
                    limpia[campo] = d[campo].strip()[:200]
            for campo, nombre in (("inicio", "inicio"), ("fin", "terminación")):
                if d.get(campo):
                    limpia[campo] = _fecha_iso(d[campo], nombre)
            if limpia.get("fin", validos[pid]["fin"]) < limpia.get("inicio", validos[pid]["inicio"]):
                raise ErrorOps("La fecha de terminación es anterior a la de inicio.")
            guardadas[pid] = limpia
            if not limpia:
                del guardadas[pid]
        decisiones["periodos"] = guardadas
    if "perfil" in cambios:
        p = cambios["perfil"] or {}
        if not p:
            decisiones.pop("perfil", None)
        else:
            def rango(minimo, maximo, que: str) -> None:
                if not isinstance(minimo, (int, float)) or minimo < 0 or (maximo is not None and (not isinstance(maximo, (int, float)) or maximo <= minimo)):
                    raise ErrorOps(f"Indique los años mínimos de {que} y, si hay tope, un máximo mayor.")

            rango(p.get("anios_minimos"), p.get("anios_maximos"), "experiencia")
            especifica_min, especifica_max = p.get("especifica_minima") or 0, p.get("especifica_maxima")
            rango(especifica_min, especifica_max, "experiencia específica")
            if p.get("posgrado") not in POSGRADOS:
                raise ErrorOps("Elija el posgrado que exige el perfil.")
            honorarios = p.get("honorarios_mensuales")
            if not isinstance(honorarios, int) or honorarios < 0:
                raise ErrorOps("Los honorarios mensuales deben ser un valor en pesos.")
            decisiones["perfil"] = {
                "anios_minimos": p["anios_minimos"], "anios_maximos": p.get("anios_maximos"), "posgrado": p["posgrado"],
                "honorarios_mensuales": honorarios, "especifica_minima": especifica_min, "especifica_maxima": especifica_max,
            }
    if "grado" in cambios:
        if cambios["grado"]:
            decisiones["grado"] = _fecha_iso(cambios["grado"], "grado")
        else:
            decisiones.pop("grado", None)
    if "posgrado" in cambios:
        if not cambios["posgrado"]:
            decisiones.pop("posgrado", None)
        elif cambios["posgrado"] in POSGRADOS:
            decisiones["posgrado"] = cambios["posgrado"]
        else:
            raise ErrorOps("Elija el posgrado que acredita la persona.")
    if "avisos_vistos" in cambios:
        validos = set(resultado.get("avisos") or [])
        decisiones["avisos_vistos"] = [a for a in dict.fromkeys(cambios["avisos_vistos"] or []) if a in validos]
    if "documentos" in cambios:
        validos = {d["clave"] for d in resultado.get("documentos") or []}
        guardadas = {**(decisiones.get("documentos") or {})}
        for clave, d in (cambios["documentos"] or {}).items():
            if clave not in validos:
                raise ErrorOps("Ese documento no está en la lista.")
            if d is None:
                guardadas.pop(clave, None)
                continue
            if d.get("estado") not in (motor_documentos.CUMPLE, motor_documentos.NO_CUMPLE, motor_documentos.NO_APLICA):
                raise ErrorOps("Indique si el documento cumple, no cumple o no aplica.")
            guardadas[clave] = {"estado": d["estado"], "nota": str(d.get("nota") or "").strip()[:500]}
        decisiones["documentos"] = guardadas
    contratacion.decisiones = decisiones
    contratacion.save(update_fields=["decisiones", "actualizada_en"])


def cambiar_datos(contratacion: ContratacionOps, datos: dict) -> bool:
    """Corrige los datos del contratista. Devuelve si hay que leer los documentos de nuevo."""
    campos: list[str] = []
    if "contratista_nombre" in datos or "contratista_cedula" in datos:
        nombre = " ".join(str(datos.get("contratista_nombre", contratacion.contratista_nombre)).split())
        cedula = re.sub(r"\D", "", str(datos.get("contratista_cedula", contratacion.contratista_cedula)))
        if len(nombre) < 5 or not 5 <= len(cedula) <= 12:
            raise ErrorOps("Escriba el nombre completo y la cédula del contratista.")
        if (nombre, cedula) != (contratacion.contratista_nombre, contratacion.contratista_cedula):
            contratacion.contratista_nombre, contratacion.contratista_cedula = nombre[:200], cedula
            campos += ["contratista_nombre", "contratista_cedula"]
    if "exige_libreta" in datos and datos["exige_libreta"] != contratacion.exige_libreta:
        contratacion.exige_libreta = datos["exige_libreta"]
        campos.append("exige_libreta")
    if datos.get("fecha_referencia"):
        fecha = date.fromisoformat(_fecha_iso(datos["fecha_referencia"], "referencia"))
        if fecha != contratacion.fecha_referencia:
            contratacion.fecha_referencia = fecha
            campos.append("fecha_referencia")
    solo_referencia = "referencia" in datos and str(datos["referencia"]).strip()[:80] != contratacion.referencia
    if solo_referencia:
        contratacion.referencia = str(datos["referencia"]).strip()[:80]
    if campos or solo_referencia:
        contratacion.save(update_fields=[*campos, "referencia", "actualizada_en"])
    return bool(campos)


# --- Certificado de idoneidad ------------------------------------------------
def certificado(contratacion: ContratacionOps) -> bytes:
    """Certificado de idoneidad y experiencia (Decreto 1082 de 2015, artículo
    2.2.1.2.1.4.9) en Word, para que la entidad lo ajuste a su formato y lo firme."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    from evaluaciones.reporte import _parrafo, _tabla, _titulo, fecha_larga, pesos

    datos = detalle(contratacion)
    perfil = datos.get("perfil")
    if perfil is None:
        raise ErrorOps("Registre primero el perfil que exige el estudio previo.")
    cuentan = [f for f in datos["experiencia"] if f["estado"] == "cuenta"]
    entidad = contratacion.entidad.nombre

    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    # Encabezado de control documental, en cada página: la entidad pone aquí
    # el código y la versión de su propio formato.
    seccion = doc.sections[0]
    control = seccion.header.add_table(rows=2, cols=4, width=seccion.page_width - seccion.left_margin - seccion.right_margin)
    control.style = "Table Grid"
    for i, (etiqueta, valor) in enumerate([
        ("Entidad", entidad), ("Código del formato", ""),
        ("Documento", "Certificado de idoneidad y experiencia para contrato de prestación de servicios"), ("Versión", ""),
    ]):
        for celda, texto, negrita in ((control.cell(i // 2, (i % 2) * 2), etiqueta, True), (control.cell(i // 2, (i % 2) * 2 + 1), valor, False)):
            celda.text = texto
            if celda.paragraphs[0].runs:
                celda.paragraphs[0].runs[0].bold = negrita
                celda.paragraphs[0].runs[0].font.size = Pt(8)

    _titulo(doc, "Certificado de idoneidad y experiencia")
    _parrafo(doc, fecha_larga(timezone.localdate()), alineacion=WD_ALIGN_PARAGRAPH.LEFT)
    _tabla(doc, ["Nombre", "Identificación"], [[contratacion.contratista_nombre.upper(), f"C.C. {contratacion.contratista_cedula}"]])
    _parrafo(doc, "")
    _parrafo(
        doc,
        f"De acuerdo con el artículo 2.2.1.2.1.4.9 del Decreto 1082 de 2015, {entidad} puede contratar directamente con la "
        "persona natural o jurídica que esté en capacidad de ejecutar el objeto del contrato, siempre que verifique la "
        "idoneidad o experiencia requerida y relacionada con el área de que se trate, sin que sea necesario obtener "
        "previamente varias ofertas, de lo cual el ordenador del gasto debe dejar constancia escrita.",
    )
    if (datos.get("estudio") or {}).get("objeto"):
        _parrafo(doc, f"Objeto a contratar: {datos['estudio']['objeto']}")
    _parrafo(doc, "Elaborados los estudios previos, se analizó la hoja de vida del contratista junto con las certificaciones de estudios y de experiencia aportadas:")

    _titulo(doc, "1. Formación académica del contratista", 2)
    minimo = perfil["descripcion"] or f"Profesional{'' if perfil['posgrado'] == SIN_POSGRADO else ' con ' + POSGRADOS[perfil['posgrado']].lower()}"
    _parrafo(doc, "1.1 Requisitos mínimos de formación académica", negrita=True)
    _parrafo(doc, f"Requisito mínimo: {minimo.rstrip('.')}.")
    _parrafo(doc, "1.2 Formación académica acreditada", negrita=True)
    _tabla(doc, ["Formación acreditada", "Fecha de grado", "Nivel"], [
        [t["nombre"] or NIVELES.get(t["nivel"], t["nivel"]), fecha_larga(date.fromisoformat(t["fecha"])) if t["fecha"] else "—",
         NIVELES.get(t["nivel"], t["nivel"])]
        for t in datos["titulos"]
    ] or [["No se leyeron títulos", "—", "—"]])

    _parrafo(doc, "")
    _titulo(doc, "2. Experiencia del contratista", 2)
    _parrafo(doc, "2.1 Experiencia mínima y relacionada con el área del servicio a prestar", negrita=True)
    maximo = perfil["anios_maximos"]
    exigida = f"entre {perfil['anios_minimos']:g} y {maximo:g} años" if maximo is not None else f"mínima de {perfil['anios_minimos']:g} años"
    if perfil["especifica_minima"] or perfil["especifica_maxima"] is not None:
        tope_especifica = perfil["especifica_maxima"]
        exigida += (
            f", y experiencia específica entre {perfil['especifica_minima']:g} y {tope_especifica:g} años" if tope_especifica is not None
            else f", y experiencia específica mínima de {perfil['especifica_minima']:g} años"
        )
    _parrafo(doc, f"Experiencia exigida: {exigida}. La experiencia se relaciona sin traslapes: el tiempo trabajado en un mismo periodo se cuenta una sola vez.")
    _parrafo(doc, "2.2 Experiencia acreditada por el contratista", negrita=True)
    _tabla(doc, ["Entidad", "Contrato o cargo", "Inicio", "Terminación", "Duración", "Tipo"], [
        [f["entidad"] or "—", f["referencia"] or "—", f"{date.fromisoformat(f['cuenta_desde']):%d/%m/%Y}", f"{date.fromisoformat(f['fin']):%d/%m/%Y}",
         f["duracion"], "Relacionada" if f["relacionada"] else "Profesional"]
        for f in cuentan
    ] or [["Sin periodos", "—", "—", "—", "—", "—"]])
    anios, meses, dias = anios_meses_dias(datos["total_dias"])
    _parrafo(doc, "")
    _parrafo(doc, f"Total de experiencia: {anios} años, {meses} meses y {dias} días.", negrita=True)
    if datos["relacionada_dias"]:
        _parrafo(doc, f"De ella, experiencia relacionada con las obligaciones del contrato: {datos['relacionada']}.")

    if perfil["honorarios_mensuales"]:
        _titulo(doc, "3. Honorarios", 2)
        texto = f"Honorarios mensuales pactados: {pesos(perfil['honorarios_mensuales'])}."
        if datos["tope"] is not None:
            texto += f" Tope de la franja {datos['franja']} según {datos['tabla_honorarios'] or 'la tabla de honorarios de la entidad'}: {pesos(datos['tope'])}."
        _parrafo(doc, texto)

    _titulo(doc, "Certificación", 2)
    _parrafo(
        doc,
        f"Analizados los aspectos establecidos en los estudios previos, certifico que {contratacion.contratista_nombre.upper()} cuenta con "
        "los requisitos de formación y experiencia determinados por esta área, que lo(a) hacen idóneo(a) para ejecutar el contrato a "
        "celebrar, y que he verificado los documentos soporte de la hoja de vida.",
    )
    for cargo in ("Jefe de la dependencia que solicita la contratación", "Ordenador del gasto o su delegado"):
        _parrafo(doc, "")
        _parrafo(doc, "Firma: ______________________________", alineacion=WD_ALIGN_PARAGRAPH.LEFT)
        _parrafo(doc, "Nombre:", alineacion=WD_ALIGN_PARAGRAPH.LEFT, espacio=0)
        _parrafo(doc, f"Cargo: {cargo}", alineacion=WD_ALIGN_PARAGRAPH.LEFT)
    quien = f" Confirmado por {datos['confirmada_por']}." if datos["confirmada_por"] else " Pendiente de confirmación por una persona."
    _parrafo(
        doc, f"Proyectado con apoyo de MiEvaluador {settings.MIEVALUADOR_VERSION} a partir de los documentos aportados.{quien}",
        tam=8, alineacion=WD_ALIGN_PARAGRAPH.LEFT,
    )
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


ESTADOS_DOCUMENTO = {
    motor_documentos.CUMPLE: "Cumple", motor_documentos.FALTA: "Falta", motor_documentos.REVISION: "Por revisar",
    motor_documentos.NO_CUMPLE: "No cumple", motor_documentos.NO_APLICA: "No aplica",
}


def lista_de_verificacion(contratacion: ContratacionOps) -> bytes:
    """Lista de verificación de los documentos del contratista, en Word, para el expediente."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    from evaluaciones.reporte import _parrafo, _tabla, _titulo, fecha_larga

    datos = detalle(contratacion)
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    _titulo(doc, "Lista de verificación de documentos")
    _tabla(doc, ["Contratista", "Identificación", "Proceso o contrato"], [[
        contratacion.contratista_nombre.upper(), f"C.C. {contratacion.contratista_cedula}", contratacion.referencia or "—",
    ]])
    _parrafo(doc, "")
    filas = []
    for i, d in enumerate(datos.get("documentos") or [], 1):
        decision = d["decision"]
        observacion = (decision["nota"] or "Verificado por una persona") if decision else d["motivo"]
        archivos = ", ".join(a for a in [d["archivo"], *(x["archivo"] for x in d["adicionales"])] if a)
        filas.append([str(i), d["nombre"], ESTADOS_DOCUMENTO.get(d["estado_final"], d["estado_final"]), archivos or "—", observacion or ""])
    _tabla(doc, ["N.º", "Documento", "Estado", "Archivo", "Observación"], filas)
    _parrafo(doc, "")
    pendientes = datos.get("documentos_pendientes") or 0
    _parrafo(doc, "Todos los documentos están al día." if not pendientes else f"Quedan {pendientes} documentos por revisar.", negrita=True)
    quien = (
        f"Confirmada por {datos['confirmada_por']} el {fecha_larga(contratacion.confirmada_en)}." if datos["confirmada_por"]
        else "Pendiente de confirmación por una persona."
    )
    _parrafo(
        doc, f"Elaborada con apoyo de MiEvaluador {settings.MIEVALUADOR_VERSION} el {fecha_larga(timezone.localdate())}. {quien}",
        tam=8, alineacion=WD_ALIGN_PARAGRAPH.LEFT,
    )
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()
