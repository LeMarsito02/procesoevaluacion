"""Pruebas de lo que hace de MiEvaluador una herramienta de apoyo y no de
decisión (expediente LEG-004 y Concepto C-1015 de 2026): muestra de control,
soporte visto antes de decidir, compromiso de uso, adopción del puntaje,
auditoría inmutable, procesos con expediente y rótulos de los informes."""
from __future__ import annotations

import io
from datetime import timedelta

from django.core.management import call_command
from django.db import DatabaseError, connection, transaction
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from cuentas.models import EventoAuditoria
from cuentas.tests import Cliente
from evaluaciones.models import (
    EstadoEvaluacion,
    EstadoExpediente,
    Evaluacion,
    Expediente,
    ItemMuestra,
    MuestraControl,
    Proponente,
    Resultado,
)
from evaluaciones.muestra import sortear
from evaluaciones.tests import DOCUMENTO_BASE, BaseEvaluaciones, correr_fila
from motor.esquemas.proceso import ResultadoRequisito

MUCHOS = [
    {"numero_orden": i, "hoja": f"P-{i:02d}", "nombre_proponente": f"Proponente {i}", "nombre_archivo": f"P{i}.zip", "drive_file_id": f"d{i}"}
    for i in range(1, 13)
]


class SorteoTests(TestCase):
    def test_misma_semilla_misma_muestra(self):
        # 30 ofertas × 5 requisitos verificados por el sistema.
        universo = [(n, f"P-{n:02d}", req, False, str(n)) for n in range(1, 31) for req in (1, 2, 3, 4, 5)]
        a = sortear(universo, 12345, 10)
        self.assertEqual(a, sortear(list(reversed(universo)), 12345, 10))
        self.assertEqual(len(a), 10)
        self.assertNotEqual(a, sortear(universo, 999, 10))
        # Repartidas: 10 ofertas distintas y los 5 requisitos cubiertos.
        self.assertEqual(len({c[1] for c in a}), 10)
        self.assertEqual({c[2] for c in a}, {1, 2, 3, 4, 5})
        # Con 10 o menos, entran todas.
        self.assertEqual(len(sortear(universo[:7], 1, 10)), 7)


class BaseFlujo(BaseEvaluaciones):
    def preparar(self, proponentes=MUCHOS):
        self.jefe = Cliente()
        self.jefe.entrar("jefe@entidad.gov.co")
        r = self.jefe.post("/api/evaluaciones/procesos", {"documento_base": DOCUMENTO_BASE, "carpeta_drive": "x", "proponentes": proponentes})
        self.assertEqual(r.status_code, 201, r.content)
        self.eid = r.json()[0]["id"]
        self.jefe.post(f"/api/evaluaciones/{self.eid}/asignar", {"responsable_id": str(self.evaluador.id)})
        self.abogado = Cliente()
        self.abogado.entrar("abogado@entidad.gov.co")
        self.evaluar_todo(self.abogado, self.eid)
        self.props = {p.hoja: p for p in Proponente.objects.filter(proceso__evaluaciones__id=self.eid)}
        # El requisito 2 quedó por revisar en todos: se resuelve con el soporte abierto.
        for hoja, p in self.props.items():
            self.ver_soporte(self.evaluador, self.eid, hoja)
            r = self.abogado.put(f"/api/evaluaciones/{self.eid}/revisiones", {"proponente_id": str(p.id), "requisito": 2, "cumple": True, "nota": "COPNIA consultado"})
            self.assertEqual(r.status_code, 200, r.content)


class MuestraTests(BaseFlujo):
    def test_sortea_diez_verificaciones_repartidas_y_solo_lo_del_sistema(self):
        self.preparar()
        r = self.abogado.post(f"/api/evaluaciones/{self.eid}/muestra")
        self.assertEqual(r.status_code, 200, r.content)
        m = r.json()["muestra"]
        # 12 ofertas × requisitos 1 y 13 verificados por el sistema (el 2 lo decidió una persona).
        self.assertEqual(m["parametros"]["universo_verificaciones"], 24)
        # Se revisan 10 verificaciones, no ofertas completas: una por oferta, de ambos requisitos.
        self.assertEqual(len(m["items"]), 10)
        self.assertEqual(len({i["hoja"] for i in m["items"]}), 10)
        self.assertEqual(len(m["ofertas_sorteadas"]), 10)
        self.assertEqual({i["requisito"] for i in m["items"]}, {1, 13})
        # Volver a pedirla devuelve la misma (no se re-sortea hasta cerrarla).
        self.assertEqual(self.abogado.post(f"/api/evaluaciones/{self.eid}/muestra").json()["muestra"]["id"], m["id"])
        self.assertTrue(EventoAuditoria.objects.filter(accion="muestra.creada").exists())

    def test_sin_ver_el_soporte_no_se_marca(self):
        self.preparar()
        m = self.abogado.post(f"/api/evaluaciones/{self.eid}/muestra").json()["muestra"]
        item = m["items"][0]
        # Sin documento soporte en el resultado, exige explicar qué consultó.
        r = self.abogado.put(f"/api/evaluaciones/{self.eid}/muestra/items/{item['id']}", {"conforme": True, "nota": "ok"})
        self.assertEqual(r.status_code, 409)
        # Con documento soporte, exige haberlo abierto después de crear la muestra.
        res = Resultado.objects.get(evaluacion_id=self.eid, proponente_id=item["proponente_id"], requisito=item["requisito"])
        res.datos = {**res.datos, "archivo_evaluado": "CARTA/carta.pdf"}
        Resultado.objects.filter(pk=res.pk).update(datos=res.datos)
        r = self.abogado.put(f"/api/evaluaciones/{self.eid}/muestra/items/{item['id']}", {"conforme": True, "nota": "Revisado contra la carta"})
        self.assertEqual(r.status_code, 409)
        self.assertIn("Abra el documento soporte", r.json()["detail"])
        self.ver_soporte(self.evaluador, self.eid, item["hoja"])
        r = self.abogado.put(f"/api/evaluaciones/{self.eid}/muestra/items/{item['id']}", {"conforme": True})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertTrue(ItemMuestra.objects.get(pk=item["id"]).soporte_visto)

    def test_lo_que_sale_del_pliego_entra_una_vez_y_su_soporte_es_el_pliego(self):
        from evaluaciones.models import AnalisisPliego, Proceso

        self.preparar()
        # El requisito 13 queda como "N.A. por el pliego" en las 12 ofertas: es
        # una sola decisión, no 12, y no tiene documento de la oferta.
        for res in Resultado.objects.filter(evaluacion_id=self.eid, requisito=13):
            Resultado.objects.filter(pk=res.pk).update(datos={**res.datos, "cumple": True, "motivo": "N.A. — el pliego no exige capacidad residual"})
        proceso = Proceso.objects.get(evaluaciones__id=self.eid)
        proceso.analisis_pliego = AnalisisPliego.objects.create(
            entidad_id=proceso.entidad_id, sha256="0" * 64, nombre_archivo="pliego.pdf", archivo="pliegos/x.pdf",
            paginas=1, extraccion={}, version=1,
        )
        proceso.save()
        m = self.abogado.post(f"/api/evaluaciones/{self.eid}/muestra").json()["muestra"]
        del_pliego = [i for i in m["items"] if i["requisito"] == 13]
        self.assertEqual(len(del_pliego), 1)
        self.assertTrue(del_pliego[0]["soporte_pliego"])
        self.assertEqual(m["parametros"]["universo_verificaciones"], 24)
        item = del_pliego[0]
        # Abrir la oferta no basta: el soporte es el pliego.
        self.ver_soporte(self.evaluador, self.eid, item["hoja"])
        r = self.abogado.put(f"/api/evaluaciones/{self.eid}/muestra/items/{item['id']}", {"conforme": True, "nota": "Visto en el pliego 3.11"})
        self.assertEqual(r.status_code, 409)
        self.assertIn("sale del pliego", r.json()["detail"])
        EventoAuditoria.objects.create(
            usuario=self.evaluador, entidad_id=self.evaluador.entidad_id, accion="pliego.visto",
            objeto_tipo="Evaluacion", objeto_id=str(self.eid), detalles={},
        )
        r = self.abogado.put(f"/api/evaluaciones/{self.eid}/muestra/items/{item['id']}", {"conforme": True})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertTrue(ItemMuestra.objects.get(pk=item["id"]).soporte_visto)

    def test_un_error_amplia_la_revision_a_todo_el_requisito(self):
        self.preparar()
        m = self.abogado.post(f"/api/evaluaciones/{self.eid}/muestra").json()["muestra"]
        item = next(i for i in m["items"] if i["requisito"] == 1)
        self.ver_soporte(self.evaluador, self.eid, item["hoja"])
        # No conforme exige explicar qué encontró.
        r = self.abogado.put(f"/api/evaluaciones/{self.eid}/muestra/items/{item['id']}", {"conforme": False, "nota": ""})
        self.assertEqual(r.status_code, 409)
        r = self.abogado.put(f"/api/evaluaciones/{self.eid}/muestra/items/{item['id']}", {"conforme": False, "nota": "La carta no está firmada"})
        self.assertEqual(r.status_code, 200, r.content)
        datos = r.json()["muestra"]
        self.assertEqual(datos["estado"], "con_hallazgos")
        self.assertEqual(datos["requisitos_ampliados"], [1])
        # El requisito 1 de los 12 proponentes pasó a revisión humana.
        avance = self.jefe.get(f"/api/evaluaciones/{self.eid}").json()["evaluacion"]["avance"]
        self.assertEqual(avance["pendientes"], 12)
        # Se revisan los demás ítems y aún no se puede cerrar: falta el requisito ampliado.
        for i in datos["items"]:
            if i["resultado"] is None:
                self.ver_soporte(self.evaluador, self.eid, i["hoja"])
                self.abogado.put(f"/api/evaluaciones/{self.eid}/muestra/items/{i['id']}", {"conforme": True})
        r = self.abogado.post(f"/api/evaluaciones/{self.eid}/muestra/cerrar")
        self.assertEqual(r.status_code, 409)
        self.assertIn("Quedan 12 requisitos por revisar", r.json()["detail"])
        for hoja, p in self.props.items():
            self.ver_soporte(self.evaluador, self.eid, hoja)
            self.abogado.put(
                f"/api/evaluaciones/{self.eid}/revisiones",
                {"proponente_id": str(p.id), "requisito": 1, "cumple": hoja != item["hoja"], "nota": "Carta revisada una por una"},
            )
        r = self.abogado.post(f"/api/evaluaciones/{self.eid}/muestra/cerrar")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["muestra"]["estado"], "cerrada")
        self.assertEqual(self.jefe.post(f"/api/evaluaciones/{self.eid}/aprobar").json()["estado"], EstadoEvaluacion.APROBADA)
        # El acta queda disponible.
        acta = self.jefe.get(f"/api/evaluaciones/{self.eid}/muestra/acta")
        self.assertEqual(acta.status_code, 200)
        self.assertTrue(acta.content.startswith(b"PK"))

    def test_un_cambio_despues_de_la_muestra_exige_otra(self):
        self.preparar(MUCHOS[:3])
        self.hacer_muestra(self.abogado, self.evaluador, self.eid)
        # Todo el universo entra cuando hay 10 verificaciones o menos.
        muestra = MuestraControl.objects.get(evaluacion_id=self.eid, estado="cerrada")
        self.assertEqual(muestra.items.count(), 6)
        p = self.props["P-01"]
        self.ver_soporte(self.evaluador, self.eid, "P-01")
        self.abogado.put(f"/api/evaluaciones/{self.eid}/revisiones", {"proponente_id": str(p.id), "requisito": 2, "cumple": False, "nota": "Se revisó de nuevo el COPNIA"})
        r = self.jefe.post(f"/api/evaluaciones/{self.eid}/aprobar")
        self.assertEqual(r.status_code, 409)
        self.assertIn("cambió después de la muestra", r.json()["detail"])
        self.hacer_muestra(self.abogado, self.evaluador, self.eid)
        self.assertEqual(self.jefe.post(f"/api/evaluaciones/{self.eid}/aprobar").json()["estado"], EstadoEvaluacion.APROBADA)
        self.assertEqual(MuestraControl.objects.filter(evaluacion_id=self.eid, estado="anulada").count(), 1)

    def test_la_muestra_se_sortea_aunque_queden_pendientes_pero_no_se_cierra(self):
        """El sorteo va primero y el cierre al final.

        Un ítem no conforme manda a revisión ese requisito en todas las
        ofertas: encontrarlo antes de revisar lo pendiente permite hacerlo
        todo en una pasada. La garantía está en el cierre, que no admite
        nada por revisar, y sin muestra cerrada no se aprueba."""
        jefe, ev = self.crear()
        abogado = Cliente()
        jefe.post(f"/api/evaluaciones/{ev['id']}/asignar", {"responsable_id": str(self.evaluador.id)})
        abogado.entrar("abogado@entidad.gov.co")
        self.evaluar_todo(abogado, ev["id"])
        # Con requisitos por revisar, la muestra se sortea igual.
        r = abogado.post(f"/api/evaluaciones/{ev['id']}/muestra")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertTrue(r.json()["muestra"]["items"])
        # Pero no se cierra: primero hay que resolver lo pendiente.
        for i in r.json()["muestra"]["items"]:
            self.ver_soporte(self.evaluador, ev["id"], i["hoja"])
            abogado.put(f"/api/evaluaciones/{ev['id']}/muestra/items/{i['id']}", {"conforme": True})
        cierre = abogado.post(f"/api/evaluaciones/{ev['id']}/muestra/cerrar")
        self.assertEqual(cierre.status_code, 409)
        self.assertIn("por revisar", cierre.json()["detail"])
        # Y sin muestra cerrada tampoco se aprueba.
        self.assertEqual(jefe.post(f"/api/evaluaciones/{ev['id']}/aprobar").status_code, 409)


class SoporteYCompromisoTests(BaseFlujo):
    def test_decidir_un_requisito_con_documento_exige_abrirlo(self):
        jefe, ev = self.crear()
        jefe.post(f"/api/evaluaciones/{ev['id']}/asignar", {"responsable_id": str(self.evaluador.id)})
        abogado = Cliente()
        abogado.entrar("abogado@entidad.gov.co")
        detalle = self.evaluar_todo(abogado, ev["id"])
        p = detalle["proponentes"][0]
        Resultado.objects.filter(evaluacion_id=ev["id"], proponente_id=p["id"], requisito=2).update(
            datos={**Resultado.objects.get(evaluacion_id=ev["id"], proponente_id=p["id"], requisito=2).datos, "archivo_evaluado": "COPNIA.pdf"}
        )
        url = f"/api/evaluaciones/{ev['id']}/revisiones"
        r = abogado.put(url, {"proponente_id": p["id"], "requisito": 2, "cumple": True, "nota": "Revisado el COPNIA"})
        self.assertEqual(r.status_code, 409)
        self.ver_soporte(self.evaluador, ev["id"], p["hoja"])
        self.assertEqual(abogado.put(url, {"proponente_id": p["id"], "requisito": 2, "cumple": True, "nota": "Revisado el COPNIA"}).status_code, 200)
        # Sin documento soporte, una justificación corta no basta.
        otro = detalle["proponentes"][1]
        r = abogado.put(url, {"proponente_id": otro["id"], "requisito": 2, "cumple": False, "nota": "No está"})
        self.assertEqual(r.status_code, 409)
        self.assertIn("qué consultó", r.json()["detail"])

    def test_sin_compromiso_de_uso_no_decide(self):
        jefe, ev = self.crear()
        jefe.post(f"/api/evaluaciones/{ev['id']}/asignar", {"responsable_id": str(self.abogado2.id)})
        abogado = Cliente()
        yo = abogado.entrar("abogado2@entidad.gov.co", aceptar_compromiso=False)["usuario"]
        self.assertTrue(yo["compromiso_pendiente"])
        self.assertIn("no traslada, atenúa ni distribuye", yo["compromiso_texto"])
        detalle = self.evaluar_todo(abogado, ev["id"])
        p = detalle["proponentes"][0]
        url = f"/api/evaluaciones/{ev['id']}/revisiones"
        r = abogado.put(url, {"proponente_id": p["id"], "requisito": 2, "cumple": True, "nota": "COPNIA consultado en línea"})
        self.assertEqual(r.status_code, 409)
        self.assertIn("compromiso de uso", r.json()["detail"])
        self.assertEqual(abogado.post("/api/auth/compromiso", {"version": "vieja", "acepto": True}).status_code, 409)
        self.assertEqual(abogado.post("/api/auth/compromiso", {"version": yo["compromiso_version"], "acepto": False}).status_code, 400)
        yo = abogado.post("/api/auth/compromiso", {"version": yo["compromiso_version"], "acepto": True}).json()
        self.assertFalse(yo["compromiso_pendiente"])
        self.assertEqual(abogado.put(url, {"proponente_id": p["id"], "requisito": 2, "cumple": True, "nota": "COPNIA consultado en línea"}).status_code, 200)
        self.assertTrue(EventoAuditoria.objects.filter(accion="compromiso.aceptado", usuario=self.abogado2).exists())


class PuntajeTests(BaseEvaluaciones):
    """El puntaje técnico lo adopta una persona (no por muestra), en un solo acto."""

    def preparar(self):
        self.jefe = Cliente()
        self.jefe.entrar("jefe@entidad.gov.co")
        from cuentas.models import TipoArea

        self.jefe_tecnico = self.jefe
        self.jefe_usuario = __import__("cuentas.models", fromlist=["Usuario"]).Usuario.objects.get(email="jefe@entidad.gov.co")
        self.jefe_usuario.areas.add(self.entidad1.areas.get(tipo=TipoArea.TECNICA))
        r = self.jefe.post(
            "/api/evaluaciones/procesos",
            {"documento_base": DOCUMENTO_BASE, "carpeta_drive": "x", "proponentes": MUCHOS[:2], "tipos": ["tecnica"]},
        )
        self.assertEqual(r.status_code, 201, r.content)
        self.eid = r.json()[0]["id"]
        self.ev = Evaluacion.objects.get(pk=self.eid)
        self.jefe.post(f"/api/evaluaciones/{self.eid}/asignar", {"responsable_id": str(self.jefe_usuario.id)})
        for p in Proponente.objects.filter(proceso=self.ev.proceso):
            factores = [(101, True, 0), (111, True, 0), (121, True, 10), (122, False, 20), (130, True, 0)]
            factores += [(n, True, 0) for n in (123, 124, 125, 126, 127, 129)]  # sin puntos en este pliego
            for numero, cumple, maximo in factores:
                datos = ResultadoRequisito(
                    hoja=p.hoja, numero_orden=p.numero_orden, nombre_proponente=p.nombre, requisito=numero, cumple=cumple,
                    motivo=None if cumple else "No aportó el plan de calidad", detalle={"puntaje_maximo": maximo},
                ).model_dump(mode="json")
                Resultado.objects.create(entidad=self.ev.entidad, evaluacion=self.ev, proponente=p, requisito=numero, datos=datos,
                                         requiere_revision=not cumple and numero in (122,))
        self.props = list(Proponente.objects.filter(proceso=self.ev.proceso).order_by("numero_orden"))

    def test_flujo_de_adopcion(self):
        self.preparar()
        url = f"/api/evaluaciones/{self.eid}/puntajes"
        estado = self.jefe.get(url).json()
        self.assertEqual([e["resuelto"] for e in estado], [False, False])
        # Con un factor pendiente no se adopta.
        r = self.jefe.post(f"{url}/{self.props[0].id}/adoptar", {})
        self.assertEqual(r.status_code, 409)
        for p in self.props:
            self.ver_soporte(self.jefe_usuario, self.eid, p.hoja)
            self.jefe.put(f"/api/evaluaciones/{self.eid}/revisiones", {"proponente_id": str(p.id), "requisito": 122, "cumple": False, "nota": "No aportó el plan"})
        estado = self.jefe.get(url).json()
        self.assertEqual([e["puntaje"] for e in estado], [10.0, 10.0])
        # La muestra no incluye el puntaje.
        m = self.jefe.post(f"/api/evaluaciones/{self.eid}/muestra").json()["muestra"]
        self.assertEqual({i["requisito"] for i in m["items"]}, {101, 111})
        for i in m["items"]:
            self.ver_soporte(self.jefe_usuario, self.eid, i["hoja"])
            self.jefe.put(f"/api/evaluaciones/{self.eid}/muestra/items/{i['id']}", {"conforme": True})
        self.assertEqual(self.jefe.post(f"/api/evaluaciones/{self.eid}/muestra/cerrar").status_code, 200)
        # Sin adoptar el puntaje de todos, no se aprueba.
        r = self.jefe.post(f"/api/evaluaciones/{self.eid}/aprobar")
        self.assertEqual(r.status_code, 409)
        self.assertIn("adoptar el puntaje de 2", r.json()["detail"])
        # Un solo acto adopta todos los puntajes de la tabla.
        estado = self.jefe.post(f"{url}/adoptar", {"nota": "Conforme con el pliego"}).json()
        self.assertEqual([e["adoptado"] for e in estado], [True, True])
        self.assertEqual(self.jefe.post(f"{url}/adoptar", {}).status_code, 409)  # ya no queda nada por adoptar
        # Hojas adoptadas: entran al orden de elegibilidad del consolidado.
        from evaluaciones.servicios import hojas_con_puntaje_adoptado

        self.assertEqual(hojas_con_puntaje_adoptado(self.ev), {"P-01", "P-02"})
        self.assertEqual(self.jefe.post(f"/api/evaluaciones/{self.eid}/aprobar").json()["estado"], EstadoEvaluacion.APROBADA)
        self.assertTrue(EventoAuditoria.objects.filter(accion="puntaje.adoptado").exists())

    def test_un_cambio_en_un_factor_invalida_la_adopcion(self):
        self.preparar()
        url = f"/api/evaluaciones/{self.eid}/puntajes"
        p = self.props[0]
        self.ver_soporte(self.jefe_usuario, self.eid, p.hoja)
        self.jefe.put(f"/api/evaluaciones/{self.eid}/revisiones", {"proponente_id": str(p.id), "requisito": 122, "cumple": False, "nota": "No aportó el plan"})
        self.jefe.post(f"{url}/{p.id}/adoptar", {})
        self.assertTrue(self.jefe.get(url).json()[0]["adoptado"])
        from evaluaciones.models import AdopcionPuntaje

        AdopcionPuntaje.objects.filter(proponente=p).update(fecha=timezone.now() - timedelta(minutes=5))
        self.jefe.put(f"/api/evaluaciones/{self.eid}/revisiones", {"proponente_id": str(p.id), "requisito": 122, "cumple": True, "nota": "Sí aportó el plan"})
        estado = self.jefe.get(url).json()[0]
        self.assertFalse(estado["adoptado"])
        self.assertTrue(estado["adopcion_desactualizada"])
        self.assertEqual(estado["puntaje"], 30.0)


class ConsolidadoPreliminarTests(TestCase):
    def test_orden_solo_con_puntaje_adoptado(self):
        from motor.consolidado import generar_informe
        from motor.tecnica.informe import ResultadoInforme

        def r(hoja, n, cumple, maximo=0):
            return ResultadoInforme(
                ResultadoRequisito(hoja=hoja, numero_orden=1, nombre_proponente=hoja, requisito=n, cumple=cumple, detalle={"puntaje_maximo": maximo}),
                False,
            )

        resultados = {"tecnica": {}}
        for hoja in ("P-01", "P-02"):
            for n in (121, 122, 123, 124, 125, 126, 127, 129):
                resultados["tecnica"][(hoja, n)] = r(hoja, n, True, 10)
            resultados["tecnica"][(hoja, 130)] = r(hoja, 130, True)
            resultados["tecnica"][(hoja, 101)] = r(hoja, 101, True)
            resultados.setdefault("juridica", {})[(hoja, 1)] = r(hoja, 1, True)
            resultados.setdefault("financiera", {})[(hoja, 201)] = r(hoja, 201, True)
        contenido = generar_informe(
            "X-1", "Objeto", [(0, "Lote único")], [("P-01", "Uno"), ("P-02", "Dos")], resultados,
            {"tecnica": {"generales": [101], "lote_0": []}, "juridica": {"generales": [1]}, "financiera": {"generales": [201]}},
            borrador=True, puntajes_adoptados={"P-01"},
        )
        hoja = load_workbook(io.BytesIO(contenido))["Consolidado"]
        valores = [[c.value for c in fila] for fila in hoja.iter_rows()]
        planos = " ".join(str(v) for fila in valores for v in fila if v is not None)
        self.assertIn("PUNTAJE PRELIMINAR", planos)
        self.assertIn("ORDEN DE ELEGIBILIDAD PRELIMINAR", planos)
        self.assertIn("PUNTAJE SIN ADOPTAR", planos)
        self.assertIn("no producen efecto hasta que el comité", planos)


class AuditoriaInmutableTests(TestCase):
    def test_la_base_de_datos_rechaza_modificar_o_borrar(self):
        e = EventoAuditoria.objects.create(accion="prueba", detalles={"a": 1})
        EventoAuditoria.objects.create(accion="prueba2")
        with self.assertRaises(DatabaseError), transaction.atomic():
            EventoAuditoria.objects.filter(pk=e.pk).update(accion="alterado")
        with self.assertRaises(DatabaseError), transaction.atomic():
            EventoAuditoria.objects.filter(pk=e.pk).delete()
        with self.assertRaises(DatabaseError), transaction.atomic(), connection.cursor() as c:
            c.execute("DELETE FROM cuentas_eventoauditoria")

    def test_la_cadena_de_huellas_detecta_manipulacion(self):
        from cuentas.management.commands.verificar_auditoria import verificar

        primero = EventoAuditoria.objects.create(accion="uno")
        EventoAuditoria.objects.create(accion="dos", detalles={"x": "y"})
        EventoAuditoria.objects.create(accion="tres")
        n, roto = verificar()
        self.assertIsNone(roto)
        self.assertGreaterEqual(n, 3)
        out = io.StringIO()
        call_command("verificar_auditoria", stdout=out)
        self.assertIn("Auditoría íntegra", out.getvalue())
        # Alguien con acceso a la base desactiva el trigger y altera un evento.
        with connection.cursor() as c:
            c.execute("SET CONSTRAINTS ALL IMMEDIATE")
            c.execute("ALTER TABLE cuentas_eventoauditoria DISABLE TRIGGER auditoria_inmutable")
            c.execute("UPDATE cuentas_eventoauditoria SET accion = 'alterado' WHERE id = %s", [primero.pk])
            c.execute("ALTER TABLE cuentas_eventoauditoria ENABLE TRIGGER auditoria_inmutable")
        n, roto = verificar()
        self.assertEqual(roto.pk, primero.pk)


class ProcesosConExpedienteTests(BaseFlujo):
    def test_no_se_elimina_aunque_se_reabra_pero_se_archiva(self):
        self.preparar(MUCHOS[:2])
        self.hacer_muestra(self.abogado, self.evaluador, self.eid)
        with self.captureOnCommitCallbacks(execute=True):
            self.assertEqual(self.jefe.post(f"/api/evaluaciones/{self.eid}/aprobar").status_code, 200)
        self.assertEqual(self.jefe.post(f"/api/evaluaciones/{self.eid}/reabrir").status_code, 200)
        ev = Evaluacion.objects.get(pk=self.eid)
        self.assertEqual(ev.version_sistema, "")
        self.assertTrue(Expediente.objects.filter(evaluacion=ev).exists())
        proceso = ev.proceso
        r = self.jefe.http.delete(
            f"/api/evaluaciones/procesos/{proceso.id}", {"confirmacion": proceso.codigo},
            content_type="application/json", headers={"X-CSRFToken": self.jefe.csrf},
        )
        self.assertEqual(r.status_code, 409)
        self.assertIn("Puede archivarlo", r.json()["detail"])
        listado = self.jefe.get("/api/evaluaciones/procesos").json()
        self.assertFalse(listado[0]["puede_eliminar"])
        self.assertTrue(listado[0]["puede_archivar"])
        archivar = lambda c, q="": c.http.post(f"/api/evaluaciones/procesos/{proceso.id}/archivar{q}", headers={"X-CSRFToken": c.csrf})  # noqa: E731
        self.assertEqual(archivar(self.jefe).status_code, 204)
        self.assertEqual(self.jefe.get("/api/evaluaciones/procesos").json(), [])
        archivados = self.jefe.get("/api/evaluaciones/procesos?archivados=true").json()
        self.assertEqual([p["codigo"] for p in archivados], [proceso.codigo])
        self.assertIsNotNone(archivados[0]["archivado_en"])
        self.assertEqual(archivar(self.jefe, "?archivar=false").status_code, 204)
        self.assertEqual(len(self.jefe.get("/api/evaluaciones/procesos").json()), 1)
        # La consulta no archiva.
        consulta = Cliente()
        consulta.entrar("control@entidad.gov.co")
        self.assertEqual(archivar(consulta).status_code, 403)


@override_settings(MEDIA_ROOT="/tmp/mievaluador-pruebas-informes")
class RotulosTests(BaseFlujo):
    def test_preinforme_y_informe_adoptado(self):
        self.preparar(MUCHOS[:2])
        r = self.jefe.get(f"/api/evaluaciones/{self.eid}/consolidado")
        self.assertEqual(r.status_code, 200, r.content)
        libro = load_workbook(io.BytesIO(r.content))
        self.assertIn("Constancia", libro.sheetnames)
        self.assertIn("Pre-informe", libro["Constancia"]["A1"].value)
        self.assertIn("Pre-informe", libro.worksheets[0].oddHeader.center.text)
        self.hacer_muestra(self.abogado, self.evaluador, self.eid)
        self.jefe.post(f"/api/evaluaciones/{self.eid}/aprobar")
        libro = load_workbook(io.BytesIO(self.jefe.get(f"/api/evaluaciones/{self.eid}/consolidado").content))
        self.assertIn("adoptado por Jefe Jurídico", libro["Constancia"]["A1"].value)
        self.assertIn("utilizó la herramienta tecnológica MiEvaluador", libro["Constancia"]["A3"].value)
        self.assertNotIn("Pre-informe", libro.worksheets[0].oddHeader.center.text)

    def test_trazabilidad_en_cada_resultado(self):
        self.preparar(MUCHOS[:1])
        r = Resultado.objects.filter(evaluacion_id=self.eid).first()
        self.assertTrue(r.trazabilidad["version_sistema"])
        self.assertNotIn("modelos_ia", r.trazabilidad)  # no usó IA


class ExpedienteEstadoTests(TestCase):
    def test_estado_listo_existe(self):
        self.assertEqual(EstadoExpediente.LISTO, "listo")


class TransparenciaTests(BaseEvaluaciones):
    def test_ficha_con_la_configuracion_real(self):
        c = Cliente()
        c.entrar("control@entidad.gov.co")
        f = c.get("/api/acerca").json()
        self.assertEqual(f["version"], __import__("django.conf").conf.settings.MIEVALUADOR_VERSION)
        self.assertTrue(any("No decide" in x for x in f["que_no_hace"]))
        self.assertEqual({m["modelo"] for m in f["modelos"]}, set(__import__("evaluaciones.cumplimiento", fromlist=["x"]).modelos_ia().values()))
        # Con Llama configurado, la licencia exige la atribución.
        self.assertEqual(f["atribucion"], "Built with Llama")
        r = c.get("/api/acerca/ficha")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"PK"))
        # Sin sesión no se ve.
        from django.test import Client

        self.assertEqual(Client().get("/api/acerca").status_code, 401)
