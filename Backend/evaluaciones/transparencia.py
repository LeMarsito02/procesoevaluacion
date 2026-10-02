"""Ficha de transparencia algorítmica de MiEvaluador.

Directiva Conjunta 007 de 2025 (Procuraduría General y Defensoría del Pueblo)
y Sentencia T-067 de 2025: las entidades que usan sistemas algorítmicos con IA
deben publicar, en lenguaje claro, su finalidad, cómo funcionan, qué datos
usan y qué efectos tienen, y responder a quien pida explicaciones.

La ficha se arma con la configuración real del servidor (versión, modelos de
IA, parámetros de la muestra de control) y con la última medición de
rendimiento: nunca queda desactualizada respecto de lo que está corriendo.
"""
from __future__ import annotations

import io

from django.conf import settings
from django.utils import timezone

from evaluaciones.cumplimiento import AVISO_INTERFAZ, modelos_ia

# Licencia de cada modelo, verificada en el archivo LICENSE oficial del punto de
# control (Hugging Face o el repositorio del fabricante), NO en la etiqueta de
# Ollama: Ollama marca Qwen2.5-VL-3B como Apache 2.0 y su licencia oficial es
# de investigación (prohíbe el uso comercial). Se recorre en orden: los
# prefijos más específicos van primero. `comercial`: True permitido, False
# prohibido, None sin verificar. Un modelo sin uso comercial verificado no se
# puede aprobar (manage.py inventario_ia).
LICENCIAS = [
    ("llama3.1", "Llama 3.1 Community License", True,
     "Uso comercial permitido; exige mostrar «Built with Llama» y el aviso de copyright de Meta",
     "https://github.com/meta-llama/llama-models/blob/main/models/llama3_1/LICENSE"),
    ("qwen3", "Apache 2.0", True, "", "https://huggingface.co/Qwen/Qwen3-4B"),
    ("qwen2.5vl:3b", "Qwen Research License", False, "Solo investigación: prohíbe el uso comercial",
     "https://huggingface.co/Qwen/Qwen2.5-VL-3B-Instruct/blob/main/LICENSE"),
    ("qwen2.5-vl-3b", "Qwen Research License", False, "Solo investigación: prohíbe el uso comercial",
     "https://huggingface.co/Qwen/Qwen2.5-VL-3B-Instruct/blob/main/LICENSE"),
    ("qwen2.5vl:7b", "Apache 2.0", True, "", "https://huggingface.co/Qwen/Qwen2.5-VL-7B-Instruct"),
    ("qwen2.5-vl-7b", "Apache 2.0", True, "", "https://huggingface.co/Qwen/Qwen2.5-VL-7B-Instruct"),
    ("granite", "Apache 2.0", True, "", "https://huggingface.co/ibm-granite"),
    ("mistral-small", "Apache 2.0", None, "Verificar la versión exacta", "https://huggingface.co/mistralai"),
]

ATRIBUCION_LLAMA = (
    "Built with Llama. Llama 3.1 is licensed under the Llama 3.1 Community License, "
    "Copyright © Meta Platforms, Inc. All Rights Reserved."
)


def licencia_completa(modelo: str) -> dict:
    """Licencia, si permite uso comercial, condiciones y fuente oficial."""
    nombre = modelo.lower()
    for prefijo, licencia, comercial, condiciones, fuente in LICENCIAS:
        if nombre.startswith(prefijo):
            return {"licencia": licencia, "comercial": comercial, "condiciones": condiciones, "fuente": fuente}
    return {"licencia": "Por verificar", "comercial": None, "condiciones": "", "fuente": ""}


USOS = {
    "extraccion_documentos": "Extraer datos de los documentos de la oferta (facultades del representante, integrantes, póliza).",
    "lectura_pliego": "Leer el pliego y proponer, con cita literal y página, los ajustes a la evaluación.",
    "vision": "Leer la fecha de expedición en la imagen de la cédula.",
}


def licencia_de(modelo: str) -> str:
    d = licencia_completa(modelo)
    if d["comercial"] is False:
        return f"{d['licencia']} (NO permite uso comercial)"
    return f"{d['licencia']} ({d['condiciones']})" if d["condiciones"] else d["licencia"]


def ficha() -> dict:
    from evaluaciones.models import MedicionRendimiento

    modelos = modelos_ia()
    medicion = MedicionRendimiento.objects.first()
    global_ = (medicion.datos or {}).get("global") if medicion else None
    usa_llama = any(m.lower().startswith("llama") for m in modelos.values())
    return {
        "sistema": "MiEvaluador",
        "responsable": "LeMarTek Labs S.A.S. · NIT 902069061-9",
        "version": settings.MIEVALUADOR_VERSION,
        "generada_en": timezone.now().isoformat(),
        "finalidad": (
            "Apoyar a los comités evaluadores de las entidades públicas en la verificación de los requisitos "
            "habilitantes (jurídicos, técnicos y financieros) de las ofertas en los procesos de contratación estatal, "
            "contrastando cada documento con las reglas del pliego del proceso."
        ),
        "que_hace": [
            "Identifica cada documento de la oferta por su contenido y lo contrasta con la regla del pliego que la entidad confirmó.",
            "Da un requisito por verificado solo si encuentra en el documento todos los datos que la regla exige; si no, lo deja pendiente de revisión humana.",
            "Muestra, para cada resultado, la regla aplicada, el dato encontrado y el documento soporte.",
            "Propone un puntaje y un orden de elegibilidad preliminares, que el evaluador revisa y adopta.",
            "Registra en una bitácora inmutable quién vio, decidió, adoptó y aprobó cada cosa.",
        ],
        "que_no_hace": [
            "No decide: no rechaza ofertas, no emite «no cumple», no requiere subsanaciones ni responde observaciones.",
            "No adjudica ni recomienda la adjudicación; no firma el informe de evaluación.",
            "No se entrena ni se afina con las ofertas de las entidades.",
            "No envía las ofertas a servicios de inteligencia artificial de terceros.",
            "No evade captchas ni controles de acceso de los portales oficiales.",
        ],
        "como_decide_el_sistema": (
            "La decisión de «verificado» o «pendiente» la toma código determinista y auditable (reglas), no el modelo de "
            "IA. Los modelos solo extraen datos, con temperatura cero, y todo dato extraído se descarta si no aparece "
            "literalmente en el documento. Si la IA no está disponible, el requisito queda pendiente de revisión humana."
        ),
        "control_humano": [
            "Todo requisito pendiente, incumplimiento, observación y subsanación lo revisa una persona, uno por uno, con justificación escrita y viendo el soporte.",
            f"Lo verificado por el sistema se adopta después de revisar, contra su soporte, una muestra de control de {settings.MUESTRA_VERIFICACIONES} verificaciones sorteadas por área, repartidas entre ofertas y requisitos distintos. Un error en la muestra envía todo ese requisito a revisión humana.",
            "El puntaje lo adopta el evaluador técnico después de revisar la tabla de puntajes preliminares.",
            "La evaluación la aprueba el jefe del área; el informe lo suscribe el comité evaluador.",
            "Quien decide acepta un compromiso de uso: la herramienta no traslada, atenúa ni distribuye su responsabilidad.",
        ],
        "datos_tratados": (
            "Documentos de las ofertas: certificados de existencia y representación, RUP, pólizas, antecedentes "
            "disciplinarios, fiscales, judiciales y de policía, cédulas, estados financieros y soportes de experiencia. "
            "La entidad es la responsable del tratamiento y LeMarTek la encargada (Ley 1581 de 2012). Las copias de las "
            f"ofertas se eliminan {settings.RETENCION_DIAS} días después de aprobada la evaluación; se conservan los "
            "resultados, las decisiones y el expediente."
        ),
        "modelos": [
            {"uso": USOS.get(clave, clave), "modelo": modelo, "licencia": licencia_de(modelo),
             "donde_se_ejecuta": "Servidores de LeMarTek (IA local); los documentos no salen a terceros"}
            for clave, modelo in modelos.items()
        ],
        "atribucion": ATRIBUCION_LLAMA if usa_llama else None,
        "medicion": (
            {
                "fecha": medicion.creada_en.isoformat(),
                "proponentes": global_.get("proponentes"),
                "decisiones": global_.get("decisiones"),
                "verificadas_por_el_sistema": global_.get("automaticas"),
                "verificaciones_contradichas_por_el_evaluador": global_.get("indebidas"),
                "limite_superior_de_error_95": global_.get("techo_error"),
            }
            if global_
            else None
        ),
        "aviso": AVISO_INTERFAZ,
        "como_pedir_explicaciones": (
            "Cada resultado se puede explicar con la regla aplicada, el dato encontrado y el documento soporte, que "
            "constan en el expediente de la evaluación. Las solicitudes de información sobre el sistema se dirigen a la "
            "entidad contratante, que cuenta con el apoyo de LeMarTek para responderlas. Las decisiones de la evaluación "
            "se controvierten por el procedimiento contractual (observaciones al informe durante el traslado), porque "
            "MiEvaluador no adopta decisiones automatizadas."
        ),
        "marco": [
            "Concepto C-1015 de 2026 de la ANCP – Colombia Compra Eficiente (uso de IA en contratación estatal)",
            "Directiva Conjunta 007 de 2025 (transparencia algorítmica) y Sentencia T-067 de 2025",
            "Circular Externa 002 de 2024 de la SIC (datos personales en sistemas de IA)",
            "CONPES 4144 de 2025 (Política Nacional de Inteligencia Artificial)",
        ],
    }


def generar_ficha_docx() -> tuple[bytes, str]:
    from docx import Document
    from docx.shared import Pt

    from evaluaciones.reporte import _parrafo, _tabla, _titulo, fecha_larga

    f = ficha()
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    _titulo(doc, "Ficha de transparencia algorítmica – MiEvaluador")
    _parrafo(doc, f"{f['responsable']} · versión {f['version']} · generada el {fecha_larga(timezone.now())}.", tam=9)
    _titulo(doc, "Finalidad", nivel=2)
    _parrafo(doc, f["finalidad"])
    _titulo(doc, "Qué hace", nivel=2)
    for x in f["que_hace"]:
        doc.add_paragraph(x, style="List Bullet")
    _titulo(doc, "Qué no hace", nivel=2)
    for x in f["que_no_hace"]:
        doc.add_paragraph(x, style="List Bullet")
    _titulo(doc, "Cómo decide el sistema", nivel=2)
    _parrafo(doc, f["como_decide_el_sistema"])
    _titulo(doc, "Control humano", nivel=2)
    for x in f["control_humano"]:
        doc.add_paragraph(x, style="List Bullet")
    _titulo(doc, "Datos que trata", nivel=2)
    _parrafo(doc, f["datos_tratados"])
    _titulo(doc, "Modelos de inteligencia artificial", nivel=2)
    _tabla(doc, ["Uso", "Modelo", "Licencia", "Dónde se ejecuta"],
           [[m["uso"], m["modelo"], m["licencia"], m["donde_se_ejecuta"]] for m in f["modelos"]], [5.0, 3.0, 4.6, 4.0])
    if f["atribucion"]:
        _parrafo(doc, f["atribucion"], negrita=True)
    if f["medicion"]:
        m = f["medicion"]
        _titulo(doc, "Desempeño medido", nivel=2)
        techo = m["limite_superior_de_error_95"]
        _parrafo(
            doc,
            f"Medición del {m['fecha'][:10]} contra evaluaciones reales: {m['proponentes']} proponentes, {m['decisiones']} "
            f"decisiones, {m['verificadas_por_el_sistema']} verificadas por el sistema y {m['verificaciones_contradichas_por_el_evaluador']} "
            f"contradichas por el evaluador" + (f" (límite superior de error del {techo:.2%} con 95 % de confianza)." if techo is not None else "."),
        )
    _titulo(doc, "Cómo pedir explicaciones", nivel=2)
    _parrafo(doc, f["como_pedir_explicaciones"])
    _titulo(doc, "Marco de referencia", nivel=2)
    for x in f["marco"]:
        doc.add_paragraph(x, style="List Bullet")
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue(), f"FICHA TRANSPARENCIA MIEVALUADOR {f['version']}.docx"
