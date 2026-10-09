"""Prestación de servicios en la plataforma: licencia, permisos, análisis en
el trabajador, decisiones de la persona y certificado. Datos inventados."""
from __future__ import annotations

import io
import shutil
import tempfile
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import override_settings
from docx import Document

from cuentas.models import EventoAuditoria
from cuentas.tests import Cliente
from evaluaciones import ops
from evaluaciones.models import ContratacionOps, EstadoOps, TablaHonorariosOps
from evaluaciones.tests import BaseEvaluaciones
from evaluaciones.tests_ops import ACTA_POSGRADO, ACTA_PREGRADO, CONTRACTUAL, ESTUDIO_COMPLETO, LABORAL, TERMINADO

NOMBRE, CEDULA = "Ana Inventada Pérez", "1000000001"


def _empleo(entidad, ingreso, retiro, cargo="CONTRATISTA"):
    (d1, m1, a1), (d2, m2, a2) = ingreso, retiro
    return (f"EMPLEO O CONTRATO ANTERIOR\nEMPRESA O ENTIDAD PÚBLICA PRIVADA PAÍS\n{entidad} X COLOMBIA\nTELÉFONOS FECHA DE INGRESO FECHA DE RETIRO\n"
            f"Día {d1:02d} Mes {m1:02d} Año {a1} Día {d2:02d} Mes {m2:02d} Año {a2}\nCARGO O CONTRATO ACTUAL DEPENDENCIA DIRECCIÓN\n{cargo} PLANEACION CALLE 1\n")


# Hoja de vida del SIGEP que relaciona toda la experiencia certificada.
HOJA_DE_VIDA = ("FORMATO ÚNICO\nHOJA DE VIDA\nPersona Natural\nANA INVENTADA PEREZ 1000000001\nEXPERIENCIA LABORAL\n"
                + _empleo("INSTITUTO DE OBRAS DE VILLA FICTICIA", (1, 2, 2024), (30, 11, 2024))
                + _empleo("CORPORACION DE RIO INVENTADO", (1, 1, 2022), (30, 6, 2022))
                + _empleo("DEPARTAMENTO DE VALLE INVENTADO", (3, 2, 2020), (30, 6, 2021))
                + _empleo("DEPARTAMENTO DE VALLE INVENTADO", (14, 2, 2012), (31, 12, 2019))
                + _empleo("CORPORACION REGIONAL DE RIO INVENTADO", (12, 2, 2009), (11, 8, 2010))
                + "TELÉFONOS FECHA DE INGRESO FECHA DE RETIRO\nDía: Mes: Año: Día: Mes: Año:\nTIEMPO TOTAL DE EXPERIENCIA\n")
PAGINAS = {
    b"estudio": [ESTUDIO_COMPLETO],
    b"cdp": ["CERTIFICADO DE DISPONIBILIDAD PRESUPUESTAL\nTOTAL CDP 90,000,000.00"],
    b"certificaciones": [CONTRACTUAL.replace("ANA INVENTADA PÉREZ", f"ANA INVENTADA PÉREZ, cédula {CEDULA},"), LABORAL, TERMINADO],
    b"titulos": [ACTA_PREGRADO, ACTA_POSGRADO],
    b"rut": [f"Formulario del Registro Único Tributario {CEDULA}"],
    b"sigep": [HOJA_DE_VIDA],
    b"otra": ["CORPORACIÓN DE RÍO INVENTADO\nCERTIFICA\nFECHA DE INICIO: 01/01/2022\nFECHA DE TERMINACIÓN: 30/06/2022"],
}


def _paginas(contenido: bytes, max_paginas: int | None = None) -> list[str]:
    return PAGINAS[contenido][:max_paginas]


def _pdf(nombre: str, contenido: bytes) -> SimpleUploadedFile:
    return SimpleUploadedFile(nombre, contenido, content_type="application/pdf")


class BaseOps(BaseEvaluaciones):
    """Una entidad con el módulo activo, un expediente inventado completo y lo que haría quien lo revisa."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.carpeta = tempfile.mkdtemp()
        cls.ajustes = override_settings(MEDIA_ROOT=cls.carpeta)
        cls.ajustes.enable()

    @classmethod
    def tearDownClass(cls):
        cls.ajustes.disable()
        shutil.rmtree(cls.carpeta, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        call_command("honorarios_iccu", nit=self.entidad1.nit, stdout=io.StringIO())
        self.entidad1.refresh_from_db()
        self.abogado = Cliente()
        self.abogado.entrar("abogado@entidad.gov.co")

    def _crear(self, c=None, **campos):
        c = c or self.abogado
        datos = {
            "contratista_nombre": NOMBRE, "contratista_cedula": "1.000.000.001", "fecha_referencia": "2026-01-22", "exige_libreta": "no",
            "documentos_entidad": [_pdf("estudio previo.pdf", b"estudio"), _pdf("cdp.pdf", b"cdp")],
            "documentos_contratista": [_pdf("14. certificaciones.pdf", b"certificaciones"), _pdf("19. titulos.pdf", b"titulos"),
                                       _pdf("9. RUT.pdf", b"rut"), _pdf("otra certificacion.pdf", b"otra"),
                                       _pdf("12. Hoja de vida SIGEP.pdf", b"sigep")],
            **campos,
        }
        return c.http.post("/api/ops", datos, headers=c._h())

    def _analizar(self):
        cumple = mock.Mock(cumple=True, archivo=None)
        with (
            mock.patch("motor.ops.expediente.paginas_de_texto", _paginas),
            mock.patch("motor.ops.documentos.paginas_de_texto", _paginas),
            mock.patch("motor.ops.documentos.antecedentes.evaluar_antecedente", return_value=cumple),
        ):
            return ops.atender_pendientes()

    def _lista(self):
        r = self._crear()
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()["estado"], "pendiente")
        self.assertEqual(self._analizar(), 1)
        return r.json()["id"]

    def _resolver(self, cid, c=None):
        """Lo que haría quien revisa: decide los documentos pendientes y confirma las lecturas dudosas."""
        c = c or self.abogado
        d = c.get(f"/api/ops/{cid}").json()
        cambios = {
            "documentos": {x["clave"]: {"estado": "cumple", "nota": "Verificado en el expediente físico"} for x in d["documentos"] if x["estado_final"] != "cumple"},
            "periodos": {f["id"]: {"confirmado": True} for f in d["experiencia"] if f["por_confirmar"]},
            "avisos_vistos": [a["texto"] for a in d["avisos"]],
        }
        r = c.put(f"/api/ops/{cid}/decisiones", cambios)
        self.assertEqual(r.status_code, 200, r.content)
        return r.json()


class OpsPlataformaTests(BaseOps):
    def test_sin_licencia_la_opcion_queda_con_candado(self):
        self.entidad1.modulo_ops = False
        self.entidad1.save(update_fields=["modulo_ops"])
        self.assertEqual(self.abogado.get("/api/ops").json()["licenciado"], False)
        r = self._crear()
        self.assertEqual((r.status_code, r.json()["detail"]), (403, "La entidad no tiene la licencia del módulo de prestación de servicios."))
        self.assertEqual(ContratacionOps.objects.count(), 0)

    def test_crear_exige_rol_nombre_cedula_y_un_pdf(self):
        control = Cliente()
        control.entrar("control@entidad.gov.co")
        self.assertEqual(self._crear(control).status_code, 403)
        self.assertEqual(self._crear(contratista_cedula="12").status_code, 400)
        r = self._crear(documentos_contratista=[SimpleUploadedFile("hoja.docx", b"texto")])
        self.assertEqual((r.status_code, r.json()["detail"]), (400, "«hoja.docx» no es un PDF ni un .zip."))
        self.assertEqual(ContratacionOps.objects.count(), 0)

    def test_analisis_lee_perfil_experiencia_y_documentos(self):
        d = self.abogado.get(f"/api/ops/{self._lista()}").json()
        self.assertEqual((d["estado"], d["contratista_cedula"], d["cdp"]["valor"], d["estudio"]["plazo_meses"]), ("lista", CEDULA, 90_000_000, 10))
        self.assertEqual((d["perfil"]["anios_minimos"], d["perfil"]["anios_maximos"], d["perfil"]["posgrado"], d["perfil"]["honorarios_mensuales"]),
                         (10, 15, "especializacion", 9_000_000))
        self.assertEqual((d["grado"], d["posgrado"]), ("2004-06-11", "especializacion"))
        # 2009-2010: 540 · 2012-2019: 2837 · 2020-2021: 508 · 2022 (archivo suelto): 180 · 2024: 285
        self.assertEqual([(f["inicio"][:4], f["estado"], f["dias"]) for f in d["experiencia"]],
                         [("2009", "cuenta", 540), ("2012", "cuenta", 2837), ("2020", "cuenta", 508), ("2022", "cuenta", 180), ("2024", "cuenta", 285)])
        self.assertEqual((d["total_dias"], d["franja"], d["tope"], d["relacionada_dias"]), (4350, "entre 10 y 15 años", 9_588_365, 285))
        self.assertEqual(d["experiencia"][-1]["obligaciones_iguales"], [1])
        estados = {x["clave"]: x["estado"] for x in d["documentos"]}
        self.assertEqual((estados["rut"], estados["titulos"], estados["certificaciones_laborales"], estados["disciplinarios"], estados["cedula"]),
                         ("cumple", "cumple", "cumple", "cumple", "falta"))
        self.assertIsNone(d["cumple"])
        self.assertGreater(d["documentos_pendientes"], 10)

    def test_la_persona_retira_un_periodo_corrige_el_perfil_y_un_analisis_nuevo_no_lo_pisa(self):
        cid = self._lista()
        d = self.abogado.get(f"/api/ops/{cid}").json()
        viejo = d["experiencia"][0]["id"]
        r = self.abogado.put(f"/api/ops/{cid}/decisiones", {
            "periodos": {viejo: {"incluir": False}}, "grado": "2013-01-01",
            "perfil": {"anios_minimos": 5, "anios_maximos": 10, "posgrado": "maestria", "honorarios_mensuales": 9_000_000},
        })
        self.assertEqual(r.status_code, 200, r.content)
        d = r.json()
        self.assertEqual([(f["inicio"][:4], f["estado"]) for f in d["experiencia"]][:2], [("2009", "retirado"), ("2012", "cuenta")])
        # Desde el grado de 2013: 2520 (2013-2019) + 508 + 180 + 285
        self.assertEqual((d["total_dias"], d["franja"], d["perfil"]["corregido"], d["grado_corregido"]), (3493, "entre 5 y 10 años", True, True))
        self.assertIn("Los honorarios ($9.000.000) superan el tope de la franja entre 5 y 10 años ($8.045.300).", d["revisiones"])
        self.assertEqual(self.abogado.put(f"/api/ops/{cid}/decisiones", {"periodos": {"p99": {"incluir": False}}}).status_code, 400)
        self.assertEqual(self.abogado.post(f"/api/ops/{cid}/reanalizar").json()["estado"], "pendiente")
        self._analizar()
        self.assertEqual(self.abogado.get(f"/api/ops/{cid}").json()["total_dias"], 3493)
        self.assertTrue(EventoAuditoria.objects.filter(accion="ops.decision", objeto_id=cid).exists())

    def test_sobra_experiencia_se_propone_retirar_y_la_persona_puede_conservarla(self):
        cid = self._lista()
        perfil = {"anios_minimos": 5, "anios_maximos": 10, "posgrado": "especializacion", "honorarios_mensuales": 8_000_000}
        d = self.abogado.put(f"/api/ops/{cid}/decisiones", {"perfil": perfil}).json()
        # Sobran los más antiguos que no dejan la experiencia por debajo del mínimo; la relacionada se conserva.
        self.assertEqual([(f["inicio"][:4], f["estado"]) for f in d["experiencia"]],
                         [("2009", "sobra"), ("2012", "cuenta"), ("2020", "sobra"), ("2022", "cuenta"), ("2024", "cuenta")])
        self.assertEqual(d["total_dias"], 4350 - 540 - 508)
        todos = {f["id"]: {"incluir": True} for f in d["experiencia"]}
        d = self.abogado.put(f"/api/ops/{cid}/decisiones", {"periodos": todos}).json()
        self.assertEqual((d["experiencia"][0]["estado"], d["experiencia"][0]["fijo"], d["total_dias"]), ("cuenta", True, 4350))
        self.assertIn("pasa de los 10 años del perfil con los periodos que se conservan.", d["revisiones"][-1])

    def test_confirmar_exige_decidir_los_documentos_y_bloquea_los_cambios(self):
        cid = self._lista()
        r = self.abogado.post(f"/api/ops/{cid}/confirmar")
        self.assertEqual(r.status_code, 409)
        self.assertIn("Decidir 17 documentos", r.json()["detail"])
        d = self._resolver(cid)
        self.assertEqual((d["documentos_pendientes"], d["por_confirmar"], d["cumple"]), (0, [], True))
        d = self.abogado.post(f"/api/ops/{cid}/confirmar").json()
        self.assertEqual((d["estado"], d["confirmada_por"]), ("confirmada", "Abogado Uno"))
        self.assertEqual(self.abogado.put(f"/api/ops/{cid}/decisiones", {"grado": "2010-01-01"}).status_code, 409)
        self.assertEqual(self.abogado.delete(f"/api/ops/{cid}").status_code, 409)
        self.assertEqual(self.abogado.post(f"/api/ops/{cid}/reabrir").json()["estado"], "lista")

    def test_certificado_en_word(self):
        cid = self._lista()
        r = self.abogado.get(f"/api/ops/{cid}/certificado")
        self.assertEqual(r.status_code, 200)
        doc = Document(io.BytesIO(b"".join(r.streaming_content)))
        texto = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Total de experiencia: 12 años, 1 meses y 0 días.", texto)
        self.assertIn("Experiencia exigida: entre 10 y 15 años.", texto)
        self.assertIn("Tope de la franja entre 10 y 15 años según Resolución 1750 de 2025: $ 9.588.365.", texto)
        filas = [[c.text for c in f.cells] for f in doc.tables[2].rows]
        self.assertEqual(filas[-1], ["INSTITUTO DE OBRAS DE VILLA FICTICIA", "IOVF-045-2024", "01/02/2024", "30/11/2024", "9 meses y 15 días", "Relacionada"])

    def test_lista_de_verificacion_en_word(self):
        cid = self._lista()
        self.abogado.put(f"/api/ops/{cid}/decisiones", {"documentos": {"cedula": {"estado": "cumple", "nota": "Vista en el expediente físico"}}})
        r = self.abogado.get(f"/api/ops/{cid}/lista-de-verificacion")
        self.assertEqual(r.status_code, 200)
        doc = Document(io.BytesIO(b"".join(r.streaming_content)))
        filas = {f.cells[1].text: [c.text for c in f.cells] for f in doc.tables[1].rows}
        self.assertEqual(filas["Cédula de ciudadanía"][2:], ["Cumple", "—", "Vista en el expediente físico"])
        self.assertEqual(filas["Registro Único Tributario (RUT)"][2:4], ["Cumple", "9. RUT.pdf"])
        self.assertTrue(any("documentos por revisar" in p.text for p in doc.paragraphs))

    def test_solo_la_ven_quien_la_creo_y_el_administrador(self):
        cid = self._lista()
        archivo = ContratacionOps.objects.get(pk=cid).documentos.first()
        for email, ve in (("admin@entidad.gov.co", True), ("abogado2@entidad.gov.co", False), ("control@entidad.gov.co", False),
                          ("admin@otraentidad.gov.co", False)):
            c = Cliente()
            c.entrar(email)
            self.assertEqual(c.get(f"/api/ops/{cid}").status_code, 200 if ve else 404, email)
            self.assertEqual(len(c.get("/api/ops").json()["contrataciones"]), 1 if ve else 0, email)
            self.assertEqual(c.get(f"/api/ops/{cid}/documentos/{archivo.id}").status_code, 200 if ve else 404, email)

    def test_tabla_de_honorarios_la_registra_el_administrador(self):
        tabla = {"vigencia": 2027, "norma": "Resolución 9 de 2026", "reconocimiento": [],
                 "profesional": [{"desde": 0, "hasta": 5, "sin_especializacion": 1, "con_especializacion": 2, "con_maestria": 3},
                                 {"desde": 5, "hasta": None, "sin_especializacion": 4, "con_especializacion": 5, "con_maestria": 6}]}
        self.assertEqual(self.abogado.put("/api/ops/honorarios", tabla).status_code, 403)
        admin = Cliente()
        admin.entrar("admin@entidad.gov.co")
        r = admin.put("/api/ops/honorarios", tabla)
        self.assertEqual((r.status_code, sorted(r.json()["vigencias"])), (200, [2026, 2027]))
        cruzada = {**tabla, "profesional": [{**tabla["profesional"][0], "hasta": None}, tabla["profesional"][1]]}
        self.assertEqual(admin.put("/api/ops/honorarios", cruzada).status_code, 400)
        self.assertEqual(TablaHonorariosOps.objects.filter(entidad=self.entidad1).count(), 2)

    def test_nombre_y_cedula_se_leen_de_los_antecedentes_si_no_se_escriben(self):
        r = self._crear(contratista_nombre="", contratista_cedula="")
        self.assertEqual(r.status_code, 201, r.content)
        cid = r.json()["id"]
        with mock.patch("evaluaciones.ops.identificar", return_value=("ANA INVENTADA PEREZ", CEDULA)):
            self._analizar()
        d = self.abogado.get(f"/api/ops/{cid}").json()
        self.assertEqual((d["estado"], d["contratista_nombre"], d["contratista_cedula"]), ("lista", "Ana Inventada Perez", CEDULA))
        self.assertEqual(self._crear(contratista_nombre="Ana", contratista_cedula="").status_code, 400)

    def test_si_no_se_pueden_leer_se_piden_y_al_corregirlos_se_vuelve_a_leer(self):
        cid = self._crear(contratista_nombre="", contratista_cedula="").json()["id"]
        self._analizar()
        d = self.abogado.get(f"/api/ops/{cid}").json()
        self.assertEqual(d["estado"], "error")
        self.assertIn("No se pudo leer el nombre y la cédula del contratista", d["error"])
        d = self.abogado.put(f"/api/ops/{cid}/datos", {"contratista_nombre": NOMBRE, "contratista_cedula": CEDULA}).json()
        self.assertEqual(d["estado"], "pendiente")
        self._analizar()
        self.assertEqual(self.abogado.get(f"/api/ops/{cid}").json()["estado"], "lista")
        # Cambiar solo el número del proceso no obliga a leer todo otra vez.
        self.assertEqual(self.abogado.put(f"/api/ops/{cid}/datos", {"referencia": "CPS-9-2026"}).json()["estado"], "lista")

    def test_antecedentes_de_otra_cedula_se_avisan(self):
        cid = self._crear().json()["id"]
        with mock.patch("evaluaciones.ops.identificar", return_value=("PEDRO OTRO RUIZ", "2000000002")):
            self._analizar()
        d = self.abogado.get(f"/api/ops/{cid}").json()
        self.assertIn("Los certificados de antecedentes son de la cédula 2000000002 (PEDRO OTRO RUIZ), no de la 1000000001", d["avisos"][0]["texto"])
        self.assertIn("Revisar 1 avisos de la lectura de los documentos.", d["por_confirmar"])

    def test_la_persona_agrega_un_periodo_corrige_fechas_y_el_posgrado(self):
        cid = self._lista()
        nuevo = {"inicio": "2005-01-01", "fin": "2005-12-31", "entidad": "Municipio de Pueblo Nuevo", "referencia": "Jefe de planeación", "relacionada": True}
        self.abogado.put(f"/api/ops/{cid}/decisiones", {"agregar": nuevo})
        d = self.abogado.put(f"/api/ops/{cid}/decisiones", {"periodos": {"m1": {"en_sigep": True}}}).json()
        fila = d["experiencia"][0]
        self.assertEqual((fila["id"], fila["manual"], fila["dias"], fila["relacionada"], d["total_dias"], d["relacionada_dias"]),
                         ("m1", True, 360, True, 4350 + 360, 285 + 360))
        self.assertEqual(self.abogado.put(f"/api/ops/{cid}/decisiones", {"agregar": {**nuevo, "fin": "2004-01-01"}}).status_code, 400)
        # Corrige la terminación de un periodo leído: ya no trae la nota de lectura dudosa.
        leido = next(f for f in d["experiencia"] if f["inicio"].startswith("2022"))
        d = self.abogado.put(f"/api/ops/{cid}/decisiones", {"periodos": {leido["id"]: {"fin": "2022-05-30"}}}).json()
        corregido = next(f for f in d["experiencia"] if f["id"] == leido["id"])
        self.assertEqual((corregido["fin"], corregido["corregido"], corregido["dias"]), ("2022-05-30", True, 150))
        self.assertEqual(self.abogado.put(f"/api/ops/{cid}/decisiones", {"periodos": {leido["id"]: {"fin": "2021-01-01"}}}).status_code, 400)
        d = self.abogado.put(f"/api/ops/{cid}/decisiones", {"quitar": "m1", "posgrado": "maestria"}).json()
        self.assertEqual((d["experiencia"][0]["id"] != "m1", d["posgrado"], d["posgrado_corregido"], d["tope"]), (True, "maestria", True, 10_234_769))
        self.assertEqual(self.abogado.put(f"/api/ops/{cid}/decisiones", {"quitar": leido["id"]}).status_code, 400)

    def test_retencion_borra_los_documentos_de_lo_confirmado_y_conserva_lo_decidido(self):
        from datetime import timedelta

        from django.utils import timezone

        from evaluaciones import retencion

        cid = self._lista()
        d = self._resolver(cid)
        self.assertEqual(self.abogado.post(f"/api/ops/{cid}/confirmar").json()["estado"], "confirmada")
        total = d["total_dias"]
        # Recién confirmada no se toca; pasado el plazo, se borran las copias.
        self.assertEqual(retencion.aplicar(dias=30).contrataciones, 0)
        ContratacionOps.objects.filter(pk=cid).update(confirmada_en=timezone.now() - timedelta(days=31))
        archivo = ContratacionOps.objects.get(pk=cid).documentos.first().archivo.path
        self.assertEqual(retencion.aplicar(dias=30, simulacro=True).contrataciones, 1)
        self.assertTrue(ContratacionOps.objects.get(pk=cid).documentos.exists())
        informe = retencion.aplicar(dias=30)
        self.assertEqual((informe.contrataciones, informe.archivos_ops), (1, 7))
        import os

        self.assertFalse(os.path.exists(archivo))
        c = ContratacionOps.objects.get(pk=cid)
        self.assertEqual((c.documentos.count(), all(p["texto"] == "" for p in c.resultado["periodos"])), (0, True))
        d = self.abogado.get(f"/api/ops/{cid}").json()
        self.assertEqual((d["total_dias"], d["estado"], bool(d["documentos_eliminados_en"]), d["archivos"]), (total, "confirmada", True, []))
        self.assertEqual(self.abogado.get(f"/api/ops/{cid}/certificado").status_code, 200)
        self.abogado.post(f"/api/ops/{cid}/reabrir")
        self.assertEqual(self.abogado.post(f"/api/ops/{cid}/reanalizar").status_code, 409)
        self.assertEqual(retencion.aplicar(dias=30).contrataciones, 0)

    def test_los_antecedentes_no_caducan(self):
        self._crear()
        cumple = mock.Mock(cumple=True, archivo=None)
        with (
            mock.patch("motor.ops.expediente.paginas_de_texto", _paginas),
            mock.patch("motor.ops.documentos.paginas_de_texto", _paginas),
            mock.patch("motor.ops.documentos.antecedentes.evaluar_antecedente", return_value=cumple) as evaluar,
        ):
            ops.atender_pendientes()
        self.assertEqual({llamada.kwargs["fecha_cierre"] for llamada in evaluar.call_args_list}, {None})

    def test_un_analisis_que_falla_queda_en_error_y_se_puede_repetir(self):
        r = self._crear()
        with mock.patch("evaluaciones.ops.analizar", side_effect=RuntimeError("PDF dañado")):
            ops.atender_pendientes()
        c = ContratacionOps.objects.get(pk=r.json()["id"])
        self.assertEqual((c.estado, c.error), (EstadoOps.ERROR, "No se pudo analizar: PDF dañado"))
        self.assertEqual(self.abogado.post(f"/api/ops/{c.id}/reanalizar").json()["estado"], "pendiente")


class NingunAprobadoIndebidoTests(BaseOps):
    """Se parte de una contratación que cumple sin nada pendiente y se le daña
    una cosa a la vez. En ningún caso el sistema puede decir «cumple», y sin
    conclusión tampoco deja confirmar."""

    def _no_aprueba(self, cid, esperado, motivo):
        d = self.abogado.get(f"/api/ops/{cid}").json()
        self.assertIs(d["cumple"], esperado, d["por_confirmar"] or d["revisiones"])
        self.assertTrue(any(motivo in x for x in [*d["por_confirmar"], *d["revisiones"]]), (motivo, d["por_confirmar"], d["revisiones"]))
        if esperado is None:
            self.assertEqual(self.abogado.post(f"/api/ops/{cid}/confirmar").status_code, 409)
        return d

    def test_control_la_contratacion_limpia_si_cumple(self):
        cid = self._lista()
        self.assertIsNone(self.abogado.get(f"/api/ops/{cid}").json()["cumple"])
        self.assertIs(self._resolver(cid)["cumple"], True)

    def test_un_documento_sin_decidir(self):
        cid = self._lista()
        self._resolver(cid)
        self.abogado.put(f"/api/ops/{cid}/decisiones", {"documentos": {"cedula": None}})
        self._no_aprueba(cid, None, "Decidir 1 documentos")

    def test_un_documento_que_la_persona_marca_no_cumple(self):
        cid = self._lista()
        self._resolver(cid)
        self.abogado.put(f"/api/ops/{cid}/decisiones", {"documentos": {"cedula": {"estado": "no_cumple", "nota": "Es de otra persona"}}})
        self._no_aprueba(cid, None, "Decidir 1 documentos")

    def test_una_certificacion_que_no_nombra_a_la_persona(self):
        cid = self._lista()
        d = self.abogado.get(f"/api/ops/{cid}").json()
        ajena = next(f for f in d["experiencia"] if f["por_confirmar"])
        self.assertIn("No se leyó en la certificación el nombre ni la cédula de la persona", ajena["nota"])
        d = self.abogado.put(f"/api/ops/{cid}/decisiones", {
            "documentos": {x["clave"]: {"estado": "cumple", "nota": ""} for x in d["documentos"] if x["estado_final"] != "cumple"}}).json()
        self._no_aprueba(cid, None, "Confirmar o corregir 1 periodos")
        # Si la persona la retira, deja de contar y ya no bloquea.
        d = self.abogado.put(f"/api/ops/{cid}/decisiones", {"periodos": {ajena["id"]: {"incluir": False}}}).json()
        self.assertEqual((d["cumple"], d["total_dias"]), (True, 4350 - 180))

    def test_sin_tabla_de_honorarios_no_hay_tope_que_verificar(self):
        cid = self._lista()
        self._resolver(cid)
        TablaHonorariosOps.objects.filter(entidad=self.entidad1).delete()
        self._no_aprueba(cid, None, "Registrar la tabla de honorarios de 2026")

    def test_honorarios_por_encima_del_tope(self):
        cid = self._lista()
        self._resolver(cid)
        perfil = {"anios_minimos": 10, "anios_maximos": 15, "posgrado": "especializacion", "honorarios_mensuales": 9_600_000}
        self.abogado.put(f"/api/ops/{cid}/decisiones", {"perfil": perfil})
        self._no_aprueba(cid, False, "superan el tope de la franja entre 10 y 15 años ($9.588.365)")

    def test_experiencia_que_no_llega_al_minimo(self):
        cid = self._lista()
        self._resolver(cid)
        perfil = {"anios_minimos": 20, "anios_maximos": 30, "posgrado": "especializacion", "honorarios_mensuales": 9_000_000}
        self.abogado.put(f"/api/ops/{cid}/decisiones", {"perfil": perfil})
        self._no_aprueba(cid, False, "no llega a los 20 años del perfil")

    def test_posgrado_exigido_que_no_acredita(self):
        cid = self._lista()
        self._resolver(cid)
        perfil = {"anios_minimos": 10, "anios_maximos": 15, "posgrado": "maestria", "honorarios_mensuales": 9_000_000}
        d = self.abogado.put(f"/api/ops/{cid}/decisiones", {"perfil": perfil}).json()
        self.assertIs(d["cumple"], False)

    def test_experiencia_especifica_exigida_sin_acreditar(self):
        cid = self._lista()
        self._resolver(cid)
        perfil = {"anios_minimos": 10, "anios_maximos": 15, "posgrado": "especializacion", "honorarios_mensuales": 9_000_000,
                  "especifica_minima": 2, "especifica_maxima": 5}
        self.abogado.put(f"/api/ops/{cid}/decisiones", {"perfil": perfil})
        self._no_aprueba(cid, False, "El perfil pide 2 años de experiencia específica")

    def test_diploma_que_no_se_leyo(self):
        with mock.patch.dict(PAGINAS, {b"titulos": ["página escaneada ilegible"]}):
            cid = self._lista()
            self._resolver(cid)
        d = self._no_aprueba(cid, None, "Confirmar la fecha de grado")
        self.assertEqual(d["posgrado"], "ninguno")
        # Con la fecha puesta por una persona, se concluye (y aquí no cumple: falta el posgrado).
        d = self.abogado.put(f"/api/ops/{cid}/decisiones", {"grado": "2004-06-11"}).json()
        self.assertEqual((d["por_confirmar"], d["cumple"]), ([], False))

    def test_antecedentes_de_otra_persona(self):
        cid = self._crear().json()["id"]
        with mock.patch("evaluaciones.ops.identificar", return_value=("PEDRO OTRO RUIZ", "2000000002")):
            self._analizar()
        d = self.abogado.get(f"/api/ops/{cid}").json()
        cambios = {
            "documentos": {x["clave"]: {"estado": "cumple", "nota": ""} for x in d["documentos"] if x["estado_final"] != "cumple"},
            "periodos": {f["id"]: {"confirmado": True} for f in d["experiencia"] if f["por_confirmar"]},
        }
        self.abogado.put(f"/api/ops/{cid}/decisiones", cambios)
        self._no_aprueba(cid, None, "Revisar 1 avisos de la lectura")

    def test_un_analisis_nuevo_con_una_lectura_dudosa_vuelve_a_pedir_confirmacion(self):
        cid = self._lista()
        self.assertIs(self._resolver(cid)["cumple"], True)
        # Llega otra certificación sin nombre: lo ya confirmado no la cubre.
        r = self.abogado.http.post(f"/api/ops/{cid}/documentos", {"origen": "contratista", "archivos": [_pdf("nueva.pdf", b"nueva")]}, headers=self.abogado._h())
        self.assertEqual(r.status_code, 200, r.content)
        with mock.patch.dict(PAGINAS, {b"nueva": ["EMPRESA INVENTADA S.A.S\nCERTIFICA\nFECHA DE INICIO: 01/01/2023\nFECHA DE TERMINACIÓN: 30/06/2023"]}):
            self._analizar()
        d = self.abogado.get(f"/api/ops/{cid}").json()
        # La confirmación anterior sigue en su periodo (2022), no saltó al nuevo (2023), y el nuevo
        # no cuenta: no está en la hoja de vida del SIGEP.
        self.assertEqual([(f["inicio"][:4], f["confirmado"], f["estado"]) for f in d["experiencia"] if f["inicio"][:4] in ("2022", "2023")],
                         [("2022", True, "cuenta"), ("2023", False, "sin_sigep")])
        self.assertEqual(d["total_dias"], 4350)

    def test_lo_decidido_sobre_un_periodo_no_cae_en_otro_al_leer_de_nuevo(self):
        cid = self._lista()
        d = self.abogado.get(f"/api/ops/{cid}").json()
        retirado = next(f for f in d["experiencia"] if f["inicio"].startswith("2012"))
        self.abogado.put(f"/api/ops/{cid}/decisiones", {"periodos": {retirado["id"]: {"incluir": False}}})
        # Aparece una certificación más antigua: todos los periodos cambian de posición.
        self.abogado.http.post(f"/api/ops/{cid}/documentos", {"origen": "contratista", "archivos": [_pdf("vieja.pdf", b"vieja")]}, headers=self.abogado._h())
        with mock.patch.dict(PAGINAS, {b"vieja": ["ANA INVENTADA PÉREZ laboró desde el 1 de enero de 2005 hasta el 30 de junio de 2005"]}):
            self._analizar()
        d = self.abogado.get(f"/api/ops/{cid}").json()
        self.assertEqual([(f["inicio"][:4], f["estado"]) for f in d["experiencia"]],
                         [("2005", "sin_sigep"), ("2009", "cuenta"), ("2012", "retirado"), ("2020", "cuenta"), ("2022", "cuenta"), ("2024", "cuenta")])


class SoloLoQueEstaEnElSigepTests(BaseOps):
    """La experiencia que cuenta debe estar relacionada en la hoja de vida del SIGEP."""

    def test_un_periodo_certificado_que_no_esta_en_el_sigep_no_cuenta(self):
        sin_2022 = HOJA_DE_VIDA.replace(_empleo("CORPORACION DE RIO INVENTADO", (1, 1, 2022), (30, 6, 2022)), "")
        with mock.patch.dict(PAGINAS, {b"sigep": [sin_2022]}):
            cid = self._lista()
        d = self.abogado.get(f"/api/ops/{cid}").json()
        fuera = next(f for f in d["experiencia"] if f["inicio"].startswith("2022"))
        self.assertEqual((fuera["estado"], fuera["motivo"], fuera["dias"]),
                         ("sin_sigep", "No está relacionado en la hoja de vida del SIGEP: no es experiencia válida.", 0))
        self.assertEqual((d["total_dias"], d["sigep"]), (4350 - 180, {"empleos": 4, "leidos": 4, "confiable": True}))
        # Quien revisa lo ve en la hoja de vida (el sistema lo leyó mal): lo marca y vuelve a contar, y queda en la auditoría.
        d = self.abogado.put(f"/api/ops/{cid}/decisiones", {"periodos": {fuera["id"]: {"en_sigep": True, "confirmado": True}}}).json()
        vuelve = next(f for f in d["experiencia"] if f["id"] == fuera["id"])
        self.assertEqual((vuelve["estado"], vuelve["sigep"], vuelve["sigep_confirmado"], d["total_dias"]), ("cuenta", "si", True, 4350))
        self.assertTrue(EventoAuditoria.objects.filter(accion="ops.decision", objeto_id=cid).exists())

    def _dudas_de_2024(self, cid):
        d = self.abogado.get(f"/api/ops/{cid}").json()
        fila = next(f for f in d["experiencia"] if f["inicio"].startswith("2024"))
        # Sin poder cruzar no se concluye, pero el periodo sigue contando mientras alguien lo confirma.
        self.assertEqual((d["cumple"], fila["estado"], fila["por_confirmar"]), (None, "cuenta", True))
        return fila["duda"]

    def test_con_la_hoja_de_vida_ilegible_no_se_concluye(self):
        ilegible = HOJA_DE_VIDA.replace("Día 01 Mes 02 Año 2024 Día 30 Mes 11 Año 2024", "D1a o1 M3s o2")
        with mock.patch.dict(PAGINAS, {b"sigep": [ilegible]}):
            cid = self._lista()
        self.assertIn("no se dejó leer completa", self._dudas_de_2024(cid))

    def test_sin_la_hoja_de_vida_no_se_concluye(self):
        sin_hoja = [_pdf("14. certificaciones.pdf", b"certificaciones"), _pdf("19. titulos.pdf", b"titulos"), _pdf("9. RUT.pdf", b"rut")]
        r = self._crear(documentos_contratista=sin_hoja)
        self._analizar()
        self.assertIn("No está la hoja de vida del SIGEP entre los documentos", self._dudas_de_2024(r.json()["id"]))

    def test_en_las_mismas_fechas_pero_con_otra_entidad_queda_en_duda(self):
        otra = HOJA_DE_VIDA.replace("INSTITUTO DE OBRAS DE VILLA FICTICIA", "MINISTERIO DE HACIENDA")
        with mock.patch.dict(PAGINAS, {b"sigep": [otra]}):
            cid = self._lista()
        d = self.abogado.get(f"/api/ops/{cid}").json()
        fila = next(f for f in d["experiencia"] if f["inicio"].startswith("2024"))
        self.assertEqual((fila["sigep"], fila["por_confirmar"]), ("si", True))
        self.assertIn("En el SIGEP, en esas fechas figura otra entidad (MINISTERIO DE HACIENDA)", fila["duda"])

    def test_un_periodo_agregado_a_mano_tambien_debe_estar_en_el_sigep(self):
        cid = self._lista()
        nuevo = {"inicio": "2005-01-01", "fin": "2005-12-31", "entidad": "Municipio de Pueblo Nuevo", "referencia": "", "relacionada": False}
        d = self.abogado.put(f"/api/ops/{cid}/decisiones", {"agregar": nuevo}).json()
        self.assertEqual((d["experiencia"][0]["estado"], d["total_dias"]), ("sin_sigep", 4350))


class LibretaDeducidaTests(BaseOps):
    def test_lo_deducido_no_da_la_libreta_por_no_aplica(self):
        # Sin indicar si es hombre menor de 50: la cédula sugiere que no aplica, pero lo confirma una persona.
        with mock.patch("motor.ops.expediente.libreta_exigible", return_value=(False, "Parece mujer. Confírmelo.")):
            cid = self._crear(exige_libreta="").json()["id"]
            self._analizar()
        d = self.abogado.get(f"/api/ops/{cid}").json()
        libreta = next(x for x in d["documentos"] if x["clave"] == "libreta_militar")
        self.assertEqual((libreta["estado"], libreta["motivo"]), ("revision", "Parece mujer. Confírmelo."))
        self.assertIsNone(d["cumple"])
        # Indicado al crear, sí es definitivo.
        cid = self._crear(exige_libreta="no").json()["id"]
        self._analizar()
        libreta = next(x for x in self.abogado.get(f"/api/ops/{cid}").json()["documentos"] if x["clave"] == "libreta_militar")
        self.assertEqual(libreta["estado"], "no_aplica")
