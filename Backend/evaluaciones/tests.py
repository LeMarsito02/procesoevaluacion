"""Pruebas de procesos, asignaciones, revisiones y aislamiento (app y RLS)."""
from __future__ import annotations

from datetime import timedelta
from unittest import mock
from urllib.parse import quote

from asgiref.sync import async_to_sync

from django.core import mail
from django.db import ProgrammingError, transaction
from django.test import TestCase
from django.utils import timezone

from cuentas.aislamiento import NINGUNA, SISTEMA, fijar_entidad
from cuentas.models import Rol, TipoArea, Usuario
from cuentas.tests import CLAVE, BaseCuentas, Cliente, crear_entidad
from evaluaciones import servicios
from evaluaciones.models import EstadoEvaluacion, EstadoTrabajo, Evaluacion, Proceso, Trabajador, Trabajo
from evaluaciones.trabajador import trabajar
from motor.esquemas.proceso import ResultadoRequisito

DOCUMENTO_BASE = {
    "codigo_proceso": "ENT-CM-037-2026",
    "fecha_cierre": "2026-08-20",
    "objeto_general": "Mantenimiento de vías",
    "lotes": [{"numero": "LOTE 1", "objeto": "Vías", "plazo_meses": 4, "valor_presupuesto": 1000000, "lugar_ejecucion": None}],
    "lote_mayor_valor": "LOTE 1",
    "presupuesto_total": 1000000,
    "garantia_seriedad": {
        "vigencia_meses": 3,
        "porcentaje": 0.1,
        "base_calculo": "lote_mayor_valor",
        "lote_base": "LOTE 1",
        "valor_base": 1000000,
        "valor_asegurado": 100000,
        "fecha_cierre": "2026-08-20",
        "fecha_vencimiento": "2026-11-20",
    },
    "advertencias": [],
}

PROPONENTES = [
    {"numero_orden": 1, "hoja": "P-01", "nombre_proponente": "Consorcio Uno", "nombre_archivo": "P1 Consorcio Uno.zip", "drive_file_id": "a1"},
    {"numero_orden": 2, "hoja": "P-02", "nombre_proponente": "Vías SAS", "nombre_archivo": "P2 Vías SAS.zip", "drive_file_id": "a2"},
]


def resultados_falsos(proponente, proceso):
    """Motor simulado: el requisito 2 queda por revisar, el resto cumple."""
    return [
        ResultadoRequisito(
            hoja=proponente.hoja,
            numero_orden=proponente.numero_orden,
            nombre_proponente=proponente.nombre_proponente,
            requisito=n,
            cumple=n != 2,
            motivo="No se encontró el COPNIA" if n == 2 else None,
        )
        for n in (1, 2, 13)
    ]


async def motor_falso(proponente, proceso):
    return resultados_falsos(proponente, proceso)


class BaseEvaluaciones(BaseCuentas):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.consulta = Usuario.objects.create_user("control@entidad.gov.co", CLAVE, nombre_completo="Control Interno", entidad=cls.entidad1, rol=Rol.CONSULTA)
        cls.abogado2 = Usuario.objects.create_user("abogado2@entidad.gov.co", CLAVE, nombre_completo="Abogado Dos", entidad=cls.entidad1, rol=Rol.EVALUADOR)
        cls.abogado2.areas.set(cls.entidad1.areas.filter(tipo=TipoArea.JURIDICA))
        cls.eval_otra = Usuario.objects.create_user("abogado@otraentidad.gov.co", CLAVE, nombre_completo="Abogado Otra", entidad=cls.otra, rol=Rol.EVALUADOR)
        cls.eval_otra.areas.set(cls.otra.areas.filter(tipo=TipoArea.JURIDICA))

    def crear(self, email="jefe@entidad.gov.co", codigo="ENT-CM-037-2026"):
        c = Cliente()
        c.entrar(email)
        doc = {**DOCUMENTO_BASE, "codigo_proceso": codigo}
        r = c.post("/api/evaluaciones/procesos", {"documento_base": doc, "carpeta_drive": "x", "proponentes": PROPONENTES})
        self.assertEqual(r.status_code, 201, r.content)
        return c, r.json()[0]

    def evaluar_todo(self, c, evaluacion_id):
        """Encola todo y deja que el trabajador real atienda la fila con el motor simulado."""
        detalle = c.get(f"/api/evaluaciones/{evaluacion_id}").json()
        r = c.post(f"/api/evaluaciones/{evaluacion_id}/evaluar", {})
        self.assertEqual(r.status_code, 200, r.content)
        correr_fila()
        return detalle


def correr_fila(motor=motor_falso):
    with mock.patch("evaluaciones.trabajador.evaluar_todos_en_proceso", motor):
        async_to_sync(trabajar)(2, una_vez=True)


class CrearYAsignarTests(BaseEvaluaciones):
    def test_jefe_crea_sin_asignar_y_asigna(self):
        jefe, ev = self.crear()
        self.assertEqual(ev["estado"], EstadoEvaluacion.SIN_ASIGNAR)
        self.assertEqual(ev["avance"]["proponentes"], 2)
        with self.captureOnCommitCallbacks(execute=True):
            r = jefe.post(f"/api/evaluaciones/{ev['id']}/asignar", {"responsable_id": str(self.evaluador.id)})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["estado"], EstadoEvaluacion.ASIGNADA)
        self.assertEqual(mail.outbox[-1].to, ["abogado@entidad.gov.co"])
        # Aparece en "mis evaluaciones" del abogado y no en la del otro abogado.
        a1 = Cliente()
        a1.entrar("abogado@entidad.gov.co")
        self.assertEqual([e["id"] for e in a1.get("/api/evaluaciones/mias").json()], [ev["id"]])
        a2 = Cliente()
        a2.entrar("abogado2@entidad.gov.co")
        self.assertEqual(a2.get("/api/evaluaciones/mias").json(), [])

    def test_evaluador_que_crea_queda_responsable(self):
        _, ev = self.crear("abogado@entidad.gov.co")
        self.assertEqual(ev["estado"], EstadoEvaluacion.ASIGNADA)
        self.assertEqual(ev["responsable"]["email"], "abogado@entidad.gov.co")

    def test_abogado_crea_su_propio_proceso_aunque_no_tenga_el_area(self):
        _, ev = self.crear("tecnico@entidad.gov.co")
        self.assertEqual(ev["estado"], EstadoEvaluacion.ASIGNADA)
        self.assertEqual(ev["responsable"]["email"], "tecnico@entidad.gov.co")
        c = Cliente()
        c.entrar("tecnico@entidad.gov.co")
        self.assertTrue(c.get(f"/api/evaluaciones/{ev['id']}").json()["evaluacion"]["puede_trabajar"])

    def test_abogado_no_asigna_a_otro_al_crear(self):
        c = Cliente()
        c.entrar("abogado@entidad.gov.co")
        r = c.post(
            "/api/evaluaciones/procesos",
            {"documento_base": DOCUMENTO_BASE, "proponentes": PROPONENTES, "responsable_id": str(self.abogado2.id)},
        )
        self.assertEqual(r.status_code, 403)

    def test_jefe_asigna_al_crear(self):
        c = Cliente()
        c.entrar("jefe@entidad.gov.co")
        with self.captureOnCommitCallbacks(execute=True):
            r = c.post(
                "/api/evaluaciones/procesos",
                {"documento_base": DOCUMENTO_BASE, "proponentes": PROPONENTES, "responsable_id": str(self.abogado2.id)},
            )
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()[0]["responsable"]["email"], "abogado2@entidad.gov.co")
        self.assertEqual(mail.outbox[-1].to, ["abogado2@entidad.gov.co"])

    def test_superadmin_crea_asigna_y_ve_equipo_de_cualquier_entidad(self):
        import pyotp

        c = Cliente()
        c.entrar("santiagopebe01@lemartek.com")
        secreto = c.post("/api/auth/2fa/configurar").json()["secreto"]
        c.post("/api/auth/2fa/verificar", {"codigo": pyotp.TOTP(secreto).now()})
        # Debe elegir la entidad.
        r = c.post("/api/evaluaciones/procesos", {"documento_base": DOCUMENTO_BASE, "proponentes": PROPONENTES})
        self.assertEqual(r.status_code, 400)
        r = c.post(
            "/api/evaluaciones/procesos",
            {"documento_base": DOCUMENTO_BASE, "proponentes": PROPONENTES, "entidad_id": str(self.otra.id), "responsable_id": str(self.eval_otra.id)},
        )
        self.assertEqual(r.status_code, 201, r.content)
        ev = r.json()[0]
        self.assertEqual(ev["entidad_nombre"], "Otra Entidad")
        self.assertEqual(ev["estado"], EstadoEvaluacion.ASIGNADA)
        # No puede asignar a alguien de otra entidad.
        self.assertEqual(c.post(f"/api/evaluaciones/{ev['id']}/asignar", {"responsable_id": str(self.evaluador.id)}).status_code, 400)
        # Reasigna, ve el equipo de esa entidad, evalúa y aprueba.
        self.assertEqual(c.post(f"/api/evaluaciones/{ev['id']}/asignar", {"responsable_id": None}).status_code, 200)
        equipo = {m["email"] for m in c.get(f"/api/evaluaciones/equipo?entidad_id={self.otra.id}").json()}
        self.assertIn("abogado@otraentidad.gov.co", equipo)
        self.assertNotIn("abogado@entidad.gov.co", equipo)
        detalle = self.evaluar_todo(c, ev["id"])
        for p in detalle["proponentes"]:
            c.http.put(
                f"/api/evaluaciones/{ev['id']}/revisiones",
                {"proponente_id": p["id"], "requisito": 2, "cumple": True, "nota": "Revisado por superadmin"},
                content_type="application/json",
                headers={"X-CSRFToken": c.csrf},
            )
        self.assertEqual(c.post(f"/api/evaluaciones/{ev['id']}/aprobar").json()["estado"], EstadoEvaluacion.APROBADA)

    def test_no_asigna_a_otra_area_ni_otra_entidad(self):
        jefe, ev = self.crear()
        self.assertEqual(jefe.post(f"/api/evaluaciones/{ev['id']}/asignar", {"responsable_id": str(self.tecnico.id)}).status_code, 400)
        self.assertEqual(jefe.post(f"/api/evaluaciones/{ev['id']}/asignar", {"responsable_id": str(self.eval_otra.id)}).status_code, 400)
        self.assertEqual(jefe.post(f"/api/evaluaciones/{ev['id']}/asignar", {"responsable_id": str(self.consulta.id)}).status_code, 400)

    def test_evaluador_y_consulta_no_asignan(self):
        _, ev = self.crear()
        for email in ("abogado@entidad.gov.co", "control@entidad.gov.co"):
            c = Cliente()
            c.entrar(email)
            r = c.post(f"/api/evaluaciones/{ev['id']}/asignar", {"responsable_id": str(self.evaluador.id)})
            self.assertEqual(r.status_code, 403, email)

    def test_consulta_no_crea_procesos(self):
        c = Cliente()
        c.entrar("control@entidad.gov.co")
        r = c.post("/api/evaluaciones/procesos", {"documento_base": DOCUMENTO_BASE, "proponentes": PROPONENTES})
        self.assertEqual(r.status_code, 403)

    def test_codigo_repetido_en_la_misma_entidad(self):
        self.crear()
        c = Cliente()
        c.entrar("jefe@entidad.gov.co")
        r = c.post("/api/evaluaciones/procesos", {"documento_base": DOCUMENTO_BASE, "proponentes": PROPONENTES})
        self.assertEqual(r.status_code, 409)

    def test_mismo_codigo_en_otra_entidad_si_se_permite(self):
        self.crear()
        self.crear("abogado@otraentidad.gov.co")

    def test_las_tres_areas_se_crean_y_asignan(self):
        c = Cliente()
        c.entrar("admin@entidad.gov.co")
        with self.captureOnCommitCallbacks(execute=True):
            r = c.post(
                "/api/evaluaciones/procesos",
                {
                    "documento_base": DOCUMENTO_BASE,
                    "proponentes": PROPONENTES,
                    "tipos": ["juridica", "tecnica", "financiera"],
                    "responsables": {"juridica": str(self.evaluador.id), "tecnica": str(self.tecnico.id), "financiera": None},
                },
            )
        self.assertEqual(r.status_code, 201, r.content)
        por_tipo = {e["tipo"]: e for e in r.json()}
        self.assertEqual(por_tipo["juridica"]["responsable"]["email"], "abogado@entidad.gov.co")
        self.assertEqual(por_tipo["tecnica"]["responsable"]["email"], "tecnico@entidad.gov.co")
        self.assertIsNone(por_tipo["financiera"]["responsable"])
        self.assertTrue(por_tipo["juridica"]["tipo_disponible"])
        self.assertTrue(por_tipo["tecnica"]["tipo_disponible"])
        self.assertTrue(por_tipo["financiera"]["tipo_disponible"])
        self.assertEqual({m.to[0] for m in mail.outbox}, {"abogado@entidad.gov.co", "tecnico@entidad.gov.co"})
        tipos = {t["clave"]: t["disponible"] for t in c.get("/api/evaluaciones/tipos").json()}
        self.assertEqual(tipos, {"juridica": True, "tecnica": True, "financiera": True})

    def test_tipo_invalido(self):
        c = Cliente()
        c.entrar("jefe@entidad.gov.co")
        r = c.post("/api/evaluaciones/procesos", {"documento_base": DOCUMENTO_BASE, "proponentes": PROPONENTES, "tipos": ["ambiental"]})
        self.assertEqual(r.status_code, 400)


class TrabajoTests(BaseEvaluaciones):
    def setUp(self):
        self.jefe, self.ev = self.crear()
        self.jefe.post(f"/api/evaluaciones/{self.ev['id']}/asignar", {"responsable_id": str(self.evaluador.id)})
        self.abogado = Cliente()
        self.abogado.entrar("abogado@entidad.gov.co")

    def test_flujo_completo(self):
        eid = self.ev["id"]
        detalle = self.evaluar_todo(self.abogado, eid)
        resumen = self.abogado.get(f"/api/evaluaciones/{eid}").json()["evaluacion"]
        self.assertEqual(resumen["estado"], EstadoEvaluacion.EN_REVISION)
        # Requisito 2 por revisar en ambos; el 13 (RUT) no cuenta; ninguno con error.
        self.assertEqual(resumen["avance"]["pendientes"], 2)
        self.assertEqual(resumen["avance"]["con_error"], 0)

        # Con pendientes no se aprueba.
        self.assertEqual(self.jefe.post(f"/api/evaluaciones/{eid}/aprobar").status_code, 409)

        for p in detalle["proponentes"]:
            r = self.abogado.http.put(
                f"/api/evaluaciones/{eid}/revisiones",
                {"proponente_id": p["id"], "requisito": 2, "cumple": p["hoja"] == "P-01", "nota": "Revisado en físico"},
                content_type="application/json",
                headers={"X-CSRFToken": self.abogado.csrf},
            )
            self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(self.abogado.get(f"/api/evaluaciones/{eid}").json()["evaluacion"]["avance"]["pendientes"], 0)

        # El evaluador no aprueba; el jefe sí.
        self.assertEqual(self.abogado.post(f"/api/evaluaciones/{eid}/aprobar").status_code, 403)
        self.assertEqual(self.jefe.post(f"/api/evaluaciones/{eid}/aprobar").json()["estado"], EstadoEvaluacion.APROBADA)

        # Aprobada: no se modifica hasta reabrir.
        r = self.abogado.http.put(
            f"/api/evaluaciones/{eid}/revisiones",
            {"proponente_id": detalle["proponentes"][0]["id"], "requisito": 2, "cumple": False},
            content_type="application/json",
            headers={"X-CSRFToken": self.abogado.csrf},
        )
        self.assertEqual(r.status_code, 409)
        self.assertEqual(self.jefe.post(f"/api/evaluaciones/{eid}/reabrir").json()["estado"], EstadoEvaluacion.EN_REVISION)

        # Informe con las decisiones aplicadas.
        r = self.abogado.get(f"/api/evaluaciones/{eid}/informe")
        self.assertEqual(r.status_code, 200)
        self.assertIn("BORRADOR", r["Content-Disposition"])
        self.assertTrue(r.content.startswith(b"PK"))

    def test_otro_evaluador_y_consulta_no_trabajan(self):
        eid = self.ev["id"]
        prop = self.abogado.get(f"/api/evaluaciones/{eid}").json()["proponentes"][0]["id"]
        for email in ("abogado2@entidad.gov.co", "control@entidad.gov.co"):
            c = Cliente()
            c.entrar(email)
            # Pueden consultar…
            self.assertEqual(c.get(f"/api/evaluaciones/{eid}").status_code, 200)
            # …pero no evaluar.
            self.assertEqual(c.post(f"/api/evaluaciones/{eid}/evaluar", {"proponente_ids": [prop]}).status_code, 403, email)

    def test_proponente_de_otro_proceso_rechazado(self):
        _, otro = self.crear(codigo="OTRO-001")
        prop_otro = self.jefe.get(f"/api/evaluaciones/{otro['id']}").json()["proponentes"][0]["id"]
        r = self.abogado.post(f"/api/evaluaciones/{self.ev['id']}/evaluar", {"proponente_ids": [prop_otro]})
        self.assertEqual(r.status_code, 404)

    def test_endpoints_de_medicion_cerrados_fuera_de_debug(self):
        r = self.abogado.post("/api/procesos/evaluar-todos/proponente", {"documento_base": DOCUMENTO_BASE, "proponente": PROPONENTES[0]})
        self.assertEqual(r.status_code, 404)


class AislamientoEvaluacionesTests(BaseEvaluaciones):
    def test_otra_entidad_no_ve_ni_modifica(self):
        jefe, ev = self.crear()
        eid = ev["id"]
        prop = jefe.get(f"/api/evaluaciones/{eid}").json()["proponentes"][0]["id"]
        c = Cliente()
        c.entrar("admin@otraentidad.gov.co")
        self.assertEqual(c.get("/api/evaluaciones/procesos").json(), [])
        self.assertEqual(c.get("/api/evaluaciones/mias").json(), [])
        self.assertEqual(c.get(f"/api/evaluaciones/{eid}").status_code, 404)
        self.assertEqual(c.get(f"/api/evaluaciones/{eid}/informe").status_code, 404)
        self.assertEqual(c.get(f"/api/evaluaciones/{eid}/proponentes/{prop}/documento?archivo=x.pdf").status_code, 404)
        self.assertEqual(c.post(f"/api/evaluaciones/{eid}/asignar", {"responsable_id": None}).status_code, 404)
        self.assertEqual(c.post(f"/api/evaluaciones/{eid}/aprobar").status_code, 404)

    def test_rls_filtra_aunque_la_consulta_no_filtre(self):
        self.crear()
        self.crear("abogado@otraentidad.gov.co", codigo="OTRA-001")
        try:
            fijar_entidad(str(self.otra.id))
            self.assertEqual(list(Proceso.objects.values_list("codigo", flat=True)), ["OTRA-001"])
            self.assertEqual(Evaluacion.objects.count(), 1)
            fijar_entidad(NINGUNA)
            self.assertEqual(Proceso.objects.count(), 0)
            # Tampoco se puede escribir una fila de otra entidad.
            fijar_entidad(str(self.otra.id))
            with self.assertRaises(ProgrammingError), transaction.atomic():
                Proceso.objects.create(
                    entidad=self.entidad1, codigo="X", fecha_cierre="2026-01-01", documento_base={}, creado_por=self.admin_otra
                )
        finally:
            fijar_entidad(SISTEMA)
        self.assertEqual(Proceso.objects.count(), 2)

    def test_superadmin_ve_todas(self):
        import pyotp

        self.crear()
        self.crear("abogado@otraentidad.gov.co", codigo="OTRA-001")
        c = Cliente()
        c.entrar("santiagopebe01@lemartek.com")
        secreto = c.post("/api/auth/2fa/configurar").json()["secreto"]
        c.post("/api/auth/2fa/verificar", {"codigo": pyotp.TOTP(secreto).now()})
        self.assertEqual(len(c.get("/api/evaluaciones/procesos").json()), 2)


class EntidadNueva(BaseEvaluaciones):
    def test_entidad_sin_procesos(self):
        crear_entidad("Vacía", "123")
        c = Cliente()
        c.entrar("admin@entidad.gov.co")
        self.assertEqual(c.get("/api/evaluaciones/procesos").json(), [])


class FilaTests(BaseEvaluaciones):
    def setUp(self):
        self.jefe, self.ev = self.crear()
        _, self.ev2 = self.crear(codigo="SEGUNDO-001")
        self.eid, self.eid2 = self.ev["id"], self.ev2["id"]

    def test_encolar_no_duplica_y_muestra_fila(self):
        r = self.jefe.post(f"/api/evaluaciones/{self.eid}/evaluar", {})
        self.assertEqual(r.json()["estado"], EstadoEvaluacion.EVALUANDO)
        self.assertEqual(r.json()["fila"]["en_fila"], 2)
        # Sin trabajadores activos no hay tiempo estimado.
        self.assertIsNone(r.json()["fila"]["eta_segundos"])
        self.jefe.post(f"/api/evaluaciones/{self.eid}/evaluar", {})
        self.assertEqual(Trabajo.objects.filter(evaluacion_id=self.eid).count(), 2)
        # Con un trabajador latiendo sí hay estimación.
        Trabajador.objects.create(id="prueba", capacidad=2, latido=timezone.now())
        fila = self.jefe.get(f"/api/evaluaciones/{self.eid}/novedades").json()["evaluacion"]["fila"]
        self.assertEqual(fila["capacidad"], 2)
        self.assertGreater(fila["eta_segundos"], 0)

    def test_reparto_por_turnos_entre_evaluaciones(self):
        self.jefe.post(f"/api/evaluaciones/{self.eid}/evaluar", {})
        self.jefe.post(f"/api/evaluaciones/{self.eid2}/evaluar", {})
        orden = []
        while (t := servicios.reclamar("prueba")) is not None:
            orden.append(str(t.evaluacion_id))
        # Primero un proponente de cada evaluación, luego el segundo de cada una.
        self.assertEqual(orden[0], self.eid)
        self.assertEqual(orden[1], self.eid2)
        self.assertEqual(sorted(orden[:2]), sorted(orden[2:]))

    def test_trabajador_evalua_guarda_y_avisa(self):
        self.jefe.post(f"/api/evaluaciones/{self.eid}/asignar", {"responsable_id": str(self.evaluador.id)})
        self.jefe.post(f"/api/evaluaciones/{self.eid}/evaluar", {})
        with self.captureOnCommitCallbacks(execute=True):
            correr_fila()
        resumen = self.jefe.get(f"/api/evaluaciones/{self.eid}").json()["evaluacion"]
        self.assertEqual(resumen["estado"], EstadoEvaluacion.EN_REVISION)
        self.assertEqual(resumen["avance"]["evaluados"], 2)
        self.assertIsNone(resumen["fila"])
        self.assertEqual(mail.outbox[-1].to, ["abogado@entidad.gov.co"])
        self.assertIn("Evaluación terminada", mail.outbox[-1].subject)
        self.assertFalse(Trabajador.objects.exists())
        # Novedades desde antes de evaluar trae los resultados guardados.
        r = self.jefe.get(f"/api/evaluaciones/{self.eid}/novedades?desde=2000-01-01T00:00:00Z").json()
        self.assertEqual(len(r["resultados"]), 6)

    def test_error_del_motor_no_detiene_la_fila(self):
        async def motor_roto(proponente, proceso):
            if proponente.hoja == "P-01":
                raise RuntimeError("zip dañado")
            return resultados_falsos(proponente, proceso)

        self.jefe.post(f"/api/evaluaciones/{self.eid}/evaluar", {})
        correr_fila(motor_roto)
        estados = dict(Trabajo.objects.filter(evaluacion_id=self.eid).values_list("proponente__hoja", "estado"))
        self.assertEqual(estados, {"P-01": EstadoTrabajo.ERROR, "P-02": EstadoTrabajo.TERMINADO})

    def test_pausar_saca_lo_que_no_empezo(self):
        self.jefe.post(f"/api/evaluaciones/{self.eid}/evaluar", {})
        servicios.reclamar("prueba")  # uno queda procesando
        r = self.jefe.post(f"/api/evaluaciones/{self.eid}/pausar").json()
        self.assertEqual(r["fila"]["en_fila"], 0)
        self.assertEqual(r["fila"]["procesando"], 1)

    def test_huerfanos_vuelven_a_la_fila_y_luego_fallan(self):
        self.jefe.post(f"/api/evaluaciones/{self.eid}/evaluar", {})
        viejo = timezone.now() - timedelta(minutes=10)
        for intento in range(3):
            t = servicios.reclamar("caido")
            Trabajo.objects.filter(pk=t.pk).update(latido=viejo)
            servicios.recuperar_huerfanos()
        t.refresh_from_db()
        self.assertEqual(t.estado, EstadoTrabajo.ERROR)

    def test_consulta_no_encola_ni_pausa(self):
        c = Cliente()
        c.entrar("control@entidad.gov.co")
        self.assertEqual(c.post(f"/api/evaluaciones/{self.eid}/evaluar", {}).status_code, 403)
        self.assertEqual(c.post(f"/api/evaluaciones/{self.eid}/pausar").status_code, 403)

    def test_otra_entidad_no_ve_novedades(self):
        c = Cliente()
        c.entrar("admin@otraentidad.gov.co")
        self.assertEqual(c.get(f"/api/evaluaciones/{self.eid}/novedades").status_code, 404)

    def test_estado_de_la_fila_por_rol(self):
        self.jefe.post(f"/api/evaluaciones/{self.eid}/evaluar", {})
        c = Cliente()
        c.entrar("admin@entidad.gov.co")
        r = c.get("/api/evaluaciones/fila/estado").json()
        self.assertEqual(r["total_en_fila"], 2)
        self.assertEqual([e["id"] for e in r["evaluaciones"]], [self.eid])
        self.assertEqual(r["trabajadores"], [])
        otra = Cliente()
        otra.entrar("admin@otraentidad.gov.co")
        self.assertEqual(otra.get("/api/evaluaciones/fila/estado").json()["evaluaciones"], [])
        self.assertEqual(self.jefe.get("/api/evaluaciones/fila/estado").status_code, 403)


PLANTILLA_EJEMPLO = "motor/plantillas/plantilla_evaluacion_juridica.xlsx"


def subir(c: "Cliente", tipo="juridica", ruta=PLANTILLA_EJEMPLO, nombre="plantilla.xlsx", entidad_id=None):
    from django.core.files.uploadedfile import SimpleUploadedFile

    with open(ruta, "rb") as f:
        archivo = SimpleUploadedFile(nombre, f.read())
    url = "/api/configuracion/plantillas" + (f"?entidad_id={entidad_id}" if entidad_id else "")
    return c.http.post(url, {"tipo": tipo, "archivo": archivo}, headers={"X-CSRFToken": c.csrf})


class PlantillasTests(BaseEvaluaciones):
    def setUp(self):
        import tempfile

        from django.test import override_settings

        self._tmp = tempfile.TemporaryDirectory()
        self._override = override_settings(MEDIA_ROOT=self._tmp.name)
        self._override.enable()
        self.admin = Cliente()
        self.admin.entrar("admin@entidad.gov.co")

    def tearDown(self):
        self._override.disable()
        self._tmp.cleanup()

    def test_subir_ajustar_mapeo_y_usar_en_el_informe(self):
        r = subir(self.admin, nombre="Plantilla de la entidad 2026.xlsx")
        self.assertEqual(r.status_code, 201, r.content)
        plantilla = r.json()
        self.assertTrue(plantilla["inspeccion"]["valida"])
        self.assertGreater(plantilla["inspeccion"]["hojas_proponente"], 0)
        self.assertIn("1", plantilla["inspeccion"]["filas"])

        mapeo = {**plantilla["mapeo"], "prefijo_codigo": "ENT3"}
        r = self.admin.http.put(
            f"/api/configuracion/plantillas/{plantilla['id']}/mapeo",
            {"mapeo": mapeo},
            content_type="application/json",
            headers={"X-CSRFToken": self.admin.csrf},
        )
        self.assertEqual(r.status_code, 200, r.content)

        listado = {t["tipo"]: t for t in self.admin.get("/api/configuracion/plantillas").json()}
        self.assertEqual(listado["juridica"]["activa"]["id"], plantilla["id"])
        self.assertFalse(listado["juridica"]["usa_plantilla_del_sistema"])
        self.assertTrue(listado["financiera"]["activa"] is None and listado["financiera"]["motor_disponible"])

        # El informe usa la plantilla y el mapeo de la entidad.
        import io

        import openpyxl

        jefe, ev = self.crear()
        self.evaluar_todo(jefe, ev["id"])
        r = jefe.get(f"/api/evaluaciones/{ev['id']}/informe")
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertIn("ENT3-ENT-CM-037", wb["P-01"]["F3"].value)

    def test_nueva_plantilla_reemplaza_la_activa_y_se_puede_volver(self):
        primera = subir(self.admin).json()
        segunda = subir(self.admin).json()
        self.assertEqual(len(self.admin.get("/api/configuracion/plantillas").json()[0]["historial"]), 2)
        r = self.admin.post(f"/api/configuracion/plantillas/{primera['id']}/activar")
        self.assertTrue(r.json()["activa"])
        activas = [p for p in self.admin.get("/api/configuracion/plantillas").json()[0]["historial"] if p["activa"]]
        self.assertEqual([p["id"] for p in activas], [primera["id"]])
        self.assertNotEqual(primera["id"], segunda["id"])

    def test_tecnica_y_financiera_aceptan_plantilla(self):
        self.assertEqual(subir(self.admin, tipo="tecnica").status_code, 201)
        self.assertEqual(subir(self.admin, tipo="financiera").status_code, 201)

    def test_rechaza_lo_que_no_es_xlsx(self):
        import tempfile

        with tempfile.NamedTemporaryFile(suffix=".xlsx") as f:
            f.write(b"esto no es un excel")
            f.flush()
            self.assertEqual(subir(self.admin, ruta=f.name).status_code, 400)
        self.assertEqual(subir(self.admin, nombre="macros.xlsm").status_code, 400)

    def test_permisos_y_aislamiento(self):
        plantilla = subir(self.admin).json()
        otra = Cliente()
        otra.entrar("admin@otraentidad.gov.co")
        self.assertEqual(otra.get(f"/api/configuracion/plantillas/{plantilla['id']}/archivo").status_code, 404)
        self.assertEqual(otra.delete(f"/api/configuracion/plantillas/{plantilla['id']}").status_code, 404)
        self.assertEqual(otra.get("/api/configuracion/plantillas").json()[0]["activa"], None)
        jefe = Cliente()
        jefe.entrar("jefe@entidad.gov.co")
        self.assertEqual(jefe.get("/api/configuracion/plantillas").status_code, 403)
        self.assertEqual(subir(jefe).status_code, 403)
        self.assertEqual(self.admin.get(f"/api/configuracion/plantillas/{plantilla['id']}/archivo").status_code, 200)

    def test_superadmin_sube_para_una_entidad(self):
        import pyotp

        c = Cliente()
        c.entrar("santiagopebe01@lemartek.com")
        secreto = c.post("/api/auth/2fa/configurar").json()["secreto"]
        c.post("/api/auth/2fa/verificar", {"codigo": pyotp.TOTP(secreto).now()})
        self.assertEqual(subir(c).status_code, 400)  # debe indicar la entidad
        self.assertEqual(subir(c, entidad_id=self.otra.id).status_code, 201)
        otra = Cliente()
        otra.entrar("admin@otraentidad.gov.co")
        self.assertIsNotNone(otra.get("/api/configuracion/plantillas").json()[0]["activa"])


def definicion_propia():
    """Base del sistema sin sanciones del RUP, con numeración propia y un requisito nuevo por bloques."""
    from motor.criterios import definicion_sistema

    base = definicion_sistema("juridica").model_dump(mode="json")
    requisitos = [r for r in base["requisitos"] if r["verificacion"] != "juridica.sanciones_rup"]
    for i, r in enumerate(requisitos, start=1):
        r["numero"] = i
    requisitos.append(
        {
            "numero": 30,
            "titulo": "Certificado de la Junta Central de Contadores",
            "corto": "JCC",
            "grupo": "antecedentes",
            "verificacion": "personalizado",
            "config": {"frases_documento": ["JUNTA CENTRAL DE CONTADORES"], "bloques": [{"tipo": "vigencia_maxima", "meses": 3}]},
        }
    )
    return {"parametros": {"copnia_meses": 6}, "requisitos": requisitos}


class PlantillaEvaluacionTests(BaseEvaluaciones):
    def setUp(self):
        self.admin = Cliente()
        self.admin.entrar("admin@entidad.gov.co")

    def publicar(self, c=None, definicion=None, nombre="Jurídica Entidad 2026", entidad_id=None):
        c = c or self.admin
        url = "/api/configuracion/evaluaciones" + (f"?entidad_id={entidad_id}" if entidad_id else "")
        return c.post(url, {"tipo": "juridica", "nombre": nombre, "definicion": definicion or definicion_propia()})

    def test_catalogo_del_motor(self):
        cat = self.admin.get("/api/configuracion/catalogo?tipo=juridica").json()
        # 17 de la evaluación base + 2 que se agregan cuando el pliego las pide.
        self.assertEqual(len(cat["verificaciones"]), 19)
        self.assertIn("copnia_meses", {p["clave"] for p in cat["parametros"]})

    def test_sin_version_propia_usa_la_base_del_sistema(self):
        juridica = self.admin.get("/api/configuracion/evaluaciones").json()[0]
        self.assertIsNone(juridica["activa"])
        self.assertEqual(len(juridica["definicion"]["requisitos"]), 17)

    def test_publicar_y_evaluar_con_la_definicion_de_la_entidad(self):
        r = self.publicar()
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()["version"], 1)

        jefe, ev = self.crear()
        self.assertEqual(ev["plantilla_version"], 1)
        detalle = jefe.get(f"/api/evaluaciones/{ev['id']}").json()
        self.assertEqual(len(detalle["catalogo"]), 17)
        self.assertEqual(detalle["catalogo"][-1]["titulo"], "Certificado de la Junta Central de Contadores")
        self.assertTrue(detalle["catalogo"][-1]["personalizado"])

        recibido = {}

        async def motor_que_mira(proponente, proceso):
            recibido["criterios"] = proceso.criterios
            return resultados_falsos(proponente, proceso)

        jefe.post(f"/api/evaluaciones/{ev['id']}/evaluar", {})
        correr_fila(motor_que_mira)
        # Además del parámetro de la entidad, el salario mínimo del año del cierre.
        self.assertEqual(recibido["criterios"]["parametros"], {"copnia_meses": 6, "smmlv": 1_750_905})
        self.assertEqual(len(recibido["criterios"]["requisitos"]), 17)

    def test_nueva_version_y_actualizar_evaluacion(self):
        self.publicar()
        jefe, ev = self.crear()
        self.evaluar_todo(jefe, ev["id"])
        from motor.criterios import definicion_sistema

        v2 = definicion_sistema("juridica").model_dump(mode="json")
        v2["requisitos"] = [r for r in v2["requisitos"] if r["numero"] in (1, 2)]
        self.assertEqual(self.publicar(definicion=v2, nombre="Solo carta y aval").json()["version"], 2)

        resumen = jefe.get(f"/api/evaluaciones/{ev['id']}").json()["evaluacion"]
        self.assertTrue(resumen["plantilla_desactualizada"])
        r = jefe.post(f"/api/evaluaciones/{ev['id']}/actualizar-plantilla")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["plantilla_version"], 2)
        self.assertFalse(r.json()["plantilla_desactualizada"])
        # Se descartan resultados de requisitos que ya no existen (13) y se vuelve a evaluar a todos.
        self.assertEqual(Trabajo.objects.filter(evaluacion_id=ev["id"], estado=EstadoTrabajo.EN_FILA).count(), 2)
        self.assertEqual(jefe.post(f"/api/evaluaciones/{ev['id']}/actualizar-plantilla").status_code, 409)
        versiones = self.admin.get("/api/configuracion/evaluaciones").json()[0]["versiones"]
        self.assertEqual([(v["version"], v["activa"], v["evaluaciones"]) for v in versiones], [(2, True, 1), (1, False, 0)])

    def test_validacion_de_la_definicion(self):
        d = definicion_propia()
        d["requisitos"][1]["numero"] = d["requisitos"][0]["numero"]
        r = self.publicar(definicion=d)
        self.assertEqual(r.status_code, 400)
        self.assertIn("repetidos", r.json()["detail"])
        d = definicion_propia()
        d["parametros"]["copnia_meses"] = 100
        self.assertEqual(self.publicar(definicion=d).status_code, 400)
        d = definicion_propia()
        d["requisitos"][0]["verificacion"] = "juridica.inventada"
        self.assertEqual(self.publicar(definicion=d).status_code, 400)
        self.assertEqual(self.publicar(definicion={"requisitos": []}).status_code, 400)

    def test_permisos_y_aislamiento(self):
        v = self.publicar().json()
        otra = Cliente()
        otra.entrar("admin@otraentidad.gov.co")
        self.assertIsNone(otra.get("/api/configuracion/evaluaciones").json()[0]["activa"])
        self.assertEqual(otra.post(f"/api/configuracion/evaluaciones/{v['id']}/activar").status_code, 404)
        jefe = Cliente()
        jefe.entrar("jefe@entidad.gov.co")
        self.assertEqual(self.publicar(c=jefe).status_code, 403)

    def test_superadmin_crea_entidad_con_base_o_copia(self):
        import pyotp

        self.publicar()
        c = Cliente()
        c.entrar("santiagopebe01@lemartek.com")
        secreto = c.post("/api/auth/2fa/configurar").json()["secreto"]
        c.post("/api/auth/2fa/verificar", {"codigo": pyotp.TOTP(secreto).now()})
        r = c.post(
            "/api/plataforma/entidades",
            {"nombre": "Tercera Entidad", "nit": "899000003", "email_admin": "admin@tercera.gov.co", "nombre_admin": "Admin Tercera", "sigla": "ent3"},
        )
        self.assertEqual(r.status_code, 201, r.content)
        tercera = r.json()["entidad"]["id"]
        juridica = c.get(f"/api/configuracion/evaluaciones?entidad_id={tercera}").json()[0]
        self.assertEqual(juridica["activa"]["version"], 1)
        self.assertEqual(juridica["definicion"]["parametros"]["prefijo_codigo"], "ENT3")
        self.assertIn("ENT3", juridica["definicion"]["parametros"]["beneficiario_claves"])

        r = c.post(
            "/api/plataforma/entidades",
            {
                "nombre": "Entidad Copiada", "nit": "123", "email_admin": "a@copia.gov.co", "nombre_admin": "Admin Copia",
                "base": "copiar", "copiar_de": str(self.entidad1.id),
            },
        )
        self.assertEqual(r.status_code, 201, r.content)
        copia = c.get(f"/api/configuracion/evaluaciones?entidad_id={r.json()['entidad']['id']}").json()[0]
        self.assertEqual(copia["activa"]["nombre"], "Jurídica Entidad 2026")
        self.assertEqual(len(copia["definicion"]["requisitos"]), 17)

    def test_probar_requisito_contra_ofertas(self):
        from concurrent.futures import ThreadPoolExecutor

        _, ev = self.crear()
        props = self.admin.get(f"/api/evaluaciones/{ev['id']}").json()["proponentes"]
        requisito = definicion_propia()["requisitos"][-1]

        def falso(proponente, proceso, req, parametros):
            from motor.esquemas.proceso import ResultadoRequisito

            return ResultadoRequisito(hoja=proponente.hoja, numero_orden=proponente.numero_orden, nombre_proponente=proponente.nombre_proponente, requisito=req["numero"], cumple=proponente.hoja == "P-01", motivo=None if proponente.hoja == "P-01" else "No se encontró el documento")

        with ThreadPoolExecutor(2) as pool, mock.patch("api.configuracion.obtener_pool", return_value=pool), mock.patch("api.configuracion.probar_requisito", falso):
            r = self.admin.post("/api/configuracion/requisitos/probar", {"evaluacion_id": ev["id"], "requisito": requisito, "proponente_ids": [p["id"] for p in props]})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(sorted((x["hoja"], x["cumple"]) for x in r.json()), [("P-01", True), ("P-02", False)])
        otra = Cliente()
        otra.entrar("admin@otraentidad.gov.co")
        r = otra.post("/api/configuracion/requisitos/probar", {"evaluacion_id": ev["id"], "requisito": requisito, "proponente_ids": [props[0]["id"]]})
        self.assertEqual(r.status_code, 404)

    def test_proponer_con_ia(self):
        propuesta = {
            "titulo": "Certificado de la Junta Central de Contadores del contador",
            "corto": "JCC",
            "grupo": "antecedentes",
            "verifica": "Sin antecedentes y máximo 3 meses",
            "config": {"frases_documento": ["JUNTA CENTRAL DE CONTADORES"], "bloques": [{"tipo": "vigencia_maxima", "meses": 3}, {"tipo": "contiene", "frases": ["NO REGISTRA ANTECEDENTES"]}]},
        }
        with mock.patch("api.configuracion.consultar_json", return_value=propuesta):
            r = self.admin.post("/api/configuracion/requisitos/proponer", {"descripcion": "Certificado de la Junta Central de Contadores del contador con vigencia de 3 meses que diga no registra antecedentes"})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["verificacion"], "personalizado")
        with mock.patch("api.configuracion.consultar_json", return_value=None):
            r = self.admin.post("/api/configuracion/requisitos/proponer", {"descripcion": "Certificado de la Junta Central de Contadores del contador"})
        self.assertEqual(r.status_code, 503)


class RetencionTests(BaseEvaluaciones):
    def test_borra_ofertas_de_procesos_aprobados_hace_mas_de_30_dias(self):
        import os
        import tempfile
        import time as tiempo
        from pathlib import Path

        from evaluaciones import retencion

        jefe, ev = self.crear()
        _, otro = self.crear(codigo="VIGENTE-001")
        with tempfile.TemporaryDirectory() as tmp:
            cache = Path(tmp)
            ofertas = cache / "drive_files"
            ocr = cache / "ocr"
            ofertas.mkdir()
            ocr.mkdir()
            for fid in ("a1", "a2"):
                (ofertas / f"{fid}.zip").write_bytes(b"zip")
                (ofertas / f"{fid}.meta.json").write_text("{}")
            viejo = ocr / "viejo.txt"
            viejo.write_text("texto")
            hace_40_dias = tiempo.time() - 40 * 86400
            os.utime(viejo, (hace_40_dias, hace_40_dias))
            (ocr / "reciente.txt").write_text("texto")

            with mock.patch.object(retencion, "CACHE", cache), mock.patch.object(retencion, "DIRECTORIO_OFERTAS", ofertas):
                # Sin aprobar: no se borra nada de ofertas.
                self.assertEqual(retencion.aplicar(30).procesos, [])
                Evaluacion.objects.filter(pk=ev["id"]).update(
                    estado=EstadoEvaluacion.APROBADA, aprobada_en=timezone.now() - timedelta(days=31)
                )
                # El otro proceso vigente usa las mismas ofertas (a1, a2): se conservan.
                self.assertEqual(retencion.aplicar(30).archivos_ofertas, 0)
                Evaluacion.objects.filter(pk=otro["id"]).update(
                    estado=EstadoEvaluacion.APROBADA, aprobada_en=timezone.now() - timedelta(days=35)
                )
                # Sin expediente permanente generado no se borra nada.
                self.assertEqual(retencion.aplicar(30).procesos, [])
                from evaluaciones.models import EstadoExpediente, Expediente

                for eid in (ev["id"], otro["id"]):
                    e = Evaluacion.objects.get(pk=eid)
                    Expediente.objects.create(entidad_id=e.entidad_id, evaluacion=e, version=1, estado=EstadoExpediente.LISTO)
                simulacro = retencion.aplicar(30, simulacro=True)
                self.assertTrue((ofertas / "a1.zip").exists())
                self.assertEqual(sorted(simulacro.procesos), ["ENT-CM-037-2026", "VIGENTE-001"])
                informe = retencion.aplicar(30)

            self.assertEqual(sorted(informe.procesos), ["ENT-CM-037-2026", "VIGENTE-001"])
            self.assertFalse((ofertas / "a1.zip").exists())
            self.assertFalse(viejo.exists())
            self.assertTrue((ocr / "reciente.txt").exists())

        # Ya no se pueden abrir los documentos, pero sí el informe y los resultados.
        prop = jefe.get(f"/api/evaluaciones/{otro['id']}").json()["proponentes"][0]["id"]
        r = jefe.get(f"/api/evaluaciones/{otro['id']}/proponentes/{prop}/documento?archivo=x.pdf")
        self.assertEqual(r.status_code, 410)
        self.assertEqual(jefe.get(f"/api/evaluaciones/{otro['id']}").status_code, 200)



def pdf_minimo(texto: str = "certificado") -> bytes:
    return b"%PDF-1.4\n1 0 obj<<>>endobj\n% " + texto.encode() + b"\n%%EOF"


class BaseHistorico(BaseEvaluaciones):
    def setUp(self):
        import tempfile

        from django.test import override_settings

        self._tmp = tempfile.TemporaryDirectory()
        self._override = override_settings(MEDIA_ROOT=self._tmp.name)
        self._override.enable()
        self.jefe, self.ev = self.crear()
        self.jefe.post(f"/api/evaluaciones/{self.ev['id']}/asignar", {"responsable_id": str(self.evaluador.id)})
        self.abogado = Cliente()
        self.abogado.entrar("abogado@entidad.gov.co")
        self.detalle = self.evaluar_todo(self.abogado, self.ev["id"])
        self.p1 = self.detalle["proponentes"][0]["id"]

    def tearDown(self):
        self._override.disable()
        self._tmp.cleanup()

    def revisar(self, c, prop, req, cumple, nota):
        return c.http.put(
            f"/api/evaluaciones/{self.ev['id']}/revisiones",
            {"proponente_id": prop, "requisito": req, "cumple": cumple, "nota": nota},
            content_type="application/json",
            headers={"X-CSRFToken": c.csrf},
        )

    def aportar(self, c, prop, requisito, persona_id=None, contenido=None, nombre="redam.pdf"):
        from django.core.files.uploadedfile import SimpleUploadedFile

        datos = {"requisito": requisito, "fecha_expedicion": "2026-08-10", "observacion": "Consultado en la página oficial", "archivo": SimpleUploadedFile(nombre, contenido or pdf_minimo())}
        if persona_id:
            datos["persona_id"] = persona_id
        return c.http.post(f"/api/evaluaciones/{self.ev['id']}/proponentes/{prop}/aportados", datos, headers={"X-CSRFToken": c.csrf})


class HistoricoTests(BaseHistorico):
    def test_justificacion_obligatoria(self):
        self.assertEqual(self.revisar(self.abogado, self.p1, 2, True, "").status_code, 400)
        self.assertEqual(self.revisar(self.abogado, self.p1, 2, True, "Se verificó el COPNIA en físico").status_code, 200)

    def test_personas_y_certificados_aportados(self):
        url = f"/api/evaluaciones/{self.ev['id']}/proponentes/{self.p1}/personas"
        empresa = self.abogado.post(url, {"rol": "integrante", "tipo": "juridica", "nombre": "Vías del Norte SAS", "documento": "900.123.456-7"}).json()
        self.assertEqual(empresa["documento"], "900123456-7")
        rep = self.abogado.post(url, {"rol": "representante_legal", "tipo": "natural", "nombre": "Pedro Pérez", "documento": "1.020.304", "fecha_expedicion_documento": "2005-03-01", "de_id": empresa["id"]})
        self.assertEqual(rep.status_code, 201, rep.content)
        # El representante debe ser persona natural y solo una jurídica tiene representantes.
        self.assertEqual(self.abogado.post(url, {"rol": "suplente", "tipo": "juridica", "nombre": "Otra SAS", "documento": "8001"}).status_code, 400)
        self.assertEqual(self.abogado.post(url, {"rol": "suplente", "tipo": "natural", "nombre": "Ana Ruiz", "documento": "5566", "de_id": rep.json()["id"]}).status_code, 400)

        r = self.aportar(self.abogado, self.p1, 2, rep.json()["id"])
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(self.aportar(self.abogado, self.p1, 2, contenido=b"no es pdf").status_code, 400)
        self.assertEqual(self.aportar(self.abogado, self.p1, 99).status_code, 400)
        datos = self.abogado.get(f"/api/evaluaciones/{self.ev['id']}/proponentes/{self.p1}/antecedentes").json()
        self.assertEqual(len(datos["personas"]), 2)
        self.assertEqual(len(datos["aportados"]), 1)
        archivo = self.abogado.get(f"/api/evaluaciones/{self.ev['id']}/aportados/{r.json()['id']}/archivo")
        self.assertEqual(b"".join(archivo.streaming_content)[:5], b"%PDF-")
        # No se quita una persona con certificados; consulta y otra entidad no modifican ni ven.
        self.assertEqual(self.abogado.delete(f"/api/evaluaciones/{self.ev['id']}/personas/{empresa['id']}").status_code, 409)
        consulta = Cliente()
        consulta.entrar("control@entidad.gov.co")
        self.assertEqual(self.aportar(consulta, self.p1, 2).status_code, 403)
        otra = Cliente()
        otra.entrar("admin@otraentidad.gov.co")
        self.assertEqual(otra.get(f"/api/evaluaciones/{self.ev['id']}/aportados/{r.json()['id']}/archivo").status_code, 404)

    def test_reporte_word_y_expediente(self):
        import io
        import zipfile

        from docx import Document

        from evaluaciones.expediente import atender_pendientes

        for p in self.detalle["proponentes"]:
            self.assertEqual(self.revisar(self.abogado, p["id"], 2, True, "COPNIA verificado en la página del Consejo").status_code, 200)
        persona = self.abogado.post(
            f"/api/evaluaciones/{self.ev['id']}/proponentes/{self.p1}/personas",
            {"rol": "representante_legal", "tipo": "natural", "nombre": "Juan Pérez", "documento": "1020304"},
        ).json()
        self.aportar(self.abogado, self.p1, 1, persona["id"])

        r = self.abogado.get(f"/api/evaluaciones/{self.ev['id']}/reporte")
        self.assertEqual(r.status_code, 200)
        self.assertIn("BORRADOR", r["Content-Disposition"])
        texto = "\n".join(c.text for t in Document(io.BytesIO(r.content)).tables for fila in t.rows for c in fila.cells)
        self.assertIn("Aprobado automáticamente por MiEvaluador", texto)
        self.assertIn("Validado manualmente por Abogado Uno", texto)
        self.assertIn("COPNIA verificado en la página del Consejo", texto)
        self.assertIn("Certificado consultado y aportado por Abogado Uno", texto)

        with self.captureOnCommitCallbacks(execute=True):
            self.assertEqual(self.jefe.post(f"/api/evaluaciones/{self.ev['id']}/aprobar").status_code, 200)
        exp = self.jefe.get(f"/api/evaluaciones/{self.ev['id']}/expedientes").json()
        self.assertEqual([(e["version"], e["estado"]) for e in exp], [(1, "pendiente")])

        oferta = io.BytesIO()
        with zipfile.ZipFile(oferta, "w") as z:
            z.writestr("CARTA/carta.pdf", pdf_minimo("carta"))
        from motor.esquemas.proceso import ResultadoRequisito

        Resultado = __import__("evaluaciones.models", fromlist=["Resultado"]).Resultado
        for res in Resultado.objects.filter(evaluacion_id=self.ev["id"], requisito=1):
            datos = ResultadoRequisito.model_validate(res.datos).model_copy(update={"archivo_evaluado": "CARTA/carta.pdf"})
            res.datos = datos.model_dump(mode="json")
            res.save()
        with mock.patch("evaluaciones.expediente.download_file_bytes", return_value=oferta.getvalue()):
            self.assertEqual(atender_pendientes(), 1)
        exp = self.jefe.get(f"/api/evaluaciones/{self.ev['id']}/expedientes").json()[0]
        self.assertEqual(exp["estado"], "listo", exp["avisos"])
        r = self.jefe.get(f"/api/evaluaciones/{self.ev['id']}/expedientes/{exp['id']}/archivo")
        zbytes = b"".join(r.streaming_content)
        with zipfile.ZipFile(io.BytesIO(zbytes)) as z:
            nombres = z.namelist()
            self.assertIn("registro.json", nombres)
            self.assertIn("MANIFIESTO.txt", nombres)
            self.assertTrue(any(n.startswith("informe/") and n.endswith(".docx") for n in nombres))
            self.assertTrue(any(n.startswith("informe/") and n.endswith(".xlsx") for n in nombres))
            self.assertTrue(any("documentos evaluados" in n for n in nombres))
            self.assertTrue(any("antecedentes aportados" in n for n in nombres))
            self.assertNotIn("BORRADOR", " ".join(nombres))
        # Si el expediente se anula mientras se arma, no revive al terminar.
        from evaluaciones.expediente import construir, solicitar
        from evaluaciones.models import EstadoExpediente, Expediente

        anulado = solicitar(Evaluacion.objects.get(pk=self.ev["id"]), None)
        Expediente.objects.filter(pk=anulado.pk).update(estado=EstadoExpediente.GENERANDO)
        anulado.estado = EstadoExpediente.GENERANDO
        Expediente.objects.filter(pk=anulado.pk).delete()
        with mock.patch("evaluaciones.expediente.download_file_bytes", return_value=oferta.getvalue()):
            construir(anulado)
        self.assertFalse(Expediente.objects.filter(pk=anulado.pk).exists())

        # Otra entidad no descarga el expediente; regenerar crea la versión 2.
        otra = Cliente()
        otra.entrar("admin@otraentidad.gov.co")
        self.assertEqual(otra.get(f"/api/evaluaciones/{self.ev['id']}/expedientes/{exp['id']}/archivo").status_code, 404)
        self.assertEqual(self.jefe.post(f"/api/evaluaciones/{self.ev['id']}/expedientes").json()["version"], 2)


class SancionesRupTests(TestCase):
    """Reglas del abogado sobre multas y sanciones del RUP (art. 58 Ley 2195/2022
    y art. 90 Ley 1474/2011)."""

    cierre = __import__("datetime").date(2026, 8, 20)

    def sancion(self, tipo, anio, contrato="C1", penal=False, incumpl=False, descripcion=""):
        from datetime import date as _date

        from motor.evaluacion.sanciones_rup import Sancion

        return Sancion(
            tipo=tipo, entidad="ENTIDAD", contrato=contrato, descripcion=descripcion, fecha=_date(anio, 3, 1),
            valor="", es_incumplimiento=incumpl, clausula_penal=penal,
        )

    def evaluar(self, sanciones):
        from motor.evaluacion.sanciones_rup import evaluar_sanciones

        return evaluar_sanciones(sanciones, self.cierre)

    def test_multa_del_ultimo_ano_rechaza_y_las_viejas_no(self):
        self.assertFalse(self.evaluar([self.sancion("multa", 2026)])[0])
        self.assertFalse(self.evaluar([self.sancion("sancion", 2026, penal=True)])[0])
        cumple, motivo = self.evaluar([self.sancion("multa", 2023)])
        self.assertTrue(cumple)
        self.assertIn("anteriores al periodo evaluado", motivo)

    def test_inhabilidad_por_incumplimiento_reiterado(self):
        cinco = [self.sancion("multa", 2023, f"C{i}") for i in range(5)]
        self.assertFalse(self.evaluar(cinco)[0])
        self.assertTrue(self.evaluar(cinco[:4])[0])
        dos_contratos = [self.sancion("declaratoria", 2023, "C1"), self.sancion("declaratoria", 2023, "C2")]
        self.assertFalse(self.evaluar(dos_contratos)[0])
        # Mismo contrato en años distintos: no es reiterado en una vigencia fiscal.
        self.assertTrue(self.evaluar([self.sancion("declaratoria", 2023, "C1"), self.sancion("declaratoria", 2022, "C1")])[0])
        mixto = [self.sancion("multa", 2023, "C1"), self.sancion("multa", 2023, "C2"), self.sancion("declaratoria", 2023, "C3")]
        self.assertFalse(self.evaluar(mixto)[0])

    def test_un_mismo_hecho_reportado_dos_veces_no_cuenta_doble(self):
        # El RUP repite el hecho en «SANCIONES» y en «DECLARATORIAS DE INCUMPLIMIENTO».
        repetido = [
            self.sancion("sancion", 2022, "1242-2018", penal=True, incumpl=True),
            self.sancion("declaratoria", 2022, "1242-2018", penal=True),
        ]
        self.assertTrue(self.evaluar(repetido)[0])

    def test_datos_dudosos_van_a_revision(self):
        contradictoria = self.sancion("sancion", 2023, descripcion="INCUMPLIMIENTO DEFINITIVO")
        cumple, motivo = self.evaluar([contradictoria])
        self.assertFalse(cumple)
        self.assertIn("se contradice", motivo)
        sin_fecha = self.sancion("multa", 2023)
        sin_fecha.fecha = None
        self.assertFalse(self.evaluar([sin_fecha])[0])

    def test_lectura_del_rup_real(self):
        """Formatos reales de tres cámaras de comercio distintas."""
        import glob

        from motor.evaluacion.camara_comercio import _norm, _texto_paginas_finales
        from motor.evaluacion.sanciones_rup import extraer_sanciones

        archivos = sorted(glob.glob(".scratch/sanciones_rup/*.pdf"))
        if not archivos:
            self.skipTest("Sin RUP de ejemplo en .scratch (solo se ejecuta en el equipo de desarrollo).")
        for ruta in archivos:
            with open(ruta, "rb") as f:
                sanciones = extraer_sanciones(_norm(_texto_paginas_finales(f.read())))
            self.assertTrue(sanciones, ruta)
            self.assertTrue(all(s.fecha and s.entidad for s in sanciones), ruta)


class HuellaCriteriosTests(TestCase):
    """La caché nunca debe servir un resultado calculado con los criterios de
    una entidad a otra que evalúa distinto."""

    def test_configuraciones_distintas_dan_huellas_distintas(self):
        from motor import criterios

        with criterios.usar({"beneficiario_claves": ["ENT"]}):
            con_entidad = criterios.huella_parametros()
        with criterios.usar({"beneficiario_claves": ["OTRA"]}):
            otra_entidad = criterios.huella_parametros()
        self.assertIsNotNone(con_entidad)
        self.assertNotEqual(con_entidad, otra_entidad)

    def test_sin_configurar_no_hay_huella(self):
        from motor import criterios

        vacios = {clave: p.defecto for clave, p in criterios.PARAMETROS.items()}
        with criterios.usar(vacios):
            self.assertIsNone(criterios.huella_parametros())

    def test_el_entorno_cuenta_como_configuracion(self):
        """Un valor puesto por variable de entorno no es "el valor por defecto":
        si no contara, un resultado calculado sin él se serviría como válido."""
        from unittest import mock

        from motor import criterios

        with mock.patch.dict(criterios._ENTORNO, {"beneficiario_claves": ["ENT"]}, clear=True):
            with criterios.usar(None):
                self.assertEqual(criterios.valor("beneficiario_claves"), ["ENT"])
                self.assertIsNotNone(criterios.huella_parametros())
        with mock.patch.dict(criterios._ENTORNO, {}, clear=True):
            with criterios.usar(None):
                self.assertIsNone(criterios.huella_parametros())

    def test_la_entidad_manda_sobre_el_entorno(self):
        from unittest import mock

        from motor import criterios

        with mock.patch.dict(criterios._ENTORNO, {"beneficiario_claves": ["ENT"]}, clear=True):
            with criterios.usar({"beneficiario_claves": ["SUYA"]}):
                self.assertEqual(criterios.valor("beneficiario_claves"), ["SUYA"])


class IdentidadAntecedentesTests(TestCase):
    """Cada entidad redacta su certificado a su manera: lo que importa es sacar
    de quién es, para no pedirle al evaluador que revise algo que ya está."""

    def identidad(self, requisito, texto):
        from motor.evaluacion import antecedentes as ant

        config = next(c for c in ant._configs() if c.requisito == requisito)
        return config.extraer_identidad(ant._norm(texto))

    def test_rnmc_con_nombre_y_cedula(self):
        texto = (
            "EL CIUDADANO CON CEDULA DE CIUDADANIA NO. 79123456 Y NOMBRE: JUAN CARLOS PEREZ GOMEZ. "
            "NO TIENE MEDIDAS CORRECTIVAS PENDIENTES POR CUMPLIR."
        )
        self.assertEqual(self.identidad(17, texto), ("JUAN CARLOS PEREZ GOMEZ", "79123456"))

    def test_rnmc_solo_con_cedula(self):
        """Variante sin nombre: la cédula sola alcanza para saber de quién es."""
        texto = (
            "QUE A LA FECHA, 01/07/2026 08:38:35 A. M. EL CIUDADANO CON CEDULA DE CIUDADANIA "
            "NO. 79123456 . NO TIENE MEDIDAS CORRECTIVAS PENDIENTES POR CUMPLIR."
        )
        self.assertEqual(self.identidad(17, texto), (None, "79123456"))

    def test_rnmc_de_persona_juridica_no_es_de_nadie(self):
        """El certificado a nombre de un NIT no se puede atribuir a una persona."""
        texto = (
            "QUE A LA FECHA, 26/05/2026 PARA - NIT, SIN DIGITO DE VERIFICACION: NO. 900123456 "
            "NO TIENE MEDIDAS CORRECTIVAS PENDIENTES POR CUMPLIR."
        )
        self.assertEqual(self.identidad(17, texto), (None, None))

    def test_la_cedula_empareja_aunque_el_certificado_no_traiga_nombre(self):
        from motor.evaluacion.antecedentes import CONFIG_RNMC, evaluar_antecedente
        from unittest import mock

        from motor.evaluacion.antecedentes import Certificado

        certificado = Certificado(
            archivo="antecedentes.pdf",
            texto=(
                "SISTEMA REGISTRO NACIONAL DE MEDIDAS CORRECTIVAS RNMC "
                "EL CIUDADANO CON CEDULA DE CIUDADANIA NO. 79123456 . "
                "NO TIENE MEDIDAS CORRECTIVAS PENDIENTES POR CUMPLIR."
            ),
            requisitos=frozenset({17}),
        )
        with mock.patch("motor.evaluacion.antecedentes.leer_certificados", return_value=[certificado]):
            resultado = evaluar_antecedente({}, CONFIG_RNMC, [("JUAN CARLOS PEREZ GOMEZ", "79.123.456")])
        self.assertTrue(resultado.cumple, resultado.motivo)


class FacultadesRepresentanteTests(TestCase):
    """Criterio del abogado: un límite de cuantía solo impide contratar si
    queda por debajo del valor del proceso."""

    CERTIFICADO = (
        "CERTIFICADO DE EXISTENCIA Y REPRESENTACION LEGAL "
        "FACULTADES Y LIMITACIONES DEL REPRESENTANTE LEGAL: EL GERENTE NO PODRA SUSCRIBIR CONTRATOS "
        "SIN PREVIA AUTORIZACION DE LA ASAMBLEA POR CUANTIA SUPERIOR A {limite}. "
    )

    def evaluar(self, limite, valor_proceso, smmlv=0):
        from motor import criterios
        from motor.evaluacion.camara_comercio import _evaluar_facultades_certificado

        with criterios.usar({"smmlv": smmlv}):
            return _evaluar_facultades_certificado(
                "certificado.pdf", self.CERTIFICADO.format(limite=limite), valor_proceso
            )

    def test_limite_por_encima_del_proceso_se_aprueba(self):
        cumple, motivo = self.evaluar("$5.000.000.000", 2_269_370_336)
        self.assertTrue(cumple, motivo)
        self.assertIn("puede suscribir", motivo)

    def test_limite_por_debajo_va_a_revision(self):
        cumple, motivo = self.evaluar("$100.000.000", 2_269_370_336)
        self.assertFalse(cumple)
        self.assertIn("autorización", motivo)

    def test_limite_en_salarios_minimos_con_el_valor_configurado(self):
        # 100 salarios de $1.500.000 = $150.000.000: no alcanza para el proceso.
        cumple, _ = self.evaluar("100 SALARIOS MINIMOS MENSUALES LEGALES VIGENTES", 2_269_370_336, smmlv=1_500_000)
        self.assertFalse(cumple)
        # 3.000 salarios sí lo cubren.
        cumple, motivo = self.evaluar("3.000 SALARIOS MINIMOS MENSUALES LEGALES VIGENTES", 2_269_370_336, smmlv=1_500_000)
        self.assertTrue(cumple, motivo)

    def test_sin_salario_minimo_configurado_queda_para_revision(self):
        """Nunca se aprueba a ciegas: sin el dato, decide una persona."""
        cumple, motivo = self.evaluar("100 SALARIOS MINIMOS MENSUALES LEGALES VIGENTES", 2_269_370_336)
        self.assertFalse(cumple)
        self.assertIn("salario mínimo", motivo)


class AntecedentesEmpresaTests(TestCase):
    """Criterio del abogado: a cada persona jurídica (por NIT) se le exigen
    Contraloría y Procuraduría; Policía, RNMC y REDAM no."""

    PROCURADURIA_EMPRESA = (
        "PROCURADURIA GENERAL DE LA NACION CERTIFICA QUE UNA VEZ CONSULTADO EL SISTEMA DE INFORMACION DE "
        "REGISTRO DE SANCIONES E INHABILIDADES (SIRI), LA PERSONA CONSTRUCTORA EJEMPLO S.A.S. IDENTIFICADO(A) "
        "CON NIT NUMERO 9001234565: NO REGISTRA SANCIONES NI INHABILIDADES VIGENTES"
    )
    PROCURADURIA_PERSONA = (
        "PROCURADURIA GENERAL DE LA NACION CERTIFICA QUE EL(LA) SEÑOR(A) JUAN CARLOS PEREZ GOMEZ IDENTIFICADO(A) "
        "CON CEDULA DE CIUDADANIA NUMERO 79123456: NO REGISTRA SANCIONES NI INHABILIDADES VIGENTES"
    )

    def evaluar(self, config, textos, empresas):
        from unittest import mock

        from motor.evaluacion.antecedentes import Certificado, evaluar_antecedente

        certificados = [Certificado(f"c{i}.pdf", t, frozenset({config.requisito})) for i, t in enumerate(textos)]
        with mock.patch("motor.evaluacion.antecedentes.leer_certificados", return_value=certificados):
            return evaluar_antecedente({}, config, [("JUAN CARLOS PEREZ GOMEZ", "79123456")], empresas)

    def test_la_empresa_se_reconoce_por_su_nit(self):
        from motor.evaluacion.antecedentes import CONFIG_PROCURADURIA
        from motor.evaluacion.camara_comercio import Empresa

        r = self.evaluar(
            CONFIG_PROCURADURIA,
            [self.PROCURADURIA_PERSONA, self.PROCURADURIA_EMPRESA],
            [Empresa("CONSTRUCTORA EJEMPLO S.A.S", "900123456")],
        )
        self.assertTrue(r.cumple, r.motivo)

    def test_nit_pegado_al_texto(self):
        from motor.evaluacion.antecedentes import _nit_del_certificado

        texto = "LA PERSONA LOS EJEMPLOS SAS IDENTIFICADO(A)CON NIT NUMERO 8909340411: NO REGISTRA SANCIONES"
        self.assertEqual(_nit_del_certificado(15, texto), "890934041")

    def test_falta_el_de_la_empresa(self):
        from motor.evaluacion.antecedentes import CONFIG_PROCURADURIA
        from motor.evaluacion.camara_comercio import Empresa

        r = self.evaluar(
            CONFIG_PROCURADURIA, [self.PROCURADURIA_PERSONA], [Empresa("CONSTRUCTORA EJEMPLO S.A.S", "900123456")]
        )
        self.assertFalse(r.cumple)
        self.assertIn("NIT 900123456", r.motivo)

    def test_resultado_de_cada_persona(self):
        """La tabla de antecedentes necesita el estado de cada persona y empresa, no solo el global."""
        from unittest import mock

        from motor.evaluacion.antecedentes import CONFIG_PROCURADURIA, Certificado, evaluar_antecedente
        from motor.evaluacion.camara_comercio import Empresa

        certificados = [Certificado("c0.pdf", self.PROCURADURIA_PERSONA, frozenset({15}))]
        personas = [("JUAN CARLOS PEREZ GOMEZ", "79123456", "representante_legal"), ("ANA MARIA RUIZ LOPEZ", "52999888", "suplente")]
        with mock.patch("motor.evaluacion.antecedentes.leer_certificados", return_value=certificados):
            r = evaluar_antecedente({}, CONFIG_PROCURADURIA, personas, [Empresa("CONSTRUCTORA EJEMPLO S.A.S", "900123456")], plural=True)
        self.assertEqual(
            [(x.nombre, x.rol, x.tipo, x.estado) for x in r.personas],
            [
                ("JUAN CARLOS PEREZ GOMEZ", "representante_legal", "natural", "cumple"),
                ("ANA MARIA RUIZ LOPEZ", "suplente", "natural", "falta"),
                ("CONSTRUCTORA EJEMPLO S.A.S", "integrante", "juridica", "falta"),
            ],
        )
        self.assertEqual(r.personas[0].archivo, "c0.pdf")

    def test_a_la_empresa_no_se_le_exige_policia_ni_rnmc(self):
        from motor.evaluacion.antecedentes import CONFIG_RNMC
        from motor.evaluacion.camara_comercio import Empresa

        rnmc_persona = (
            "SISTEMA REGISTRO NACIONAL DE MEDIDAS CORRECTIVAS RNMC EL CIUDADANO CON CEDULA DE CIUDADANIA NO. "
            "79123456 . NO TIENE MEDIDAS CORRECTIVAS PENDIENTES POR CUMPLIR."
        )
        r = self.evaluar(CONFIG_RNMC, [rnmc_persona], [Empresa("CONSTRUCTORA EJEMPLO S.A.S", "900123456")])
        self.assertTrue(r.cumple, r.motivo)


class IntegrantesNaturalesTests(TestCase):
    """Criterio del abogado: al integrante persona natural de un consorcio se
    le piden todos los antecedentes menos el REDAM."""

    def test_se_agregan_sin_repetir_al_representante(self):
        from motor.evaluacion.antecedentes import _con_integrantes_naturales
        from motor.evaluacion.proponente_plural import Integrante

        personas = [("JUAN CARLOS PEREZ GOMEZ", "79123456")]
        integrantes = [
            Integrante("CONSTRUCTORA EJEMPLO S.A.S", "900123456-5", False),
            Integrante("MARIA LUISA TORRES DIAZ", "52111222", True),
            Integrante("JUAN CARLOS PEREZ GOMEZ", "79.123.456", True),  # también es el representante
        ]
        self.assertEqual(
            _con_integrantes_naturales(personas, integrantes),
            [("JUAN CARLOS PEREZ GOMEZ", "79123456"), ("MARIA LUISA TORRES DIAZ", "52111222", "integrante")],
        )

    def test_el_redam_no_se_le_pide_al_integrante(self):
        from motor.evaluacion.antecedentes import REQUISITOS_INTEGRANTE_NATURAL

        self.assertNotIn(5, REQUISITOS_INTEGRANTE_NATURAL)

    def test_forma_societaria(self):
        from motor.evaluacion.proponente_plural import MARCA_PERSONA_JURIDICA_RE

        for empresa in ("INJERR S.A.S", "PC PROYECCION E INGENIERIA SAS", "SESAC S.A", "CODIPRO LTDA.", "UG21 SL SUCURSAL EN COLOMBIA"):
            self.assertIsNotNone(MARCA_PERSONA_JURIDICA_RE.search(empresa), empresa)
        for persona in ("GERMAN ALONSO PUENTES GORDO", "MARTA EUGENIA GARCIA BETANCUR", "ARCELIA ARIAS DIAZ"):
            self.assertIsNone(MARCA_PERSONA_JURIDICA_RE.search(persona), persona)


# --- El pliego de cada proceso ---
def _pagina(numero: int, *lineas: str) -> "object":
    from motor.pliego.lectura import Pagina

    encabezado = ["DOCUMENTO BASE", f"Código CCE-EJEMPLO-01 Página {numero} de 4"]
    return Pagina(numero, "\n".join([*encabezado, *lineas, str(numero), "Código CCE-EJEMPLO-01 Versión 1"]))


PLIEGO_DE_PRUEBA = [
    _pagina(1, "CAPÍTULO I. INFORMACIÓN GENERAL ..................... 1", "2.3 LIMITACIÓN A MIPYME ................ 2",
            "3.3.2 PERSONAS JURÍDICAS ................ 3", "3.4 SEGURIDAD SOCIAL .................... 3",
            "8.1 GARANTÍA DE SERIEDAD ................ 4"),
    _pagina(2, "CAPÍTULO II. ELABORACIÓN Y PRESENTACIÓN DE LA OFERTA", "2.3 LIMITACIÓN A MIPYME",
            "El presente proceso se limita a Mipyme: únicamente podrán participar Mipyme domiciliadas en Colombia."),
    _pagina(3, "CAPÍTULO III. REQUISITOS HABILITANTES Y SU VERIFICACIÓN", "3.2 CAPACIDAD JURÍDICA",
            "E. Presentar el certificado del Registro de Deudores Alimentarios Morosos – REDAM en los términos de la Ley",
            "2097 de 2021",
            "La Entidad debe consultar los Antecedentes Judiciales en línea en los registros de las bases de datos.",
            "3.3.2 PERSONAS JURÍDICAS",
            "a. Fecha de expedición del certificado no mayor a treinta (30) días calendario anteriores a la fecha de cierre.",
            "c. Las personas jurídicas deberán acreditar que su duración no será inferior a la del plazo del contrato y un año más.",
            "III. Fotocopia del documento de identificación del representante legal."),
    _pagina(4, "3.4 CERTIFICACIÓN DE PAGOS AL SISTEMA DE SEGURIDAD SOCIAL",
            "El proponente acreditará el pago de aportes de seguridad social.",
            "CAPÍTULO XI. LISTA DE ANEXOS", "11.2 FORMATOS", "1. Formato 1 – Carta de presentación de la oferta",
            "2. Formato 3 – Experiencia", "3. Formato 14 – Acreditación de Mipyme"),
]


class LecturaPliegoTests(TestCase):
    def test_secciones_sin_indice_ni_encabezados(self):
        from motor.pliego.lectura import secciones

        s = {x.numero: x for x in secciones(PLIEGO_DE_PRUEBA)}
        self.assertIn("3.3.2", s)
        self.assertEqual(s["3.3.2"].pagina, 3)
        self.assertNotIn("Página", s["3.2"].texto)  # encabezado repetido fuera
        self.assertEqual(sum(1 for x in secciones(PLIEGO_DE_PRUEBA) if x.numero == "2.3"), 1)  # el índice no cuenta

    def test_no_borra_contenido_que_parece_numero_de_pagina(self):
        """"2097 de 2021" es parte de "Ley 2097 de 2021", no un "página 2 de 4"."""
        from motor.pliego.lectura import secciones

        s = {x.numero: x for x in secciones(PLIEGO_DE_PRUEBA)}
        self.assertIn("2097 de 2021", s["3.2"].texto)


class AnalisisPliegoTests(TestCase):
    def hallazgos(self, definicion=None):
        from motor import criterios
        from motor.pliego.analisis import comparar, extraer

        return {h.id: h for h in comparar(extraer(PLIEGO_DE_PRUEBA), definicion or criterios.definicion_sistema("juridica"))}

    def test_detecta_lo_que_cambia_la_evaluacion(self):
        h = self.hallazgos()
        self.assertEqual(h["vigencia_existencia"].tipo, "ajuste_parametro")
        self.assertEqual((h["vigencia_existencia"].parametro, h["vigencia_existencia"].valor_pliego), ("camara_dias", 30))
        self.assertEqual(h["vigencia_existencia"].pagina, 3)
        self.assertIn("treinta (30) días", h["vigencia_existencia"].cita)
        for requisito in ("duracion_sociedad", "identidad_representante", "limitacion_mipyme"):
            self.assertEqual(h[requisito].tipo, "requisito_nuevo", requisito)
            self.assertTrue(h[requisito].requiere_decision)
        self.assertEqual(h["consulta_antecedentes_entidad"].tipo, "aclaracion")
        self.assertFalse(h["consulta_antecedentes_entidad"].requiere_decision)

    def test_formatos_y_alcance(self):
        h = self.hallazgos()
        tipos = {x.titulo: x.tipo for x in h.values() if x.id.startswith("formato_")}
        self.assertEqual(tipos.get("Formato 3 – Experiencia"), "fuera_de_alcance")
        self.assertNotIn("Formato 1 – Carta de presentación de la oferta", tipos)  # ya lo revisa la plantilla

    def test_sin_limitacion_mipyme_no_pide_nada(self):
        from motor.pliego.analisis import comparar, extraer
        from motor import criterios

        paginas = list(PLIEGO_DE_PRUEBA)
        paginas[1] = _pagina(2, "CAPÍTULO II. ELABORACIÓN", "2.3 LIMITACIÓN A MIPYME",
                             "El presente procedimiento de selección no es susceptible de limitarse a Mipyme.")
        h = {x.id: x for x in comparar(extraer(paginas), criterios.definicion_sistema("juridica"))}
        self.assertEqual(h["limitacion_mipyme"].tipo, "informativo")

    def test_si_la_plantilla_ya_usa_30_dias_no_hay_ajuste(self):
        from motor import criterios

        definicion = criterios.definicion_sistema("juridica")
        definicion.parametros = {"camara_dias": 30}
        self.assertNotIn("vigencia_existencia", self.hallazgos(definicion))

    def test_obligacion_desconocida_queda_a_la_vista(self):
        """Una exigencia que ningún detector conoce no se pierde: se muestra con su cita."""
        from motor import criterios
        from motor.pliego.analisis import comparar, extraer

        paginas = list(PLIEGO_DE_PRUEBA)
        paginas[3] = _pagina(4, "3.4 CERTIFICACIÓN DE PAGOS AL SISTEMA DE SEGURIDAD SOCIAL",
                             "El proponente deberá aportar el paz y salvo ambiental expedido por la corporación autónoma regional.")
        h = [x for x in comparar(extraer(paginas), criterios.definicion_sistema("juridica")) if x.tipo == "obligacion"]
        self.assertEqual(len(h), 1)
        self.assertIn("paz y salvo ambiental", h[0].cita)
        self.assertEqual(h[0].pagina, 4)

    def test_pliego_sin_estructura_avisa(self):
        from motor import criterios
        from motor.pliego.analisis import comparar, extraer
        from motor.pliego.lectura import Pagina

        paginas = [Pagina(i, "texto escaneado ilegible sin títulos") for i in range(1, 6)]
        h = comparar(extraer(paginas), criterios.definicion_sistema("juridica"))
        self.assertEqual(h[0].id, "estructura_no_reconocida")

    def test_aplicar_ajustes_aceptados(self):
        from evaluaciones.pliego import aplicar_ajustes
        from motor import criterios

        h = self.hallazgos()
        ajustes = [
            {"id": "vigencia_existencia", "decision": "aceptado", "hallazgo": h["vigencia_existencia"].model_dump()},
            {"id": "duracion_sociedad", "decision": "aceptado", "hallazgo": h["duracion_sociedad"].model_dump()},
            {"id": "identidad_representante", "decision": "rechazado", "nota": "No aplica", "hallazgo": h["identidad_representante"].model_dump()},
        ]
        base = criterios.definicion_sistema("juridica")
        ajustada = aplicar_ajustes(base, ajustes)
        self.assertEqual(ajustada.parametros["camara_dias"], 30)
        nuevos = ajustada.requisitos[len(base.requisitos):]
        self.assertEqual([(r.titulo, r.verificacion) for r in nuevos], [("Duración de la sociedad", "juridica.duracion")])
        self.assertEqual(nuevos[0].numero, max(r.numero for r in base.requisitos) + 1)

    def test_la_verificacion_manual_queda_pendiente(self):
        from datetime import date

        from motor import criterios
        from motor.esquemas.proceso import Proponente
        from motor.evaluacion.todos import evaluar_requisito

        req = criterios.RequisitoDefinicion(numero=19, titulo="Duración de la sociedad", corto="Duración",
                                            verificacion=criterios.MANUAL, verifica="Plazo del contrato y un año más")
        proponente = Proponente(numero_orden=1, hoja="P-01", nombre_proponente="Ejemplo", nombre_archivo="p1.zip", drive_file_id="x")
        r = evaluar_requisito(proponente, None, req)
        self.assertFalse(r.cumple)
        self.assertIn("Verificación manual", r.motivo)
        self.assertEqual(r.requisito, 19)

    def test_vigencia_en_dias(self):
        from datetime import date

        from motor import criterios
        from motor.evaluacion.camara_comercio import limite_expedicion_camara

        with criterios.usar({"camara_dias": 30}):
            self.assertEqual(limite_expedicion_camara(date(2026, 8, 20))[0], date(2026, 7, 21))
        self.assertEqual(limite_expedicion_camara(date(2026, 8, 20))[0], date(2026, 7, 20))  # 1 mes


class PliegoProcesoTests(BaseEvaluaciones):
    """El proceso se crea con el pliego analizado y lo que la persona decidió."""

    def analisis(self, entidad=None):
        from motor.pliego.analisis import extraer

        with mock.patch("evaluaciones.pliego.leer", return_value=extraer(PLIEGO_DE_PRUEBA)):
            from evaluaciones import pliego

            a, _ = pliego.analizar((entidad or self.entidad1).id, b"%PDF-prueba", "pliego.pdf", self.jefe)
        return a

    def crear_con(self, analisis_id, decisiones, email="jefe@entidad.gov.co"):
        c = Cliente()
        c.entrar(email)
        datos = {"documento_base": DOCUMENTO_BASE, "carpeta_drive": "x", "proponentes": PROPONENTES,
                 "analisis_pliego_id": str(analisis_id), "decisiones_pliego": decisiones}
        return c, c.post("/api/evaluaciones/procesos", datos)

    def test_no_se_crea_sin_decidir_todo(self):
        _, r = self.crear_con(self.analisis().id, {"vigencia_existencia": {"decision": "aceptado"}})
        self.assertEqual(r.status_code, 400)
        self.assertIn("Falta decidir", r.json()["detail"])

    def test_rechazar_exige_razon(self):
        a = self.analisis()
        decisiones = {h: {"decision": "aceptado"} for h in ("vigencia_existencia", "duracion_sociedad", "limitacion_mipyme")}
        decisiones["identidad_representante"] = {"decision": "rechazado", "nota": ""}
        _, r = self.crear_con(a.id, decisiones)
        self.assertEqual(r.status_code, 400)
        self.assertIn("Explique", r.json()["detail"])

    def test_los_ajustes_llegan_a_la_evaluacion_y_al_reporte(self):
        from evaluaciones.models import Evaluacion
        from evaluaciones.reporte import generar_reporte

        a = self.analisis()
        decisiones = {h: {"decision": "aceptado"} for h in ("vigencia_existencia", "duracion_sociedad", "limitacion_mipyme")}
        decisiones["identidad_representante"] = {"decision": "rechazado", "nota": "La cédula se revisa con el RUP."}
        c, r = self.crear_con(a.id, decisiones)
        self.assertEqual(r.status_code, 201, r.content)
        evaluacion = Evaluacion.objects.get(pk=r.json()[0]["id"])
        definicion = servicios.definicion_de(evaluacion)
        self.assertEqual(definicion.parametros["camara_dias"], 30)
        self.assertIn("Duración de la sociedad", [x.titulo for x in definicion.requisitos])
        self.assertNotIn("Documento de identidad del representante legal", [x.titulo for x in definicion.requisitos])
        detalle = c.get(f"/api/evaluaciones/{evaluacion.id}").json()
        self.assertEqual(detalle["pliego"]["nombre_archivo"], "pliego.pdf")
        self.assertIn("Duración de la sociedad", [x["titulo"] for x in detalle["catalogo"]])
        contenido, _ = generar_reporte(evaluacion)
        self.assertGreater(len(contenido), 1000)
        from evaluaciones.reporte import notas_del_pliego

        notas = notas_del_pliego(evaluacion.proceso, definicion)
        existencia = next(x.numero for x in definicion.requisitos if x.verificacion == "juridica.existencia")
        duracion = next(x.numero for x in definicion.requisitos if x.titulo == "Duración de la sociedad")
        self.assertIn("Debido al pliego", notas[existencia])
        self.assertIn("30 días", notas[existencia])
        self.assertIn("pág.", notas[duracion])

    def test_mismo_pdf_se_reutiliza_y_otra_entidad_no_lo_ve(self):
        from evaluaciones import pliego

        a = self.analisis()
        with mock.patch("evaluaciones.pliego.leer", side_effect=AssertionError("no debía releerse")):
            b, reutilizado = pliego.analizar(self.entidad1.id, b"%PDF-prueba", "otro-nombre.pdf", self.jefe)
        self.assertTrue(reutilizado)
        self.assertEqual(a.id, b.id)
        _, r = self.crear_con(a.id, {}, email="abogado@otraentidad.gov.co")
        self.assertEqual(r.status_code, 404)


    def test_el_pliego_queda_en_el_expediente(self):
        import io
        import json
        import zipfile

        from evaluaciones.expediente import construir
        from evaluaciones.models import EstadoExpediente, Evaluacion, Expediente

        a = self.analisis()
        decisiones = {h: {"decision": "aceptado"} for h in ("vigencia_existencia", "duracion_sociedad", "limitacion_mipyme", "identidad_representante")}
        _, r = self.crear_con(a.id, decisiones)
        self.assertEqual(r.status_code, 201, r.content)
        evaluacion = Evaluacion.objects.get(pk=r.json()[0]["id"])
        exp = Expediente.objects.create(entidad=evaluacion.entidad, evaluacion=evaluacion, version=1, estado=EstadoExpediente.GENERANDO)
        construir(exp)
        exp.refresh_from_db()
        with exp.archivo.open("rb") as f, zipfile.ZipFile(io.BytesIO(f.read())) as z:
            self.assertIn("pliego/pliego.pdf", z.namelist())
            registro = json.loads(z.read("registro.json"))
        self.assertEqual(registro["pliego"]["sha256"], a.sha256)
        self.assertEqual(len(registro["pliego"]["ajustes"]), 4)


class TablaAntecedentesTests(BaseHistorico):
    """La tabla de antecedentes se arma sola con todas las personas que exige la
    regla (las que sean) y el estado del certificado de cada una."""

    def guardar(self, personas_por_requisito):
        from evaluaciones.models import Evaluacion, Proponente
        from motor.esquemas.proceso import PersonaAntecedente, ResultadoRequisito

        ev = Evaluacion.objects.get(pk=self.ev["id"])
        prop = Proponente.objects.get(pk=self.p1)
        resultados = [
            ResultadoRequisito(
                hoja=prop.hoja, numero_orden=prop.numero_orden, nombre_proponente=prop.nombre, requisito=req,
                cumple=all(p[3] == "cumple" for p in personas), motivo=None,
                personas_antecedente=[
                    PersonaAntecedente(nombre=n, documento=d, tipo=t, rol=r, estado=e) for n, d, r, e, t in
                    [(p[0], p[1], p[2], p[3], p[4] if len(p) > 4 else "natural") for p in personas]
                ],
            )
            for req, personas in personas_por_requisito.items()
        ]
        servicios.guardar_resultados(ev, prop, resultados)

    def tabla(self):
        return self.abogado.get(f"/api/evaluaciones/{self.ev['id']}/proponentes/{self.p1}/antecedentes").json()

    def test_varias_personas_con_su_estado(self):
        rep = ("ANA MARIA RUIZ LOPEZ", "52999888", "representante_legal")
        sup = ("CARLOS PEREZ GOMEZ", "79123456", "suplente")
        emp = ("CONSTRUCTORA EJEMPLO S.A.S", "900123456", "integrante")
        self.guardar({
            14: [(*rep, "cumple"), (*sup, "falta"), (*emp, "cumple", "juridica")],
            16: [(*rep, "cumple"), (*sup, "con_novedad")],
        })
        t = self.tabla()
        personas = {p["nombre"]: p for p in t["personas"]}
        self.assertEqual(set(personas), {"ANA MARIA RUIZ LOPEZ", "CARLOS PEREZ GOMEZ", "CONSTRUCTORA EJEMPLO S.A.S"})
        self.assertTrue(all(p["detectada"] for p in personas.values()))
        estados = t["estados"]
        self.assertEqual(estados[personas["CARLOS PEREZ GOMEZ"]["id"]]["14"]["estado"], "falta")
        self.assertEqual(estados[personas["CARLOS PEREZ GOMEZ"]["id"]]["16"]["estado"], "con_novedad")
        # A la empresa no se le exige Policía.
        self.assertEqual(estados[personas["CONSTRUCTORA EJEMPLO S.A.S"]["id"]]["16"]["estado"], "no_requerido")

    def test_reevaluar_no_toca_lo_manual_ni_lo_aportado(self):
        self.guardar({14: [("ANA MARIA RUIZ LOPEZ", "52999888", "representante_legal", "falta")]})
        manual = self.abogado.post(
            f"/api/evaluaciones/{self.ev['id']}/proponentes/{self.p1}/personas",
            {"rol": "integrante", "tipo": "natural", "nombre": "Pedro Pablo León", "documento": "1020304050"},
        ).json()
        ana = next(p for p in self.tabla()["personas"] if p["nombre"] == "ANA MARIA RUIZ LOPEZ")
        self.aportar(self.abogado, self.p1, 14, ana["id"])
        # Al reevaluar, el programa ya no la exige (otro representante): se conserva por tener un certificado aportado.
        self.guardar({14: [("JORGE LUIS DIAZ", "80111222", "representante_legal", "cumple")]})
        nombres = {p["nombre"] for p in self.tabla()["personas"]}
        self.assertEqual(nombres, {"ANA MARIA RUIZ LOPEZ", "PEDRO PABLO LEÓN", "JORGE LUIS DIAZ"})
        self.assertIn(manual["id"], {p["id"] for p in self.tabla()["personas"]})


class SalarioMinimoTests(BaseEvaluaciones):
    def test_la_evaluacion_usa_el_del_ano_del_cierre(self):
        from evaluaciones.models import Evaluacion, SalarioMinimo

        _, ev = self.crear()
        evaluacion = Evaluacion.objects.get(pk=ev["id"])
        ano = evaluacion.proceso.fecha_cierre.year
        SalarioMinimo.objects.update_or_create(ano=ano, defaults={"valor": 1_750_905})
        self.assertEqual(servicios.definicion_de(evaluacion).parametros["smmlv"], 1_750_905)
        SalarioMinimo.objects.filter(ano=ano).delete()
        self.assertNotIn("smmlv", servicios.definicion_de(evaluacion).parametros)

    def test_los_valores_oficiales_vienen_cargados(self):
        from evaluaciones.models import SalarioMinimo

        self.assertEqual(SalarioMinimo.objects.get(ano=2025).valor, 1_423_500)
        self.assertEqual(SalarioMinimo.objects.get(ano=2026).valor, 1_750_905)

    def test_solo_el_superadmin_lo_actualiza(self):
        import pyotp

        c = Cliente()
        c.entrar("admin@entidad.gov.co")
        self.assertEqual(c.http.put("/api/plataforma/salarios-minimos/2027", {"valor": 1900000}, content_type="application/json",
                                    headers={"X-CSRFToken": c.csrf}).status_code, 403)
        s = Cliente()
        s.entrar("santiagopebe01@lemartek.com")
        secreto = s.post("/api/auth/2fa/configurar").json()["secreto"]
        s.post("/api/auth/2fa/verificar", {"codigo": pyotp.TOTP(secreto).now()})
        r = s.http.put("/api/plataforma/salarios-minimos/2027", {"valor": 1900000, "norma": "Decreto de prueba"},
                       content_type="application/json", headers={"X-CSRFToken": s.csrf})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(s.get("/api/plataforma/salarios-minimos").json()[0]["ano"], 2027)


class NombresOfertasDriveTests(TestCase):
    """Las ofertas llegan nombradas "P1 Nombre" o, como las descarga el SECOP II
    y las numera la entidad, "110. CONSORCIO ASF.zip"."""

    def test_formatos_reconocidos(self):
        from motor.integrations.drive import _parse_nombre

        self.assertEqual(_parse_nombre("P1 JJAB SAS.zip"), (1, "JJAB SAS"))
        self.assertEqual(_parse_nombre("P-03 - Consorcio Ejemplo.rar"), (3, "Consorcio Ejemplo"))
        self.assertEqual(_parse_nombre("110. CONSORCIO ASF..zip"), (110, "CONSORCIO ASF"))
        self.assertEqual(_parse_nombre("7) OBRAS SAS.zip"), (7, "OBRAS SAS"))

    def test_un_pdf_numerado_no_es_una_oferta(self):
        from motor.integrations.drive import _parse_nombre

        self.assertIsNone(_parse_nombre("2026 informe.pdf"))
        self.assertIsNone(_parse_nombre("Mostrar mensaje.pdf"))

    def test_las_subcarpetas_no_se_mezclan_con_las_ofertas_de_la_raiz(self):
        from unittest import mock

        from motor.integrations import drive

        archivos = [
            {"id": "a", "name": "1. OBRAS SAS.zip", "profundidad": 0, "carpeta": ""},
            {"id": "b", "name": "2. CONSORCIO X.zip", "profundidad": 0, "carpeta": ""},
            {"id": "c", "name": "1. OBRAS SAS.zip", "profundidad": 1, "carpeta": "SOBRE ECONOMICO"},
        ]
        import pathlib
        import tempfile

        with tempfile.TemporaryDirectory() as tmp, mock.patch.object(drive, "CACHE_LISTADOS_DIR", pathlib.Path(tmp)), \
                mock.patch.object(drive, "get_drive_service"), mock.patch.object(drive, "_listar_archivos_recursivo", return_value=archivos), \
                mock.patch.object(drive, "_solo_cache", return_value=False):
            r = drive.list_proponentes("1k6QaO1v0bEbnhKnT2edkavswCpK3zzYc_prueba")
        self.assertEqual([p.drive_file_id for p in r.proponentes], ["a", "b"])
        self.assertIn("SOBRE ECONOMICO/1. OBRAS SAS.zip (subcarpeta: no se usa)", r.no_reconocidos)


class DocumentoBaseObjetoUnicoTests(TestCase):
    """Pliegos de obra de un solo objeto, sin tabla de lotes."""

    def test_objeto_unico(self):
        from unittest import mock

        from motor.parsers.documento_base import _objeto_unico

        tabla = [
            ["Objeto del proyecto", "Plazo del contrato", "", "Valor presupuesto", "", "", "Lugar(es) de"],
            ["", "", "", "oficial (pesos incluido IVA)", "", "", "ejecución"],
            ["CONSTRUCCIÓN DE LA INSTITUCIÓN EDUCATIVA DEL MUNICIPIO", "DOCE (12) MESES", "$5.458.954.545,00", "", "", "Municipio de Ejemplo", ""],
        ]
        pagina = mock.Mock()
        pagina.extract_tables.return_value = [tabla]
        pdf = mock.Mock(pages=[pagina])
        lote = _objeto_unico(pdf, 0)
        self.assertEqual((lote.numero, lote.plazo_meses, lote.valor_presupuesto), ("ÚNICO", 12, 5_458_954_545.0))
        self.assertEqual(lote.objeto, "CONSTRUCCIÓN DE LA INSTITUCIÓN EDUCATIVA DEL MUNICIPIO")
        self.assertEqual(lote.lugar_ejecucion, "Municipio de Ejemplo")


class CertificadoSeguridadSocialTests(TestCase):
    """Si el pliego dice que "bastará el certificado suscrito por el revisor
    fiscal o el representante legal", una certificación propia vale como el formato."""

    PROPIA = "CERTIFICACION DE PAGOS DE SEGURIDAD SOCIAL Y APORTES PARAFISCALES - ARTICULO 50 DE LA LEY 789 DE 2002"

    def test_solo_con_el_ajuste_del_pliego(self):
        from motor import criterios
        from motor.evaluacion.seguridad_social import _es_certificado

        self.assertTrue(_es_certificado("FORMATO 6 – PAGOS DE SEGURIDAD SOCIAL Y APORTES LEGALES"))
        self.assertFalse(_es_certificado(self.PROPIA))
        with criterios.usar({"seguridad_social_certificado": 1}):
            self.assertTrue(_es_certificado(self.PROPIA))
            # El formulario del SECOP nunca cuenta.
            self.assertFalse(_es_certificado(self.PROPIA + " ESTA PREGUNTA REQUIERE ANEXAR DOCUMENTOS"))

    def test_el_pliego_lo_propone(self):
        from motor import criterios
        from motor.pliego.analisis import comparar, extraer

        paginas = list(PLIEGO_DE_PRUEBA)
        paginas[3] = _pagina(4, "3.4 CERTIFICACIÓN DE PAGOS AL SISTEMA DE SEGURIDAD SOCIAL",
                             "La Entidad no exigirá las planillas de pago. Bastará el certificado suscrito por el revisor fiscal, "
                             "en los casos requeridos por la ley, o por el representante legal que así lo acredite.")
        h = {x.id: x for x in comparar(extraer(paginas), criterios.definicion_sistema("juridica"))}
        self.assertEqual((h["certificado_seguridad_social"].parametro, h["certificado_seguridad_social"].valor_pliego),
                         ("seguridad_social_certificado", 1))
        self.assertTrue(h["certificado_seguridad_social"].requiere_decision)


class LecturaPolizaTests(TestCase):
    """Garantía de seriedad (Req 11): la fila del amparo en los formatos
    reales de las aseguradoras, y los errores del OCR sin aprobar de más."""

    EXIGIDO = 545_895_454.50

    def leer(self, texto):
        from motor.evaluacion.formato1 import _norm
        from motor.evaluacion.garantia import leer_vigencia_y_valor

        return leer_vigencia_y_valor(_norm(texto), 2026, self.EXIGIDO)

    def test_valores_con_separadores_del_ocr(self):
        from motor.evaluacion.garantia import _parsear_valor_pesos

        self.assertEqual(_parsear_valor_pesos("545.895.454,50"), 545_895_454.50)
        self.assertEqual(_parsear_valor_pesos("$545,895,454.50"), 545_895_454.50)
        self.assertEqual(_parsear_valor_pesos("545.,895.454,50"), 545_895_454.50)
        self.assertEqual(_parsear_valor_pesos("545,895"), 545_895)

    def test_fechas_del_ocr(self):
        from datetime import date

        from motor.evaluacion.garantia import _fechas_posibles

        self.assertEqual(_fechas_posibles("24", "11", "2026", 2026), {date(2026, 11, 24)})
        # 0 leído como 8, 2 como 7, una cifra de más: solo queda el año razonable.
        self.assertEqual(_fechas_posibles("24", "11", "2826", 2026), {date(2026, 11, 24)})
        self.assertEqual(_fechas_posibles("24", "11", "7076", 2026), {date(2026, 11, 24)})
        self.assertEqual(_fechas_posibles("24", "11", "28726", 2026), {date(2026, 11, 24)})
        self.assertEqual(_fechas_posibles("31", "18", "2826", 2026), {date(2026, 10, 31)})
        # Si hubo que corregir, el día también puede estar mal: todas las lecturas.
        self.assertEqual(_fechas_posibles("18", "11", "2826", 2026), {date(2026, 11, 18), date(2026, 11, 10)})

    def test_formatos_de_aseguradoras(self):
        from datetime import date

        casos = {
            "SERIEDAD DE LA OFERTA ---------- COP 545.895.454,50 VIGENCIA DE LA COBERTURA : DESDE LAS 0 HS DEL "
            "24/07/2026, HASTA LAS 0 HS DEL 24/11/2026 PRIMA DE LA COBERTURA : COP 545.895,00": date(2026, 11, 24),
            "SERTEDAD OFERTA CO 2026/07/24 2026/11/24 545,895,454.50 COL$ 545,895.45 COL$": date(2026, 11, 24),
            "SERIEDAD DE LA OFERTA 24-07-2026 11-11-2026 545,895,455.00 545,895,455.00 545,895.00": date(2026, 11, 11),
            "SERIEDAD DE LA OFERTA 00:00 Horas Del 24/07/2026 24:00 Horas Del 30/12/28026 545.895.454,58": date(2026, 12, 30),
        }
        for texto, hasta in casos.items():
            lectura = self.leer(texto)
            self.assertEqual(lectura.hasta, {hasta}, texto)
            self.assertGreaterEqual(lectura.valor, self.EXIGIDO)

    def test_la_fecha_de_adjudicacion_no_es_la_vigencia(self):
        # El OCR dañó la vigencia hasta ("202€6"): no se toma la fecha que sigue.
        texto = "SERIEDAD DE LA OFERTA 24/07/2026 24/11/202€6 $545,898,454.50 FECHA ADJUDICACION : 13/08/2026"
        self.assertIsNone(self.leer(texto))
        texto = "SERIEDAD DE LA OFERTA 24/07/2076 24/11/2026 $545,895,455.00 FECHA ADJUDICACION : 23/08/2026"
        self.assertEqual({str(f) for f in self.leer(texto).hasta}, {"2026-11-24"})

    def test_no_toma_la_prima_ni_el_clausulado(self):
        self.assertIsNone(self.leer("SERIEDAD DE LA OFERTA 24/07/2026 24/11/2026 545.895,00"))
        self.assertIsNone(self.leer("LA GARANTIA DE SERIEDAD DE LA OFERTA CUBRIRA LA SANCION DERIVADA DEL INCUMPLIMIENTO"))

    def test_encabezado_de_mundial(self):
        texto = ("VIGENCIA DESDE VIGENCIA HASTA VIGENCIA DEL CERTIFICADO DESDE VIGENCIA DEL CERTIFICADO HASTA "
                 "00:08 HORAS DEL | 24/07/2826 |24:08 HORAS DEL | 01/12/2826 CONDICIONES PARTICULARES "
                 "TOTAL ASEGURADO $ 545.895.454,50")
        lectura = self.leer(texto)
        self.assertEqual({str(f) for f in lectura.hasta}, {"2026-12-01"})

    def test_beneficiario_con_etiqueta_asegurado(self):
        from motor import criterios
        from motor.evaluacion.formato1 import _norm
        from motor.evaluacion.garantia import _beneficiario_re

        with criterios.usar({"beneficiario_claves": ["ICCU", "INSTITUTO DE CAMINOS"]}):
            patron = _beneficiario_re()
            self.assertTrue(patron.search(_norm("ASEGURADO: INSTITUTO DE CAMINOS Y CONSTRUCCIONES DE CUNDINAMARCA - ICCU")))
            self.assertTrue(patron.search(_norm("BENEFICIARIO INSTITUTO DE CAMINOS Y CONSTRUCCIONES")))
            # El clausulado nombra al asegurado sin ser la entidad.
            self.assertFalse(patron.search(_norm(
                "EL TOMADOR Y/O ASEGURADO SEGÚN CORRESPONDA, SE COMPROMETE A PAGAR LA PRIMA DENTRO DE LOS 30 DÍAS")))
            # Póliza de otra entidad (se vio en una oferta real).
            self.assertFalse(patron.search(_norm("ASEGURADO MUNICIPIO DE MANIZALES BENEFICIARIO MUNICIPIO DE MANIZALES")))


class OneDriveTests(TestCase):
    """Carpetas de ofertas compartidas por OneDrive con enlace público."""

    ENLACE = "https://1drv.ms/f/c/d09ede0cd2e6118e/IgBXSGkuRU2QR6ag50US4Rkb?e=hPgliQ"

    def respuesta(self, datos, estado=200):
        r = mock.Mock(status_code=estado)
        r.json.return_value = datos
        r.raise_for_status.return_value = None
        return r

    def test_reconoce_enlaces(self):
        from motor.integrations import onedrive

        self.assertTrue(onedrive.es_enlace(self.ENLACE))
        self.assertTrue(onedrive.es_enlace("https://onedrive.live.com/?id=ABC"))
        self.assertFalse(onedrive.es_enlace("https://drive.google.com/drive/folders/abc"))
        self.assertTrue(onedrive._token_de_enlace(self.ENLACE).startswith("u!aHR0cHM6Ly8xZHJ2"))

    def test_lista_ofertas_como_drive(self):
        import tempfile
        from pathlib import Path

        from motor.integrations import drive, onedrive

        zip_ = {"id": "D09!s1", "name": "1. GARANS SAS.zip", "size": 10, "file": {"hashes": {"quickXorHash": "abc="}}}
        otro = {"id": "D09!s2", "name": "LEAME.txt", "size": 1, "file": {"hashes": {"quickXorHash": "x"}}}
        raiz = {"id": "D09!raiz", "name": "OFERTAS", "folder": {}, "children": [zip_, otro]}
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(onedrive, "_DIR", Path(tmp)), \
                mock.patch.object(onedrive, "_ENLACES", Path(tmp) / "e.json"), \
                mock.patch.object(drive, "CACHE_LISTADOS_DIR", Path(tmp)), \
                mock.patch.object(onedrive, "_token_anonimo", return_value="T"), \
                mock.patch.object(onedrive.requests, "get", return_value=self.respuesta(raiz)):
            r = drive.list_proponentes(self.ENLACE)
            self.assertEqual([(p.hoja, p.nombre_proponente) for p in r.proponentes], [("P-01", "GARANS SAS")])
            self.assertEqual(r.proponentes[0].drive_file_id, "onedrive!D09!s1")
            self.assertEqual(r.no_reconocidos, ["LEAME.txt"])
            # La descarga sabe de qué enlace es cada archivo.
            self.assertEqual(onedrive._enlace_de("onedrive!D09!s1"), self.ENLACE)

    def test_enlace_sin_permiso(self):
        from motor.integrations import onedrive

        with mock.patch.object(onedrive, "_token_anonimo", return_value="T"), \
                mock.patch.object(onedrive.requests, "get", return_value=self.respuesta({}, 403)):
            with self.assertRaisesMessage(onedrive.OneDriveError, "Cualquier persona con el vínculo"):
                onedrive._redimir("https://1drv.ms/f/c/otro")


class ContenidoCartaTests(TestCase):
    """Carta de presentación frente al Formato 1 del pliego (casos reales
    que el abogado rechazó y el programa aprobaba)."""

    NUMERAL_5 = ("5. Tengo conocimiento acerca de las características y condiciones del sitio de ejecución del proyecto, "
                 "por lo que asumo la responsabilidad de su revisión con la presentación de esta oferta.")
    NUMERAL_5_CAMBIADO = ("5. Tengo conocimiento acerca de las características y condiciones del sitio de ejecución del "
                          "proyecto y asumo los Riesgos previsibles inherentes al mismo, así como aquellos asignados en el "
                          "Pliego de Condiciones.")

    def test_modalidad_del_pliego(self):
        from motor.evaluacion.formato1_contenido import modalidad_de

        self.assertEqual(modalidad_de("DOCUMENTO BASE LICITACIÓN DE OBRA PÚBLICA DE INFRAESTRUCTURA SOCIAL"), "licitacion_social")
        self.assertEqual(modalidad_de("LICITACIÓN DE INFRAESTRUCTURA DE TRANSPORTE (VERSIÓN 4)"), "licitacion_transporte")
        self.assertEqual(modalidad_de("SELECCIÓN ABREVIADA DE MENOR CUANTÍA"), "menor_cuantia")
        self.assertEqual(modalidad_de("SELECCIÓN ABREVIADA DE MENOR CUANTÍA", "…DE INFRAESTRUCTURA SOCIAL…"), "menor_cuantia_social")
        self.assertEqual(modalidad_de("Código CCE-EICP-GI-11 INTERVENTORÍA DE OBRA PÚBLICA"), "interventoria_transporte")
        self.assertEqual(modalidad_de("INFRAESTRUCTURA DE AGUA, SANEAMIENTO BÁSICO … MEDIANTE LICITACIÓN PÚBLICA"), "licitacion_agua")
        self.assertIsNone(modalidad_de("CONCURSO DE MÉRITOS"))

    def test_numeral_cambiado(self):
        from motor.evaluacion.formato1_contenido import clausulas_faltantes

        def falta_revision(texto):
            return any("RESPONSABILIDAD DE SU REVISI" in c.upper() for c in clausulas_faltantes(texto, "licitacion_transporte"))

        self.assertFalse(falta_revision(self.NUMERAL_5))
        self.assertTrue(falta_revision(self.NUMERAL_5_CAMBIADO))

    def test_composicion_accionaria_vacia(self):
        from motor.evaluacion.formato1_contenido import composicion_accionaria_vacia

        encabezado = ("Composición de la persona jurídica: Porcentaje NIT, Cédula o Nombre o participación Documento de "
                      "Razón social Identificación del Accionista\n")
        self.assertTrue(composicion_accionaria_vacia(encabezado + "21. La oferta contiene información reservada"))
        self.assertFalse(composicion_accionaria_vacia(encabezado + "51% CC 1.088.245.241 Marcela Ruiz\n21. La oferta"))
        self.assertFalse(composicion_accionaria_vacia("Carta sin cuadro de composición"))

    def test_aval_del_mismo_representante(self):
        from motor.evaluacion.formato1_contenido import avalista_del_parrafo, representante_de_la_carta

        carta = ("Nombre del representante legal: Andrés Felipe García Ávila\nC. C. No. 1’069.725.868\n"
                 "“De acuerdo con lo expresado en la Ley 842 de 2003 y debido a que el suscriptor de la presente propuesta "
                 "no es ingeniero matriculado, yo Andrés Felipe García Ávila ingeniero con matrícula profesional No. "
                 "25202-251630 CND, avalo la presente propuesta”.")
        self.assertEqual(avalista_del_parrafo(carta), "ANDRES FELIPE GARCIA AVILA")
        self.assertEqual(representante_de_la_carta(carta), "ANDRES FELIPE GARCIA AVILA")


class RevisorFiscalSeguridadSocialTests(TestCase):
    """Si la sociedad tiene revisor fiscal, el formato de seguridad social
    debe venir certificado por él (caso real rechazado por el abogado)."""

    def test_revisor_designado_en_el_certificado(self):
        from motor.evaluacion.camara_comercio import REVISOR_FISCAL_DESIGNADO_RE

        m = REVISOR_FISCAL_DESIGNADO_RE.search(
            "SE DESIGNO A: CARGO NOMBRE IDENTIFICACION REVISOR FISCAL ANGELICA MARIA BARRIOS C.C. NO. 1140873167")
        self.assertEqual(m.group(1), "ANGELICA MARIA BARRIOS")
        self.assertIsNone(REVISOR_FISCAL_DESIGNADO_RE.search(
            "O NOMBRAMIENTO DE REPRESENTANTES LEGALES, ADMINISTRADORES O REVISORES FISCALES, QUE MODIFIQUEN"))

    def test_falta_la_firma_del_revisor(self):
        from motor.evaluacion.seguridad_social import _falta_revisor_fiscal

        revisores = {"901332677": "ANGELICA MARIA BARRIOS"}
        firmado_por_rl = "YO JUAN PEREZ EN CALIDAD DE REPRESENTANTE LEGAL DE EFR SAS NIT 901.332.677-1 CERTIFICO"
        self.assertEqual(_falta_revisor_fiscal(firmado_por_rl, "901332677", revisores), "ANGELICA MARIA BARRIOS")
        self.assertIsNone(_falta_revisor_fiscal(firmado_por_rl + " ANGELICA MARIA BARRIOS REVISOR FISCAL", "901332677", revisores))
        self.assertIsNone(_falta_revisor_fiscal(firmado_por_rl, "900111222", revisores))  # otra sociedad, sin revisor
        self.assertIsNone(_falta_revisor_fiscal(firmado_por_rl, "901332677", {}))


class VigenciaAntecedentesTests(TestCase):
    """REDAM vencido al cierre (caso real rechazado por el abogado) y
    antigüedad máxima de los demás certificados de antecedentes."""

    def test_redam(self):
        from datetime import date

        from motor.evaluacion.antecedentes import CONFIG_REDAM, problema_de_vigencia

        texto = "SE EXPIDE EN BOGOTA EL 28/04/2026 10:53 AM CODIGO VERIFICACION: Z81TN3L4PD VALIDA HASTA: 27/07/2026"
        self.assertIn("venció el 27/07/2026", problema_de_vigencia(CONFIG_REDAM, texto, date(2026, 8, 3)))
        self.assertIsNone(problema_de_vigencia(CONFIG_REDAM, texto, date(2026, 7, 20)))
        # Sin "válida hasta": tres meses desde la expedición.
        self.assertIsNone(problema_de_vigencia(CONFIG_REDAM, "SE EXPIDE EN BOGOTA EL 17/07/2026", date(2026, 8, 3)))
        self.assertIsNotNone(problema_de_vigencia(CONFIG_REDAM, "SIN FECHAS", date(2026, 8, 3)))

    def test_demas_certificados(self):
        from datetime import date

        from motor import criterios
        from motor.evaluacion.antecedentes import CONFIG_PROCURADURIA, problema_de_vigencia

        reciente = "BOGOTA DC, 31 DE JULIO DEL 2026 LA PROCURADURIA GENERAL DE LA NACION CERTIFICA"
        viejo = "BOGOTA DC, 15 DE MAYO DEL 2026 LA PROCURADURIA GENERAL DE LA NACION CERTIFICA"
        self.assertIsNone(problema_de_vigencia(CONFIG_PROCURADURIA, reciente, date(2026, 8, 3)))
        self.assertIn("más de 1 mes", problema_de_vigencia(CONFIG_PROCURADURIA, viejo, date(2026, 8, 3)))
        with criterios.usar({"antecedentes_meses": 0}):
            self.assertIsNone(problema_de_vigencia(CONFIG_PROCURADURIA, viejo, date(2026, 8, 3)))


class ObjetoSocialTests(TestCase):
    def test_objeto_social(self):
        from motor.evaluacion.camara_comercio import PROPORCION_MINIMA_OBJETO_SOCIAL, _proporcion_objeto_relacionado_laxo as p

        base = "MANTENIMIENTO DE LA VÍA DESDE EL SECTOR LA PLAYA HACIA EL CASCO URBANO DEL MUNICIPIO DE SUESCA"
        self.assertGreaterEqual(p("CONSTRUCCION Y MANTENIMIENTO DE VIAS Y OBRAS CIVILES", base), PROPORCION_MINIMA_OBJETO_SOCIAL)
        self.assertEqual(p("PODRA REALIZAR CUALQUIER ACTIVIDAD COMERCIAL O CIVIL LICITA", base), 1.0)
        # Comparte "mantenimiento" pero no es del sector de obras.
        self.assertEqual(p("VENTA Y MANTENIMIENTO DE EQUIPOS DE COMPUTO", base), 0.0)


class RequisitosDelPliegoTests(TestCase):
    """Verificaciones que el pliego agrega y el motor revisa solo."""

    def test_duracion(self):
        from datetime import date

        from motor.evaluacion.camara_comercio import DURACION_INDEFINIDA_RE, DISUELTA_RE, _duracion_hasta

        self.assertTrue(DURACION_INDEFINIDA_RE.search("LA PERSONA JURIDICA NO SE ENCUENTRA DISUELTA Y SU DURACION ES INDEFINIDA"))
        self.assertIsNone(DISUELTA_RE.search("LA PERSONA JURIDICA NO SE ENCUENTRA DISUELTA Y SU DURACION ES INDEFINIDA"))
        self.assertTrue(DISUELTA_RE.search("LA SOCIEDAD SE ENCUENTRA DISUELTA Y EN ESTADO DE LIQUIDACION"))
        self.assertEqual(_duracion_hasta("NO SE ENCUENTRA DISUELTA Y SU DURACION ES HASTA EL 30 DE MAYO DE 2063."), date(2063, 5, 30))
        self.assertEqual(_duracion_hasta("QUE LA SOCIEDAD NO SE HALLA DISUELTA. DURACION HASTA EL 30 DE JULIO DE 2045"), date(2045, 7, 30))

    def test_numero_de_cedula_con_ocr(self):
        from motor.evaluacion.identidad import _numero_en

        self.assertTrue(_numero_en("1069725868", "CEDULA DE CIUDADANIA NUMERO 1.069.725.868 GARCIA AVILA"))
        self.assertTrue(_numero_en("1069725868", "NUMERO 1.069.725.863"))  # un dígito mal leído
        self.assertFalse(_numero_en("1069725868", "NUMERO 1.069.735.863"))
        self.assertFalse(_numero_en("1069725868", "NUMERO 79.446.297"))

    def test_sociedad_anonima_por_razon_social(self):
        from motor.evaluacion.camara_comercio import ABIERTA_O_CERRADA_RE, es_sociedad_anonima

        self.assertFalse(es_sociedad_anonima(
            "RAZON SOCIAL: KA S.A.S. NIT: 830141859 LA SOCIEDAD SE TRANSFORMO DE SOCIEDAD ANONIMA A SOCIEDAD POR ACCIONES SIMPLIFICADA"))
        self.assertTrue(es_sociedad_anonima("RAZON SOCIAL: MOVITIERRA CONSTRUCCIONES S.A. NIT: 800128984"))
        self.assertTrue(es_sociedad_anonima("RAZON SOCIAL: EMPRESAS PUBLICAS S.A. E.S.P. NIT: 890904996"))
        self.assertFalse(es_sociedad_anonima("RAZON SOCIAL: INGENIEROS ASOCIADOS LTDA NIT: 800000000"))
        self.assertTrue(ABIERTA_O_CERRADA_RE.search("LA EMPRESA MOVITIERRA CONSTRUCCIONES S.A. ES UNA SOCIEDAD ANONIMA CERRADA"))
        self.assertIsNone(ABIERTA_O_CERRADA_RE.search("LA CUENTA ABIERTA EN EL BANCO"))

    def test_limitacion_mipyme(self):
        from motor.pliego.analisis import _MIPYME_NO_RE, _MIPYME_SI_RE

        no = "LA ENTIDAD NO LIMITA EL PROCESO DE CONTRATACION A LAS MIPYME COLOMBIANAS POR NO HABERSE CUMPLIDO LAS CONDICIONES"
        self.assertTrue(_MIPYME_NO_RE.search(no))
        como_pedir = "LOS INTERESADOS MANIFESTARAN SU INTENCION DE LIMITAR LAS CONVOCATORIAS A MIPYME EN LA SECCION MENSAJES"
        self.assertIsNone(_MIPYME_SI_RE.search(como_pedir))
        self.assertTrue(_MIPYME_SI_RE.search("EL PRESENTE PROCESO SE LIMITA A MIPYME COLOMBIANAS"))


def _requisito_pliego(**campos):
    from motor.pliego.lector_ia import RequisitoPliego

    base = {"id": "x", "requisito": "Requisito de prueba", "cita": "cita de prueba del pliego", "seccion": "3.1 Prueba", "pagina": 5}
    return RequisitoPliego(**{**base, **campos})


class LectorPliegoIATests(TestCase):
    """La lectura del pliego con la IA local: nada que no esté en el pliego,
    sin repetidos, y cada requisito con su forma de verificarse."""

    def trozo(self, texto):
        from motor.pliego.lector_ia import Trozo

        return Trozo(seccion="3.4 Capacidad jurídica", titulo="Capacidad jurídica", texto=texto, pagina=12)

    def test_la_cita_debe_estar_en_el_pliego(self):
        from motor.pliego.lector_ia import _limpiar

        trozo = self.trozo("El proponente debe aportar el certificado de antecedentes fiscales expedido por la Contraloría.")
        bueno = {"requisito": "Certificado de antecedentes fiscales", "documento": "Certificado de la Contraloría",
                 "cita": "debe aportar el certificado de antecedentes fiscales expedido por la Contraloría",
                 "aplica_a": ["persona_juridica", "inventado"], "vigencia_dias": "30"}
        r = _limpiar(bueno, trozo, lambda cita, pagina: pagina)
        self.assertIsNotNone(r)
        self.assertEqual(r.aplica_a, ["persona_juridica"])
        self.assertEqual(r.vigencia_dias, 30)
        self.assertEqual(r.pagina, 12)
        inventado = {**bueno, "cita": "el proponente debe aportar la licencia ambiental vigente del proyecto"}
        self.assertIsNone(_limpiar(inventado, trozo, lambda cita, pagina: pagina))

    def test_se_descarta_lo_que_no_es_requisito_juridico(self):
        from motor.pliego.lector_ia import _limpiar

        texto = "El proponente debe acreditar una capacidad residual igual o superior a la del proceso."
        crudo = {"requisito": "Capacidad residual del proponente", "cita": "debe acreditar una capacidad residual igual o superior"}
        self.assertIsNone(_limpiar(crudo, self.trozo(texto), lambda cita, pagina: pagina))

    def test_une_repetidos(self):
        from motor.pliego.lector_ia import _sin_repetidos

        a = _requisito_pliego(id="a", requisito="Certificado de existencia y representación legal", condiciones=["expedido por la Cámara"])
        b = _requisito_pliego(id="b", requisito="Certificado de existencia y representación legal vigente", vigencia_dias=30,
                              condiciones=["no mayor a 30 días"])
        c = _requisito_pliego(id="c", requisito="Garantía de seriedad de la oferta")
        unicos = _sin_repetidos([a, b, c])
        self.assertEqual([u.id for u in unicos], ["a", "c"])
        self.assertEqual(unicos[0].vigencia_dias, 30)
        self.assertEqual(unicos[0].condiciones, ["expedido por la Cámara", "no mayor a 30 días"])

    def test_trozos_solo_juridicos(self):
        from motor.pliego.lector_ia import trozos
        from motor.pliego.lectura import Seccion

        secciones = [
            Seccion(numero="3.1", titulo="Capacidad jurídica", texto="x " * 3000, pagina=10),
            Seccion(numero="4.1", titulo="Experiencia", texto="y " * 500, pagina=20),
            Seccion(numero="1.2", titulo="Garantía de seriedad de la oferta", texto="z " * 200, pagina=4),
        ]
        ambitos = {"3.1": "juridica", "4.1": "tecnica", "1.2": "general"}
        lista = trozos(secciones, ambitos, largo=2000)
        self.assertEqual({t.seccion for t in lista}, {"3.1 Capacidad jurídica", "1.2 Garantía de seriedad de la oferta"})
        self.assertGreaterEqual(sum(t.seccion.startswith("3.1") for t in lista), 3)  # la sección larga se parte

    def test_si_la_ia_no_responde_no_queda_como_leido(self):
        import requests

        from motor.pliego import lector_ia
        from motor.pliego.lectura import Seccion

        secciones = [Seccion(numero="3.1", titulo="Capacidad jurídica", texto="El proponente debe aportar " * 50, pagina=10)]
        with mock.patch.object(lector_ia, "_preguntar", side_effect=requests.ConnectionError("apagado")):
            with self.assertRaises(RuntimeError):
                lector_ia.leer(secciones, {"3.1": "juridica"}, [])

    def test_lectura_completa_con_respuesta_de_la_ia(self):
        import json

        from motor.pliego import lector_ia
        from motor.pliego.lectura import Pagina, Seccion

        texto = ("Cada integrante debe presentar el certificado del Registro de Deudores Alimentarios Morosos REDAM "
                 "expedido dentro de los 30 días anteriores al cierre.")
        respuesta = json.dumps({"requisitos": [{
            "requisito": "Certificado del REDAM de cada integrante", "documento": "Certificado REDAM",
            "cita": "debe presentar el certificado del Registro de Deudores Alimentarios Morosos REDAM", "vigencia_dias": 30,
        }]})
        secciones = [Seccion(numero="3.2", titulo="Documentos jurídicos", texto=texto, pagina=8)]
        with mock.patch.object(lector_ia, "_preguntar", return_value=respuesta):
            requisitos = lector_ia.leer(secciones, {"3.2": "juridica"}, [Pagina(numero=8, texto=texto)])
        self.assertEqual(len(requisitos), 1)
        self.assertEqual(requisitos[0].verificacion, "juridica.redam")


class CatalogoPliegoTests(TestCase):
    """Cómo se verifica cada requisito que el pliego exige."""

    def test_verificacion_por_lo_que_exige(self):
        from motor.pliego.catalogo import verificacion_de

        casos = {
            "Certificado de inscripción en el Registro Único de Proponentes RUP": "juridica.rup",
            "Certificado de antecedentes disciplinarios de la Procuraduría": "juridica.procuraduria",
            "No estar incurso en causales de inhabilidad o incompatibilidad": "juridica.carta",
            "Fotocopia de la cédula de ciudadanía del representante legal": "juridica.identidad",
            "Término de duración de la sociedad": "juridica.duracion",
            "Garantía de seriedad de la oferta": "juridica.garantia",
        }
        for requisito, esperado in casos.items():
            with self.subTest(requisito=requisito):
                self.assertEqual(verificacion_de(_requisito_pliego(requisito=requisito)), esperado)
        self.assertIsNone(verificacion_de(_requisito_pliego(requisito="Poder otorgado al apoderado del consorcio")))
        self.assertIsNone(verificacion_de(_requisito_pliego(requisito="Certificado de tamaño empresarial MIPYME")))

    def test_parametros_que_fija_el_pliego(self):
        from motor.pliego.catalogo import parametros_de

        rup = _requisito_pliego(requisito="RUP", verificacion="juridica.rup", vigencia_dias=30)
        self.assertEqual(parametros_de(rup), {"camara_dias": 30})
        policia = _requisito_pliego(requisito="Antecedentes judiciales", verificacion="juridica.policia", vigencia_dias=90)
        self.assertEqual(parametros_de(policia), {"antecedentes_meses": 3})
        extranjero = _requisito_pliego(requisito="Certificado de existencia", verificacion="juridica.existencia",
                                       vigencia_dias=90, aplica_a=["extranjero"])
        self.assertEqual(parametros_de(extranjero), {})

    def test_config_desde_el_pliego(self):
        from motor import criterios
        from motor.pliego.catalogo import config_desde_pliego

        r = _requisito_pliego(requisito="Certificado de la ARL", titulo_documento=["CERTIFICADO DE AFILIACION ARL"],
                              vigencia_dias=45, condiciones=["que cubra a todo el personal"], aplica_a=["plural", "persona_juridica"],
                              cita="cada integrante del consorcio o unión temporal debe aportar el certificado de la ARL")
        config = config_desde_pliego(r)
        criterios.ConfigPersonalizado.model_validate(config)
        self.assertEqual(config["bloques"][0], {"tipo": "vigencia_maxima", "meses": 2})
        self.assertEqual(config["bloques"][1]["tipo"], "confirmar")
        self.assertEqual(config["aplica_a"], ["consorcio", "union_temporal"])
        self.assertIsNone(config_desde_pliego(_requisito_pliego(requisito="Manifestar la intención de participar")))

    def test_confirmar_nunca_aprueba_solo(self):
        from datetime import date

        from motor import criterios
        from motor.evaluacion import personalizado

        config = criterios.ConfigPersonalizado(
            frases_documento=["CERTIFICADO DE AFILIACION"], bloques=[{"tipo": "confirmar", "frases": ["que cubra a todo el personal"]}]
        )
        with mock.patch.object(personalizado, "extraer_texto", return_value="CERTIFICADO DE AFILIACION ARL"):
            cumple, motivo, archivo = personalizado.evaluar_config({"arl.pdf": b"x"}, config, date(2026, 5, 25), None, [], "ACME")
        self.assertFalse(cumple)
        self.assertIn("confirma lo que exige el pliego", motivo)
        self.assertEqual(archivo, "arl.pdf")


class ComparacionConLecturaIATests(TestCase):
    """Lo que la IA leyó del pliego frente a la plantilla de la entidad."""

    def comparar(self, requisitos):
        from motor import criterios
        from motor.pliego.analisis import Extraccion, comparar

        extraccion = Extraccion(secciones=[], exigencias=[], paginas=40, documento_tipo=None)
        return comparar(extraccion, criterios.definicion_sistema("juridica"), requisitos)

    def test_propone_lo_que_falta_y_lo_que_cambia(self):
        requisitos = [
            _requisito_pliego(id="a", requisito="RUP", verificacion="juridica.rup", vigencia_dias=15),
            _requisito_pliego(id="b", requisito="Término de duración de la sociedad", verificacion="juridica.duracion"),
            _requisito_pliego(id="c", requisito="Certificado de la ARL", documento="Certificado de afiliación a la ARL"),
            _requisito_pliego(id="d", requisito="Visita a la obra"),
            _requisito_pliego(id="e", requisito="Documentos apostillados", aplica_a=["extranjero"]),
        ]
        hallazgos = {h.id: h for h in self.comparar(requisitos)}
        self.assertEqual(hallazgos["ia_param_camara_dias"].valor_pliego, 15)
        self.assertEqual(hallazgos["ia_falta_juridica.duracion"].requisito_propuesto["verificacion"], "juridica.duracion")
        self.assertEqual(hallazgos["ia_c"].requisito_propuesto["verificacion"], "personalizado")
        self.assertNotIn("verificacion", hallazgos["ia_d"].requisito_propuesto)
        self.assertEqual(hallazgos["ia_extranjeros"].tipo, "aclaracion")
        self.assertFalse(any(h.tipo == "requisito_no_exigido" for h in hallazgos.values()))  # lectura corta: no se afirma

    def test_no_exigido_solo_con_lectura_suficiente(self):
        from motor.pliego.analisis import MINIMO_REQUISITOS_PARA_NO_EXIGIDOS

        requisitos = [_requisito_pliego(id=str(i), requisito=f"Requisito {i}", verificacion="juridica.existencia")
                      for i in range(MINIMO_REQUISITOS_PARA_NO_EXIGIDOS)]
        no_exigidos = {h.verificacion for h in self.comparar(requisitos) if h.tipo == "requisito_no_exigido"}
        self.assertIn("juridica.rup", no_exigidos)
        self.assertNotIn("juridica.carta", no_exigidos)  # la carta siempre se exige
        self.assertNotIn("juridica.existencia", no_exigidos)

    def test_aplicar_ajustes_quita_y_agrega(self):
        from evaluaciones.pliego import aplicar_ajustes
        from motor import criterios

        definicion = criterios.definicion_sistema("juridica")
        config = {"frases_documento": ["CERTIFICADO ARL"], "bloques": [{"tipo": "confirmar", "frases": ["todo el personal"]}]}
        ajustes = [
            {"decision": "aceptado", "hallazgo": {"tipo": "requisito_no_exigido", "verificacion": "juridica.rup"}},
            {"decision": "aceptado", "hallazgo": {"tipo": "requisito_nuevo", "requisito_propuesto": {
                "verificacion": "personalizado", "titulo": "Certificado de la ARL", "corto": "ARL", "config": config}}},
            {"decision": "rechazado", "hallazgo": {"tipo": "requisito_no_exigido", "verificacion": "juridica.redam"}},
        ]
        nueva = aplicar_ajustes(definicion, ajustes)
        verificaciones = [r.verificacion for r in nueva.requisitos]
        self.assertNotIn("juridica.rup", verificaciones)
        self.assertIn("juridica.redam", verificaciones)
        arl = next(r for r in nueva.requisitos if r.titulo == "Certificado de la ARL")
        self.assertEqual(arl.config.bloques[0].tipo, "confirmar")


class LecturaIAProcesoTests(BaseEvaluaciones):
    """La lectura con IA la hace el trabajador y guarda lo que encontró."""

    analisis = PliegoProcesoTests.analisis

    def test_trabajador_lee_y_guarda(self):
        from evaluaciones import pliego
        from evaluaciones.models import AnalisisPliego

        a = self.analisis()
        self.assertEqual(a.estado_ia, "pendiente")
        requisitos = [_requisito_pliego(id="b", requisito="Duración de la sociedad", verificacion="juridica.duracion")]
        with mock.patch("motor.pliego.lector_ia.disponible", return_value=True), \
             mock.patch("motor.pliego.lector_ia.leer", return_value=requisitos), \
             mock.patch("evaluaciones.pliego.lectura.leer_paginas", return_value=[]):
            self.assertEqual(pliego.atender_lecturas_pendientes(), 1)
        a = AnalisisPliego.objects.get(pk=a.pk)
        self.assertEqual((a.estado_ia, a.progreso_ia), ("listo", 100))
        self.assertEqual(pliego.mapa(a)[0]["estado"], "motor_nuevo")

    def test_sin_ia_queda_no_disponible(self):
        from evaluaciones import pliego
        from evaluaciones.models import AnalisisPliego

        a = self.analisis()
        with mock.patch("motor.pliego.lector_ia.disponible", return_value=False):
            pliego.atender_lecturas_pendientes()
        self.assertEqual(AnalisisPliego.objects.get(pk=a.pk).estado_ia, "no_disponible")


class FiltrosLecturaIATests(TestCase):
    """Lo que se aprendió de la prueba 4 (documento tipo de menor cuantía)."""

    def test_nacional_o_extranjero_es_para_todos(self):
        from motor.pliego.lector_ia import es_de_extranjeros

        todos = _requisito_pliego(requisito="Ser persona natural o jurídica nacional o extranjera domiciliada en Colombia",
                                  cita="personas naturales nacionales o extranjeras")
        self.assertFalse(es_de_extranjeros(todos))
        self.assertTrue(es_de_extranjeros(_requisito_pliego(requisito="Documentos públicos otorgados en el exterior apostillados")))

    def test_condicionales_no_se_exigen_a_todos(self):
        from motor import criterios
        from motor.pliego.analisis import Extraccion, comparar

        requisitos = [
            _requisito_pliego(id="a", requisito="El proponente debe acreditar que el apoderado que firma la oferta está facultado",
                              documento="Poder", titulo_documento=["PODER ESPECIAL AMPLIO Y SUFICIENTE"]),
            _requisito_pliego(id="b", requisito="La persona natural que reúna los requisitos para acceder a la pensión de vejez",
                              verificacion="juridica.seguridad_social"),
        ]
        extraccion = Extraccion(secciones=[], exigencias=[], paginas=40, documento_tipo=None)
        h = {x.id: x for x in comparar(extraccion, criterios.definicion_sistema("juridica"), requisitos)}
        self.assertNotIn("ia_a", h)
        self.assertEqual(h["ia_condicionales"].tipo, "aclaracion")

    def test_lo_que_el_pliego_menciona_no_se_propone_quitar(self):
        from motor import criterios
        from motor.pliego.analisis import MINIMO_REQUISITOS_PARA_NO_EXIGIDOS, Extraccion, comparar
        from motor.pliego.catalogo import temas_en

        temas = temas_en("F. La Entidad debe consultar los antecedentes judiciales en línea y el certificado de antecedentes "
                         "disciplinarios y el Registro Nacional de Medidas Correctivas")
        self.assertTrue({"juridica.policia", "juridica.procuraduria", "juridica.rnmc"} <= set(temas))
        requisitos = [_requisito_pliego(id=str(i), requisito=f"Requisito {i}", verificacion="juridica.existencia")
                      for i in range(MINIMO_REQUISITOS_PARA_NO_EXIGIDOS)]
        extraccion = Extraccion(secciones=[], exigencias=[], paginas=40, documento_tipo=None, temas=temas)
        no_exigidos = {x.verificacion for x in comparar(extraccion, criterios.definicion_sistema("juridica"), requisitos)
                       if x.tipo == "requisito_no_exigido"}
        self.assertFalse({"juridica.policia", "juridica.procuraduria", "juridica.rnmc"} & no_exigidos)
        self.assertIn("juridica.rup", no_exigidos)

    def test_se_lee_la_seriedad_y_lo_juridico_de_paso(self):
        from motor.pliego.lector_ia import _se_lee
        from motor.pliego.lectura import Seccion

        self.assertTrue(_se_lee(Seccion(numero="7.1", titulo="GARANTÍA DE SERIEDAD DE LA OFERTA", pagina=83, texto="x"), "garantias"))
        self.assertFalse(_se_lee(Seccion(numero="7.2", titulo="GARANTÍA DE CUMPLIMIENTO", pagina=84, texto="x"), "garantias"))
        rup = "D. Los Proponentes obligados a estar inscritos en el Registro Único de Proponentes (RUP), deben aportar certificado"
        self.assertTrue(_se_lee(Seccion(numero="3.1", titulo="GENERALIDADES", pagina=25, texto=rup), "puntaje"))
        self.assertFalse(_se_lee(Seccion(numero="3.5", titulo="EXPERIENCIA", pagina=32, texto=rup), "tecnica"))


class AQuienSeExigeTests(TestCase):
    def test_solo_se_restringe_si_el_pliego_lo_dice(self):
        from motor.pliego.catalogo import aplica_a_de

        general = _requisito_pliego(requisito="El plazo ofrecido no debe ser superior al de la Entidad", aplica_a=["persona_juridica"])
        self.assertEqual(aplica_a_de(general), [])
        natural = _requisito_pliego(requisito="El proponente persona natural debe acreditar la afiliación", aplica_a=["persona_natural"])
        self.assertEqual(aplica_a_de(natural), ["persona_natural"])
        ambos = _requisito_pliego(requisito="Ser persona natural o jurídica", aplica_a=["persona_natural", "persona_juridica"])
        self.assertEqual(aplica_a_de(ambos), [])

    def test_mipyme_en_titulos_vecinos_no_tapa_la_carta(self):
        from motor.pliego.catalogo import verificacion_de

        carta = _requisito_pliego(requisito="Presentar la carta de presentación de la oferta", documento="Carta de presentación de la oferta",
                                  titulo_documento=["Carta de presentación de la oferta", "Acreditación de Mipyme"])
        self.assertEqual(verificacion_de(carta), "juridica.carta")

    def test_extranjero_junto_a_otros_no_es_exclusivo(self):
        from motor.pliego.lector_ia import es_de_extranjeros

        self.assertFalse(es_de_extranjeros(_requisito_pliego(requisito="Tener capacidad jurídica",
                                                            aplica_a=["persona_natural", "persona_juridica", "extranjero"])))
        self.assertTrue(es_de_extranjeros(_requisito_pliego(requisito="Tener capacidad jurídica", aplica_a=["extranjero"])))


class SinFalsosPositivosDelPliegoTests(TestCase):
    """Nada que la IA leyó del pliego puede aprobarse solo."""

    def test_el_requisito_armado_desde_el_pliego_siempre_lo_confirma_una_persona(self):
        from datetime import date

        from motor import criterios
        from motor.evaluacion import personalizado
        from motor.pliego.catalogo import config_desde_pliego

        req = _requisito_pliego(requisito="Presentar la autorización de tratamiento de datos personales",
                                documento="Autorización de tratamiento de datos personales",
                                titulo_documento=["AUTORIZACION DE TRATAMIENTO DE DATOS PERSONALES"])
        config = criterios.ConfigPersonalizado.model_validate(config_desde_pliego(req))
        self.assertTrue(any(b.tipo == "confirmar" for b in config.bloques))
        with mock.patch.object(personalizado, "extraer_texto", return_value="AUTORIZACION DE TRATAMIENTO DE DATOS PERSONALES firmada"):
            cumple, motivo, _ = personalizado.evaluar_config({"x.pdf": b"y"}, config, date(2026, 5, 25), None, [], "ACME")
        self.assertFalse(cumple)
        self.assertIn("confirma", motivo)


class DocumentoInventadoPorLaIATests(TestCase):
    """La IA a veces nombra un documento que no corresponde a lo exigido."""

    def requisito(self):
        return _requisito_pliego(
            requisito="El proponente debe presentar una garantía de seriedad de la oferta que cumpla el Decreto 1082 de 2015",
            documento="Certificado de existencia y representación legal",
            titulo_documento=["CERTIFICADO DE EXISTENCIA Y REPRESENTACION LEGAL"],
        )

    def test_manda_lo_que_se_exige_no_el_documento(self):
        from motor.pliego.catalogo import verificacion_de

        self.assertEqual(verificacion_de(self.requisito()), "juridica.garantia")

    def test_no_se_busca_un_documento_que_no_corresponde(self):
        from motor.pliego.catalogo import config_desde_pliego

        incoherente = _requisito_pliego(requisito="Presentar la autorización de tratamiento de datos personales",
                                        documento="Carta de presentación de la oferta",
                                        titulo_documento=["CARTA DE PRESENTACION DE LA OFERTA"])
        self.assertIsNone(config_desde_pliego(incoherente))


class IntegrantesJuridicosSinCertificadoTests(TestCase):
    """A todo integrante jurídico del consorcio se le exigen antecedentes,
    haya aportado o no su certificado de existencia."""

    def test_se_agrega_el_integrante_que_no_aporto_certificado(self):
        from motor.evaluacion.antecedentes import _con_integrantes_juridicos
        from motor.evaluacion.camara_comercio import Empresa
        from motor.evaluacion.proponente_plural import Integrante

        con_certificado = [Empresa("CONSTRUCTORA ALFA S.A.S.", "900123456")]
        integrantes = [
            Integrante(nombre="CONSTRUCTORA ALFA SAS", identificacion="900.123.456-1", persona_natural=False),
            Integrante(nombre="INGENIERIA BETA LTDA", identificacion="830.987.654-3", persona_natural=False),
            Integrante(nombre="PEDRO PEREZ", identificacion="79446297", persona_natural=True),
        ]
        empresas = _con_integrantes_juridicos(con_certificado, integrantes)
        self.assertEqual([e.nombre for e in empresas], ["CONSTRUCTORA ALFA S.A.S.", "INGENIERIA BETA LTDA"])
        self.assertEqual(empresas[1].nit, "830987654")

    def test_el_certificado_sin_razon_social_toma_el_nombre_del_integrante(self):
        """Un certificado de existencia que no dice la razón social no puede
        hacer que la misma empresa aparezca dos veces (una con su nombre y
        otra como "la empresa con NIT…"), con una falsa alarma."""
        from motor.evaluacion.antecedentes import _con_integrantes_juridicos
        from motor.evaluacion.camara_comercio import Empresa
        from motor.evaluacion.proponente_plural import Integrante

        empresas = [Empresa(None, "900574741"), Empresa("9D SOLUCIONES INTEGRALES S.A.S", "901799409")]
        integrantes = [
            Integrante(nombre="MSING S.A.S", identificacion=None, persona_natural=False),
            Integrante(nombre="9D SOLUCIONES INTEGRALES SAS", identificacion=None, persona_natural=False),
        ]
        resultado = _con_integrantes_juridicos(empresas, integrantes)
        self.assertEqual([(e.nombre, e.nit) for e in resultado],
                         [("MSING S.A.S", "900574741"), ("9D SOLUCIONES INTEGRALES S.A.S", "901799409")])

    def test_sin_nit_queda_como_falta_y_no_aprueba(self):
        from datetime import date

        from motor.evaluacion.antecedentes import _configs, evaluar_antecedente
        from motor.evaluacion.camara_comercio import Empresa

        config = next(c for c in _configs() if c.requisito == 14)
        resultado = evaluar_antecedente(
            {}, config, personas=[], empresas=[Empresa("INGENIERIA BETA LTDA", "")], plural=True, fecha_cierre=date(2026, 5, 25)
        )
        self.assertFalse(resultado.cumple)
        self.assertEqual([p.estado for p in resultado.personas], ["falta"])
        self.assertNotIn("NIT )", resultado.motivo or "")


class ConsultaEnLineaTests(TestCase):
    """Certificados que el programa trae solo de la página oficial."""

    def test_rnmc_de_persona_exige_la_fecha_de_expedicion(self):
        from motor.consultas.linea import ConsultaError, consultar_rnmc

        with self.assertRaises(ConsultaError) as caso:
            consultar_rnmc("79446297", tipo="cedula")
        self.assertIn("fecha de expedición", str(caso.exception))

    def test_tipo_de_documento_no_soportado(self):
        from motor.consultas.linea import ConsultaError, consultar_rnmc

        with self.assertRaises(ConsultaError):
            consultar_rnmc("79446297", tipo="pasaporte_extranjero")

    def test_copnia_necesita_matricula_o_cedula(self):
        from motor.consultas.linea import ConsultaError, consultar_copnia

        with self.assertRaises(ConsultaError):
            consultar_copnia("  ")

    def test_solo_se_consultan_las_fuentes_sin_captcha(self):
        from api.historico import FUENTES_EN_LINEA

        self.assertEqual(set(FUENTES_EN_LINEA), {"juridica.rnmc", "juridica.copnia_antecedentes", "juridica.aval_ingeniero"})
        for con_captcha in ("juridica.procuraduria", "juridica.contraloria", "juridica.policia", "juridica.redam"):
            self.assertNotIn(con_captcha, FUENTES_EN_LINEA)


class FechaExpedicionCedulaTests(TestCase):
    """La fecha de expedición se lee del reverso de la cédula (la pide el RNMC)."""

    def test_reversos_reales(self):
        from datetime import date

        from motor.evaluacion.identidad import fecha_expedicion_en

        casos = {
            "LUGAR DE NACIMIENTO: E I 1.69 O+ M ESTATURA GS.RH SEXO N 22-MAY-2002 PEREIRA FECHA Y LUGAR DE EXPEDICION": date(2002, 5, 22),
            "ALVEAR APELLIDOS LUGAR DE NACIMIENTO 1.71 O+ ESTATURA G.S. AH 05-JUN-2001 CARTAGENA FECHA Y LUGAR DE EXPEDICION P-05": date(2001, 6, 5),
            "LUGAR DE NACIMIENTO 1.60 B+ ESTATURA G.S.RH SEXO 30-ENE-1989 BOGOTA D.C. FECHA Y LUGAR DE EXPEDICIONF.20 447": date(1989, 1, 30),
            "SANTA MARTA (MAGDALENA) LUGAR DE NACIMIENTO ESTATURA G.S. RH * 09-MAR-1992 SANTA MARTA FECHA Y LUGAR DE EXPEDICION A IW": date(1992, 3, 9),
        }
        for texto, esperada in casos.items():
            with self.subTest(texto=texto[:40]):
                self.assertEqual(fecha_expedicion_en(texto), esperada)

    def test_no_confunde_la_fecha_de_nacimiento(self):
        from motor.evaluacion.identidad import fecha_expedicion_en

        solo_nacimiento = "FECHA DE NACIMIENTO 19-MAY-1984 PEREIRA (RISARALDA) LUGAR DE NACIMIENTO 1.69 ESTATURA"
        self.assertIsNone(fecha_expedicion_en(solo_nacimiento))

    def test_no_inventa_cuando_hay_dos_fechas_distintas(self):
        from motor.evaluacion.identidad import fecha_expedicion_en

        confuso = "05-JUN-2001 CARTAGENA FECHA Y LUGAR DE EXPEDICION ... 30-ENE-1989 BOGOTA EXPEDICION"
        self.assertIsNone(fecha_expedicion_en(confuso))

    def test_descarta_fechas_imposibles(self):
        from motor.evaluacion.identidad import fecha_expedicion_en

        self.assertIsNone(fecha_expedicion_en("30-FEB-1999 BOGOTA FECHA Y LUGAR DE EXPEDICION"))
        self.assertIsNone(fecha_expedicion_en("05-JUN-2045 BOGOTA FECHA Y LUGAR DE EXPEDICION"))


class FechaSoloDeSuCedulaTests(TestCase):
    """Una carpeta puede traer las cédulas de varias personas: la fecha de
    expedición de otro no sirve para consultar nada."""

    def test_no_toma_la_fecha_de_la_cedula_de_otra_persona(self):
        from unittest.mock import patch

        from motor.evaluacion import identidad

        pagina = ("REPUBLICA DE COLOMBIA CEDULA DE CIUDADANIA NUMERO 43.001.767 GARCIA BETANCUR MARTA EUGENIA "
                  "01-MAR-1979 MEDELLIN FECHA Y LUGAR DE EXPEDICION")
        with patch.object(identidad, "cedula_de", return_value="doc.pdf"), \
             patch.object(identidad, "paginas_cedula", return_value=[("doc.pdf", pagina)]), \
             patch.object(identidad, "texto_ocr_reforzado", return_value=""), \
             patch.object(identidad, "lecturas_de_paginas", side_effect=lambda *a, **k: ({1: pagina}, {})):
            propia = identidad.fecha_expedicion_cedula({"doc.pdf": b"x"}, "MARTA EUGENIA GARCIA BETANCUR", "43001767")
            ajena = identidad.fecha_expedicion_cedula({"doc.pdf": b"x"}, "ADRIANA MARCELA ROJAS PRIETO", "52371321")
        self.assertEqual(propia.isoformat(), "1979-03-01")
        self.assertIsNone(ajena)


class CertificadoMasRecienteTests(TestCase):
    """Si el evaluador aportó un certificado nuevo, ese manda sobre el de la
    oferta: si no, adjuntarlo no serviría de nada."""

    def test_gana_el_expedido_mas_tarde(self):
        from motor.evaluacion.antecedentes import _mas_reciente

        viejo = ("oferta/contraloria.pdf", "SE EXPIDE EL 26 DE MAYO DE 2026")
        nuevo = ("aportados/Req 14 - RNMC.pdf", "SE EXPIDE EL 20 DE SEPTIEMBRE DE 2026")
        self.assertEqual(_mas_reciente([viejo, nuevo])[0], nuevo[0])
        self.assertEqual(_mas_reciente([nuevo, viejo])[0], nuevo[0])
        self.assertEqual(_mas_reciente([viejo])[0], viejo[0])
        self.assertIsNone(_mas_reciente([]))

    def test_los_aportados_entran_como_documentos_del_proponente(self):
        from motor.esquemas.proceso import Proponente
        from motor.procesamiento.zip_utils import pdfs_con_aportados

        proponente = Proponente(
            numero_orden=1, hoja="P-01", nombre_proponente="ACME", nombre_archivo="1. ACME.zip", drive_file_id="x",
            documentos_aportados=[("Req 17 - RNMC 43001767.pdf", b"%PDF-nuevo")],
        )
        with mock.patch("motor.procesamiento.zip_utils.extraer_pdfs", return_value={"oferta/carta.pdf": b"%PDF-viejo"}):
            pdfs = pdfs_con_aportados(b"zip", proponente)
        self.assertEqual(sorted(pdfs), ["aportados/Req 17 - RNMC 43001767.pdf", "oferta/carta.pdf"])


class CertificadoAportadoSeVeTests(BaseHistorico):
    """El certificado que se aportó o se consultó en línea se abre desde la
    pantalla y entra a la evaluación como un documento más de la oferta."""

    def test_se_abre_por_su_ruta_de_aportado(self):
        r = self.aportar(self.abogado, self.p1, 14, nombre="contraloria.pdf")
        self.assertEqual(r.status_code, 201, r.content)
        ruta = f"aportados/Req 14 - {r.json()['nombre_original']}"
        respuesta = self.abogado.get(
            f"/api/evaluaciones/{self.ev['id']}/proponentes/{self.p1}/documento?archivo={quote(ruta)}"
        )
        self.assertEqual(respuesta.status_code, 200, respuesta.content[:200])
        self.assertEqual(respuesta["Content-Type"], "application/pdf")
        self.assertTrue(bytes(respuesta.content).startswith(b"%PDF"))

    def test_aportar_vuelve_a_poner_al_proponente_en_la_fila(self):
        from evaluaciones.models import Trabajo

        Trabajo.objects.filter(evaluacion_id=self.ev["id"]).delete()
        self.aportar(self.abogado, self.p1, 14, nombre="contraloria.pdf")
        self.assertTrue(Trabajo.objects.filter(evaluacion_id=self.ev["id"], proponente_id=self.p1).exists())

    def test_el_motor_recibe_el_certificado_aportado(self):
        from evaluaciones import servicios
        from evaluaciones.models import Evaluacion, Proponente

        self.aportar(self.abogado, self.p1, 14, nombre="contraloria.pdf")
        evaluacion = Evaluacion.objects.get(pk=self.ev["id"])
        proponente = servicios.proponente_motor(Proponente.objects.get(pk=self.p1), evaluacion)
        self.assertEqual([n for n, _ in proponente.documentos_aportados], ["Req 14 - contraloria.pdf"])
        self.assertTrue(proponente.documentos_aportados[0][1].startswith(b"%PDF"))


class ErroresDeConsultaClarosTests(TestCase):
    """Cuando de verdad no se puede, hay que decir por qué."""

    def test_el_aviso_de_la_pagina_llega_tal_cual(self):
        from motor.consultas import linea

        pagina = ("Policía Nacional de Colombia × Error La fecha de expedición de la Cedula de Ciudadania no es correcta, "
                  "por favor verifique. Aceptar Sistema Registro Nacional de Medidas Correctivas RNMC")
        aviso = linea._AVISO_RNMC_RE.search(linea._norm(pagina))
        self.assertIsNotNone(aviso)
        self.assertIn("NO ES CORRECTA", aviso.group(1))

    def test_el_titulo_de_la_pagina_no_se_confunde_con_un_resultado(self):
        from motor.consultas import linea

        # "Medidas Correctivas" está en el título: solo la frase completa del
        # resultado significa que la persona está limpia.
        titulo = linea._norm("Sistema Registro Nacional de Medidas Correctivas RNMC")
        self.assertNotIn("NO TIENE MEDIDAS CORRECTIVAS PENDIENTES", titulo)


class CedulaConIATests(TestCase):
    """Cuando el OCR no puede, la IA local mira la imagen de la cédula. Nunca
    se acepta una fecha que no sea de esa persona ni una imposible."""

    def leer(self, respuesta, cedula="52371321"):
        """Las dos miradas (recorte y página) leen lo mismo."""
        from motor.llm import vision

        with mock.patch.object(vision, "disponible", return_value=True), \
             mock.patch.object(vision, "_imagenes", return_value=[b"recorte", b"pagina"]), \
             mock.patch.object(vision, "_preguntar", return_value=respuesta):
            return vision.leer_cedula(b"%PDF", cedula)

    def test_lee_la_fecha_de_expedicion(self):
        from datetime import date

        fecha = self.leer({"numero": "52.371.321", "fecha_expedicion": "2002-05-22", "fecha_nacimiento": "1984-05-19"})
        self.assertEqual(fecha, date(2002, 5, 22))

    def test_no_acepta_la_cedula_de_otra_persona(self):
        self.assertIsNone(self.leer({"numero": "43001767", "fecha_expedicion": "1979-03-01"}))

    def test_no_acepta_fecha_sin_saber_de_quien_es(self):
        self.assertIsNone(self.leer({"numero": None, "fecha_expedicion": "2002-05-22"}))

    def test_descarta_fechas_imposibles(self):
        self.assertIsNone(self.leer({"numero": "52371321", "fecha_expedicion": "2045-01-01"}))
        self.assertIsNone(self.leer({"numero": "52371321", "fecha_expedicion": "1930-01-01"}))

    def test_descarta_la_expedicion_anterior_al_nacimiento(self):
        self.assertIsNone(self.leer({"numero": "52371321", "fecha_expedicion": "1980-01-01", "fecha_nacimiento": "1984-05-19"}))

    def test_sin_modelo_de_vision_no_falla(self):
        from motor.llm import vision

        with mock.patch.object(vision, "disponible", return_value=False):
            self.assertIsNone(vision.leer_cedula(b"%PDF", "52371321"))


class OcrDudosoVaALaIATests(TestCase):
    """Una fecha mal leída gasta una consulta a una página del Estado para
    nada: si el OCR no quedó limpio, decide la IA."""

    def test_confianza_segun_la_etiqueta(self):
        from datetime import date

        from motor.evaluacion.identidad import lectura_ocr

        limpio = lectura_ocr("LUGAR DE NACIMIENTO 1.69 ESTATURA 22-MAY-2002 PEREIRA FECHA Y LUGAR DE EXPEDICION")
        self.assertEqual((limpio.fecha, limpio.confiable, limpio.fuente), (date(2002, 5, 22), True, "ocr"))
        roto = lectura_ocr("ESTATURA 10-JUL-2015 FECHA Y LUEGAR DE EXPEDICION")
        self.assertEqual(roto.fecha, date(2015, 7, 10))
        self.assertFalse(roto.confiable)

    def test_la_ia_resuelve_lo_que_el_ocr_dejo_dudoso(self):
        from datetime import date

        from motor.evaluacion import identidad

        pagina = "REPUBLICA DE COLOMBIA CEDULA NUMERO 52.371.321 ROJAS PRIETO 10-JUL-2015 FECHA Y LUEGAR DE EXPEDICION"
        with mock.patch.object(identidad, "cedula_de", return_value="doc.pdf"), \
             mock.patch.object(identidad, "paginas_cedula", return_value=[("doc.pdf", pagina)]), \
             mock.patch.object(identidad, "paginas_cedula_numeradas", return_value=[("doc.pdf", 1, pagina)]), \
             mock.patch.object(identidad, "texto_ocr_reforzado", return_value=""), \
             mock.patch.object(identidad, "lecturas_de_paginas", side_effect=lambda *a, **k: ({1: pagina}, {})), \
             mock.patch("motor.llm.vision.leer_cedula_detallado",
                        return_value=__import__("motor.llm.vision", fromlist=["LecturaVision"]).LecturaVision(date(2002, 5, 22), 2)) as ia:
            lectura = identidad.leer_fecha_expedicion({"doc.pdf": b"x"}, "ADRIANA MARCELA ROJAS PRIETO", "52371321")
        self.assertTrue(ia.called)
        self.assertEqual((lectura.fecha, lectura.confiable, lectura.fuente), (date(2002, 5, 22), True, "ia"))

    def test_sin_ia_la_fecha_dudosa_no_se_da_por_buena(self):
        from datetime import date

        from motor.evaluacion import identidad

        pagina = "CEDULA NUMERO 52.371.321 ROJAS PRIETO 10-JUL-2015 FECHA Y LUEGAR DE EXPEDICION"
        with mock.patch.object(identidad, "cedula_de", return_value="doc.pdf"), \
             mock.patch.object(identidad, "paginas_cedula", return_value=[("doc.pdf", pagina)]), \
             mock.patch.object(identidad, "paginas_cedula_numeradas", return_value=[("doc.pdf", 1, pagina)]), \
             mock.patch.object(identidad, "texto_ocr_reforzado", return_value=""), \
             mock.patch.object(identidad, "lecturas_de_paginas", side_effect=lambda *a, **k: ({1: pagina}, {})), \
             mock.patch("motor.llm.vision.leer_cedula_detallado",
                        return_value=__import__("motor.llm.vision", fromlist=["LecturaVision"]).LecturaVision(None)):
            lectura = identidad.leer_fecha_expedicion({"doc.pdf": b"x"}, "ADRIANA MARCELA ROJAS PRIETO", "52371321")
            guardable = identidad.fecha_expedicion_cedula({"doc.pdf": b"x"}, "ADRIANA MARCELA ROJAS PRIETO", "52371321")
        self.assertEqual(lectura.fecha, date(2015, 7, 10))
        self.assertFalse(lectura.confiable)
        self.assertIsNone(guardable)


class FechaDeExpedicionEnLaFichaTests(BaseHistorico):
    """El abogado la necesita a la vista: las páginas de consulta la piden."""

    def persona(self):
        return self.abogado.post(
            f"/api/evaluaciones/{self.ev['id']}/proponentes/{self.p1}/personas",
            {"rol": "representante_legal", "tipo": "natural", "nombre": "Pedro Pérez", "documento": "1020304"},
        ).json()

    def test_se_guarda_desde_la_ficha_de_la_persona(self):
        persona = self.persona()
        self.assertIsNone(persona["fecha_expedicion_documento"])
        r = self.abogado.http.patch(
            f"/api/evaluaciones/{self.ev['id']}/personas/{persona['id']}",
            {"fecha_expedicion_documento": "2002-05-22"},
            content_type="application/json",
            headers={"X-CSRFToken": self.abogado.csrf},
        )
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["fecha_expedicion_documento"], "2002-05-22")

    def test_no_acepta_una_fecha_futura(self):
        r = self.abogado.http.patch(
            f"/api/evaluaciones/{self.ev['id']}/personas/{self.persona()['id']}",
            {"fecha_expedicion_documento": "2099-01-01"},
            content_type="application/json",
            headers={"X-CSRFToken": self.abogado.csrf},
        )
        self.assertEqual(r.status_code, 400)


class LaIADebeVerDosVecesLoMismoTests(TestCase):
    """En una prueba real el modelo leyó «24» donde decía «13»: una sola
    lectura no puede dar por buena una fecha."""

    def lectura(self, respuestas):
        from motor.llm import vision

        with mock.patch.object(vision, "disponible", return_value=True), \
             mock.patch.object(vision, "_imagenes", return_value=[b"a", b"b"]), \
             mock.patch.object(vision, "_preguntar", side_effect=respuestas):
            return vision.leer_cedula_detallado(b"%PDF", "52371321")

    def test_dos_lecturas_iguales_se_dan_por_buenas(self):
        from datetime import date

        misma = {"numero": "52371321", "fecha_expedicion": "2002-05-22"}
        lectura = self.lectura([misma, dict(misma)])
        self.assertEqual((lectura.fecha, lectura.veces, lectura.confirmada), (date(2002, 5, 22), 2, True))

    def test_dos_lecturas_distintas_no_se_confirman(self):
        lectura = self.lectura([
            {"numero": "52371321", "fecha_expedicion": "1996-05-13"},
            {"numero": "52371321", "fecha_expedicion": "1996-05-24"},
        ])
        self.assertFalse(lectura.confirmada)

    def test_una_sola_lectura_se_sugiere_pero_no_se_usa_sola(self):
        from datetime import date

        from motor.llm import vision

        lectura = self.lectura([{"numero": "52371321", "fecha_expedicion": "2002-05-22"}, {}])
        self.assertEqual(lectura.fecha, date(2002, 5, 22))
        self.assertFalse(lectura.confirmada)
        with mock.patch.object(vision, "leer_cedula_detallado", return_value=lectura):
            self.assertIsNone(vision.leer_cedula(b"%PDF", "52371321"))

    def test_la_ia_confirma_lo_que_el_ocr_leyo_a_medias(self):
        from datetime import date

        from motor.evaluacion import identidad
        from motor.llm.vision import LecturaVision

        pagina = "CEDULA NUMERO 52.371.321 ROJAS PRIETO 22-MAY-2002 BOGOTA FECHA Y LUEGAR DE EXPEDICION"
        with mock.patch.object(identidad, "cedula_de", return_value="doc.pdf"), \
             mock.patch.object(identidad, "paginas_cedula", return_value=[("doc.pdf", pagina)]), \
             mock.patch.object(identidad, "paginas_cedula_numeradas", return_value=[("doc.pdf", 1, pagina)]), \
             mock.patch.object(identidad, "texto_ocr_reforzado", return_value=""), \
             mock.patch.object(identidad, "lecturas_de_paginas", side_effect=lambda *a, **k: ({1: pagina}, {})), \
             mock.patch("motor.llm.vision.leer_cedula_detallado", return_value=LecturaVision(date(2002, 5, 22), 1)):
            lectura = identidad.leer_fecha_expedicion({"doc.pdf": b"x"}, "ADRIANA MARCELA ROJAS PRIETO", "52371321")
        self.assertEqual(lectura.fecha, date(2002, 5, 22))
        self.assertTrue(lectura.confiable)  # una lectura de la IA + el OCR dicen lo mismo

    def test_si_la_ia_y_el_ocr_no_coinciden_decide_una_persona(self):
        from datetime import date

        from motor.evaluacion import identidad
        from motor.llm.vision import LecturaVision

        pagina = "CEDULA NUMERO 52.371.321 ROJAS PRIETO 13-MAY-1996 CUCUTA FECHA Y LUEGAR DE EXPEDICION"
        with mock.patch.object(identidad, "cedula_de", return_value="doc.pdf"), \
             mock.patch.object(identidad, "paginas_cedula", return_value=[("doc.pdf", pagina)]), \
             mock.patch.object(identidad, "paginas_cedula_numeradas", return_value=[("doc.pdf", 1, pagina)]), \
             mock.patch.object(identidad, "texto_ocr_reforzado", return_value=""), \
             mock.patch.object(identidad, "lecturas_de_paginas", side_effect=lambda *a, **k: ({1: pagina}, {})), \
             mock.patch("motor.llm.vision.leer_cedula_detallado", return_value=LecturaVision(date(1996, 5, 24), 1)):
            lectura = identidad.leer_fecha_expedicion({"doc.pdf": b"x"}, "ADRIANA MARCELA ROJAS PRIETO", "52371321")
        self.assertFalse(lectura.confiable)


class FechaSoloSeGuardaSiSirveTests(BaseHistorico):
    """Una fecha que la página rechazó no puede quedarse pegada a la persona:
    todas las consultas siguientes fallarían por el mismo dato malo."""

    def setUp(self):
        super().setUp()
        # Una pausa por falla de la página no puede pasar de una prueba a otra.
        from django.core.cache import cache

        cache.clear()

    def persona(self):
        return self.abogado.post(
            f"/api/evaluaciones/{self.ev['id']}/proponentes/{self.p1}/personas",
            {"rol": "representante_legal", "tipo": "natural", "nombre": "Pedro Pérez", "documento": "1020304",
             "fecha_expedicion_documento": "2005-03-01"},
        ).json()

    def consultar(self, persona_id, fecha=None):
        from evaluaciones.models import PlantillaEvaluacion  # noqa: F401  (asegura la plantilla cargada)

        datos = {"requisito": 17, "persona_id": persona_id}
        if fecha:
            datos["fecha_expedicion_documento"] = fecha
        return self.abogado.post(f"/api/evaluaciones/{self.ev['id']}/proponentes/{self.p1}/consultar", datos)

    def test_la_fecha_rechazada_por_la_pagina_se_borra(self):
        from motor.consultas.linea import FechaRechazada
        from evaluaciones.models import PersonaVerificada

        persona = self.persona()
        with mock.patch("motor.consultas.linea.consultar_rnmc",
                        side_effect=FechaRechazada("La página de la Policía responde: «La fecha de expedición de la cedula "
                                                  "de ciudadania no es correcta, por favor verifique.»")):
            r = self.consultar(persona["id"], "1999-09-09")
        self.assertEqual(r.status_code, 400)
        self.assertIsNone(PersonaVerificada.objects.get(pk=persona["id"]).fecha_expedicion_documento)


class FallasDeLaPaginaNoSonDeLosDatosTests(FechaSoloSeGuardaSiSirveTests):
    """Si la página de la Policía falla por su cuenta, no se toca la fecha
    guardada y no se le vuelve a insistir por un rato."""

    def fecha_guardada(self, persona_id):
        from evaluaciones.models import PersonaVerificada

        return PersonaVerificada.objects.get(pk=persona_id).fecha_expedicion_documento

    def test_un_error_de_la_pagina_no_borra_la_fecha_y_pausa_las_consultas(self):
        from motor.consultas.linea import PaginaNoDisponible

        persona = self.persona()
        falla = PaginaNoDisponible("La página de la Policía respondió con un error propio: «Servicio no disponible».")
        with mock.patch("motor.consultas.linea.consultar_rnmc", side_effect=falla) as rnmc:
            primera = self.consultar(persona["id"])
            segunda = self.consultar(persona["id"])
        self.assertEqual(primera.status_code, 503)
        self.assertIsNotNone(self.fecha_guardada(persona["id"]))  # la fecha buena sigue ahí
        self.assertEqual(segunda.status_code, 503)
        self.assertIn("no insistirle", segunda.json()["detail"])
        self.assertEqual(rnmc.call_count, 1)  # la segunda no llegó a la Policía

    def test_una_falla_inesperada_tampoco_borra_la_fecha(self):
        persona = self.persona()
        with mock.patch("motor.consultas.linea.consultar_rnmc", side_effect=TimeoutError("chromium")):
            r = self.consultar(persona["id"])
        self.assertEqual(r.status_code, 502)
        self.assertIsNotNone(self.fecha_guardada(persona["id"]))

    def test_el_aviso_de_la_pagina_se_clasifica(self):
        from motor.consultas import linea

        fecha_mala = linea._AVISO_RNMC_RE.search(linea._norm(
            "× Error La fecha de expedición de la Cedula de Ciudadania no es correcta, por favor verifique. Aceptar"))
        otro = linea._AVISO_RNMC_RE.search(linea._norm("× Error El servicio no está disponible en este momento. Aceptar"))
        self.assertIn("FECHA DE EXPEDICI", fecha_mala.group(1))
        self.assertNotIn("FECHA DE EXPEDICI", otro.group(1))


class NadaDeConsultasInnecesariasTests(FechaSoloSeGuardaSiSirveTests):
    """A la página de la Policía solo se le pregunta cuando de verdad hace falta."""

    def test_lo_ya_consultado_hoy_no_se_vuelve_a_consultar(self):
        from datetime import date

        from motor.consultas.linea import CertificadoEnLinea

        persona = self.persona()
        certificado = CertificadoEnLinea(
            fuente="rnmc", texto="NO TIENE MEDIDAS CORRECTIVAS PENDIENTES POR CUMPLIR", pdf=b"%PDF-1.4 rnmc",
            nombre_archivo="RNMC 1020304.pdf", fecha_expedicion=date.today(), sin_novedades=True,
        )
        with mock.patch("motor.consultas.linea.consultar_rnmc", return_value=certificado) as rnmc:
            primera = self.consultar(persona["id"])
            segunda = self.consultar(persona["id"])
        self.assertEqual(primera.status_code, 201, primera.content)
        self.assertEqual(segunda.status_code, 201)
        self.assertEqual(primera.json()["id"], segunda.json()["id"])
        self.assertEqual(rnmc.call_count, 1)  # la segunda vez no se llamó a la Policía

    def test_quien_ya_cumple_no_se_consulta(self):
        from evaluaciones import servicios
        from evaluaciones.models import Resultado

        persona = self.persona()
        from evaluaciones.models import Evaluacion

        evaluacion = Evaluacion.objects.get(pk=self.ev["id"])
        ya_cumple = [{"nombre": persona["nombre"], "documento": persona["documento"], "tipo": "natural",
                      "rol": "representante_legal", "estado": "cumple"}]
        Resultado.objects.update_or_create(
            evaluacion=evaluacion, proponente_id=self.p1, requisito=17,
            defaults={"entidad_id": evaluacion.entidad_id, "requiere_revision": False, "datos": {
                "hoja": "P-01", "numero_orden": 1, "nombre_proponente": "x", "requisito": 17, "cumple": True,
                "personas_antecedente": ya_cumple}},
        )
        self.assertTrue(servicios.clave_persona(persona["nombre"], persona["documento"]))
        with mock.patch("motor.consultas.linea.consultar_rnmc") as rnmc:
            r = self.consultar(persona["id"])
        self.assertEqual(r.status_code, 409)
        self.assertIn("no hace falta", r.json()["detail"])
        rnmc.assert_not_called()


class EliminarProcesoTests(BaseEvaluaciones):
    """Eliminar un proceso es irreversible: solo quien puede, con confirmación,
    nunca uno aprobado, y sin dejar nada suelto."""

    def eliminar(self, c, proceso_id, confirmacion):
        return c.http.delete(
            f"/api/evaluaciones/procesos/{proceso_id}",
            data={"confirmacion": confirmacion},
            content_type="application/json",
            headers={"X-CSRFToken": c.csrf},
        )

    def proceso_de(self, ev):
        from evaluaciones.models import Evaluacion

        return Evaluacion.objects.get(pk=ev["id"]).proceso

    def test_se_elimina_con_todo_lo_suyo(self):
        from evaluaciones.models import Evaluacion, Proceso, Proponente, Resultado

        jefe, ev = self.crear()
        self.evaluar_todo(jefe, ev["id"])
        proceso = self.proceso_de(ev)
        self.assertTrue(Resultado.objects.filter(evaluacion_id=ev["id"]).exists())
        admin = Cliente()
        admin.entrar("admin@entidad.gov.co")
        r = self.eliminar(admin, proceso.id, proceso.codigo)
        self.assertEqual(r.status_code, 204, r.content)
        self.assertFalse(Proceso.objects.filter(pk=proceso.id).exists())
        self.assertFalse(Evaluacion.objects.filter(pk=ev["id"]).exists())
        self.assertFalse(Proponente.objects.filter(proceso_id=proceso.id).exists())
        self.assertFalse(Resultado.objects.filter(evaluacion_id=ev["id"]).exists())

    def test_hay_que_escribir_el_codigo(self):
        jefe, ev = self.crear()
        proceso = self.proceso_de(ev)
        admin = Cliente()
        admin.entrar("admin@entidad.gov.co")
        r = self.eliminar(admin, proceso.id, "otro código")
        self.assertEqual(r.status_code, 400)
        self.assertIn(proceso.codigo, r.json()["detail"])

    def test_quien_no_lo_creo_ni_administra_no_puede(self):
        jefe, ev = self.crear()
        proceso = self.proceso_de(ev)
        otro = Cliente()
        otro.entrar("abogado2@entidad.gov.co")
        self.assertEqual(self.eliminar(otro, proceso.id, proceso.codigo).status_code, 403)
        ajeno = Cliente()
        ajeno.entrar("abogado@otraentidad.gov.co")
        self.assertIn(self.eliminar(ajeno, proceso.id, proceso.codigo).status_code, (403, 404))

    def test_quien_lo_creo_si_puede(self):
        jefe, ev = self.crear()
        proceso = self.proceso_de(ev)
        self.assertEqual(self.eliminar(jefe, proceso.id, proceso.codigo.lower()).status_code, 204)

    def test_un_proceso_con_evaluacion_aprobada_no_se_elimina(self):
        from evaluaciones.models import EstadoEvaluacion, Evaluacion

        jefe, ev = self.crear()
        Evaluacion.objects.filter(pk=ev["id"]).update(estado=EstadoEvaluacion.APROBADA)
        proceso = self.proceso_de(ev)
        admin = Cliente()
        admin.entrar("admin@entidad.gov.co")
        r = self.eliminar(admin, proceso.id, proceso.codigo)
        self.assertEqual(r.status_code, 409)
        self.assertIn("aprobadas", r.json()["detail"])
        listado = {p["id"]: p for p in admin.get("/api/evaluaciones/procesos").json()}
        self.assertFalse(listado[str(proceso.id)]["puede_eliminar"])

    def test_queda_en_la_auditoria(self):
        from cuentas.models import EventoAuditoria

        jefe, ev = self.crear()
        proceso = self.proceso_de(ev)
        codigo = proceso.codigo
        self.eliminar(jefe, proceso.id, codigo)
        self.assertTrue(EventoAuditoria.objects.filter(accion="proceso.eliminado", detalles__codigo=codigo).exists())


class EliminarProcesoConCertificadosTests(BaseHistorico):
    def test_se_eliminan_tambien_los_certificados_aportados_y_sus_archivos(self):
        from evaluaciones.models import DocumentoAportado, Evaluacion, PersonaVerificada, Proceso

        persona = self.abogado.post(
            f"/api/evaluaciones/{self.ev['id']}/proponentes/{self.p1}/personas",
            {"rol": "representante_legal", "tipo": "natural", "nombre": "Pedro Pérez", "documento": "1020304"},
        ).json()
        self.assertEqual(self.aportar(self.abogado, self.p1, 14, persona["id"]).status_code, 201)
        doc = DocumentoAportado.objects.get(evaluacion_id=self.ev["id"])
        archivo = doc.archivo.name
        self.assertTrue(doc.archivo.storage.exists(archivo))

        proceso = Evaluacion.objects.get(pk=self.ev["id"]).proceso
        admin = Cliente()
        admin.entrar("admin@entidad.gov.co")
        with self.captureOnCommitCallbacks(execute=True):
            r = admin.http.delete(
                f"/api/evaluaciones/procesos/{proceso.id}", data={"confirmacion": proceso.codigo},
                content_type="application/json", headers={"X-CSRFToken": admin.csrf},
            )
        self.assertEqual(r.status_code, 204, r.content)
        self.assertFalse(Proceso.objects.filter(pk=proceso.id).exists())
        self.assertFalse(DocumentoAportado.objects.filter(pk=doc.pk).exists())
        self.assertFalse(PersonaVerificada.objects.filter(pk=persona["id"]).exists())
        self.assertFalse(doc.archivo.storage.exists(archivo))  # el PDF tampoco queda en disco


class CedulaDeOtraPersonaTests(TestCase):
    """Caso real (P-06): la carpeta «DOC LEGAL/MEGB» traía la cédula de la
    suplente, y por decir «LEGAL» se le asignaba a la representante legal."""

    REVERSO_DE_MARTA = (
        "REPUBLICA DE COLOMBIA CEDULA DE CIUDADANIA FECHA DE NACIMIENTO 09-SEP-1960 MEDELLIN (ANTIOQUIA) LUGAR DE "
        "NACIMIENTO 1.60 A+ F ESTATURA G.S. RH SEXO 01-MAR-1979 MEDELLIN FECHA Y LUGAR DE EXPEDICION INDICE DERECHO "
        "REGISTRADOR NACIONAL A-0100100-00157053-F-0043001767-20090520 0011606230A"
    )

    def test_no_se_asigna_por_el_nombre_del_archivo(self):
        from motor.evaluacion import identidad

        with mock.patch.object(identidad, "paginas_cedula",
                               return_value=[("DOC LEGAL/MEGB/Doc Legal MEGB.pdf", self.REVERSO_DE_MARTA)]), \
             mock.patch.object(identidad, "texto_ocr_reforzado", return_value=""):
            adriana = identidad.cedula_de({"x": b""}, "ADRIANA MARCELA ROJAS PRIETO", "52371321", principal=True)
            marta = identidad.cedula_de({"x": b""}, "MARTA EUGENIA GARCIA BETANCUR", "43001767")
        self.assertIsNone(adriana)
        self.assertEqual(marta, "DOC LEGAL/MEGB/Doc Legal MEGB.pdf")

    def test_la_cedula_de_otro_se_reconoce_por_el_numero(self):
        from motor.evaluacion.identidad import _de_otra_persona

        self.assertTrue(_de_otra_persona("52371321", self.REVERSO_DE_MARTA))
        self.assertFalse(_de_otra_persona("43001767", self.REVERSO_DE_MARTA))
        self.assertFalse(_de_otra_persona("52371321", "REPUBLICA DE COLOMBIA FECHA Y LUGAR DE EXPEDICION"))  # sin números


class PliegoSeAutoadaptaTests(TestCase):
    """Lo que el pliego pide y el motor ya sabe revisar se revisa solo."""

    def test_revisor_fiscal_de_sociedad_anonima_va_al_motor(self):
        from motor.pliego.catalogo import verificacion_de

        req = _requisito_pliego(requisito="Certificación del revisor fiscal para sociedades anónimas colombianas")
        self.assertEqual(verificacion_de(req), "juridica.revisor_fiscal")  # el 18: N.A. para las S.A.S.

    def test_formatos_de_puntaje_e_implicitos_no_son_requisitos(self):
        from motor.pliego.lector_ia import _NO_ES_REQUISITO_RE, _norm

        for texto in ("Formato 6 – Vinculación de personas en condición de discapacidad",
                      "Autorización para el tratamiento de datos personales",
                      "Formato 7 – Puntaje de industria nacional"):
            with self.subTest(texto=texto):
                self.assertTrue(_NO_ES_REQUISITO_RE.search(_norm(texto)))

    def test_debe_ser_firmado_se_revisa_solo(self):
        from motor.pliego.catalogo import config_desde_pliego

        req = _requisito_pliego(requisito="Compromiso anticorrupción", titulo_documento=["COMPROMISO ANTICORRUPCION"],
                                condiciones=["Debe ser firmado por el representante legal"])
        tipos = [b["tipo"] for b in config_desde_pliego(req)["bloques"]]
        self.assertEqual(tipos, ["firmado", "menciona_representante"])

    def test_lo_que_no_sabe_revisar_sigue_para_una_persona(self):
        from motor.pliego.catalogo import config_desde_pliego

        req = _requisito_pliego(requisito="Certificado de la ARL", titulo_documento=["CERTIFICADO DE AFILIACION ARL"],
                                condiciones=["que cubra a todo el personal del contrato"])
        self.assertIn("confirmar", [b["tipo"] for b in config_desde_pliego(req)["bloques"]])

    def test_el_bloque_de_firma(self):
        from datetime import date

        from motor import criterios
        from motor.evaluacion import personalizado

        config = criterios.ConfigPersonalizado(frases_documento=["COMPROMISO ANTICORRUPCION"], bloques=[{"tipo": "firmado"}])
        with mock.patch.object(personalizado, "extraer_texto", return_value="COMPROMISO ANTICORRUPCION"):
            with mock.patch.object(personalizado, "esta_firmado", return_value=True):
                cumple, motivo, _ = personalizado.evaluar_config({"c.pdf": b"%PDF"}, config, date(2026, 5, 25), None, [], "ACME")
            self.assertTrue(cumple, motivo)
            with mock.patch.object(personalizado, "esta_firmado", return_value=False):
                cumple, motivo, _ = personalizado.evaluar_config({"c.pdf": b"%PDF"}, config, date(2026, 5, 25), None, [], "ACME")
        self.assertFalse(cumple)
        self.assertIn("firma", motivo)


class AuditoriaDeHoyTests(TestCase):
    """Riesgos encontrados en la auditoría."""

    def test_dos_palabras_en_comun_no_hacen_suya_una_cedula(self):
        from motor.evaluacion.identidad import _es_suyo

        cedula_de_marta = "REPUBLICA DE COLOMBIA CEDULA GARCIA BETANCUR MARTA EUGENIA FECHA Y LUGAR DE EXPEDICION"
        self.assertTrue(_es_suyo("MARTA EUGENIA GARCIA BETANCUR", "", cedula_de_marta))
        self.assertFalse(_es_suyo("EUGENIA GARCIA LOPEZ", "", cedula_de_marta))

    def test_palabras_de_archivo_no_son_iniciales(self):
        from motor.evaluacion.identidad import _nombre_en_archivo

        self.assertFalse(_nombre_en_archivo("DIANA ORTEGA CASTRO", "DOC LEGAL MEGB"))
        self.assertTrue(_nombre_en_archivo("MARTA EUGENIA GARCIA BETANCUR", "DOC LEGAL MEGB"))

    def test_varias_empresas_sin_razon_social_no_se_cruzan(self):
        from motor.evaluacion.antecedentes import _con_integrantes_juridicos
        from motor.evaluacion.camara_comercio import Empresa
        from motor.evaluacion.proponente_plural import Integrante

        empresas = [Empresa(None, "900000001"), Empresa(None, "900000002")]
        integrantes = [Integrante("ALFA SAS", None, False), Integrante("BETA SAS", None, False)]
        resultado = _con_integrantes_juridicos(empresas, integrantes)
        self.assertEqual([e.nit for e in resultado], ["900000001", "900000002"])
        self.assertTrue(all(e.razon_social is None for e in resultado))  # sin nombres cruzados ni duplicados


class RequisitosDelPliegoRevalidadosTests(TestCase):
    """Lo aceptado del pliego con reglas anteriores se revalida con las de hoy,
    sin mover los números de los demás requisitos."""

    def ajuste(self, titulo, verificacion="personalizado", cita="cita del pliego"):
        propuesto = {"titulo": titulo, "corto": titulo[:20], "verificacion": verificacion}
        if verificacion == "personalizado":
            propuesto["config"] = {"frases_documento": [titulo.upper()[:40]], "bloques": [{"tipo": "confirmar", "frases": ["x"]}]}
        return {"decision": "aceptado", "hallazgo": {
            "tipo": "requisito_nuevo", "titulo": titulo, "cita": cita, "seccion": "11.2 FORMATOS", "pagina": 96,
            "requisito_propuesto": propuesto}}

    def test_se_retiran_los_que_sobran_y_los_numeros_no_se_mueven(self):
        from evaluaciones.pliego import aplicar_ajustes
        from motor import criterios

        definicion = criterios.definicion_sistema("juridica")
        ajustes = [
            self.ajuste("Duración de la sociedad", "juridica.duracion"),
            self.ajuste("Certificación del revisor fiscal para sociedades anónimas colombianas"),
            self.ajuste("Formato para vinculación de personas en condición de discapacidad"),
            self.ajuste("Autorización para el tratamiento de datos personales"),
            self.ajuste("Compromiso anticorrupción firmado"),
        ]
        antes = {r.titulo: r.numero for r in aplicar_ajustes(definicion, ajustes, revalidar=False).requisitos}
        despues = {r.titulo: r.numero for r in aplicar_ajustes(definicion, ajustes).requisitos}
        self.assertNotIn("Formato para vinculación de personas en condición de discapacidad", despues)
        self.assertNotIn("Autorización para el tratamiento de datos personales", despues)
        self.assertNotIn("Certificación del revisor fiscal para sociedades anónimas colombianas", despues)
        # Lo que sigue conserva su número: su resultado guardado sigue siendo suyo.
        self.assertEqual(despues["Compromiso anticorrupción firmado"], antes["Compromiso anticorrupción firmado"])
        self.assertEqual(despues["Duración de la sociedad"], antes["Duración de la sociedad"])


class RequisitosRetiradosNoCuentanTests(BaseEvaluaciones):
    def test_una_evaluacion_aprobada_no_cambia(self):
        from evaluaciones import servicios
        from evaluaciones.models import EstadoEvaluacion, Evaluacion

        jefe, ev = self.crear()
        evaluacion = Evaluacion.objects.get(pk=ev["id"])
        evaluacion.estado = EstadoEvaluacion.APROBADA
        evaluacion.save(update_fields=["estado"])
        self.assertEqual(servicios.depurar_requisitos_retirados(evaluacion), 0)


class CedulaConCerosDeLaCamaraTests(TestCase):
    """La Cámara de Comercio escribe la cédula del suplente con ceros a la
    izquierda ("C.C. 000000009990000006"): sigue siendo la misma cédula."""

    def test_los_ceros_no_impiden_reconocer_la_cedula(self):
        from motor.evaluacion.identidad import SUPLENTE_RE, _digitos, _es_suyo

        certificado = "SUPLENTE DEL GERENTE CASTRO MEJIA PAULA ANDREA C.C. 000000009990000006"
        m = SUPLENTE_RE.search(certificado)
        self.assertEqual(_digitos(m.group(2)), "9990000006")
        cedula = "REPUBLICA DE COLOMBIA IDENTIFICACION PERSONAL CEDULA DE CIUDADANIA NUMERO 9.990.000.006 CASTRO MEJIA PAULA ANDREA"
        self.assertTrue(_es_suyo("CASTRO MEJIA PAULA ANDREA", _digitos(m.group(2)), cedula))


class PersonasConCedulasConsecutivasTests(TestCase):
    """Dos personas con cédulas de 10 dígitos consecutivas (hermanos que la
    sacaron el mismo día) no pueden fundirse en una: los antecedentes de una
    taparían los de la otra."""

    def test_la_cedula_completa_distingue_a_las_personas(self):
        from evaluaciones.servicios import clave_persona

        self.assertNotEqual(clave_persona("LAURA MENDEZ", "1020304050", "natural"), clave_persona("JORGE PAREDES", "1020304051", "natural"))
        self.assertEqual(clave_persona("LAURA MENDEZ", "1.020.304.050", "natural"), clave_persona("LAURA MENDEZ", "1020304050", "natural"))
        # El NIT con o sin dígito de verificación es la misma empresa.
        self.assertEqual(clave_persona("ACME SAS", "900123456-7", "juridica"), clave_persona("ACME SAS", "900123456", "juridica"))


class CedulasConsecutivasNoSeConfundenTests(TestCase):
    def test_la_cedula_del_hermano_no_es_suya(self):
        from motor.evaluacion.identidad import _es_suyo

        reverso_de_jorge = "21-JUN-2005 BOGOTA FECHA Y LUGAR DE EXPEDICION INDICE DERECHO A-0000000-00000000-M-1999000002-20050621"
        self.assertFalse(_es_suyo("LAURA CAMILA MENDEZ RIOS", "1999000001", reverso_de_jorge))
        self.assertTrue(_es_suyo("JORGE ANDRES PAREDES LUNA", "1999000002", reverso_de_jorge))

    def test_un_digito_mal_leido_con_su_nombre_si_es_suya(self):
        from motor.evaluacion.identidad import _es_suyo

        anverso_ocr = "CEDULA DE CIUDADANIA NUMERO 1.999.000.007 MENDEZ RIOS LAURA CAMILA"  # el OCR leyó 7 por 1
        self.assertTrue(_es_suyo("LAURA CAMILA MENDEZ RIOS", "1999000001", anverso_ocr))


class CopniaMasRecienteTests(TestCase):
    """Si se adjunta un COPNIA nuevo porque el de la oferta estaba vencido,
    el nuevo es el que cuenta."""

    def test_gana_el_mas_reciente_del_mismo_profesional(self):
        from motor.evaluacion import copnia

        viejo = ("oferta/copnia.pdf", "COPNIA 1. Que MARIA JOSE RESTREPO DIAZ, identificado(a) con CEDULA DE CIUDADANIA 1999000009 se expide a los veinte (20) días del mes de Abril del año dos mil veintiseis (2026).")
        nuevo = ("aportados/Req 3 - COPNIA.pdf", "COPNIA 1. Que MARIA JOSE RESTREPO DIAZ, identificado(a) con CEDULA DE CIUDADANIA 1999000009 se expide a los veinte (20) días del mes de Septiembre del año dos mil veintiseis (2026).")
        with mock.patch.object(copnia, "encontrar_copnias", return_value=[viejo, nuevo]):
            elegido = copnia._elegir_copnia_del_profesional({}, ["MARIA JOSE RESTREPO DIAZ"])
        self.assertEqual(elegido[0], "aportados/Req 3 - COPNIA.pdf")


class TarjetaProfesionalAvalTests(TestCase):
    """La copia de la tarjeta profesional del ingeniero que avala se exige
    solo si la sección del aval del pliego la pide y no deja suplirla."""

    def test_aval_solo_con_copnia_no_exige_tarjeta(self):
        from motor.parsers.documento_base import _aval_pide_tarjeta

        texto = ("la oferta tendrá que ser avalada por un Ingeniero, para lo cual adjuntará el certificado de vigencia "
                 "de matrícula profesional expedida por el Copnia. El aval del ingeniero hace parte integral del Formato 1. "
                 "Para el equipo de trabajo: copia de la tarjeta profesional y certificado de antecedentes.")
        self.assertFalse(_aval_pide_tarjeta(texto))

    def test_aval_que_pide_tarjeta(self):
        from motor.parsers.documento_base import _aval_pide_tarjeta

        texto = ("la oferta tendrá que ser avalada por un ingeniero, para lo cual debe adjuntar copia de la tarjeta "
                 "profesional y copia del certificado de vigencia de matrícula profesional. El aval hace parte integral del Formato 1.")
        self.assertTrue(_aval_pide_tarjeta(texto))

    def test_aval_que_pide_tarjeta_suplible(self):
        from motor.parsers.documento_base import _aval_pide_tarjeta

        texto = ("la oferta tendrá que ser avalada por un ingeniero, para lo cual debe adjuntar copia de la tarjeta "
                 "profesional y del certificado de vigencia. El requisito de la tarjeta profesional se puede suplir con el "
                 "registro de que trata el artículo 18 del Decreto 2106 de 2019. El aval hace parte integral del Formato 1.")
        self.assertFalse(_aval_pide_tarjeta(texto))

    def test_documento_base_viejo_conserva_la_regla_anterior(self):
        from motor.esquemas.proceso import ProcesoDocumentoBase

        campos = ProcesoDocumentoBase.model_fields
        self.assertIsNone(campos["tarjeta_exigida"].default)
        doc = ProcesoDocumentoBase.model_construct(tarjeta_suplible=False, tarjeta_exigida=None)
        self.assertTrue(doc.exige_tarjeta_profesional)
        # Aunque el aval del pliego no la mencione, se exige salvo que el
        # pliego deje suplirla (por prudencia: sin ella hubo un aval aprobado
        # que el abogado rechazó).
        doc = ProcesoDocumentoBase.model_construct(tarjeta_suplible=False, tarjeta_exigida=False)
        self.assertTrue(doc.exige_tarjeta_profesional)
        doc = ProcesoDocumentoBase.model_construct(tarjeta_suplible=True, tarjeta_exigida=False)
        self.assertFalse(doc.exige_tarjeta_profesional)


class GravamenesExistenciaTests(TestCase):
    """Embargos, órdenes judiciales e insolvencia en el certificado de
    existencia mandan el requisito a revisión; las frases de rutina no."""

    def test_detecta_embargo_y_reorganizacion(self):
        from motor.evaluacion.camara_comercio import gravamenes_del_certificado

        casos = [
            "** ORDENES DE AUTORIDADES COMPETENTES: POR OFICIO NO. 0816 DEL 22 DE MAYO DE 2026 DEL JUZGADO PRIMERO LABORAL "
            "SE DECRETO EMBARGO DE ESTABLECIMIENTO DE COMERCIO.",
            "SE DECRETO EL EMBARGO DE LAS CUOTAS SOCIALES QUE POSEA EL DEMANDADO EN LA SOCIEDAD DE LA REFERENCIA.",
            "LA SUPERINTENDENCIA DE SOCIEDADES ORDENO LA ADMISION AL PROCESO DE REORGANIZACION DE LA SOCIEDAD DE LA REFERENCIA.",
            "RAZON SOCIAL: CANO JIMENEZ ESTUDIOS S A - EN REORGANIZACION NIT: 800.000.000-1",
            "SE NOMBRO PROMOTOR(A) DENTRO DEL TRAMITE DE REORGANIZACION EMPRESARIAL DE LA SOCIEDAD",
        ]
        for texto in casos:
            self.assertTrue(gravamenes_del_certificado(texto), texto)

    def test_frases_de_rutina_no_cuentan(self):
        from motor.evaluacion.camara_comercio import gravamenes_del_certificado

        casos = [
            "LOS BIENES SUJETOS A REGISTRO MERCANTIL RELACIONADOS EN EL PRESENTE CERTIFICADO, SE ENCUENTRAN LIBRES DE EMBARGOS.",
            "COMPRAR Y VENDER TODA CLASE DE BIENES, SIN EMBARGO SE PROHIBE EXPRESAMENTE A LA SOCIEDAD CONSTITUIRSE EN GARANTE.",
            "LA PERSONA JURIDICA NO SE ENCUENTRA DISUELTA Y SU DURACION ES INDEFINIDA.",
            "D) SOLICITAR LA ADMISION DE LA SOCIEDAD A UN PROCESO DE REORGANIZACION O DE LIQUIDACION EN LOS TERMINOS DEL "
            "REGIMEN DE INSOLVENCIA QUE CONSAGRA LA LEY 1116 DE 2006",
            "F) CUMPLIR LAS DISPOSICIONES LEGALES Y DEMAS ORDENES DE AUTORIDAD COMPETENTE",
            "CONTRALORIA DE EMPRESAS EN LIQUIDACION, REALIZACION DE ESTUDIOS",
        ]
        for texto in casos:
            self.assertEqual(gravamenes_del_certificado(texto), [], texto)


class FechaExpedicionCedulaTests(TestCase):
    """La fecha de expedición se lee de la cédula antigua y de la digital,
    solo de páginas de la persona."""

    def test_etiqueta_danada_por_el_ocr(self):
        from motor.evaluacion.identidad import fecha_expedicion_en
        from datetime import date

        texto = "1.85 O+ M ESTATURA AS AH SEXO 14-DIC-1987 VALLEDUPAR FECHA Y LUGAR DE EXPED&CION REGISTRADOR"
        self.assertEqual(fecha_expedicion_en(texto), date(1987, 12, 14))

    def test_cedula_digital_con_etiqueta(self):
        from motor.evaluacion.identidad import fecha_expedicion_en
        from datetime import date

        texto = "FECHA DE NACIMIENTO 09 OCT 1981 FECHA Y LUGAR DE EXPEDICION 26 ENE 2000, BUCARAMANGA"
        self.assertEqual(fecha_expedicion_en(texto), date(2000, 1, 26))

    def test_cedula_digital_sin_etiquetas_por_la_mrz(self):
        from motor.evaluacion.identidad import fecha_expedicion_por_mrz
        from datetime import date

        texto = ("GOMEZ FALLA JORGE 27 MAR 1976 NEIVA (HUILA) 07 MAYO 1997 SARRANQUILLA 13 SEPT 2032 "
                 "ICCOLOD3931792403001<<<<<<<<<< 7903270M3209136COL72007216<<<7 GOMEZ<FALLA<<JORGE<<<<")
        self.assertEqual(fecha_expedicion_por_mrz(texto, "72007216"), date(1997, 5, 7))
        # La MRZ es de otra persona: no se afirma nada.
        self.assertIsNone(fecha_expedicion_por_mrz(texto, "91505590"))

    def test_mrz_con_digito_de_control_malo_no_sirve(self):
        from motor.evaluacion.identidad import fecha_expedicion_por_mrz

        texto = "07 MAYO 1997 13 SEPT 2032 7903271M3209136COL72007216<<<7"
        self.assertIsNone(fecha_expedicion_por_mrz(texto, "72007216"))

    def test_mes_danado(self):
        from motor.evaluacion.identidad import fecha_expedicion_por_mrz
        from datetime import date

        texto = "LORA MUNO MARIA ISABE 31 00T 2012 BARRANQUILLA 19 DIC 2033 9407288F3312197C0L1140872459<"
        self.assertEqual(fecha_expedicion_por_mrz(texto, "1140872459"), date(2012, 10, 31))


class LecturasQueNoCoincidenTests(TestCase):
    """Cédula real: el texto decía 28-MAR-2017 y el OCR reforzado leyó
    20-MAR-2017 con la etiqueta limpia. Si las dos lecturas no coinciden,
    ninguna se usa sola."""

    def test_dos_lecturas_distintas_no_son_confiables(self):
        from datetime import date

        from motor.evaluacion import identidad

        normal = "CEDULA NUMERO 1.098.817.511 JIMENEZ CORREA DANIELA ANDREA 28-MAR-2017 BUCARAMANGA FECHA Y LUGAR DE EXPEDICION"
        reforzada = "CEDULA 1.098.817.511 JIMENEZ CORREA DANIELA ANDREA 20-MAR-2017 BUCARAMANGA FECHA Y LUGAR DE EXPEDICION"
        with mock.patch.object(identidad, "cedula_de", return_value="doc.pdf"), \
             mock.patch.object(identidad, "lecturas_de_paginas", side_effect=lambda *a, **k: ({1: normal}, {1: reforzada})):
            lectura = identidad.leer_fecha_expedicion({"doc.pdf": b"x"}, "DANIELA ANDREA JIMENEZ CORREA", "1098817511", con_ia=False)
        self.assertFalse(lectura.confiable)
        self.assertIn(lectura.fecha, {date(2017, 3, 28), date(2017, 3, 20)})

    def test_dos_lecturas_iguales_con_etiqueta_danada_son_confiables(self):
        from datetime import date

        from motor.evaluacion import identidad

        normal = "NUMERO 77.031.208 ARAQUE BLANCO JUAN JOSE 14-DIC-1987 VALLEDUPAR FECHA Y LUGAR DE EXPED&CION"
        reforzada = "77033208 ARAQUE BLAQQ 14-DIC-1987 VALLEDUPAR FECHA Y LUGAR DE EXPEDIC|ON"
        with mock.patch.object(identidad, "cedula_de", return_value="doc.pdf"), \
             mock.patch.object(identidad, "lecturas_de_paginas", side_effect=lambda *a, **k: ({1: normal}, {1: reforzada})):
            lectura = identidad.leer_fecha_expedicion({"doc.pdf": b"x"}, "JUAN JOSE ARAQUE BLANCO", "77031208", con_ia=False)
        self.assertEqual((lectura.fecha, lectura.confiable), (date(1987, 12, 14), True))


class NumeroDelReversoTests(TestCase):
    """La línea de abajo del reverso trae el número de la cédula: con dos
    dígitos mal leídos sigue siendo suya; muy distinto, es de otra persona."""

    def test_numero_del_reverso(self):
        from motor.evaluacion.identidad import _reverso_es_suyo

        self.assertTrue(_reverso_es_suyo("52371321", "A-1500150-01172910-F-0052871821-20201023 00721698982"))
        self.assertFalse(_reverso_es_suyo("52371321", "A-0100100-00157053-F-0043001767-20090520 0011606230A"))
        self.assertIsNone(_reverso_es_suyo("52371321", "FECHA Y LUGAR DE EXPEDICION"))

    def test_archivo_con_su_nombre(self):
        from motor.evaluacion.identidad import _archivo_con_su_nombre

        self.assertTrue(_archivo_con_su_nombre("DOBLE R/CC ADRIANA ROJAS RL.pdf", "ADRIANA MARCELA ROJAS PRIETO"))
        self.assertFalse(_archivo_con_su_nombre("CEDULAS/CC LFRM.pdf", "LUIS FELIPE RUIZ MEJIA"))
        self.assertFalse(_archivo_con_su_nombre("DOC LEGAL/MEGB/Doc Legal MEGB.pdf", "ADRIANA MARCELA ROJAS PRIETO"))


class NumeroParecidoTests(TestCase):
    def test_numero_con_un_digito_de_mas(self):
        from motor.evaluacion.identidad import _numero_parecido

        self.assertTrue(_numero_parecido("19275138", "A-1500150-00184486-M-00192715138-20091009"))
        self.assertTrue(_numero_parecido("52371321", "F-0052871821-20201023"))
        self.assertFalse(_numero_parecido("52371321", "F-0043001767-20090520"))


class FechasDanadasPorElOcrTests(TestCase):
    def test_letras_por_digitos_y_ano_pegado(self):
        from datetime import date

        from motor.evaluacion.identidad import _fecha_de_la_franja

        self.assertEqual(_fecha_de_la_franja("ESTATURA G.S. RH SEXO OB-FEB-2013VALLEDUPAR — FECHA Y LUGAR DE EXPEDICION"), date(2013, 2, 8))
        self.assertEqual(_fecha_de_la_franja("GS. AH SEXO 19-AG0-1976 BOGOTA D.C. FECHA Y LUGAR DE"), date(1976, 8, 19))
        self.assertEqual(_fecha_de_la_franja("SEXO 03-0CT-1997 BOGOTA D.C FECHA Y LUGAR DE EXPEDICION"), date(1997, 10, 3))
