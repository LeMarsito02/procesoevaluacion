"""Lo que hace de MiEvaluador una herramienta de apoyo y no de decisión.

Reúne los textos y las reglas que exige el expediente jurídico-técnico LEG-004
(Concepto C-1015 de 2026 de la ANCP-CCE) para que el juicio y la decisión
queden, de forma verificable, en las personas:

- el lenguaje: el sistema *verifica* y *propone*; la persona *adopta*;
- los rótulos de los informes antes y después de la adopción;
- la trazabilidad de cada resultado (versión del sistema y modelo de IA);
- los indicadores que muestran que la revisión humana no fue nominal.
"""
from __future__ import annotations

from django.conf import settings
from django.utils import timezone

# --- Textos -----------------------------------------------------------------
# Numeral 6.3 del LEG-004: aviso permanente en la interfaz.
AVISO_INTERFAZ = (
    "Los resultados de MiEvaluador son verificaciones preliminares que usted adopta, corrige o descarta. "
    "La decisión es del comité evaluador."
)
# Numeral 3, paso 2: el pre-informe va rotulado mientras no se adopte.
LEYENDA_PREINFORME = "Pre-informe – documento de apoyo, pendiente de adopción por el comité evaluador."
NOTA_PRELIMINAR = "No produce efecto hasta que el comité lo adopte y lo motive en el informe que suscribe."

FORMA_VERIFICADO = "Verificado por MiEvaluador"
FORMA_NO_APLICA = "No aplica según verificación de MiEvaluador"

# Numeral 6.4: compromiso de uso del evaluador. Cambiar el texto obliga a
# subir la versión: todos lo vuelven a aceptar.
COMPROMISO_VERSION = "2026-09-27"
COMPROMISO_TEXTO = (
    "Declaro que conozco el alcance y las limitaciones de MiEvaluador; que revisaré de manera sustantiva los "
    "incumplimientos, observaciones y subsanaciones de mi área, así como la muestra de control de los ítems "
    "verificados por la herramienta, antes de incorporarlos a mi evaluación; que verificaré la información contra "
    "las fuentes oficiales; y que entiendo que el uso de la herramienta no traslada, atenúa ni distribuye mi "
    "responsabilidad por el contenido del informe que suscribo."
)


def forma_verificado(evaluacion) -> str:
    """«Verificado por MiEvaluador – adoptado por …» o «– pendiente de adopción»."""
    from evaluaciones.models import EstadoEvaluacion
    from evaluaciones.reporte import fecha_larga

    if evaluacion.estado == EstadoEvaluacion.APROBADA and evaluacion.aprobada_por_id:
        return f"{FORMA_VERIFICADO} – adoptado por {evaluacion.aprobada_por.nombre_completo} el {fecha_larga(evaluacion.aprobada_en)}"
    return f"{FORMA_VERIFICADO} – pendiente de adopción"


def titulo_adoptado(evaluacion) -> str:
    """Título del informe aprobado: es el informe del comité, hecho con MiEvaluador."""
    from evaluaciones.reporte import fecha_larga

    area = evaluacion.get_tipo_display().lower()
    quien = evaluacion.aprobada_por.nombre_completo if evaluacion.aprobada_por_id else "el comité evaluador"
    return f"Informe de evaluación {area} – adoptado por {quien} el {fecha_larga(evaluacion.aprobada_en)}"


# --- Trazabilidad -------------------------------------------------------------
def modelos_ia() -> dict[str, str]:
    """Modelos de IA configurados en este momento, por uso."""
    from motor.llm import cliente, vision
    from motor.pliego import lector_ia

    return {
        "extraccion_documentos": cliente.LLM_MODELO,
        "lectura_pliego": lector_ia.MODELO,
        "vision": vision.MODELO,
    }


def usa_ia(datos: dict) -> bool:
    """El resultado usó datos extraídos con IA (y comprobados contra el texto)."""
    return "IA local" in (datos.get("motivo") or "")


def trazabilidad_resultado(datos: dict) -> dict:
    traza: dict = {"version_sistema": settings.MIEVALUADOR_VERSION, "registrado_en": timezone.now().isoformat()}
    if usa_ia(datos):
        traza["modelos_ia"] = modelos_ia()
    return traza


def es_puntaje(definicion, numero: int) -> bool:
    return any(r.numero == numero and r.grupo == "puntaje" for r in definicion.requisitos)


def numeros_puntaje(definicion) -> set[int]:
    return {r.numero for r in definicion.requisitos if r.grupo == "puntaje"}


# --- Indicadores de revisión (numerales 2, 8 y 9.2) ---------------------------
def indicadores(evaluacion) -> dict:
    """Cuánto verificó el sistema y cuánto revisó, corrigió o descartó una
    persona. Prueba que la supervisión no fue nominal."""
    from cuentas.models import EventoAuditoria
    from evaluaciones.models import MuestraControl, Resultado, Revision

    resultados = {(r.proponente_id, r.requisito): r for r in Resultado.objects.filter(evaluacion=evaluacion)}
    revisiones = list(Revision.objects.filter(evaluacion=evaluacion))
    decididas = {(r.proponente_id, r.requisito): r for r in revisiones}
    automaticas = [
        k for k, r in resultados.items()
        if r.datos.get("cumple") is True and not r.datos.get("error") and k not in decididas
    ]
    corregidas = 0  # la persona decidió distinto de lo que decía el sistema
    confirmadas = 0
    for k, rev in decididas.items():
        r = resultados.get(k)
        if r is None:
            continue
        if r.datos.get("cumple") is True and not r.datos.get("error"):
            if rev.cumple:
                confirmadas += 1
            else:
                corregidas += 1
    pendientes_resueltas = sum(
        1 for k, rev in decididas.items() if k in resultados and resultados[k].requiere_revision
    )
    no_cumple_humano = sum(1 for rev in revisiones if not rev.cumple)
    documentos_vistos = EventoAuditoria.objects.filter(
        accion="documento.visto", objeto_tipo="Evaluacion", objeto_id=str(evaluacion.pk)
    ).count()
    muestra = MuestraControl.objects.filter(evaluacion=evaluacion).exclude(estado="anulada").first()
    datos_muestra = None
    if muestra is not None:
        items = list(muestra.items.all())
        datos_muestra = {
            "ofertas": len(muestra.ofertas_sorteadas),
            "verificaciones": muestra.items.count(),
            "items": len(items),
            "conformes": sum(1 for i in items if i.resultado == "conforme"),
            "no_conformes": sum(1 for i in items if i.resultado == "no_conforme"),
            "con_soporte_visto": sum(1 for i in items if i.soporte_visto),
            "requisitos_ampliados": list(muestra.requisitos_ampliados),
            "estado": muestra.get_estado_display(),
        }
    total = len(resultados)
    return {
        "resultados": total,
        "verificados_por_el_sistema_sin_revision": len(automaticas),
        "revisados_por_una_persona": len(decididas),
        "pendientes_resueltos_por_una_persona": pendientes_resueltas,
        "verificaciones_del_sistema_confirmadas": confirmadas,
        "verificaciones_del_sistema_corregidas": corregidas,
        "no_cumple_decididos_por_una_persona": no_cumple_humano,
        "documentos_soporte_abiertos": documentos_vistos,
        "muestra_de_control": datos_muestra,
    }


def constancia(evaluacion) -> str:
    """Numeral 6.2 del LEG-004 (versión 1.2), con los datos reales de la evaluación."""
    from evaluaciones.models import MuestraControl

    version = evaluacion.version_sistema or settings.MIEVALUADOR_VERSION
    area = evaluacion.get_tipo_display().lower()
    muestra = MuestraControl.objects.filter(evaluacion=evaluacion, estado="cerrada").first()
    if muestra is not None:
        texto_muestra = (
            f"previa revisión de una muestra de control de {muestra.items.count()} verificaciones sorteadas "
            f"en {len(muestra.ofertas_sorteadas)} ofertas, revisadas contra su soporte, cuyo resultado consta en el acta anexa"
        )
    else:
        texto_muestra = "previa revisión de la muestra de control prevista en el procedimiento"
    return (
        f"En el presente proceso, el comité evaluador utilizó la herramienta tecnológica MiEvaluador (versión {version}) "
        "como apoyo para contrastar la información de las ofertas con las reglas del pliego de condiciones, previamente "
        "confirmadas por la entidad. La herramienta verificó con evidencia documental parte de los requisitos y dejó los "
        f"demás pendientes de revisión humana; no emitió resultados de incumplimiento ni adoptó decisión alguna. El "
        f"evaluador del área {area} revisó individualmente todos los requisitos pendientes, los incumplimientos, las "
        f"observaciones y las subsanaciones, y adoptó los resultados verificados por la herramienta {texto_muestra}. "
        "Las conclusiones, puntajes y recomendaciones de este informe son del comité evaluador, que las adopta y motiva "
        "bajo su responsabilidad. La bitácora de uso reposa en el expediente del proceso."
    )


# --- Rótulo de los informes en Excel (numerales 3 y 6.2) ----------------------
def rotular_excel(contenido: bytes, *, adoptado: bool, titulo: str, texto_constancia: str) -> bytes:
    """Encabezado en cada hoja y una hoja «Constancia» al final.

    Antes de la adopción, todo el libro dice que es un pre-informe. Después,
    dice quién lo adoptó y cuándo: ese libro es el informe del comité."""
    import io

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font

    libro = load_workbook(io.BytesIO(contenido))
    encabezado = titulo if adoptado else LEYENDA_PREINFORME
    for hoja in libro.worksheets:
        hoja.oddHeader.center.text = encabezado[:250]
        hoja.oddHeader.center.size = 8
    if "Constancia" in libro.sheetnames:
        del libro["Constancia"]
    hoja = libro.create_sheet("Constancia")
    hoja.column_dimensions["A"].width = 120
    hoja["A1"] = titulo if adoptado else LEYENDA_PREINFORME
    hoja["A1"].font = Font(bold=True, size=12)
    hoja["A3"] = texto_constancia
    hoja["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    hoja.row_dimensions[3].height = 150
    hoja["A5"] = f"Generado por MiEvaluador, versión {settings.MIEVALUADOR_VERSION}, el {timezone.localtime():%d/%m/%Y %H:%M}."
    hoja["A5"].font = Font(size=9, italic=True)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def rotular_informe_de(evaluacion, contenido: bytes) -> bytes:
    from evaluaciones.models import EstadoEvaluacion

    adoptado = evaluacion.estado == EstadoEvaluacion.APROBADA
    return rotular_excel(
        contenido,
        adoptado=adoptado,
        titulo=titulo_adoptado(evaluacion) if adoptado else LEYENDA_PREINFORME,
        texto_constancia=constancia(evaluacion),
    )
