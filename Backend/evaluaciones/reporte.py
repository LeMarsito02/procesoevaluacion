"""Reporte formal de evaluación en Word.

Deja constancia de quién creó, asignó, evaluó y aprobó, de cómo se verificó
cada requisito de cada proponente y del porqué: «Aprobado automáticamente por
MiEvaluador» (con la razón y el documento soporte) o «Validado manualmente por»
(con la persona, la fecha y su justificación).
"""
from __future__ import annotations

from datetime import datetime

import io
from collections import defaultdict
from dataclasses import dataclass

from django.utils import timezone
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from evaluaciones import servicios
from evaluaciones.models import DocumentoAportado, EstadoEvaluacion, Evaluacion, PersonaVerificada, Resultado, Revision
from motor import criterios
from motor.esquemas.proceso import ProcesoDocumentoBase, ResultadoRequisito

AZUL = RGBColor(0x00, 0x14, 0x3C)
AZUL_MARCA = RGBColor(0x00, 0x48, 0xDC)
GRIS = RGBColor(0x55, 0x60, 0x75)
COPYRIGHT = "© {anio} LeMarTek Labs S.A.S. · NIT 902069061-9"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
TIPOS_PROPONENTE = {
    "persona_natural": "Persona natural",
    "persona_juridica": "Persona jurídica",
    "consorcio": "Consorcio",
    "union_temporal": "Unión temporal",
}


def fecha_larga(valor) -> str:
    if valor is None:
        return "—"
    if hasattr(valor, "hour"):
        valor = timezone.localtime(valor)
        return f"{valor.day} de {MESES[valor.month - 1]} de {valor.year}, {valor:%H:%M}"
    return f"{valor.day} de {MESES[valor.month - 1]} de {valor.year}"


def pesos(valor: float) -> str:
    return "$ " + f"{valor:,.0f}".replace(",", ".")


def _frase(texto: str | None) -> str:
    """Primera letra en mayúscula y un solo punto final."""
    texto = (texto or "").strip().rstrip(".").strip()
    return (texto[:1].upper() + texto[1:] + ".") if texto else ""


@dataclass
class Validacion:
    estado: str  # Cumple / No cumple / No aplica / Pendiente
    forma: str
    detalle: str


def explicar_automatico(r: ResultadoRequisito, verifica: str) -> str:
    """Por qué MiEvaluador dio por cumplido el requisito, con los datos que leyó."""
    partes = [f"Se verificó: {verifica.rstrip('.')}." if verifica else "Se verificaron las condiciones del requisito."]
    if r.representante_legal:
        partes.append(f"Representante legal identificado: {r.representante_legal}.")
    if r.profesion_certificada or r.matricula_profesional:
        partes.append(
            "Profesional certificado"
            + (f": {r.profesion_certificada}" if r.profesion_certificada else "")
            + (f", matrícula {r.matricula_profesional}" if r.matricula_profesional else "")
            + "."
        )
    if r.copnia_fecha_expedicion:
        partes.append(f"Certificado expedido el {fecha_larga(r.copnia_fecha_expedicion)}.")
    if r.lotes_encontrados:
        partes.append(f"Lotes a los que se presenta: {', '.join(r.lotes_encontrados)}.")
    if r.motivo:
        partes.append(_frase(r.motivo))
    if r.archivo_evaluado:
        partes.append(f"Documento soporte: {r.archivo_evaluado}.")
    return " ".join(partes)


def validacion(r: ResultadoRequisito, revision: Revision | None, verifica: str) -> Validacion:
    if revision is not None:
        quien = revision.usuario.nombre_completo if revision.usuario_id else "el evaluador"
        justificacion = revision.nota.strip() or "Sin justificación registrada."
        soporte = f" Documento revisado: {r.archivo_evaluado}." if r.archivo_evaluado else ""
        motor = f" Observación del sistema: {_frase(r.error or r.motivo)}" if (r.error or r.motivo) else ""
        return Validacion(
            "Cumple" if revision.cumple else "No cumple",
            f"Validado manualmente por {quien} el {fecha_larga(revision.fecha)}",
            f"Justificación: {_frase(justificacion)}{soporte}{motor}",
        )
    if r.error:
        return Validacion("Pendiente", "No se pudo evaluar automáticamente", _frase(r.error))
    if (r.motivo or "").startswith("N.A."):
        return Validacion("No aplica", "Determinado automáticamente por MiEvaluador", _frase(r.motivo.removeprefix("N.A.").lstrip(" —-")))
    if r.cumple:
        return Validacion("Cumple", "Aprobado automáticamente por MiEvaluador", explicar_automatico(r, verifica))
    return Validacion("Pendiente", "Requiere revisión del evaluador", _frase(r.motivo or "El sistema no pudo confirmar el requisito"))


# --- Utilidades de formato ---
def _sombrear(celda, color: str) -> None:
    tc = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc.append(shd)


def _tabla(doc, encabezados: list[str], filas: list[list[str]], anchos: list[float] | None = None):
    tabla = doc.add_table(rows=1, cols=len(encabezados))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, texto in enumerate(encabezados):
        celda = tabla.rows[0].cells[i]
        celda.text = ""
        run = celda.paragraphs[0].add_run(texto)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _sombrear(celda, "00143C")
    for fila in filas:
        celdas = tabla.add_row().cells
        for i, texto in enumerate(fila):
            celdas[i].text = ""
            run = celdas[i].paragraphs[0].add_run(str(texto))
            run.font.size = Pt(9)
    # Repetir encabezado en cada página.
    tr_pr = tabla.rows[0]._tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")
    repetir.set(qn("w:val"), "true")
    tr_pr.append(repetir)
    if anchos:
        # Diseño fijo: Word y LibreOffice respetan los anchos solo así.
        tabla.autofit = False
        tbl_pr = tabla._tbl.tblPr
        diseno = OxmlElement("w:tblLayout")
        diseno.set(qn("w:type"), "fixed")
        tbl_pr.append(diseno)
        for i, ancho in enumerate(anchos):
            tabla.columns[i].width = Cm(ancho)
        for fila in tabla.rows:
            for i, ancho in enumerate(anchos):
                fila.cells[i].width = Cm(ancho)
    return tabla


def _campo_pagina(parrafo) -> None:
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        run = parrafo.add_run()
        if tipo:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), tipo)
            run._r.append(f)
        else:
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = texto
            run._r.append(t)


def _parrafo(doc, texto: str, *, negrita=False, tam=10.5, color=None, alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY, espacio=6):
    p = doc.add_paragraph()
    p.alignment = alineacion
    p.paragraph_format.space_after = Pt(espacio)
    run = p.add_run(texto)
    run.bold = negrita
    run.font.size = Pt(tam)
    if color is not None:
        run.font.color.rgb = color
    return p


def _titulo(doc, texto: str, nivel: int = 1):
    h = doc.add_heading(texto, level=nivel)
    for run in h.runs:
        run.font.color.rgb = AZUL if nivel == 1 else AZUL_MARCA
        run.font.name = "Arial"
    return h


# --- Reporte ---
def _seccion_pliego(doc, proceso) -> None:
    """Qué se ajustó de la evaluación por lo que dice el pliego de este proceso,
    con la cita, la página y quién lo decidió."""
    analisis = proceso.analisis_pliego
    _titulo(doc, "Ajustes derivados del pliego", nivel=2)
    if analisis is None:
        _parrafo(doc, "El proceso se creó sin análisis del pliego: la evaluación aplica la plantilla de la entidad sin ajustes.")
        return
    tipo = f" (documento tipo {analisis.documento_tipo})" if analisis.documento_tipo else ""
    _parrafo(
        doc,
        f"El sistema leyó completo el pliego del proceso, «{analisis.nombre_archivo}», de {analisis.paginas} páginas{tipo}, "
        "y lo comparó con la plantilla de evaluación de la entidad. Cada diferencia fue presentada al evaluador con la cita "
        "literal y la página del pliego, y se aplicó solo lo que este aceptó.",
    )
    ajustes = proceso.ajustes_pliego or []
    if not ajustes:
        _parrafo(doc, "El pliego no exigió ajustes a la evaluación de la entidad.")
    else:
        filas = []
        for a in ajustes:
            h = a["hallazgo"]
            if a["decision"] == "aceptado":
                decision = "Aplicado"
                if h.get("parametro"):
                    decision += f": {h.get('valor_plantilla')} → {h.get('valor_pliego_texto') or h.get('valor_pliego')}"
                elif h.get("tipo") == "requisito_nuevo":
                    decision += ": requisito agregado"
            else:
                decision = f"No aplicado. Razón: {a.get('nota') or '—'}"
            filas.append([
                h["titulo"],
                f"«{h['cita']}» ({h['seccion']}, pág. {h['pagina']})",
                decision,
                f"{a['por']}, {fecha_larga(datetime.fromisoformat(a['en']))}",
            ])
        _tabla(doc, ["Exigencia del pliego", "Texto del pliego", "Decisión", "Decidió"], filas, [3.4, 7.0, 3.6, 2.6])


def notas_del_pliego(proceso, definicion) -> dict[int, str]:
    """Por requisito, lo que cambió debido al pliego (aceptado por el
    evaluador), con la sección y la página: "Debido al pliego (3.4, pág.
    12): 1 mes → 30 días"."""
    notas: dict[int, list[str]] = defaultdict(list)
    for a in proceso.ajustes_pliego or []:
        if a.get("decision") != "aceptado":
            continue
        h = a["hallazgo"]
        donde = f"{h.get('seccion') or 'pliego'}, pág. {h.get('pagina')}"
        propuesto = h.get("requisito_propuesto") or {}
        if h.get("parametro"):
            cambio = f"{h.get('valor_plantilla')} → {h.get('valor_pliego_texto') or h.get('valor_pliego')}"
            numeros = [r.numero for r in definicion.requisitos if r.verificacion == h.get("verificacion")]
        elif h.get("tipo") == "requisito_nuevo":
            cambio = "requisito exigido por el pliego"
            numeros = [r.numero for r in definicion.requisitos
                       if r.titulo == (propuesto.get("titulo") or "")[:200] or
                       (propuesto.get("verificacion") not in (None, criterios.MANUAL, criterios.PERSONALIZADO)
                        and r.verificacion == propuesto.get("verificacion"))][:1]
        else:
            continue
        for n in numeros:
            notas[n].append(f"Debido al pliego ({donde}): {cambio}.")
    return {n: " ".join(dict.fromkeys(v)) for n, v in notas.items()}


def generar_reporte(evaluacion: Evaluacion) -> tuple[bytes, str]:
    evaluacion = Evaluacion.objects.select_related(
        "entidad", "proceso__creado_por", "responsable", "asignada_por", "aprobada_por", "plantilla"
    ).get(pk=evaluacion.pk)
    proceso = evaluacion.proceso
    documento_base = ProcesoDocumentoBase.model_validate(proceso.documento_base)
    definicion = servicios.definicion_de(evaluacion)
    catalogo = servicios.catalogo(definicion)
    info = {c["numero"]: c for c in catalogo}
    por_pliego = notas_del_pliego(proceso, definicion)
    orden = [c["numero"] for c in catalogo]
    tipo = evaluacion.get_tipo_display().lower()
    caracter = {"juridica": "jurídico", "tecnica": "técnico", "financiera": "financiero"}.get(evaluacion.tipo, tipo)
    aprobada = evaluacion.estado == EstadoEvaluacion.APROBADA
    ahora = timezone.now()

    resultados: dict = defaultdict(dict)
    for r in Resultado.objects.filter(evaluacion=evaluacion):
        resultados[r.proponente_id][r.requisito] = ResultadoRequisito.model_validate(r.datos)
    revisiones = {(r.proponente_id, r.requisito): r for r in Revision.objects.filter(evaluacion=evaluacion).select_related("usuario")}
    personas = defaultdict(list)
    for p in PersonaVerificada.objects.filter(evaluacion=evaluacion).order_by("creada_en"):
        personas[p.proponente_id].append(p)
    aportados = defaultdict(list)
    for d in DocumentoAportado.objects.filter(evaluacion=evaluacion).select_related("persona", "subido_por"):
        aportados[(d.proponente_id, d.requisito)].append(d)

    doc = Document()
    seccion = doc.sections[0]
    seccion.orientation = WD_ORIENT.PORTRAIT
    seccion.left_margin = seccion.right_margin = Cm(2.2)
    seccion.top_margin = seccion.bottom_margin = Cm(2)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(10.5)

    encabezado = seccion.header.paragraphs[0]
    encabezado.text = f"MiEvaluador · Reporte de evaluación {tipo} · {proceso.codigo}"
    encabezado.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in encabezado.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = GRIS
    pie = seccion.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pie.add_run(f"{COPYRIGHT.format(anio=ahora.year)} · Página ")
    run.font.size = Pt(8)
    run.font.color.rgb = GRIS
    _campo_pagina(pie)

    # Portada
    _parrafo(doc, "MiEvaluador", negrita=True, tam=13, color=AZUL_MARCA, alineacion=WD_ALIGN_PARAGRAPH.LEFT, espacio=0)
    _parrafo(doc, "by LeMarTek", tam=9, color=GRIS, alineacion=WD_ALIGN_PARAGRAPH.LEFT, espacio=18)
    titulo = f"REPORTE DE EVALUACIÓN {evaluacion.get_tipo_display().upper()}"
    _parrafo(doc, titulo + ("" if aprobada else " — BORRADOR"), negrita=True, tam=18, color=AZUL, alineacion=WD_ALIGN_PARAGRAPH.LEFT, espacio=2)
    _parrafo(doc, f"Proceso de selección No. {proceso.codigo}", negrita=True, tam=13, alineacion=WD_ALIGN_PARAGRAPH.LEFT, espacio=2)
    _parrafo(doc, evaluacion.entidad.nombre, tam=12, color=GRIS, alineacion=WD_ALIGN_PARAGRAPH.LEFT, espacio=16)

    # 1. Objeto
    _titulo(doc, "1. Objeto del reporte")
    _parrafo(
        doc,
        f"El presente reporte consolida el resultado de la verificación de los requisitos habilitantes de carácter {caracter} "
        f"de las ofertas presentadas dentro del proceso de selección No. {proceso.codigo} de {evaluacion.entidad.nombre}, "
        f"cuyo objeto es: «{documento_base.objeto_general.strip().rstrip('.')}», con fecha de cierre del "
        f"{fecha_larga(documento_base.fecha_cierre)}.",
    )
    if documento_base.lotes:
        _tabla(
            doc,
            ["Lote", "Objeto", "Presupuesto oficial"],
            [[l.numero, l.objeto, pesos(l.valor_presupuesto)] for l in documento_base.lotes],
            [2.5, 10.5, 3.6],
        )
        g = documento_base.garantia_seriedad
        _parrafo(
            doc,
            f"Garantía de seriedad exigida: {g.porcentaje:.0%} del valor base, con vigencia mínima hasta el "
            f"{fecha_larga(g.fecha_vencimiento)} ({g.vigencia_meses} meses contados desde el cierre).",
            tam=9.5,
        )

    # 2. Responsables
    _titulo(doc, "2. Responsables")
    _tabla(
        doc,
        ["Actuación", "Persona", "Fecha"],
        [
            ["Creación del proceso en MiEvaluador", proceso.creado_por.nombre_completo, fecha_larga(proceso.creado_en)],
            [
                "Asignación de la evaluación",
                evaluacion.asignada_por.nombre_completo if evaluacion.asignada_por else "—",
                fecha_larga(evaluacion.asignada_en),
            ],
            ["Evaluador responsable", evaluacion.responsable.nombre_completo if evaluacion.responsable else "Sin asignar", "—"],
            [
                "Aprobación de la evaluación",
                evaluacion.aprobada_por.nombre_completo if evaluacion.aprobada_por else "Pendiente de aprobación",
                fecha_larga(evaluacion.aprobada_en),
            ],
        ],
        [6.2, 6.2, 4.2],
    )

    # 3. Metodología
    _titulo(doc, "3. Metodología")
    version = (
        f"la versión {evaluacion.plantilla.version} de la plantilla de evaluación «{evaluacion.plantilla.nombre}» de la entidad"
        if evaluacion.plantilla_id
        else "la plantilla de evaluación base del sistema"
    )
    for texto in (
        f"La verificación se realizó con apoyo de MiEvaluador, aplicando {version}, que comprende {len(catalogo)} requisitos. "
        "Para cada requisito, el sistema identificó el documento correspondiente dentro de la oferta de cada proponente y "
        "verificó las condiciones exigidas.",
        "Los requisitos que el sistema confirmó por sí solo se identifican como «Aprobado automáticamente por MiEvaluador», "
        "junto con la razón y el documento soporte. Los casos que el sistema no pudo confirmar con certeza fueron revisados "
        "por el evaluador, quien registró su decisión y justificación; se identifican como «Validado manualmente por», con el "
        "nombre de quien validó y la fecha.",
        "Cuando el proponente no aportó un certificado de antecedentes, el evaluador lo consultó en la fuente oficial y lo "
        "incorporó al expediente, indicando la persona consultada y la fecha de expedición del certificado.",
    ):
        _parrafo(doc, texto)
    _seccion_pliego(doc, proceso)

    # 4. Resumen
    proponentes = list(proceso.proponentes.all().order_by("numero_orden"))
    veredictos = {}
    filas_resumen = []
    for p in proponentes:
        estados = []
        no_cumple = []
        for n in orden:
            r = resultados[p.id].get(n)
            if r is None:
                estados.append("Pendiente")
                continue
            v = validacion(r, revisiones.get((p.id, n)), info[n]["verifica"])
            estados.append(v.estado)
            if v.estado == "No cumple":
                no_cumple.append(info[n]["corto"])
        if no_cumple:
            veredicto = "NO HABILITADO"
        elif "Pendiente" in estados:
            veredicto = "PENDIENTE"
        else:
            veredicto = "HABILITADO"
        veredictos[p.id] = veredicto
        filas_resumen.append([p.hoja, p.nombre, veredicto, ", ".join(no_cumple) or "—"])
    _titulo(doc, "4. Resumen de resultados")
    conteo = {k: sum(1 for v in veredictos.values() if v == k) for k in ("HABILITADO", "NO HABILITADO", "PENDIENTE")}
    _parrafo(
        doc,
        f"Se evaluaron {len(proponentes)} proponentes: {conteo['HABILITADO']} habilitados, {conteo['NO HABILITADO']} no "
        f"habilitados" + (f" y {conteo['PENDIENTE']} con requisitos pendientes de verificación." if conteo["PENDIENTE"] else "."),
    )
    _tabla(doc, ["No.", "Proponente", "Resultado", "Requisitos que no cumple"], filas_resumen, [1.6, 7.4, 3.2, 4.4])

    # 5. Detalle
    _titulo(doc, "5. Detalle por proponente")
    for p in proponentes:
        _titulo(doc, f"{p.hoja} · {p.nombre}", nivel=2)
        lista = resultados[p.id]
        tipo_proponente = next((r.tipo_proponente for r in lista.values() if r.tipo_proponente), None)
        representante = next((r.representante_legal for r in lista.values() if r.representante_legal), None)
        _parrafo(
            doc,
            f"Tipo de proponente: {TIPOS_PROPONENTE.get(tipo_proponente, 'no determinado')}. "
            f"Representante legal: {representante or 'no identificado'}. Resultado: {veredictos[p.id]}.",
            negrita=True,
            tam=10,
        )
        if personas[p.id]:
            _parrafo(doc, "Personas verificadas (antecedentes):", tam=9.5, espacio=2)
            _tabla(
                doc,
                ["Rol", "Nombre", "Documento", "Expedición"],
                [
                    [
                        per.get_rol_display() + (f" de {per.de.nombre}" if per.de_id else ""),
                        per.nombre,
                        ("NIT " if per.tipo == "juridica" else "C.C. ") + per.documento,
                        fecha_larga(per.fecha_expedicion_documento) if per.fecha_expedicion_documento else "—",
                    ]
                    for per in personas[p.id]
                ],
                [4.6, 5.6, 3.4, 3.0],
            )
        filas = []
        for n in orden:
            r = lista.get(n)
            if r is None:
                filas.append([str(n), info[n]["titulo"], "Pendiente", "No evaluado."])
                continue
            v = validacion(r, revisiones.get((p.id, n)), info[n]["verifica"])
            detalle = f"{v.forma}. {v.detalle}"
            if n in por_pliego:
                detalle += f" {por_pliego[n]}"
            for d in aportados[(p.id, n)]:
                detalle += (
                    f" Certificado consultado y aportado por {d.subido_por.nombre_completo if d.subido_por else 'el evaluador'}"
                    f" el {fecha_larga(d.subido_en)}"
                    + (f" para {d.persona.nombre} ({'NIT' if d.persona.tipo == 'juridica' else 'C.C.'} {d.persona.documento})" if d.persona else "")
                    + f", expedido el {fecha_larga(d.fecha_expedicion)}."
                    + (f" {d.observacion}" if d.observacion else "")
                )
            filas.append([str(n), info[n]["titulo"], v.estado, detalle])
        _tabla(doc, ["No.", "Requisito", "Resultado", "Forma de validación y justificación"], filas, [1.0, 4.4, 2.2, 9.0])

    # 6. Constancia
    _titulo(doc, "6. Constancia")
    _parrafo(
        doc,
        f"Reporte generado por MiEvaluador el {fecha_larga(ahora)}. "
        + (
            f"La evaluación fue aprobada por {evaluacion.aprobada_por.nombre_completo} el {fecha_larga(evaluacion.aprobada_en)}."
            if aprobada and evaluacion.aprobada_por
            else "La evaluación aún no ha sido aprobada: este documento es un borrador y no constituye el informe definitivo."
        ),
    )
    doc.add_paragraph()
    firmas = doc.add_table(rows=2, cols=2)
    for i, (rol, persona) in enumerate(
        (
            ("Elaboró", evaluacion.responsable.nombre_completo if evaluacion.responsable else ""),
            ("Aprobó", evaluacion.aprobada_por.nombre_completo if evaluacion.aprobada_por else ""),
        )
    ):
        firmas.rows[0].cells[i].text = "\n\n______________________________"
        firmas.rows[1].cells[i].text = f"{rol}: {persona}"

    buffer = io.BytesIO()
    doc.save(buffer)
    nombre = f"REPORTE EVALUACION {evaluacion.tipo.upper()} {proceso.codigo}{'' if aprobada else ' (BORRADOR)'}.docx"
    return buffer.getvalue(), nombre
