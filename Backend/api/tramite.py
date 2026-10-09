"""Traslado del informe de evaluación: su término en días hábiles y la matriz
de observaciones con sus respuestas (RF-16).

Lo ve quien ve la evaluación; lo registra y responde el comité o quien la
gestiona. Cada cambio queda en la auditoría con quién y cuándo, y la persona
que decidió queda en el propio registro."""
from __future__ import annotations

import io
from datetime import date
from uuid import UUID

from django.db import transaction
from django.db.models import Max
from django.http import FileResponse, HttpRequest
from django.shortcuts import get_object_or_404
from django.utils import timezone
from ninja import Router, Schema
from ninja.errors import HttpError

from api.evaluaciones import _evaluacion
from cuentas.models import Usuario
from cuentas.seguridad import auditar, sesion_activa
from evaluaciones.calendario import dias_habiles_entre, sumar_dias_habiles
from evaluaciones.models import (
    DecisionObservacion,
    Evaluacion,
    ObservacionInforme,
    Traslado,
)
from evaluaciones.permisos import exigir_compromiso, puede_trabajar

router = Router(auth=sesion_activa, tags=["tramite"])


class TrasladoIn(Schema):
    publicado_en: date
    dias_habiles: int


class ObservacionIn(Schema):
    observante: str
    recibida_en: date
    texto: str
    proponente_id: UUID | None = None
    requisito: int | None = None


class RespuestaObservacionIn(Schema):
    respuesta: str
    decision: str
    modifica_resultado: bool = False


def _para_registrar(request: HttpRequest, evaluacion_id: UUID) -> Evaluacion:
    evaluacion = _evaluacion(request.auth, evaluacion_id)
    if not puede_trabajar(request.auth, evaluacion):
        raise HttpError(403, "Solo el comité evaluador o quien gestiona la evaluación registra el traslado.")
    exigir_compromiso(request.auth)
    return evaluacion


def _nombre(u: Usuario | None) -> str | None:
    return u.nombre_completo if u else None


def _observacion_out(o: ObservacionInforme, traslado: Traslado | None) -> dict:
    return {
        "id": str(o.id), "consecutivo": o.consecutivo, "observante": o.observante,
        "proponente_id": str(o.proponente_id) if o.proponente_id else None, "proponente": o.proponente.nombre if o.proponente_id else None,
        "requisito": o.requisito, "recibida_en": o.recibida_en.isoformat(), "texto": o.texto,
        "extemporanea": bool(traslado and o.recibida_en > traslado.vence_en),
        "respuesta": o.respuesta, "decision": o.decision, "modifica_resultado": o.modifica_resultado,
        "respondida_por": _nombre(o.respondida_por), "respondida_en": o.respondida_en.isoformat() if o.respondida_en else None,
        "registrada_por": _nombre(o.registrada_por), "registrada_en": o.registrada_en.isoformat(),
    }


def _detalle(evaluacion: Evaluacion, usuario: Usuario) -> dict:
    traslado = Traslado.objects.filter(evaluacion=evaluacion).first()
    hoy = timezone.localdate()
    observaciones = list(evaluacion.observaciones.select_related("proponente", "respondida_por", "registrada_por"))
    return {
        "puede_registrar": puede_trabajar(usuario, evaluacion),
        "traslado": None if traslado is None else {
            "publicado_en": traslado.publicado_en.isoformat(), "dias_habiles": traslado.dias_habiles, "vence_en": traslado.vence_en.isoformat(),
            "dias_habiles_restantes": dias_habiles_entre(hoy, traslado.vence_en), "vencido": hoy > traslado.vence_en,
        },
        "proponentes": [{"id": str(p.id), "nombre": p.nombre, "hoja": p.hoja} for p in evaluacion.proceso.proponentes.all()],
        "observaciones": [_observacion_out(o, traslado) for o in observaciones],
        "resumen": {
            "observaciones": len(observaciones),
            "sin_responder": sum(1 for o in observaciones if o.decision == DecisionObservacion.PENDIENTE),
        },
    }


def _texto(valor: str, que: str, largo: int = 8000) -> str:
    limpio = (valor or "").strip()
    if not limpio:
        raise HttpError(400, f"Escriba {que}.")
    return limpio[:largo]


@router.get("/{evaluacion_id}")
def ver(request: HttpRequest, evaluacion_id: UUID) -> dict:
    return _detalle(_evaluacion(request.auth, evaluacion_id), request.auth)


@router.put("/{evaluacion_id}/traslado")
def fijar_traslado(request: HttpRequest, evaluacion_id: UUID, datos: TrasladoIn) -> dict:
    evaluacion = _para_registrar(request, evaluacion_id)
    if not 1 <= datos.dias_habiles <= 30:
        raise HttpError(400, "El traslado debe durar entre 1 y 30 días hábiles.")
    vence = sumar_dias_habiles(datos.publicado_en, datos.dias_habiles)
    Traslado.objects.update_or_create(
        evaluacion=evaluacion,
        defaults={"entidad_id": evaluacion.entidad_id, "publicado_en": datos.publicado_en, "dias_habiles": datos.dias_habiles,
                  "vence_en": vence, "actualizado_por": request.auth},
    )
    auditar(request, "tramite.traslado", entidad_id=evaluacion.entidad_id, objeto=evaluacion,
            publicado_en=datos.publicado_en.isoformat(), dias_habiles=datos.dias_habiles, vence_en=vence.isoformat())
    return _detalle(evaluacion, request.auth)


@router.post("/{evaluacion_id}/observaciones")
def crear_observacion(request: HttpRequest, evaluacion_id: UUID, datos: ObservacionIn) -> dict:
    evaluacion = _para_registrar(request, evaluacion_id)
    proponente = get_object_or_404(evaluacion.proceso.proponentes, pk=datos.proponente_id) if datos.proponente_id else None
    with transaction.atomic():
        siguiente = (ObservacionInforme.objects.select_for_update().filter(evaluacion=evaluacion).aggregate(m=Max("consecutivo"))["m"] or 0) + 1
        o = ObservacionInforme.objects.create(
            entidad_id=evaluacion.entidad_id, evaluacion=evaluacion, consecutivo=siguiente,
            observante=_texto(datos.observante, "quién presenta la observación", 300), proponente=proponente,
            requisito=datos.requisito, recibida_en=datos.recibida_en, texto=_texto(datos.texto, "la observación"),
            registrada_por=request.auth,
        )
    auditar(request, "tramite.observacion", entidad_id=evaluacion.entidad_id, objeto=o, observante=o.observante,
            recibida_en=o.recibida_en.isoformat())
    return _detalle(evaluacion, request.auth)


@router.put("/{evaluacion_id}/observaciones/{observacion_id}")
def responder_observacion(request: HttpRequest, evaluacion_id: UUID, observacion_id: UUID, datos: RespuestaObservacionIn) -> dict:
    evaluacion = _para_registrar(request, evaluacion_id)
    o = get_object_or_404(ObservacionInforme, pk=observacion_id, evaluacion=evaluacion)
    if datos.decision not in DecisionObservacion.values or datos.decision == DecisionObservacion.PENDIENTE:
        raise HttpError(400, "Indique si la observación se acoge, se acoge parcialmente o no se acoge.")
    anterior = {"respuesta": o.respuesta, "decision": o.decision} if o.respondida_en else None
    o.respuesta = _texto(datos.respuesta, "la respuesta")
    o.decision, o.modifica_resultado = datos.decision, datos.modifica_resultado
    o.respondida_por, o.respondida_en = request.auth, timezone.now()
    o.save()
    # Si la respuesta se cambia, la versión anterior queda en la auditoría.
    auditar(request, "tramite.observacion_respuesta", entidad_id=evaluacion.entidad_id, objeto=o, decision=o.decision,
            modifica_resultado=o.modifica_resultado, respuesta=o.respuesta, anterior=anterior)
    return _detalle(evaluacion, request.auth)


@router.get("/{evaluacion_id}/matriz")
def matriz_observaciones(request: HttpRequest, evaluacion_id: UUID):
    """Matriz de observaciones al informe, en Excel."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    evaluacion = _evaluacion(request.auth, evaluacion_id)
    datos = _detalle(evaluacion, request.auth)
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Observaciones"
    titulo = f"Proceso {evaluacion.proceso.codigo} · evaluación {evaluacion.get_tipo_display().lower()}"
    nombres = {"pendiente": "Pendiente", "acoge": "Se acoge", "acoge_parcial": "Se acoge parcialmente", "no_acoge": "No se acoge"}

    def tabla(h, encabezado, filas, anchos):
        h.append([titulo])
        h["A1"].font = Font(bold=True, size=12)
        h.append([])
        h.append(encabezado)
        for celda in h[3]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="00143C")
        for fila in filas:
            h.append(fila)
        for i, ancho in enumerate(anchos):
            h.column_dimensions[chr(65 + i)].width = ancho
        for fila in h.iter_rows(min_row=4):
            for celda in fila:
                celda.alignment = Alignment(wrap_text=True, vertical="top")

    tabla(hoja, ["N.º", "Recibida", "Observante", "Proponente", "Requisito", "Observación", "Extemporánea", "Respuesta", "Decisión",
                 "Modifica el resultado", "Respondida por", "Respondida el"], [
        [o["consecutivo"], o["recibida_en"], o["observante"], o["proponente"] or "", o["requisito"] or "", o["texto"],
         "Sí" if o["extemporanea"] else "No", o["respuesta"], nombres[o["decision"]], "Sí" if o["modifica_resultado"] else "No",
         o["respondida_por"] or "", (o["respondida_en"] or "")[:10]]
        for o in datos["observaciones"]
    ], [6, 12, 26, 26, 10, 60, 13, 60, 18, 12, 24, 13])
    salida = io.BytesIO()
    libro.save(salida)
    auditar(request, "tramite.matriz", entidad_id=evaluacion.entidad_id, objeto=evaluacion)
    return FileResponse(
        io.BytesIO(salida.getvalue()), as_attachment=True, filename=f"Matriz de observaciones - {evaluacion.proceso.codigo}.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
