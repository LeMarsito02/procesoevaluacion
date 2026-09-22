"""Informe Excel de la evaluación técnica, con la forma del consolidado que
usan las entidades: por lote, la experiencia de cada proponente y su
puntaje (factor de calidad, industria nacional, discapacidad, mujeres,
MIPYME y reducción), más el detalle de los contratos evaluados.

Lo que ni el programa verificó ni una persona decidió sale como
"PENDIENTE": el informe nunca presenta como definitivo algo sin revisar.
"""
from __future__ import annotations

import io
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from motor.esquemas.proceso import ResultadoRequisito

PENDIENTE = "PENDIENTE"
# Columnas de puntaje: (título, números internos que suman en la columna).
COLUMNAS_PUNTAJE: list[tuple[str, tuple[int, ...]]] = [
    ("FACTOR DE CALIDAD", (121, 122, 123)),
    ("APOYO A LA INDUSTRIA NACIONAL", (124,)),
    ("PERSONAL CON DISCAPACIDAD", (125,)),
    ("E y E DE MUJERES", (126,)),
    ("MIPYME", (127,)),
]
OBRAS_INCONCLUSAS = 130

_BORDE = Border(*(Side(style="thin", color="999999"),) * 4)
_ENCABEZADO = PatternFill("solid", fgColor="1F3864")
_PENDIENTE = PatternFill("solid", fgColor="FFF2CC")


@dataclass
class ResultadoInforme:
    resultado: ResultadoRequisito
    # Una persona tomó la decisión (prevalece sobre el programa).
    decidido: bool


def _puntos(r: ResultadoInforme | None) -> float | str:
    if r is None:
        return PENDIENTE
    maximo = float((r.resultado.detalle or {}).get("puntaje_maximo") or 0)
    if r.resultado.cumple:
        return maximo
    return 0.0 if r.decidido else PENDIENTE


def _experiencia(r: ResultadoInforme | None) -> str:
    if r is None:
        return PENDIENTE
    if r.resultado.cumple:
        return "CUMPLE"
    return "NO CUMPLE" if r.decidido else PENDIENTE


def _reduccion(r: ResultadoInforme | None) -> float | str:
    if r is None:
        return PENDIENTE
    if r.resultado.cumple:
        return 0.0
    return -1.0 if r.decidido else PENDIENTE


def _encabezado(hoja, fila: int, titulos: list[str]) -> None:
    for j, titulo in enumerate(titulos, 1):
        celda = hoja.cell(row=fila, column=j, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = _ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = _BORDE


def _fila(hoja, fila: int, valores: list) -> None:
    for j, valor in enumerate(valores, 1):
        celda = hoja.cell(row=fila, column=j, value=valor)
        celda.border = _BORDE
        celda.alignment = Alignment(vertical="top", wrap_text=j == 2)
        if valor == PENDIENTE:
            celda.fill = _PENDIENTE


def generar_informe(
    codigo: str,
    objeto: str,
    lotes: list[tuple[int, str]],
    proponentes: list[tuple[str, str]],
    resultados: dict[tuple[str, int], ResultadoInforme],
    numero_por_lote: dict[int, int],
    borrador: bool,
    titulos: dict[int, str] | None = None,
) -> bytes:
    """`lotes`: [(índice, nombre)]; `proponentes`: [(hoja, nombre)];
    `resultados`: {(hoja, número de requisito): resultado};
    `numero_por_lote`: índice del lote -> número del requisito de experiencia."""
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    fila = 1
    resumen.cell(row=fila, column=1, value=f"EVALUACIÓN TÉCNICA - {codigo}" + (" (BORRADOR)" if borrador else "")).font = Font(bold=True, size=13)
    fila += 1
    resumen.cell(row=fila, column=1, value=objeto).alignment = Alignment(wrap_text=True)
    fila += 2
    encabezados = ["PROP.", "NOMBRE DEL PROPONENTE", "EXPERIENCIA DEL PROPONENTE", *(t for t, _ in COLUMNAS_PUNTAJE), "REDUCCIÓN DE PUNTAJE", "PUNTAJE PRELIMINAR"]
    for indice, nombre in lotes:
        resumen.cell(row=fila, column=1, value=nombre.upper()).font = Font(bold=True)
        fila += 1
        _encabezado(resumen, fila, encabezados)
        fila += 1
        for hoja, nombre_proponente in proponentes:
            experiencia = _experiencia(resultados.get((hoja, numero_por_lote.get(indice, -1))))
            puntos = []
            for _, numeros in COLUMNAS_PUNTAJE:
                partes = [_puntos(resultados.get((hoja, n))) for n in numeros]
                puntos.append(PENDIENTE if PENDIENTE in partes else sum(partes))
            reduccion = _reduccion(resultados.get((hoja, OBRAS_INCONCLUSAS)))
            numericos = [p for p in (*puntos, reduccion) if p != PENDIENTE]
            total = sum(numericos)
            completo = len(numericos) == len(puntos) + 1
            _fila(resumen, fila, [hoja, nombre_proponente, experiencia, *puntos, reduccion, total if completo else f"{total:g} + {PENDIENTE}"])
            fila += 1
        fila += 1
    anchos = [8, 46, 16, 14, 16, 14, 12, 10, 14, 16]
    for j, ancho in enumerate(anchos, 1):
        resumen.column_dimensions[get_column_letter(j)].width = ancho

    detalle = libro.create_sheet("Detalle experiencia")
    columnas = ["PROP.", "LOTE", "ORDEN", "CONSECUTIVO RUP", "CONTRATANTE", "No. CONTRATO", "VALOR RUP (SMMLV)",
                "PARTICIPACIÓN", "VALOR APORTADO (SMMLV)", "UNSPSC", "LONGITUD (km)", "OBSERVACIONES"]
    _encabezado(detalle, 1, columnas)
    fila = 2
    for hoja, _ in proponentes:
        for indice, nombre in lotes:
            r = resultados.get((hoja, numero_por_lote.get(indice, -1)))
            for c in ((r.resultado.detalle or {}).get("contratos", []) if r else []):
                _fila(detalle, fila, [
                    hoja, nombre, c.get("orden"), ", ".join(c.get("consecutivos") or []), c.get("contratante"),
                    c.get("numero_contrato"), c.get("valor_smmlv"),
                    None if c.get("participacion") is None else round(c["participacion"] * 100, 2),
                    None if c.get("valor_aportado") is None else round(c["valor_aportado"], 2),
                    {True: "SI", False: "NO"}.get(c.get("unspsc"), ""),
                    c.get("longitud_km"), "; ".join(c.get("problemas") or []),
                ])
                fila += 1
    for j, ancho in enumerate([8, 10, 7, 14, 36, 22, 14, 12, 16, 9, 12, 60], 1):
        detalle.column_dimensions[get_column_letter(j)].width = ancho

    observaciones = libro.create_sheet("Observaciones")
    _encabezado(observaciones, 1, ["PROP.", "REQUISITO", "RESULTADO", "OBSERVACIÓN"])
    fila = 2
    for (hoja, numero), r in sorted(resultados.items()):
        texto = r.resultado.motivo or r.resultado.error
        if not texto:
            continue
        estado = "CUMPLE" if r.resultado.cumple else ("NO CUMPLE" if r.decidido else PENDIENTE)
        _fila(observaciones, fila, [hoja, (titulos or {}).get(numero, str(numero)), estado, texto])
        fila += 1
    for j, ancho in enumerate([8, 40, 13, 110], 1):
        observaciones.column_dimensions[get_column_letter(j)].width = ancho

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
