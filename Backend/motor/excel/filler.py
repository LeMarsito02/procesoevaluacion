from __future__ import annotations

import io
import re

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from pydantic import BaseModel, Field

from motor.esquemas.proceso import ProcesoDocumentoBase, Proponente, ResultadoRequisito

DATOS_SHEET_NAME = "DATOS PROCESO"

# Fila (en cada hoja P-XX) donde va la columna L (SI/NO/REVISAR) y N (archivo/
# motivo) de cada requisito, según la plantilla jurídica de ejemplo.
FILA_POR_REQUISITO: dict[int, int] = {
    1: 29,
    2: 33,
    3: 37,
    4: 41,
    5: 45,
    6: 49,
    7: 53,
    8: 57,
    9: 61,
    10: 65,
    11: 69,
    12: 73,
    13: 77,
    14: 81,
    15: 85,
    16: 89,
    17: 93,
    18: 97,
}


class MapeoPlantilla(BaseModel):
    """Dónde escribe el informe dentro de una plantilla de Excel. Cada entidad
    sube su propia plantilla y ajusta este mapeo; los valores por defecto son
    los de la plantilla jurídica de ejemplo."""

    prefijo_codigo: str = Field("", description="Sigla de la entidad antepuesta al código del proceso")
    titulo: str = Field("EVALUACIÓN {tipo} - PROCESO {codigo}{de_anio}", description="Título de cada hoja de proponente")
    hoja_proponente_patron: str = Field(r"^P-\d+$", description="Expresión regular del nombre de las hojas por proponente")
    hoja_modelo: str | None = "MODELO"
    celda_titulo: str | None = "F3"
    celda_objeto: str | None = "C6"
    celda_nombre_proponente: str | None = "E10"
    columna_resultado: str = "L"
    columna_detalle: str = "N"
    filas_por_requisito: dict[int, int] = Field(default_factory=lambda: dict(FILA_POR_REQUISITO))
    hoja_resumen: str | None = "RESUMEN"
    celda_resumen: str | None = "B4"
    texto_cumple: str = "SI"
    texto_no_cumple: str = "NO"
    texto_revisar: str = "REVISAR"
    agregar_hoja_datos: bool = True


MAPEO_POR_DEFECTO = MapeoPlantilla()


def _codigo_y_anio(codigo_proceso: str) -> tuple[str, str]:
    match = re.match(r"^(.*)-(\d{4})$", codigo_proceso.strip())
    if match:
        return match.group(1), match.group(2)
    return codigo_proceso.strip(), ""


def _codigo_completo(codigo_proceso: str, prefijo: str = "") -> str:
    codigo_sin_anio, _ = _codigo_y_anio(codigo_proceso)
    return f"{prefijo}-{codigo_sin_anio}" if prefijo else codigo_sin_anio


def _titulo_proceso(codigo_proceso: str, mapeo: MapeoPlantilla, tipo: str) -> str:
    _, anio = _codigo_y_anio(codigo_proceso)
    return mapeo.titulo.format(
        tipo=tipo.upper(),
        codigo=_codigo_completo(codigo_proceso, mapeo.prefijo_codigo),
        anio=anio,
        de_anio=f" DE {anio}" if anio else "",
    )


def _objeto_con_lotes(proceso: ProcesoDocumentoBase) -> str:
    partes = []
    objeto_general = proceso.objeto_general.strip()
    if objeto_general:
        partes.append(objeto_general.rstrip(". ") + ":")
    for lote in proceso.lotes:
        partes.append(f"{lote.numero}: {lote.objeto}")
    return "\n\n".join(partes)


def _formato_pesos(valor: float) -> str:
    return f"$ {valor:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def _write_datos_proceso_sheet(
    wb: openpyxl.Workbook, proceso: ProcesoDocumentoBase, proponentes: list[Proponente], prefijo: str = ""
) -> None:
    if DATOS_SHEET_NAME in wb.sheetnames:
        del wb[DATOS_SHEET_NAME]
    ws = wb.create_sheet(DATOS_SHEET_NAME, 0)

    bold = Font(bold=True)
    title_font = Font(bold=True, size=13)
    wrap = Alignment(wrap_text=True, vertical="top")

    ws["A1"] = "DATOS DEL PROCESO (extraídos del Documento Base)"
    ws["A1"].font = title_font

    ws["A3"] = "Código del proceso"
    ws["B3"] = _codigo_completo(proceso.codigo_proceso, prefijo)
    ws["A4"] = "Fecha de cierre"
    ws["B4"] = proceso.fecha_cierre.strftime("%d/%m/%Y")
    ws["A5"] = "Objeto general"
    ws["B5"] = proceso.objeto_general
    ws["B5"].alignment = wrap

    for coord in ("A3", "A4", "A5"):
        ws[coord].font = bold

    headers = ["Lote", "Objeto", "Plazo (meses)", "Valor Presupuesto Oficial", "Lugar de ejecución"]
    header_row = 7
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = bold

    row = header_row + 1
    for lote in proceso.lotes:
        ws.cell(row=row, column=1, value=lote.numero)
        objeto_cell = ws.cell(row=row, column=2, value=lote.objeto)
        objeto_cell.alignment = wrap
        ws.cell(row=row, column=3, value=lote.plazo_meses)
        valor_cell = ws.cell(row=row, column=4, value=lote.valor_presupuesto)
        valor_cell.number_format = '"$" #,##0.00'
        ws.cell(row=row, column=5, value=lote.lugar_ejecucion or "")
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Lote de mayor valor").font = bold
    ws.cell(row=row, column=2, value=proceso.lote_mayor_valor)
    row += 1
    ws.cell(row=row, column=1, value="Presupuesto Oficial total").font = bold
    total_cell = ws.cell(row=row, column=2, value=proceso.presupuesto_total)
    total_cell.number_format = '"$" #,##0.00'

    row += 2
    ws.cell(row=row, column=1, value="GARANTÍA DE SERIEDAD DE LA OFERTA").font = title_font
    row += 1
    g = proceso.garantia_seriedad
    campos = [
        ("Base de cálculo", "Lote de mayor valor" if g.base_calculo == "lote_mayor_valor" else "Presupuesto Oficial total"),
        ("Lote base", g.lote_base or "-"),
        ("Valor base", g.valor_base),
        ("Porcentaje asegurado", f"{g.porcentaje:.0%}"),
        ("Valor asegurado requerido", g.valor_asegurado),
        ("Fecha de cierre", g.fecha_cierre.strftime("%d/%m/%Y")),
        ("Vigencia requerida", f"{g.vigencia_meses} meses contados desde la fecha de cierre"),
        ("Fecha de vencimiento mínima de la garantía", g.fecha_vencimiento.strftime("%d/%m/%Y")),
    ]
    for label, value in campos:
        ws.cell(row=row, column=1, value=label).font = bold
        cell = ws.cell(row=row, column=2, value=value)
        if label in ("Valor base", "Valor asegurado requerido"):
            cell.number_format = '"$" #,##0.00'
        row += 1

    if proponentes:
        row += 1
        ws.cell(row=row, column=1, value="PROPONENTES (desde Google Drive)").font = title_font
        row += 1
        for header, col in (("No.", 1), ("Hoja", 2), ("Proponente", 3), ("Archivo en Drive", 4)):
            ws.cell(row=row, column=col, value=header).font = bold
        row += 1
        for proponente in proponentes:
            ws.cell(row=row, column=1, value=proponente.numero_orden)
            ws.cell(row=row, column=2, value=proponente.hoja)
            ws.cell(row=row, column=3, value=proponente.nombre_proponente)
            ws.cell(row=row, column=4, value=proponente.nombre_archivo)
            row += 1

    if proceso.advertencias:
        row += 1
        ws.cell(row=row, column=1, value="ADVERTENCIAS DE EXTRACCIÓN").font = title_font
        row += 1
        for advertencia in proceso.advertencias:
            cell = ws.cell(row=row, column=1, value=f"- {advertencia}")
            cell.alignment = wrap
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            row += 1

    column_widths = [28, 55, 16, 24, 22]
    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def fill_template(
    template_path: str,
    proceso: ProcesoDocumentoBase,
    proponentes: list[Proponente] | None = None,
    resultados: list[ResultadoRequisito] | None = None,
    mapeo: MapeoPlantilla | None = None,
    tipo: str = "Jurídica",
) -> bytes:
    mapeo = mapeo or MAPEO_POR_DEFECTO
    wb = openpyxl.load_workbook(template_path)
    hoja_re = re.compile(mapeo.hoja_proponente_patron)

    titulo = _titulo_proceso(proceso.codigo_proceso, mapeo, tipo)
    objeto_texto = _objeto_con_lotes(proceso)

    hojas_a_actualizar = [name for name in wb.sheetnames if name == mapeo.hoja_modelo or hoja_re.match(name)]
    for name in hojas_a_actualizar:
        ws = wb[name]
        if mapeo.celda_titulo:
            ws[mapeo.celda_titulo] = titulo
        if mapeo.celda_objeto:
            ws[mapeo.celda_objeto] = objeto_texto

    if mapeo.hoja_resumen and mapeo.celda_resumen and mapeo.hoja_resumen in wb.sheetnames:
        resumen = wb[mapeo.hoja_resumen]
        codigo_completo = _codigo_completo(proceso.codigo_proceso, mapeo.prefijo_codigo)
        _, anio = _codigo_y_anio(proceso.codigo_proceso)
        actual = str(resumen[mapeo.celda_resumen].value or "")
        match = re.search(r" - (.*)$", actual)
        sufijo = match.group(1) if match else f"PRIMER INFORME DE EVALUACIÓN {tipo.upper()}"
        resumen[mapeo.celda_resumen] = f"PROCESO DE SELECCIÓN No. {codigo_completo} DE {anio} - {sufijo}"

    if mapeo.celda_nombre_proponente:
        for proponente in proponentes or []:
            if proponente.hoja in wb.sheetnames:
                wb[proponente.hoja][mapeo.celda_nombre_proponente] = proponente.nombre_proponente

    for resultado in resultados or []:
        if resultado.hoja not in wb.sheetnames:
            continue
        fila = mapeo.filas_por_requisito.get(resultado.requisito)
        if fila is None:
            continue
        ws = wb[resultado.hoja]
        col_l, col_n = f"{mapeo.columna_resultado}{fila}", f"{mapeo.columna_detalle}{fila}"
        if resultado.error:
            ws[col_l] = mapeo.texto_revisar
            ws[col_n] = f"No se pudo evaluar automáticamente: {resultado.error}"
        elif resultado.cumple:
            ws[col_l] = mapeo.texto_cumple
            ws[col_n] = resultado.archivo_evaluado or resultado.motivo or ""
        else:
            ws[col_l] = mapeo.texto_no_cumple
            archivo = resultado.archivo_evaluado or "no se encontró el documento"
            ws[col_n] = f"{archivo} — NO CUMPLE: {resultado.motivo}"

    if mapeo.agregar_hoja_datos:
        _write_datos_proceso_sheet(wb, proceso, proponentes or [], mapeo.prefijo_codigo)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def inspeccionar_plantilla(contenido: bytes, mapeo: MapeoPlantilla | None = None) -> dict:
    """Revisa una plantilla subida: que abra, cuántas hojas de proponente tiene
    y qué texto hay en la fila de cada requisito (para confirmar el mapeo)."""
    mapeo = mapeo or MAPEO_POR_DEFECTO
    problemas: list[str] = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True)
    except Exception as exc:  # noqa: BLE001
        return {"valida": False, "problemas": [f"No se pudo abrir como Excel (.xlsx): {exc}"], "hojas": [], "hojas_proponente": 0, "filas": {}}
    try:
        hoja_re = re.compile(mapeo.hoja_proponente_patron)
    except re.error as exc:
        return {"valida": False, "problemas": [f"El patrón de hojas no es válido: {exc}"], "hojas": wb.sheetnames, "hojas_proponente": 0, "filas": {}}
    hojas_proponente = [n for n in wb.sheetnames if hoja_re.match(n)]
    if not hojas_proponente:
        problemas.append(
            f"No hay hojas de proponente que coincidan con «{mapeo.hoja_proponente_patron}». Hojas encontradas: "
            + ", ".join(wb.sheetnames[:12])
        )
    filas: dict[int, str] = {}
    if hojas_proponente:
        ws = wb[hojas_proponente[0]]
        maxima = max(mapeo.filas_por_requisito.values(), default=0)
        textos: dict[int, list[str]] = {}
        for fila in ws.iter_rows(min_row=1, max_row=maxima):
            for celda in fila:
                if getattr(celda, "row", None) in set(mapeo.filas_por_requisito.values()) and celda.value not in (None, ""):
                    textos.setdefault(celda.row, []).append(str(celda.value).strip())
        for requisito, fila in sorted(mapeo.filas_por_requisito.items()):
            etiqueta = " · ".join(t for t in textos.get(fila, []) if t)[:160]
            filas[requisito] = etiqueta
            if not etiqueta:
                problemas.append(f"La fila {fila} (requisito {requisito}) está vacía en la hoja {hojas_proponente[0]}.")
    wb.close()
    return {
        "valida": bool(hojas_proponente),
        "problemas": problemas,
        "hojas": wb.sheetnames,
        "hojas_proponente": len(hojas_proponente),
        "filas": filas,
    }
