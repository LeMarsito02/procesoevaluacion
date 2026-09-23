"""Capacidad residual del proponente (pliego 3.11.2):

    CRP = CO × (E + CT + CF) / 100 − SCE

por integrante; la del plural es la suma de la de cada uno, sin ponderar por
participación (una negativa resta).

- E (experiencia): valor de los contratos del segmento 72 inscritos en el
  RUP (en SMMLV × participación, liquidado con el SMMLV del año) ÷
  (presupuesto de los lotes a los que se presenta × participación del
  integrante). Se calcula con el RUP, no con lo que declara el formato.
- CF (capacidad financiera): liquidez del RUP. También se calcula.
- CO (capacidad de organización): mayor ingreso operacional de los últimos
  cinco años (Formato 5D o equivalente).
- CT (capacidad técnica): socios y profesionales de arquitectura, ingeniería
  y geología del Formato 5B; si no se leen, cuenta 0 (lo desfavorable).
- SCE (saldos de contratos en ejecución): del Formato 5C; si dice que no tiene
  contratos en ejecución, 0; si no se lee el total, a revisión.

Solo se aprueba si con esta estimación prudente alcanza.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from motor.financiera.capacidad import Revision
from motor.financiera.integrantes import (
    _agrupado, _del_integrante, _nit_escrito, _nombre_escrito, estados_del_integrante, integrante_del_titular,
    titulares_del_documento,
)
from motor.financiera.parametros import LoteFinanciero
from motor.llm.cliente import consultar_json
from motor.tecnica.experiencia import IntegranteTecnico
from motor.tecnica.rup import normalizar, numero

# La capacidad residual se aprueba sola solo si supera en este factor la
# exigida. La entidad recalcula los saldos de los contratos en ejecución y
# revisa el ingreso operacional con sus soportes; en la evaluación de
# referencia (LP-022) su K llegó a ser menos de la mitad de la calculada con
# lo que declara el proponente. Con menos margen, a revisión.
HOLGURA = 1.5
_E_TABLA = [(3, 60), (6, 80), (10, 100)]  # (≤ tope, puntaje); más de 10: 120
_CF_TABLA = [(0.50, 20), (0.75, 25), (1.00, 30), (1.50, 35)]  # (≤ tope, puntaje); más: 40
_CT_TABLA = [(0, 0), (5, 20), (10, 30)]  # (≤ tope, puntaje); más: 40

_FORMATO5_RE = re.compile(r"CAPACIDAD\s+RESIDUAL|FORMATO\s*(?:N[O°º]\.?\s*)?5(?:\.\d|\s*[A-D])?\b")
# Ingresos operacionales del estado de resultados: la fila trae, tras la nota
# opcional ("13"), el valor de cada año comparado (y a veces la variación).
# Primero las filas de total (las de arriba pueden ser de detalle); "TOTAL
# INGRESOS" a secas no, porque puede sumar los no operacionales.
_VALORES_FILA = r"(?:(?:\(\s*|\b)(?:NOTA\s*)?\d{1,2}(?:\s\d)?\s*\)?\s+)?((?:\$?\s*\(?\d{1,3}(?:[.,]\d{3}){2,}(?:[.,]\d{1,2})?\)?\s*){1,2})"
_INGRESOS_TOTAL_RE = re.compile(
    # "ACTIVI…": también "ACTIVIDAD ORDINARIA" y "ACTIVIADES" (errores de digitación).
    r"(?:TOTAL\s+INGRESOS\s+(?:OPERACIONALES|ORDINARIOS|NETOS|(?:(?:DE|POR)\s+)?(?:LAS\s+)?ACTIVI\w*\s+ORDINARIAS?)"
    r"|INGRESOS\s+(?:(?:DE|POR)\s+(?:LAS\s+)?ACTIVI\w*\s+ORDINARIAS?|OPERACIONALES|ORDINARIOS)\s+NETOS"
    r"|VENTAS\s+NETAS|INGRESOS\s+NETOS\s+OPERACIONALES)[\s:]{0,5}" + _VALORES_FILA
)
_INGRESOS_RE = re.compile(
    r"(?:INGRESOS?\s+(?:(?:DE|POR)\s+(?:LAS\s+)?ACTIVI\w*\s+ORDINARIAS?|OPERACIONALES|ORDINARIOS)|VENTAS\s+NETAS"
    # "VENTAS 10 10.510.259.000": solo con nota o "$" enseguida, y no "COSTO DE VENTAS".
    r"|(?<!COSTO\s)(?<!COSTOS\s)(?<!DE\s)\bVENTAS(?=\s+(?:\(?\s*(?:NOTA\s*)?\d{1,2}\s*\)?\s+|\$)))[^\d$\n]{0,60}?"
    + _VALORES_FILA
)
_CIFRA_RE = re.compile(r"\(?\d{1,3}(?:[.,]\d{3}){2,}(?:[.,]\d{1,2})?\)?")


def _cifra(texto: str) -> float | None:
    """Cifra de un estado financiero, con miles a la colombiana
    ("16.017.439.164") o a la inglesa ("40,850,976,491.79")."""
    t = texto.strip("() ")
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d{1,2})?", t):
        t = t.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d{1,2})?", t):
        t = t.replace(",", "")
    else:
        return None
    return float(t)


def _ingreso_operacional(texto: str) -> float | None:
    """Mayor ingreso operacional de los años que compara el estado de
    resultados (el pliego pide el del año de mayor ingreso)."""
    # El PDF a veces separa dígitos: "$ 1 54,066,247,835", "1 .918.119.166",
    # "$ 3 2 . 7 6 0 . 9 7 9 . 3 4 8".
    texto = re.sub(r"(?<![\d.,])((?:[\d.,] ){3,}[\d.,]+)", lambda m: m.group(1).replace(" ", ""), texto)
    texto = re.sub(r"(\$\s*)(\d)\s(\d{1,2}[.,]\d{3}[.,]\d{3})", r"\1\2\3", texto)
    texto = re.sub(r"(?<![\d.,])(\d{1,2})\s(\.\d{3}\.\d{3})", r"\1\2", texto)
    m = _INGRESOS_TOTAL_RE.search(texto) or _INGRESOS_RE.search(texto)
    if m is None:
        return None
    valores = [v for c in _CIFRA_RE.findall(m.group(1)) if not c.startswith("(") and (v := _cifra(c))]
    return max(valores) if valores else None


_CO_RE = re.compile(r"MAYOR\s+INGRESO\s+OPERACIONAL[^$]{0,220}?\$\s*([\d.,]{5,})")
_SIN_CONTRATOS_RE = re.compile(
    r"NO\s+(?:CUENTA|CUENTO|CONTAMOS|TIENE|TENGO|TENEMOS|POSEE|POSEEMOS|REGISTRA|REGISTRAMOS)\s+(?:CON\s+)?(?:NINGUN\s+)?"
    r"CONTRATOS?\s+(?:DE\s+OBRA\s+)?EN\s+EJECUCION"
    r"|SIN\s+CONTRATOS\s+EN\s+EJECUCION|NO\s+EXISTEN\s+CONTRATOS\s+EN\s+EJECUCION"
)
_PROFESION_RE = re.compile(r"\b(INGENIER[OA]\s+[A-Z]{4,}|ARQUITECT[OA]\b|GEOLOG[OA]\b)")
_INSTRUCCION_SCE = (
    "Este texto es el Formato 5C (o 5.1/5.3) de capacidad residual: el listado de contratos de obra en ejecución de un "
    "proponente, con el saldo por ejecutar en los próximos 12 meses de cada contrato. Di el TOTAL del saldo de los "
    "contratos en ejecución (SCE) en pesos: si el formato trae la fila de total, ese valor; si no, la suma de la "
    "columna del saldo. Si dice que no tiene contratos en ejecución, 0. Responde JSON: "
    '{"sce": número o null, "cita": "texto copiado tal cual donde aparece el total"}'
)


def _puntaje(valor: float | None, tabla: list[tuple[float, int]], maximo: int, indeterminado: int | None = None) -> int:
    if valor is None:
        return indeterminado if indeterminado is not None else 0
    for tope, puntos in tabla:
        if valor <= tope:
            return puntos
    return maximo


@dataclass
class ResidualIntegrante:
    nombre: str
    co: float | None = None
    e: float | None = None
    ct_profesionales: int = 0
    cf_liquidez: float | None = None
    sce: float | None = None
    puntos_e: int = 0
    puntos_ct: int = 0
    puntos_cf: int = 0
    crp: float | None = None
    faltas: list[str] = field(default_factory=list)
    uso_ia: bool = False


def _texto_formato(contenido: bytes, paginas: int = 25, tablas: bool = False) -> str:
    """Texto del formato. La lectura normal primero; la de tablas (OCR fila
    por fila) solo cuando la normal no dejó leer el listado, porque hacerlo
    con todos los documentos de una oferta cuesta el doble de tiempo."""
    from motor.procesamiento.pdf_utils import extraer_texto, texto_completo

    leer = texto_completo if tablas else extraer_texto
    return leer(contenido, max_paginas=paginas)


def _textos_formato5(pdfs: dict[str, bytes], excels: dict[str, bytes], tablas: bool = False) -> dict[str, str]:
    textos = {}
    for archivo, contenido in pdfs.items():
        # "Formato 5 - Capacidad residual", "OBRAS SAS 5C.pdf", "FORMA 5.3".
        if not re.search(r"RESIDUAL|FORMA(?:TO)?\s*5|CAPACIDAD|\b5\s*[A-D]\b|\b5\.[1-4]\b|\bK\b", normalizar(archivo)):
            continue
        try:
            texto = normalizar(" ".join(_texto_formato(contenido, tablas=tablas).split()))
        except Exception:  # noqa: BLE001
            continue
        if _FORMATO5_RE.search(texto):
            textos[archivo] = texto
    for archivo, contenido in excels.items():
        try:
            import openpyxl

            libro = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        except Exception:  # noqa: BLE001
            continue
        for hoja in libro.worksheets:
            filas = [" ".join(str(c) for c in f if c not in (None, "")) for f in hoja.iter_rows(values_only=True, max_row=300)]
            texto = normalizar(" ".join(" ".join(filas).split()))
            if _FORMATO5_RE.search(texto) or re.search(r"CONTRATOS\s+EN\s+EJECUCION|CAPACIDAD\s+TECNICA", texto):
                textos[f"{archivo} ({hoja.title})"] = texto
    return textos


def _numero_escrito(valor: float, texto: str) -> re.Match | None:
    """El número escrito tal cual en el texto, como una cifra completa
    ("441.253.132" o "441,253,132" o "441253132"), no como pedazo de otra."""
    digitos = f"{valor:.0f}"
    grupos = []
    while digitos:
        grupos.insert(0, digitos[-3:])
        digitos = digitos[:-3]
    patron = r"(?<![\d.,])" + r"[.,\s]?".join(grupos) + r"(?:[.,]\d{1,2})?(?![\d.,]*\d)"
    return re.search(patron, texto)


# Título del listado; no una mención dentro de las notas ("…debe incluirse
# en el Formato 5C con el valor total…").
_SECCION_SCE_RE = re.compile(
    r"(?<!EN\sEL\s)(?<!DEL\s)(?<!AL\s)(?<!EL\s)(?<!FORMATO\s)"
    r"(?:FORMATO\s*(?:N[O°º]\.?\s*)?5\s*(?:\.\s*3|-?\s*C)\b|LISTADO\s+DE\s+(?:LOS\s+)?CONTRATOS\s+EN\s+EJECUCION"
    r"|SALDOS?\s+(?:DE\s+)?CONTRATOS\s+EN\s+EJECUCION)"
    # Una frase de las notas sigue con coma o con "cuando…", "y se encuentren…".
    r"(?!\s*,)(?!\s+(?:Y\s+SE\s|CUANDO\b|CON\s+EL\s+VALOR|DEBE\b|SE\s+DEBE\b|DE\s+LA\s+A\b))"
)
_FIN_SCE_RE = re.compile(r"EN\s+CONSTANCIA|NOTA\s*1\s*:|FIRMA\s+REPRESENTANTE|FORMATO\s*(?:N[O°º]\.?\s*)?5\s*(?:\.\s*[124]|-?\s*[ABD])\b")
# La fila de total del listado: "SUMATORIA COLUMNA (F) $2.935.598.357", "TOTAL $ 4.412.531.329,82".
_TOTAL_SCE_RE = re.compile(
    r"(?:SUMATORIA\s+(?:DE\s+LA\s+)?COLUMNA\s*\(?\s*(?:F|\d{1,2})\s*\)?\s*[-:]?\s*\$?"
    r"|(?<!VALOR\s)(?<!VALOR\s\s)\bTOTAL(?:\s+(?:SCE|SALDOS?|GENERAL)[A-Z\s()]{0,40}?)?\s*:?\s*\$"
    r"|(?<!\()\bSCE\s*:?\s*\$"
    r"|VALOR\s+TOTAL\s+(?:DE\s+LOS\s+)?CONTRATOS\s+EN\s+EJECUCION[^$\d]{0,40}\$"
    # Resumen del aplicativo de Colombia Compra: "SALDO DE CONTRATOS EN EJECUCION $ 81.408.551".
    r"|SALDO\s+DE\s+(?:LOS\s+)?CONTRATOS\s+EN\s+EJECUCION\s*\$)"
    r"\s*(\d{1,3}(?:[.,]\d{3})+(?:[.,]\d{1,2})?)(?![\d.,]*\d)"
)


_MONTO = r"\d{1,3}(?:[.,]\d{3})+(?:[.,]\d{1,2})?|\d{4,}(?:[.,]\d{1,2})?"
# Inicio de fila: valor del contrato, plazo en meses y fecha de inicio.
_INICIO_FILA_RE = re.compile(r"\$\s*(?:" + _MONTO + r")\s+\d{1,3}(?:[.,]\d{1,2})?\s+\d{1,2}/\d{1,2}/\d{2,4}")
_FILA_SCE_RE = re.compile(
    r"\$\s*(?P<valor>" + _MONTO + r")\s+(?P<plazo>\d{1,3}(?:[.,]\d{1,2})?)\s+\d{1,2}/\d{1,2}/\d{2,4}\s+"
    r"(?P<part>\d{1,3}(?:[.,]\d{1,2})?)\s*%(?:\s+(?:SI|NO)){0,2}\s*\$?\s*(?P<diario>" + _MONTO + r")\s+"
    # Sin "$", el saldo tiene que traer separadores de miles (un "2025" suelto es otra columna).
    r"(?:\$\s*(?P<saldo>" + _MONTO + r")|(?P<saldo_sin>\d{1,3}(?:[.,]\d{3})+(?:[.,]\d{1,2})?))(?![\d.,]*\d)"
)

# La tabla parte el número del contrato entre columnas y lo deja en medio de
# la fila: "LOP $2.618.545.675,68 5 6/04/2026 50% NO NO 003 DE $17.456.971,17
# $881.577.044,15 2025". Solo se dejan pasar trozos cortos (ni montos ni
# porcentajes) y el saldo diario tiene que venir con "$".
_FILA_SCE_PARTIDA_RE = re.compile(
    r"\$\s*(?P<valor>" + _MONTO + r")\s+(?P<plazo>\d{1,3}(?:[.,]\d{1,2})?)\s+\d{1,2}/\d{1,2}/\d{2,4}\s+"
    r"(?P<part>\d{1,3}(?:[.,]\d{1,2})?)\s*%(?:\s+(?:SI|NO|[A-ZÑ]{1,8}|\d{1,4})){1,6}\s+"
    r"\$\s*(?P<diario>" + _MONTO + r")\s+\$\s*(?P<saldo>" + _MONTO + r")(?![\d.,]*\d)"
)

# Otra lectura del PDF saca primero el saldo: "$ 761.466.589,72 NO $ 14.952.434.852,59 18 11/11/2026 10% $ 27.689.694,17".
_FILA_SCE_SALDO_PRIMERO_RE = re.compile(
    r"\$\s*(?P<saldo>" + _MONTO + r")\s+(?:SI|NO)\s+\$\s*(?P<valor>" + _MONTO + r")\s+(?P<plazo>\d{1,3}(?:[.,]\d{1,2})?)\s+"
    r"\d{1,2}/\d{1,2}/\d{2,4}\s+(?P<part>\d{1,3}(?:[.,]\d{1,2})?)\s*%"
)


def _monto(texto: str) -> float | None:
    return _cifra(texto) if re.search(r"\d[.,]\d{3}", texto) else numero(texto)


def _suma_filas_sce(tramo: str) -> float | None:
    """Suma de la columna de saldo cuando el listado no trae la fila de
    total. Solo si se leyeron todas las filas y cada saldo cuadra con el
    valor del contrato por la participación."""
    # El PDF a veces separa el primer dígito: "$ 1 24.428.807,90", "$ 9 .505.190,77".
    tramo = re.sub(r"\$\s*(\d)\s+(\d{0,2}[.,]\d{3})", r"$ \1\2", tramo)
    tramo = re.sub(r"\$\s*(\d)\s+(\.\d{3})", r"$ \1\2", tramo)
    inicios = len(_INICIO_FILA_RE.findall(tramo))
    filas = list(_FILA_SCE_RE.finditer(tramo))
    if len(filas) != inicios:
        filas = list(_FILA_SCE_PARTIDA_RE.finditer(tramo))
    if len(filas) != inicios:
        filas = list(_FILA_SCE_SALDO_PRIMERO_RE.finditer(tramo))
    if not filas or len(filas) != inicios:
        return None
    total = 0.0
    for f in filas:
        grupos = f.groupdict()
        saldo = _monto(grupos.get("saldo") or grupos.get("saldo_sin") or "")
        valor, part = _monto(f["valor"]), numero(f["part"])
        diario = _monto(grupos["diario"]) if grupos.get("diario") else None
        if valor is None or saldo is None or part is None or not 0 < part <= 100 or saldo > valor * part / 100 * 1.01:
            return None
        # El saldo es el saldo diario por los días que faltan (y la
        # participación): uno menor que un día de obra está mal leído.
        if diario is not None and 0 < saldo < diario * part / 100 * 0.99:
            return None
        total += saldo
    return total

# Sin la palabra TOTAL (se pierde al leer la tabla): la celda del total va
# justo antes de la nota del formato "SE DEBE DILIGENCIAR UNICAMENTE…".
_TOTAL_ANTES_DE_NOTA_RE = re.compile(
    r"\$?\s*(\d{1,3}(?:[.,]\d{3})+(?:[.,]\d{1,2})?)\s*\(?\s*SE\s+DEBE\s+DILIGENCIAR"
)


def _tramos_sce(texto: str) -> list[str]:
    """Cada listado de contratos en ejecución del texto, hasta su cierre."""
    tramos, hasta = [], -1
    for m in _SECCION_SCE_RE.finditer(texto):
        if m.start() < hasta:
            continue  # el título repetido ("FORMATO 5.3 - LISTADO DE…") es el mismo listado
        fin = _FIN_SCE_RE.search(texto, m.end() + 20)
        hasta = fin.start() if fin and fin.start() - m.end() < 12000 else m.end() + 12000
        tramos.append(texto[m.start(): hasta])
    return tramos


def _listado_vacio(tramo: str) -> bool:
    """El listado trae sus columnas pero ninguna fila: ni valores ni fechas
    (si la tabla fuera una imagen, tampoco estarían las columnas)."""
    encabezado = re.search(r"OFERENTE|PROPONENTE|INTEGRANTE", tramo[:600])
    columnas = encabezado and re.search(r"VALOR\s+(?:TOTAL\s+)?DEL\s+CONTRATO", tramo) and re.search(r"SALDO", tramo)
    return bool(columnas) and not re.search(
        r"\d[.,]\d{3}[.,]\d{3}|\d{5,}|\d{1,2}/\d{1,2}/\d{2,4}|\d{4}-\d{2}-\d{2}|\$\s*\d", tramo
    )


def _total_del_tramo(tramo: str, permite_vacio: bool) -> float | None:
    """Saldo de un listado: 0 si dice que no hay contratos o está vacío; si
    no, el mayor entre la fila de total y la suma de las filas."""
    tramo = re.sub(r"\$\s*(\d)\s+(\d{0,2}[.,]\d{3})", r"$ \1\2", tramo)  # "$ 8 .411.673.217"
    tramo = re.sub(r"\$\s*(\d)\s+(\.\d{3})", r"$ \1\2", tramo)
    totales = [v for c in (_TOTAL_SCE_RE.findall(tramo) or _TOTAL_ANTES_DE_NOTA_RE.findall(tramo)) if (v := _cifra(c)) is not None]
    suma = _suma_filas_sce(tramo)
    if totales or suma is not None:
        return max(totales + ([suma] if suma is not None else []))
    if _SIN_CONTRATOS_RE.search(tramo) or (permite_vacio and _listado_vacio(tramo)):
        return 0.0
    return None


def _sce(texto: str, permite_vacio: bool = True) -> tuple[float | None, bool]:
    """(saldo, se usó IA) del listado de contratos en ejecución. Si el
    integrante trae varios listados (copias), el mayor (lo prudente), y
    todos tienen que poder leerse. `permite_vacio`: False con hojas de
    cálculo, donde una celda con fórmula sin valor guardado se lee vacía."""
    tramos = _tramos_sce(texto)
    if tramos:
        saldos = [_total_del_tramo(t, permite_vacio) for t in tramos]
        if all(v is not None for v in saldos):
            return max(saldos), False
    inicio = texto.find("EJECUCION")
    tramo = texto[max(0, inicio - 200): inicio + 6000] if inicio >= 0 else texto[:6000]
    if _SIN_CONTRATOS_RE.search(tramo):
        return 0.0, False
    respuesta = consultar_json(_INSTRUCCION_SCE, tramo)
    if not isinstance(respuesta, dict) or respuesta.get("sce") is None:
        return None, False
    try:
        valor = float(str(respuesta["sce"]).replace(",", ""))
    except ValueError:
        return None, False
    if valor == 0:
        return None, False  # un cero solo lo acepta la frase explícita de arriba
    # Anti-invención: el total tiene que estar escrito en el formato, como
    # cifra completa y junto a la palabra TOTAL (un saldo parcial inflaría la
    # capacidad residual).
    escrito = _numero_escrito(valor, tramo)
    if escrito is None:
        return None, False
    alrededor = tramo[max(0, escrito.start() - 120): escrito.end() + 40]
    return (valor, True) if "TOTAL" in alrededor else (None, False)


def _bloques_sce_hoja(contenido: bytes) -> list[tuple[str, float | None]]:
    """Formato 5C en Excel: cada listado de la hoja (una hoja puede traer el
    de varios integrantes, uno debajo del otro) con el texto que lo precede
    (dice de qué integrante es) y la suma de la columna del saldo del
    contrato en ejecución (F), o la fila de total si la trae. El valor es
    None si no hay filas o alguna celda es una fórmula sin valor guardado
    (se leería vacía)."""
    import openpyxl

    try:
        valores = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        formulas = openpyxl.load_workbook(io.BytesIO(contenido), data_only=False, read_only=True)
    except Exception:  # noqa: BLE001
        return []
    bloques = []
    for hoja in valores.worksheets:
        filas = list(hoja.iter_rows(values_only=True, max_row=600))
        crudas = list(formulas[hoja.title].iter_rows(values_only=True, max_row=600))
        contexto: list[str] = []
        i = 0
        while i < len(filas):
            textos = [normalizar(str(c or "")) for c in filas[i]]
            columna = next((j for j, t in enumerate(textos)
                            # Celda de encabezado, no una nota que la nombra.
                            if re.search(r"SALDO\s+DEL\s+CONTRATO\s+EN\s+EJECUCION", t) and "DIARIO" not in t
                            and len(t) < 160 and not t.lstrip().startswith("NOTA")), None)
            if columna is None:
                contexto.append(" ".join(t for t in textos if t))
                i += 1
                continue
            suma, n, total, valido, con_valor = 0.0, 0, None, True, False
            i += 1
            while i < len(filas):
                fila, cruda = filas[i], crudas[i] if i < len(crudas) else ()
                texto = normalizar(" ".join(str(c) for c in fila if c is not None))
                valor = fila[columna] if columna < len(fila) else None
                formula = cruda[columna] if columna < len(cruda) else None
                if isinstance(formula, str) and formula.startswith("=") and valor is None:
                    valido = False
                if re.search(r"\bTOTAL|SUMATORIA", texto):
                    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
                        total = float(valor)
                    i += 1
                    break
                if re.search(r"^\s*NOTA\b|SE\s+DEBE\s+DILIGENCIAR", texto):
                    break
                if isinstance(valor, (int, float)) and not isinstance(valor, bool):
                    if valor < 0:
                        valido = False
                    suma += float(valor)
                    n += 1
                if any(isinstance(c, (int, float)) and not isinstance(c, bool) and c >= 1_000_000
                       for j, c in enumerate(fila) if j != columna):
                    con_valor = True  # la fila trae un contrato (su valor en pesos)
                i += 1
            dato = max(suma, total or 0.0) if valido and (n or total is not None) else None
            if dato == 0 and con_valor:
                dato = None  # contratos con valor y saldo cero: fórmulas sin calcular o mal leídas
            bloques.append((" ".join(contexto[-15:]), dato))
            contexto = []
    return bloques


def _monto_celda(celda: str | None) -> float | None:
    """Cifra de una celda de tabla, aunque el renglón la parta ("$3.059.701.53 5,00")."""
    t = re.sub(r"[\s$]", "", celda or "").strip("()")
    if not re.fullmatch(r"\d[\d.,]*", t):
        return None
    return _cifra(t) if re.search(r"\d[.,]\d{3}", t) else numero(t)


def _duenio_de_la_tabla(contexto: str, integrantes: list[IntegranteTecnico]) -> IntegranteTecnico | None:
    """ "…INTEGRANTES: A (90 %), B (10 %) PARA A (90 %):": la tabla es del
    integrante nombrado después del último "PARA"."""
    ultimo = None
    for m in re.finditer(r"\bPARA\s+(.{4,90})", contexto):
        ultimo = m
    if ultimo is None:
        return None
    compacto = re.sub(r"\s+", "", ultimo.group(1))
    duenios = [i for i in integrantes if _nombre_escrito(i, compacto)]
    return duenios[0] if len(duenios) == 1 else None


def _bloques_sce_pdf(contenido: bytes) -> list[tuple[str, float | None]]:
    """Formato 5C en PDF con tabla dibujada: cada tabla del listado leída por
    celdas, con el texto que tiene justo encima (dice de qué integrante es) y
    la suma de la columna del saldo del contrato en ejecución, o su fila de
    total. None si una fila no se puede leer o su saldo no cuadra con el
    valor del contrato por la participación."""
    from motor.procesamiento.pdf_utils import abrir_pdf

    bloques = []
    with abrir_pdf(contenido) as pdf:
        for page in pdf.pages[:25]:
            if "EJECUCION" not in normalizar(page.extract_text() or ""):
                page.flush_cache()
                continue
            for tabla in page.find_tables():
                filas = [[" ".join((c or "").split()) for c in f] for f in tabla.extract()]
                encabezado = next((i for i, f in enumerate(filas[:4])
                                   if any(re.search(r"SALDO\s*(?:DEL\s*)?CONTRATO", normalizar(c)) for c in f)), None)
                if encabezado is None:
                    continue
                cab = [normalizar(c) for c in filas[encabezado]]
                col_saldo = next((j for j, c in enumerate(cab) if re.search(r"SALDO", c) and not re.search(r"DIARIO", c)
                                  and re.search(r"EJECU|CONTRATO", c)), None)
                col_valor = next((j for j, c in enumerate(cab) if re.search(r"VALOR", c)), None)
                col_part = next((j for j, c in enumerate(cab) if re.search(r"%|PARTICI", c)), None)
                if col_saldo is None:
                    continue
                arriba = page.crop((0, max(0, tabla.bbox[1] - 110), page.width, tabla.bbox[1])).extract_text() or ""
                suma, n, total, valido = 0.0, 0, None, True
                for f in filas[encabezado + 1:]:
                    texto = normalizar(" ".join(f))
                    if not texto.strip() or "FORMULA" in texto:
                        continue
                    saldo = _monto_celda(f[col_saldo] if col_saldo < len(f) else None)
                    if re.search(r"\bTOTAL|SUMATORIA", texto):
                        total = saldo if saldo is not None else next(
                            (v for c in reversed(f) if (v := _monto_celda(c)) is not None), None)
                        break
                    if saldo is None:
                        valido = False
                        continue
                    valor = _monto_celda(f[col_valor]) if col_valor is not None and col_valor < len(f) else None
                    part = numero(re.sub(r"[%\s]", "", f[col_part])) if col_part is not None and col_part < len(f) else None
                    if valor is None or part is None or not 0 < part <= 100 or saldo > valor * part / 100 * 1.01:
                        valido = False
                    suma += saldo
                    n += 1
                dato = max(suma, total or 0.0) if valido and (n or total is not None) else None
                bloques.append((normalizar(" ".join(arriba.split())), dato))
            page.flush_cache()
    return bloques


def _es_su_listado(integrante: IntegranteTecnico, tramo: str, del_archivo: bool, integrantes: list[IntegranteTecnico]) -> bool:
    """El listado es de este integrante: lo dice el propio formato ("PROPONENTE
    O INTEGRANTE: …"), lo nombra o trae su NIT; si no nombra a ninguno, es del
    dueño del archivo."""
    if len(integrantes) <= 1:
        return True
    cabeza = tramo[:700]
    # En las hojas de cálculo el rótulo puede venir después de las notas del
    # bloque anterior: se busca en todo el contexto.
    titulares = titulares_del_documento(cabeza) or titulares_del_documento(tramo[:4000])
    if titulares:
        elegido = integrante_del_titular(titulares, integrantes)
        if elegido is not None:
            return elegido is integrante
    if _del_integrante(integrante, cabeza, integrantes):
        return True
    nombra = any(_nit_escrito(i, cabeza) or _nombre_escrito(i, re.sub(r"\s+", "", cabeza)) for i in integrantes)
    return del_archivo and not nombra


def _tramos_sce_de(pdfs: dict[str, bytes], excels: dict[str, bytes], integrante: IntegranteTecnico,
                   integrantes: list[IntegranteTecnico], memoria: dict[str, str]) -> list[str]:
    """Listados del integrante releyendo los formatos fila por fila (solo se
    hace cuando la lectura normal no dejó ningún saldo)."""
    if not memoria:
        memoria.update(_textos_formato5(pdfs, excels, tablas=True))
    return [tramo for a, t in memoria.items() if a in pdfs for tramo in _tramos_sce(t)
            if _es_su_listado(integrante, tramo, True, integrantes)]


def residual_del_proponente(
    pdfs: dict[str, bytes], excels: dict[str, bytes], integrantes: list[IntegranteTecnico],
    lotes: list[LoteFinanciero], smmlv: float, plural: bool, estados: dict[str, str] | None = None,
    completos: dict[str, str] | None = None,
) -> Revision:
    """CRP frente a la capacidad residual de los lotes a los que se presenta."""
    exigida = sum(l.capacidad_residual_del_proceso or 0 for l in lotes)
    presupuesto = sum(l.presupuesto or 0 for l in lotes)
    if not lotes or not exigida or not presupuesto:
        return Revision(False, ["no se pudo calcular la capacidad residual del proceso (presupuesto, plazo o anticipo del lote)"])
    textos = _textos_formato5(pdfs, excels)
    if not textos:
        # Escaneado: se vuelve a leer fila por fila antes de darlo por perdido.
        textos = _textos_formato5(pdfs, excels, tablas=True)
    if not textos:
        return Revision(False, ["no se encontró el Formato 5 – Capacidad residual"])
    resultados: list[ResidualIntegrante] = []
    bloques_hoja = [b for contenido in excels.values() for b in _bloques_sce_hoja(contenido)]
    textos_tablas: dict[str, str] = {}
    bloques_pdf = []
    for archivo in textos:
        if archivo in pdfs and "EJECUCION" in textos[archivo]:
            try:
                bloques_pdf += [(archivo, c, v) for c, v in _bloques_sce_pdf(pdfs[archivo])]
            except Exception:  # noqa: BLE001
                pass
    for integrante in integrantes:
        r = ResidualIntegrante(integrante.nombre)
        suyos = {a: t for a, t in textos.items() if _del_integrante(integrante, t, integrantes)}
        todo = " ".join(suyos.values())
        # CO: mayor ingreso operacional, del estado de resultados de cada año
        # (lo que exige el pliego); si el Formato 5D declara menos, lo menor.
        suyos_estados = estados_del_integrante(integrante, estados or {}, integrantes, completos)
        de_estados = [v for t in suyos_estados.values() if (v := _ingreso_operacional(t))]
        declarados = [v for v in (numero(m) for m in _CO_RE.findall(todo)) if v]
        if de_estados:
            r.co = min(max(de_estados), min(declarados)) if declarados else max(de_estados)
        elif declarados:
            # No se leyó la fila del estado de resultados (tablas en columnas):
            # vale lo que declara el Formato 5D solo si esa cifra está escrita
            # tal cual en los estados financieros del integrante.
            verificados = [v for v in declarados if any(_numero_escrito(v, t) for t in suyos_estados.values())]
            if verificados:
                r.co = min(verificados)
        if r.co is None:
            r.faltas.append("no se leyó el ingreso operacional en sus estados financieros (CO)")
        # E: con el RUP (segmento 72), liquidado con el SMMLV del año.
        rup = integrante.rup
        participacion = (integrante.participacion if plural else 1.0) or None
        if rup is None or participacion is None:
            r.faltas.append("no se pudo calcular la experiencia (E): falta el RUP o la participación")
        else:
            valor = sum((e.valor_smmlv or 0) * (e.participacion if e.participacion is not None else 1.0)
                        for e in rup.experiencias.values() if any(c.startswith("72") for c in e.clases))
            r.e = valor * smmlv / (presupuesto * participacion)
            r.puntos_e = _puntaje(r.e, _E_TABLA, 120)
        # CF: liquidez del RUP (sin pasivo corriente: el mayor puntaje).
        f = rup.financiera if rup else None
        if f is None or f.activo_corriente is None or f.pasivo_corriente is None:
            r.faltas.append("no se leyó la liquidez del RUP (CF)")
        else:
            r.cf_liquidez = f.activo_corriente / f.pasivo_corriente if f.pasivo_corriente else None
            r.puntos_cf = _puntaje(r.cf_liquidez, _CF_TABLA, 40, indeterminado=40)
        # CT: profesionales del Formato 5B (si no se leen, 0: lo desfavorable).
        tecnica = re.search(r"CAPACIDAD\s+TECNICA(.{0,5000}?)(?:CONTRATOS\s+EN\s+EJECUCION|FORMATO\s*(?:NO\.?\s*)?5\.?\s*[C3D]|$)", todo)
        r.ct_profesionales = len(_PROFESION_RE.findall(tecnica.group(1))) if tecnica else 0
        r.puntos_ct = _puntaje(r.ct_profesionales, _CT_TABLA, 40)
        # SCE.
        # Cada listado del PDF es de quien lo encabeza ("INTEGRANTE: …"); si
        # no nombra a ningún integrante, del dueño del archivo.
        de_pdf = [tramo for a, t in textos.items() if a in pdfs for tramo in _tramos_sce(t)
                   if _es_su_listado(integrante, tramo, a in suyos, integrantes)]
        # Separados por el cierre del formato, para que no se fundan en uno.
        saldos = [_sce(" EN CONSTANCIA ".join(de_pdf))] if de_pdf else []
        # Tablas dibujadas del PDF leídas por celdas (el texto corrido las revuelve).
        saldos += [(v, False) for a, contexto, v in bloques_pdf
                   if _duenio_de_la_tabla(contexto, integrantes) is integrante
                   or _duenio_de_la_tabla(contexto, integrantes) is None and _del_integrante(integrante, contexto, integrantes)
                   or (a in suyos and not any(_nit_escrito(i, contexto) or _nombre_escrito(i, re.sub(r"\s+", "", contexto))
                                              for i in integrantes))]
        # En hojas de cálculo, cada listado por su columna y de quien lo encabece.
        saldos += [(v, False) for contexto, v in bloques_hoja
                   if _es_su_listado(integrante, contexto, False, integrantes)]
        # El PDF firmado y su copia en Excel son el mismo listado: vale el
        # mayor de los que se pudieron leer (lo prudente).
        if not any(v is not None for v, _ in saldos):
            # Con la lectura normal no apareció: se reintenta fila por fila.
            for tramo in _tramos_sce_de(pdfs, excels, integrante, integrantes, textos_tablas):
                saldos.append(_sce(tramo))
        leidos = [(v, ia) for v, ia in saldos if v is not None]
        if len(leidos) < len(saldos) and not de_pdf:
            leidos = []  # un listado en Excel ilegible: no se sabe si es el único
        r.sce, r.uso_ia = (max(v for v, _ in leidos), any(ia for _, ia in leidos)) if leidos else (None, False)
        if r.sce is None:
            r.faltas.append("no se leyó el saldo de los contratos en ejecución (SCE, Formato 5C)")
        cf_leida = f is not None and f.activo_corriente is not None and f.pasivo_corriente is not None
        if r.co is not None and r.e is not None and r.sce is not None and cf_leida:
            r.crp = r.co * (r.puntos_e + r.puntos_ct + r.puntos_cf) / 100 - r.sce
        resultados.append(r)
    detalle = {
        "exigida": exigida,
        "integrantes": [
            {"nombre": r.nombre, "co": r.co, "e": r.e, "puntos_e": r.puntos_e, "profesionales": r.ct_profesionales,
             "puntos_ct": r.puntos_ct, "liquidez": r.cf_liquidez, "puntos_cf": r.puntos_cf, "sce": r.sce, "crp": r.crp}
            for r in resultados
        ],
    }
    faltas = [f"{r.nombre}: {f}" for r in resultados for f in r.faltas]
    if any(r.crp is None for r in resultados):
        return Revision(False, faltas or ["no se pudo calcular la capacidad residual"], detalle)
    total = sum(r.crp for r in resultados)
    texto = f"capacidad residual del proponente ${total:,.0f}; se exige ${exigida:,.0f} para los lotes a los que se presenta"
    if any(r.uso_ia for r in resultados):
        texto += " (saldo de contratos en ejecución leído con IA local y verificado en el formato)"
    return decidir_residual(total, exigida, lotes, texto, detalle)


def decidir_residual(total: float, exigida: float, lotes: list[LoteFinanciero], texto: str, detalle: dict) -> Revision:
    """Cumple si la capacidad residual supera la exigida con la holgura. Si
    no, se habilita solo en los lotes de mayor valor que alcance a cubrir
    con esa holgura (3.11); el resto, a revisión."""
    detalle["crp"] = total
    if total >= exigida * HOLGURA:
        detalle["lotes_cubiertos"] = [l.nombre for l in lotes]
        return Revision(True, [texto], detalle)
    cubiertos, acumulado = [], 0.0
    for l in sorted(lotes, key=lambda x: -(x.presupuesto or 0)):
        if acumulado + (l.capacidad_residual_del_proceso or 0) <= total / HOLGURA:
            acumulado += l.capacidad_residual_del_proceso or 0
            cubiertos.append(l.nombre)
    detalle["lotes_cubiertos"] = cubiertos
    alcanza = f"; alcanzaría para: {', '.join(cubiertos)}" if cubiertos else ""
    if total >= exigida:
        # La entidad recalcula los saldos de los contratos en ejecución con
        # sus propias fechas (Nota 10 del Formato 5C) y puede darle menos.
        return Revision(False, [
            f"{texto}: alcanza por menos de {100 * (HOLGURA - 1):.0f} % de margen; confirma los saldos de los contratos "
            f"en ejecución (la entidad los recalcula) antes de habilitar{alcanza}"
        ], detalle)
    return Revision(False, [f"no alcanza: {texto}{alcanza}"], detalle)
