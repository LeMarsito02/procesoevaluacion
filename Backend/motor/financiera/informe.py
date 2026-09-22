"""Informe Excel de la evaluación financiera, con la forma del resumen que
usan las entidades: por proponente y lote, CUMPLE / NO CUMPLE / N/A, más
los indicadores, el capital de trabajo y la capacidad residual.

Un lote cumple si cumplen todos los requisitos que le aplican (los
generales y los de ese lote). Lo que ni el programa verificó ni una persona
decidió sale como "PENDIENTE": el informe nunca presenta como definitivo
algo sin revisar.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from motor.tecnica.informe import PENDIENTE, ResultadoInforme, _encabezado, _fila

NO_APLICA = "N/A"


def _estado(r: ResultadoInforme | None) -> str:
    if r is None:
        return PENDIENTE
    if ((r.resultado.detalle or {}).get("financiera") or {}).get("no_aplica") and r.resultado.cumple:
        return NO_APLICA
    if r.resultado.cumple:
        return "CUMPLE"
    return "NO CUMPLE" if r.decidido else PENDIENTE


def _lote(estados: list[str]) -> str:
    """Resultado del lote a partir de sus requisitos."""
    if NO_APLICA in estados:
        return NO_APLICA
    if "NO CUMPLE" in estados:
        return "NO CUMPLE"
    if PENDIENTE in estados:
        return PENDIENTE
    return "CUMPLE"


def _num(valor, dec: int = 3):
    return None if valor is None else round(float(valor), dec)


def generar_informe(
    codigo: str,
    objeto: str,
    lotes: list[tuple[int, str]],
    proponentes: list[tuple[str, str]],
    resultados: dict[tuple[str, int], ResultadoInforme],
    generales: list[int],
    por_lote: dict[int, list[int]],
    borrador: bool,
    titulos: dict[int, str] | None = None,
) -> bytes:
    """`lotes`: [(índice, nombre)]; `proponentes`: [(hoja, nombre)];
    `resultados`: {(hoja, número de requisito): resultado}; `generales`:
    requisitos que valen para todos los lotes; `por_lote`: índice del lote ->
    sus requisitos (capital de trabajo y capacidad residual)."""
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    resumen.cell(row=1, column=1, value=f"RESUMEN EVALUACIÓN FINANCIERA - {codigo}" + (" (BORRADOR)" if borrador else "")).font = Font(bold=True, size=13)
    resumen.cell(row=2, column=1, value=objeto).alignment = Alignment(wrap_text=True)
    _encabezado(resumen, 4, ["PROP.", "NOMBRE DEL PROPONENTE", *(nombre.upper() for _, nombre in lotes)])
    fila = 5
    for hoja, nombre in proponentes:
        comunes = [_estado(resultados.get((hoja, n))) for n in generales]
        por = []
        for indice, _ in lotes:
            propios = [_estado(resultados.get((hoja, n))) for n in por_lote.get(indice, [])]
            por.append(NO_APLICA if NO_APLICA in propios else _lote(comunes + propios))
        _fila(resumen, fila, [hoja, nombre, *por])
        fila += 1
    for j, ancho in enumerate([8, 50, *([16] * len(lotes))], 1):
        resumen.column_dimensions[get_column_letter(j)].width = ancho

    detalle = libro.create_sheet("Indicadores")
    columnas = ["PROP.", "LIQUIDEZ", "ENDEUDAMIENTO", "COBERTURA", "RENT. ACTIVO", "RENT. PATRIMONIO", "VALIDEZ CO"]
    for _, nombre in lotes:
        columnas += [f"CAPITAL DE TRABAJO {nombre.upper()}", f"K RESIDUAL {nombre.upper()}"]
    _encabezado(detalle, 1, columnas)
    fila = 2
    numeros_generales = sorted(generales)
    for hoja, _ in proponentes:
        def dato(numero: int) -> dict:
            r = resultados.get((hoja, numero))
            return ((r.resultado.detalle or {}).get("financiera") or {}) if r else {}

        indicadores = dato(numeros_generales[0]) if numeros_generales else {}
        organizacional = dato(numeros_generales[1]) if len(numeros_generales) > 1 else {}
        validez = _estado(resultados.get((hoja, numeros_generales[2]))) if len(numeros_generales) > 2 else PENDIENTE
        valores = [hoja, _num(indicadores.get("liquidez")), _num(indicadores.get("endeudamiento")),
                   _num(indicadores.get("cobertura"), 2), _num(organizacional.get("roa")), _num(organizacional.get("roe")), validez]
        for indice, _ in lotes:
            ct, k = (por_lote.get(indice, []) + [None, None])[:2]
            d_ct, d_k = (dato(ct) if ct else {}), (dato(k) if k else {})
            valores += [NO_APLICA if d_ct.get("no_aplica") else _num(d_ct.get("capital_de_trabajo"), 0),
                        NO_APLICA if d_k.get("no_aplica") else _num(d_k.get("crp"), 0)]
        _fila(detalle, fila, valores)
        fila += 1
    for j, ancho in enumerate([8, 11, 14, 12, 13, 16, 12, *([20] * (2 * len(lotes)))], 1):
        detalle.column_dimensions[get_column_letter(j)].width = ancho

    observaciones = libro.create_sheet("Observaciones")
    _encabezado(observaciones, 1, ["PROP.", "REQUISITO", "RESULTADO", "OBSERVACIÓN"])
    fila = 2
    for (hoja, numero), r in sorted(resultados.items()):
        texto = r.resultado.motivo or r.resultado.error
        if not texto:
            continue
        _fila(observaciones, fila, [hoja, (titulos or {}).get(numero, str(numero)), _estado(r), texto])
        fila += 1
    for j, ancho in enumerate([8, 40, 13, 110], 1):
        observaciones.column_dimensions[get_column_letter(j)].width = ancho

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
