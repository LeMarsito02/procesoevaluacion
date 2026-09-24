"""Informe consolidado del proceso: junta las tres áreas por lote.

Por cada proponente y lote dice si quedó habilitado en jurídica, técnica y
financiera, el puntaje que asigna el programa (los puntos que el pliego da a
cada factor, menos la reducción por obras civiles inconclusas) y el orden de
elegibilidad entre los habilitados.

Lo que ni el programa verificó ni una persona decidió sale como "PENDIENTE",
y un proponente con algo pendiente no recibe orden: el consolidado nunca
presenta como definitivo algo sin revisar.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from motor.financiera.informe import NO_APLICA, _estado as _estado_financiero, _lote as _resultado_del_lote
from motor.tecnica.informe import (
    COLUMNAS_PUNTAJE, OBRAS_INCONCLUSAS, OPCIONALES, PENDIENTE, ResultadoInforme, _encabezado, _fila, _puntos, _reduccion,
)


def _estado(r: ResultadoInforme | None) -> str:
    """CUMPLE / NO CUMPLE / N/A / PENDIENTE de un requisito."""
    return _estado_financiero(r)


def _area(resultados: dict[tuple[str, int], ResultadoInforme], hoja: str, generales: list[int], del_lote: list[int]) -> str:
    """Resultado del área en el lote. No aplica solo cuando los requisitos
    propios del lote no aplican (el proponente no se presentó a él); un
    requisito general que no aplica —el patrimonio, por ejemplo— no puede
    volver "N/A" todo el lote y tapar un "no cumple"."""
    if not generales and not del_lote:
        return NO_APLICA
    propios = [_estado(resultados.get((hoja, n))) for n in del_lote]
    if propios and all(e == NO_APLICA for e in propios):
        return NO_APLICA
    return _resultado_del_lote(propios + [_estado(resultados.get((hoja, n))) for n in generales])


def _puntaje(resultados: dict[tuple[str, int], ResultadoInforme], hoja: str) -> float | str:
    """El puntaje que asigna el programa: los factores del pliego menos la
    reducción por obras inconclusas. PENDIENTE si algún factor no está
    resuelto."""
    partes: list[float | str] = []
    for _, numeros in COLUMNAS_PUNTAJE:
        for n in numeros:
            if n in OPCIONALES and (hoja, n) not in resultados:
                partes.append(0.0)
            else:
                partes.append(_puntos(resultados.get((hoja, n))))
    partes.append(_reduccion(resultados.get((hoja, OBRAS_INCONCLUSAS))))
    if any(p == PENDIENTE for p in partes):
        return PENDIENTE
    return round(sum(p for p in partes if isinstance(p, (int, float))), 2)


def generar_informe(
    codigo: str,
    objeto: str,
    lotes: list[tuple[int, str]],
    proponentes: list[tuple[str, str]],
    resultados: dict[str, dict[tuple[str, int], ResultadoInforme]],
    requisitos: dict[str, dict[str, list[int]]],
    borrador: bool,
) -> bytes:
    """`resultados` y `requisitos` van por área ("juridica", "tecnica",
    "financiera"). `requisitos[area]` trae "generales" (valen para todos los
    lotes) y "lote_0", "lote_1"… con los de cada lote."""
    libro = Workbook()
    hoja_resumen = libro.active
    hoja_resumen.title = "Consolidado"
    hoja_resumen.cell(row=1, column=1, value=f"CONSOLIDADO DE LA EVALUACIÓN - {codigo}" + (" (BORRADOR)" if borrador else "")).font = Font(bold=True, size=13)
    hoja_resumen.cell(row=2, column=1, value=objeto).alignment = Alignment(wrap_text=True)
    fila = 4
    areas = [("juridica", "JURÍDICA"), ("tecnica", "TÉCNICA"), ("financiera", "FINANCIERA")]
    for indice, nombre in lotes:
        hoja_resumen.cell(row=fila, column=1, value=nombre.upper()).font = Font(bold=True)
        fila += 1
        _encabezado(hoja_resumen, fila, ["PROP.", "NOMBRE DEL PROPONENTE", *(t for _, t in areas), "HABILITADO",
                                         "PUNTAJE", "ORDEN DE ELEGIBILIDAD"])
        fila += 1
        filas = []
        for hoja, nombre_proponente in proponentes:
            estados = []
            for area, _ in areas:
                reparto = requisitos.get(area, {})
                estados.append(_area(resultados.get(area, {}), hoja, reparto.get("generales", []),
                                     reparto.get(f"lote_{indice}", [])))
            habilitado = NO_APLICA if NO_APLICA in estados else _resultado_del_lote(estados)
            puntaje = _puntaje(resultados.get("tecnica", {}), hoja) if habilitado == "CUMPLE" else ""
            filas.append([hoja, nombre_proponente, *estados, habilitado, puntaje])
        # El orden de elegibilidad solo entre los habilitados con puntaje en
        # firme. Los que empatan comparten puesto y se marcan: el desempate lo
        # decide la entidad con los criterios del pliego, no el programa.
        elegibles = sorted(
            [f for f in filas if f[5] == "CUMPLE" and isinstance(f[6], (int, float))],
            key=lambda f: -f[6],
        )
        orden: dict[int, object] = {}
        puesto = 0
        for i, f in enumerate(elegibles):
            if i == 0 or f[6] != elegibles[i - 1][6]:
                puesto = i + 1
            empatados = sum(1 for otro in elegibles if otro[6] == f[6])
            orden[id(f)] = f"{puesto} (empate entre {empatados})" if empatados > 1 else puesto
        for f in filas:
            _fila(hoja_resumen, fila, [*f, orden.get(id(f), PENDIENTE if f[5] == PENDIENTE else "")])
            fila += 1
        fila += 1
    for j, ancho in enumerate([8, 46, 14, 14, 14, 14, 12, 20], 1):
        hoja_resumen.column_dimensions[get_column_letter(j)].width = ancho

    detalle = libro.create_sheet("Pendientes")
    _encabezado(detalle, 1, ["PROP.", "ÁREA", "REQUISITO", "ESTADO", "MOTIVO"])
    fila = 2
    for area, titulo in areas:
        for (hoja, numero), r in sorted(resultados.get(area, {}).items()):
            estado = _estado(r)
            if estado != PENDIENTE:
                continue
            _fila(detalle, fila, [hoja, titulo, str(numero), estado, r.resultado.motivo or r.resultado.error or ""])
            fila += 1
    for j, ancho in enumerate([8, 14, 12, 13, 110], 1):
        detalle.column_dimensions[get_column_letter(j)].width = ancho

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
