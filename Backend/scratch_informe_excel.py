"""Veredictos de un informe de evaluación en Excel (plantilla del ICCU: una
hoja por proponente, un requisito cada 4 filas desde la 29, veredicto en la
columna L). SOLO para medir el acierto del programa; no alimenta su lógica."""
import json
import re
import sys

import openpyxl

FILA_INICIAL, PASO, REQUISITOS = 29, 4, 18


def main(ruta: str, salida: str) -> None:
    wb = openpyxl.load_workbook(ruta, data_only=True)
    proponentes = []
    for hoja in wb.sheetnames:
        if not re.fullmatch(r"P-\d+", hoja):
            continue
        ws = wb[hoja]
        requisitos = {}
        # La fila del requisito 1 cambia entre versiones de la plantilla (29 o 28).
        inicio = next((r for r in range(20, 40) if str(ws[f"B{r}"].value or "").strip() in ("1", "1.0")), FILA_INICIAL)
        for n in range(REQUISITOS):
            fila = inicio + n * PASO
            numero = ws[f"B{fila}"].value
            veredicto = str(ws[f"L{fila}"].value or "").strip().upper().replace("N.A", "N.A.").replace("N.A..", "N.A.")
            nota = " ".join(
                str(ws.cell(row=r, column=c).value).strip()
                for r in range(fila, fila + PASO) for c in range(14, 24)
                if ws.cell(row=r, column=c).value not in (None, "") and str(ws.cell(row=r, column=c).value).strip() not in ("SI", "NO", "N.A.")
            )
            requisitos[str(int(float(numero)) if re.fullmatch(r"\d+(\.0)?", str(numero or "")) else n + 1)] = {"veredicto": veredicto or "?", "nota": nota[:300]}
        proponentes.append({"hoja": hoja, "nombre": hoja, "requisitos": requisitos})
    json.dump(proponentes, open(salida, "w"), ensure_ascii=False, indent=1)
    import collections
    conteo = collections.Counter(r["veredicto"] for p in proponentes for r in p["requisitos"].values())
    print(f"{len(proponentes)} proponentes · veredictos {dict(conteo)}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
